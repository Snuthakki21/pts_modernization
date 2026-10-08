"""Editable, bounded component diagrams shared by management and framework decks."""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from .domain import require

RED=RGBColor(215,30,40)
INK=RGBColor(59,51,49)


def label(slide,text,x,y,w,h,size=16,bold=False):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    box.text_frame.word_wrap=True;box.text_frame.text=text
    for p in box.text_frame.paragraphs:
        p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=INK
    return box


def component_slide(prs,process_id):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(255,255,255)
    label(slide,'One process · one evidence trail',.6,.35,12.1,.7,27,True)
    label(slide,'Process '+process_id+' · source access remains read-only',.6,1.1,12.1,.5,15)
    nodes=[('Mainframe + Db2','COBOL · JCL · BMS\nDDL · exported records'),
           ('Copilot retrieves','Zowe CLI source\nApproved Db2 MCP'),
           ('Process folder','Frozen source + lineage\nSaved requirements Markdown'),
           ('Claude develops','Local analysis + code\nTests + independent review'),
           ('Coordinator verifies','One SME return\nImmutable tests + report gates'),
           ('Target + report','Python / SQLite candidate\nSource comparison + specific gaps')]
    positions=[(.6,2),(4.6,2),(8.6,2),(8.6,4.1),(4.6,4.1),(.6,4.1)]
    for (title,body),(x,y) in zip(nodes,positions):
        box=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(3.35),Inches(1.55))
        box.fill.solid();box.fill.fore_color.rgb=RGBColor(249,246,243);box.line.color.rgb=RED
        label(slide,title,x+.12,y+.1,3.1,.42,17,True)
        label(slide,body,x+.12,y+.56,3.1,.82,15)
    endpoints=[(3.95,2.76,4.5,2.76),(7.95,2.76,8.5,2.76),(10.27,3.55,10.27,4),
               (8.5,4.86,7.95,4.86),(4.5,4.86,3.95,4.86)]
    for x1,y1,x2,y2 in endpoints:
        connector=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
        connector.line.color.rgb=RED;connector.line.width=Pt(2)
        # Editable OOXML arrow; no raster diagram or external resources.
        line=connector.line._get_or_add_ln()
        from lxml import etree
        tail=etree.SubElement(line,'{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd');tail.set('type','triangle')
    label(slide,'Analyst: follow the guided next action. SME: open one review file, answer, Save and return it.\nExecutives: see verified scope, exclusions, remaining gaps, effort and actual AI credits.',.65,6.05,12,.92,14)
    return slide


def inspect_deck(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            require(shape.left>=0 and shape.top>=0 and shape.left+shape.width<=prs.slide_width and shape.top+shape.height<=prs.slide_height,'Deck geometry exceeds canvas')


def framework_deck(path):
    """Create an operator/executive briefing, without fabricated migration metrics."""
    from pptx import Presentation
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    component_slide(prs,'framework overview')
    def page(title,rows,note):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        label(slide,title,.6,.35,12.1,.75,27,True)
        for index,(heading,body) in enumerate(rows):
            y=1.55+index*.96
            label(slide,heading,.65,y,3.1,.82,17,True)
            label(slide,body,3.9,y,8.7,.82,16)
        label(slide,note,.65,6.42,12,.68,13)
    page('The analyst job aid',[
        ('1 · Save setup','Enter nonsecret MCP and Zowe choices. Save creates the workspace instructions and shows Add process.'),
        ('2 · Define process','Provide the ordered process Markdown. No exports yet starts a specific retrieval checkpoint.'),
        ('3 · Copilot retrieves','Copy the exact prompt; Copilot writes source and Db2 observations into that request inbox. Continue validates it.'),
        ('4 · Scope + Claude','Save default Yes or explicit No selections. Copy the local analysis, development and test prompt to Claude.'),
        ('5 · SME + results','SME opens one local review file, answers and Saves a returned file. Import with their name; inspect gaps and PPT.')],
        'Missing identities and unsupported semantics stop at named gates. The single human review cannot be supplied by an agent.')
    page('Why the factory matters',[
        ('Explain the process','One process folder preserves jobs, programs, screens, dependencies, source and current requirements.'),
        ('Simplify with evidence','Consolidate implementations only when each original rule maps to verified target behavior.'),
        ('Report honest progress','Original, selected No, verified and unverified rules remain separate. LOC reduction is a size metric.'),
        ('Plan remaining work','Measure development, validation, waiting and actual Copilot credits; forecast only from comparable completed pilots.'),
        ('Support decisions','Executives get a one-slide component view, legacy/target metrics, precise gaps and the next action.')],
        'Real application conversion, live access and observed mainframe parity are not certified by this framework briefing.')
    page('A stable source boundary, replaceable target',[
        ('Fixed source access','Read-only Zowe CLI for mainframe source; approved Db2 MCP for catalogs and exported records.'),
        ('Implemented today','Bounded Python record and BMS layout adapters; SQLite comparison storage and qualified schema candidates.'),
        ('Future targets','Python/Oracle, Python/BigQuery, Java and .NET are extension candidates, not implemented backends.'),
        ('Qualification contract','Versioned source model, selected-rule mappings, native target comparisons, randomized linked states and adversarial review.'),
        ('Architectural decisions','Use transaction, ownership, workload and operational evidence. A warehouse or another service is not an automatic replacement.')],
        'Native CICS controllers, business SQL, production transactions and future languages remain unverified until their adapters pass the existing gates.')
    inspect_deck(prs);prs.save(path)
    return path
