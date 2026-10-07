"""One frozen metric model drives CSV, workbook and editable PowerPoint."""
from .domain import path_is_link
import csv
import json
from pathlib import Path
from io import StringIO
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from .domain import encode, atomic_json, sha, require
from .ledger import now
from .coverage import build_coverage, write_coverage
from .executive import executive_summary, render_executive, inspect_executive, PRIMARY_REPORT
from .inventory import report_inventory
from .rule_inventory import build_rule_inventory, add_workbook, write_inventory, render_rule_summary

def metrics(ledger, doc, coverage=None, portfolio_model=None):
    coverage=coverage or build_coverage(doc,ledger.root)
    summary=coverage['summary']
    final_status='COMPLETED' if summary['completion_eligible'] and not doc.get('blockers') and not doc.get('cancel_requested') else 'COMPLETED_WITH_BLOCKERS'
    projected=portfolio_model if portfolio_model is not None else report_portfolio(ledger,doc,final_status)
    a=doc.get('analysis') or {'programs':{},'assets':[],'rules':[]}
    assets=[x for x in a['assets'] if x.get('selected',True)];run=doc['runs'][-1] if doc['runs'] else {'programs':{}}
    verified=0;matches=0;cases=0
    answers=(doc.get('answers') or {}).get('items',{})
    source_status={}
    for row in coverage['rows']:
        if row['disposition'] in ('non_executable','out_of_scope'):continue
        status=source_status.setdefault(row['source_path'],{'applicable':0,'verified':True})
        status['applicable']+=1;status['verified'] &= row['disposition'] in ('mapped_verified','platform_replaced_verified')
    for name,p in a['programs'].items():
        r=run['programs'].get(name)
        if not r:continue
        cases+=len(r['actual'])
        program_status=source_status.get(p['path'],{})
        if program_status.get('applicable',0)>0 and program_status.get('verified') and not r['differences'] and r['coverage']['complete'] and not p['blockers']:
            matches+=len(r['actual'])-len(r['differences'])
            verified+=sum(answers.get(rule['id'],{}).get('answer')=='Yes' and not answers.get(rule['id'],{}).get('correction') for rule in p['rules'])
    versions=doc.get('program_versions',{})
    target_loc=sum(sum(bool(x.strip()) and not x.lstrip().startswith('#') for x in ((ledger.root/'shared/target/python'/f'{v}.py').read_text(encoding='utf-8').splitlines() if (ledger.root/'shared/target/python'/f'{v}.py').is_file() else [])) for v in set(versions.values()))
    result={'process_id':doc['id'],'demo':doc['demo'],'report_final_status':final_status,
        'portfolio_completed_processes':projected['completed_processes'],
        'portfolio_basis':'Includes current report outcome upon atomic artifact acceptance; excludes demonstrations',
        'source_programs':len(a['programs']),
        'source_copybooks':sum(x['kind']=='copybook' for x in assets),
        'source_bms_screens':sum(len(x.get('screens',[])) for x in assets),'source_cics_transactions':None,
        'source_db2_table_references':len({t.upper() for x in assets for t in x.get('tables',[])}),
        'source_vsam_files':None,'inbound_interfaces':None,'outbound_interfaces':None,
        'target_python_programs':len(versions),'target_react_business_screens':0,'target_business_rest_apis':0,
        'source_code_loc':sum(x['loc']['code'] for x in assets),'source_physical_loc':sum(x['loc']['physical'] for x in assets),
        'target_program_code_loc':target_loc,'rules_documented':len(a['rules']),'rules_verified':verified,
        'known_rule_verification_percent':round(100*verified/len(a['rules']),2) if a['rules'] else None,
        'synthetic_cases':cases,'matching_cases':matches,'matching_cases_basis':'Matching cases credited only for intact, SME-confirmed whole-program evidence','unresolved_blockers':len(doc['blockers']),
        'unsupported_source_lines':sum(x['disposition']=='unsupported' for p in a['programs'].values() for x in p['coverage']),
        'source_inventory_files':summary['source_files'],'source_accounted_lines':summary['source_lines'],
        'source_in_scope_lines':summary['in_scope_lines'],'source_out_of_scope_lines':summary['out_of_scope_lines'],
        'source_applicable_lines':summary['applicable_lines'],'source_verified_applicable_lines':summary['verified_applicable_lines'],
        'source_applicable_line_verification_percent':summary['line_verification_percent'],
        'source_semantic_units':summary['semantic_units']['total'],'source_applicable_semantic_units':summary['semantic_units']['applicable'],
        'source_verified_semantic_units':summary['semantic_units']['verified'],
        'source_semantic_unit_verification_percent':summary['semantic_units']['verification_percent'],
        'source_blocked_lines':summary['dispositions']['blocked'],'source_non_executable_lines':summary['non_executable_lines'],
        'source_unverified_mapped_lines':summary['dispositions']['mapped_unverified']+summary['dispositions']['platform_replaced_unverified'],
        'source_platform_replaced_verified_lines':summary['dispositions']['platform_replaced_verified'],
        'source_full_accounting':summary['fully_accounted'],'source_completion_eligible':summary['completion_eligible'],
        'source_integrity_errors':len(summary['integrity_errors']),
        'observed_mainframe_parity':False,'evidence_basis':'SOURCE_DERIVED_EXPECTED',
        'verification_percent_denominator':'Extracted known rules only; unsupported and unknown rules are not silently excluded from completion gates.',
        'target_environment':'Non-production Python / SQLite; JSON record adapter'}
    if doc.get('transactions'):
        result['declared_online_transactions']=len(doc['transactions'])
        result['target_online_api_candidates']=len((doc.get('online_delivery') or {}).get('transactions',{}))
        result['verified_native_cics_replacements']=0
    with ledger.lock:documents=ledger.list()
    result['estate_inventory_json']=json.dumps(report_inventory(doc,documents,coverage),ensure_ascii=False,sort_keys=True)
    result['adapter_priorities_json']=json.dumps([{key:value for key,value in group.items() if key!='source_paths'} | {'source_files':len(group['source_paths']),'evidence':'coverage.json summary.adapter_groups'} for group in summary.get('adapter_groups',[])],ensure_ascii=False,sort_keys=True)
    validation=doc.get('logic_validation')
    if doc.get('logic_validation_min_records'):
        records=[value.get('coverage',{}).get('logic_validation') for value in run['programs'].values()]
        records=[value for value in records if isinstance(value,dict)]
        meeting=sum(value.get('logic_meeting_minimum',0) for value in records) if not summary['integrity_errors'] else 0
        validation={**(validation or {}),'contract_version':doc.get('fixture_contract_version',3),'minimum_distinct_records_per_logic':doc['logic_validation_min_records'],
                    'known_supported_logic':len(a['rules']),'logic_meeting_minimum':meeting,
                    'known_logic_missing_minimum':max(0,len(a['rules'])-meeting),
                    'unresolved_record_gaps':sum(value.get('gaps',0) for value in records),
                    'unsupported_source_lines':result['unsupported_source_lines'],
                    'unknown_legacy_logic_count':None,
                    'complete':bool(records) and not summary['integrity_errors'] and (validation or {}).get('complete',True) and all(value.get('complete') for value in records) and not a.get('blockers'),
                    'basis':'Distinct source predicate input states at execution; duplicates and unused-field padding do not count. Unsupported logic remains unresolved; records do not prove complete legacy parity.'}
    if validation is not None:result['logic_validation_json']=json.dumps(validation,ensure_ascii=False,sort_keys=True)
    return result

def portfolio(ledger):
    with ledger.lock:
        p=ledger.list()
    # Scope belongs to a process membership, not to the first globally registered
    # asset document. The same source version can be excluded by one process and
    # selected by another. Keep discovered exports separate from selected scope.
    discovered={};selected={}
    for process in p:
        for asset in (process.get('analysis') or {}).get('assets',[]):
            discovered[asset['id']]=asset
            if asset.get('selected',True):selected[asset['id']]=asset
    return {'processes':len(p),'unique_program_versions':sum(x['kind']=='cobol_program' for x in selected.values()),
        'unique_copybook_versions':sum(x['kind']=='copybook' for x in selected.values()),
        'discovered_program_versions':sum(x['kind']=='cobol_program' for x in discovered.values()),
        'discovered_copybook_versions':sum(x['kind']=='copybook' for x in discovered.values()),
        'source_scope_basis':'Unique versions selected by at least one non-demo process; discovered versions include exported out-of-scope assets.',
        'program_memberships':sum(len((x.get('analysis') or {}).get('programs',{})) for x in p),
        'completed_processes':sum(x['status']=='COMPLETED' for x in p),'demo_excluded':True}

def report_portfolio(ledger,doc,final_status):
    """Project only this report's accepted outcome; keep live portfolio unchanged.

    The coordinator commits these frozen metrics with terminal status only after
    mandatory artifact inspection. Failed or interrupted output remains a
    projection and never adds accepted history or live completion credit.
    """
    with ledger.lock:
        result=portfolio(ledger)
        if not doc.get('demo'):
            current=next((p for p in ledger.list() if p['id']==doc['id']),None)
            if current:
                result['completed_processes']+=int(final_status=='COMPLETED')-int(current['status']=='COMPLETED')
    result['basis']='Includes current report outcome upon atomic artifact acceptance; excludes demonstrations'
    return result


def generate_reports(ledger,doc,root,checkpoint=None,coverage=None):
    root=Path(root)
    names=('metrics.json','metrics.csv','metrics.xlsx','management.pptx','inspection.json',
           'program-insights.json','program-insights.html','factory.json','factory.html','economics.json','economics.html','economics.csv','coverage.json','coverage.csv','coverage.xlsx','coverage.html','rules.json','rules.csv','rules.html',PRIMARY_REPORT)
    require(not path_is_link(root) and not any(path_is_link(p) for p in root.parents),'Unsafe report output path')
    require(not any((root/name).exists() or path_is_link(root/name) for name in names),
            'Report evidence already exists; create a new report version')
    root.mkdir(parents=True,exist_ok=True)
    coverage=coverage if coverage is not None else build_coverage(doc,ledger.root,checkpoint=checkpoint)
    coverage_paths=write_coverage(coverage,root)
    factory_paths=[]
    factory=None
    if doc.get('factory_contract_version'):
        from .factory import factory_view, render_factory
        factory=factory_view(doc,coverage)
        atomic_json(root/'factory.json',factory)
        (root/'factory.html').write_text(render_factory(factory),encoding='utf-8')
        from .program_insights import program_insights, render_program_insights
        insights=program_insights(doc,coverage,complete=True)
        atomic_json(root/'program-insights.json',insights)
        (root/'program-insights.html').write_text(render_program_insights(insights),encoding='utf-8')
        factory_paths=[root/'factory.json',root/'factory.html',root/'program-insights.json',root/'program-insights.html']
    from .economics import economics_view, render_economics
    economics=economics_view(ledger,doc['id'])
    atomic_json(root/'economics.json',economics)
    (root/'economics.html').write_text(render_economics(economics),encoding='utf-8')
    economics_paths=[root/'economics.json',root/'economics.html',root/'economics.csv']
    pilot=economics['processes'][0]
    forecast=economics.get('forecast') or {}
    economics_summary={'as_of':economics['as_of'],'pilot_effort_hours':pilot['work']['total_effort_hours'],
        'service_hours':pilot['timing']['service_hours'],'framework_hours':economics['framework']['recorded_detail_hours'],
        'remaining_base_hours':forecast.get('remaining_effort_hours',{}).get('base'),
        'capacity_base_weeks':forecast.get('capacity_weeks',{}).get('base')}
    economics_rows=[['Measure','Value','Basis']]
    economics_rows.extend([[key,'Unknown' if value is None else value,'Frozen as of report generation; current reporting attempt may still be open'] for key,value in economics_summary.items()])
    economics_rows.extend([['AI '+r['provider']+' / '+r['account']+' / '+r['model']+' / '+r['unit'],r['quantity'],r['provenance']] for r in economics['usage']['quantities']])
    economics_rows.extend([['Forecast '+r['id']+' '+k,'Unknown' if v is None else v,r['basis']] for r in forecast.get('cohorts',[]) for k,v in r['remaining_effort_hours'].items()])
    # Imported labels are inert in spreadsheet applications.
    def cell(value):return "'"+value if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')) else value
    economics_rows=[[cell(v) for v in row] for row in economics_rows]
    out=StringIO();csv.writer(out).writerows(economics_rows);(root/'economics.csv').write_text(out.getvalue(),encoding='utf-8')
    rule_inventory=build_rule_inventory(doc,coverage,ledger.root)
    rule_paths=write_inventory(rule_inventory,root)
    from .comparison import render_comparison_report
    (root/'rules.html').write_text(render_comparison_report(rule_inventory),encoding='utf-8')
    rule_paths.append(root/'rules.html')
    final_status='COMPLETED' if coverage['summary']['completion_eligible'] and not doc.get('blockers') and not doc.get('cancel_requested') else 'COMPLETED_WITH_BLOCKERS'
    pf=report_portfolio(ledger,doc,final_status)
    m=metrics(ledger,doc,coverage,portfolio_model=pf);model={'created':now(),'metrics':m,'portfolio':pf,'blockers':doc['blockers'],'lineage':doc.get('lineage',(doc.get('analysis') or {}).get('graph',[]))}
    if factory is not None:model['factory']=factory
    model['inventory']=json.loads(m['estate_inventory_json'])
    model['economics']=economics
    m.update({'pilot_effort_hours':economics_summary['pilot_effort_hours'],'observed_service_hours':economics_summary['service_hours'],'framework_effort_hours':economics_summary['framework_hours'],'forecast_remaining_base_hours':economics_summary['remaining_base_hours']})
    model['executive_context']={key:doc.get(key) for key in ('id','name','status','demo','fixture_only','cancel_requested','packet_imported','verification_finished','manifest_hash','blockers')}
    model['executive_context']['economics_summary']=economics_summary
    if factory is not None:model['executive_context']['factory_report']=True
    model['executive_context']['rule_inventory']={k:v for k,v in rule_inventory.items() if k not in ('rules','requirements_comparison')}
    for category,counts in rule_inventory['summary'].items():
        for key in ('total','selected','converted_verified','blocked','implemented_unverified','excluded_by_requirements'):
            m[category+'_'+key]=counts[key]
    model['executive_context']['analysis']={'source_snapshot':(doc.get('analysis') or {}).get('source_snapshot')}
    model['executive_context']['artifacts']=[name for name in doc.get('artifacts',[]) if name=='analysis/source-analysis.json']
    model['executive']=executive_summary(model['executive_context'],m,coverage)
    atomic_json(root/'metrics.json',model)
    csvout=StringIO();writer=csv.writer(csvout);writer.writerow(['Metric','Value'])
    for k,v in m.items():writer.writerow([k,'Unknown' if v is None else v])
    (root/'metrics.csv').write_text(csvout.getvalue(), encoding='utf-8')
    book=Workbook();sheet=book.active;sheet.title='Metrics';sheet.append(['Metric','Value'])
    for k,v in m.items():sheet.append([k,'Unknown' if v is None else v])
    sheet.column_dimensions['A'].width=48;sheet.column_dimensions['B'].width=90;sheet.freeze_panes='B2'
    portfolio_sheet=book.create_sheet('Portfolio');portfolio_sheet.append(['Measure','Value'])
    for key,value in model['portfolio'].items():portfolio_sheet.append([key,value])
    portfolio_sheet.column_dimensions['A'].width=42;portfolio_sheet.column_dimensions['B'].width=95
    estate=book.create_sheet('Estate inventory')
    estate.append(['Category','Declared baseline (unverified)','Observed process','Converted process POC','Converted cumulative POC','Remaining vs baseline','Unique versions','Memberships'])
    for row in model['inventory']['rows']:
        label=row['label'];label="'"+label if label.lstrip().startswith(('=','+','-','@')) else label
        estate.append([label,row['baseline_count'],'Unknown' if row['process_observed_count'] is None else row['process_observed_count'],row['process_converted_count'],row['cumulative_converted_count'],row['remaining_vs_baseline'],row['portfolio_unique_versions'],row['portfolio_memberships']])
    estate.append(['Declared total',model['inventory']['declared_total']]);estate.append(['Category sum',model['inventory']['category_total']]);estate.append(['Unreconciled',model['inventory']['unreconciled_count']])
    estate.append(['Baseline status',model['inventory']['baseline_status']]);estate.append(['Count basis',model['inventory']['scope_boundary']])
    estate.freeze_panes='B2'
    estate.column_dimensions['A'].width=32
    history=book.create_sheet('History');history.append(['Process','Created','Programs','Verified rules','Acceptance'])
    accepted_history=ledger.history()
    for h in accepted_history:
        if h['document'].get('demo'):continue
        history.append([h['process_id'],h['created'],h['document']['source_programs'],h['document']['rules_verified'],'Accepted snapshot'])
    # Live effort snapshots change with reporting time; that is not another
    # accepted conversion. Compare stable conversion metrics for history rows.
    timing_metrics={'pilot_effort_hours','observed_service_hours','framework_effort_hours','forecast_remaining_base_hours'}
    conversion_metrics=lambda values:{k:v for k,v in values.items() if k not in timing_metrics}
    if not doc['demo'] and not any(h['process_id']==doc['id'] and encode(conversion_metrics(h['document']))==encode(conversion_metrics(m)) for h in accepted_history):
        history.append([doc['id'],model['created'],m['source_programs'],m['rules_verified'],'Current report upon atomic acceptance'])
    economics_sheet=book.create_sheet('Effort and forecast')
    for row in economics_rows:economics_sheet.append(row)
    economics_sheet.freeze_panes='B2';economics_sheet.column_dimensions['A'].width=60;economics_sheet.column_dimensions['B'].width=30;economics_sheet.column_dimensions['C'].width=85
    add_workbook(book,rule_inventory)
    if factory is not None:
        sheet=book.create_sheet('Factory capabilities');sheet.append(['Capability','State','Source count','Gap count'])
        for item in factory['capabilities']:sheet.append([item['label'],item['state'],item['source_count'],item['gap_count']])
        sheet=book.create_sheet('Online transactions');sheet.append(['Transaction','Program','Mapset','Map','API candidate','State','Native CICS verified'])
        for item in factory['transactions']:sheet.append([item['id'],item['program'],item['mapset'],item['map'],item['api'],item['state'],False])
    book.save(root/'metrics.xlsx');book.close()
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    def slide(title,rows,note,headers=None):
        s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(248,250,252)
        box=s.shapes.add_textbox(Inches(.6),Inches(.35),Inches(12.1),Inches(.7))
        p=box.text_frame.paragraphs[0];p.text=title;p.font.size=Pt(27);p.font.bold=True;p.font.color.rgb=RGBColor(15,35,60)
        headers=headers or ['Measure','Evidence / value']
        table=s.shapes.add_table(len(rows)+1,len(headers),Inches(.65),Inches(1.4),Inches(12),Inches(min(4.5,(len(rows)+1)*.52))).table
        if len(headers)==2:table.columns[0].width=Inches(7.6);table.columns[1].width=Inches(4.4)
        else:
            table.columns[0].width=Inches(4.2)
            for column in list(table.columns)[1:]:column.width=Inches(7.8/(len(headers)-1))
        for i,row in enumerate([headers]+rows):
            for j,value in enumerate(row):
                cell=table.cell(i,j);cell.text=str(value);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor(15,35,60) if i==0 else RGBColor(255,255,255)
                for par in cell.text_frame.paragraphs:
                    par.font.size=Pt(16 if len(headers)==2 else 13);par.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(25,45,65)
        b=s.shapes.add_textbox(Inches(.65),Inches(6.35),Inches(12),Inches(.8));b.text_frame.word_wrap=True
        b.text_frame.text=('Fictional test fixture. ' if doc.get('fixture_only') else '')+note
        for par in b.text_frame.paragraphs:par.font.size=Pt(13);par.font.color.rgb=RGBColor(70,85,105)
    inventory=model['inventory']
    inventory_rows=[[row['label'],row['baseline_count'],'Unknown' if row['process_observed_count'] is None else row['process_observed_count'],row['process_converted_count'],row['remaining_vs_baseline']] for row in inventory['rows']]
    slide('Estate inventory · '+doc['id'],inventory_rows,
          f"User-reported, unverified: declared {inventory['declared_total']:,}; category sum {inventory['category_total']:,}; unreconciled {inventory['unreconciled_count']:,}. Delta = baseline minus cumulative verified local POC assets. CICS screens are separate from transactions; staging locations provide no conversion credit.",
          ['Category','Baseline','Process','Converted','Delta'])
    slide('Before → after',[['BMS screens → React business screens',f"{m['source_bms_screens']} → 0"],['Local online API candidates (unverified native replacement)',m.get('target_online_api_candidates',0)],['All selected source code LOC → program Python LOC',f"{m['source_code_loc']} → {m['target_program_code_loc']}"],['CICS / VSAM / inbound / outbound counts','Unknown until evidenced'],['Target environment','Python / SQLite (non-production)']],'Workbench UI and its control endpoints are excluded from modernized business-screen/API counts. LOC is a size metric, not a parity metric.')
    slide('Rule verification',[['Business rules: verified / total',str(m['business_rule_converted_verified'])+' / '+str(m['business_rule_total'])],['Technical logic: verified / total',str(m['technical_logic_converted_verified'])+' / '+str(m['technical_logic_total'])],['Unclassified spans (rule count unknown)',m['unclassified_total']],['Unsupported source lines',m['unsupported_source_lines']],['Observed mainframe parity','NOT established']],'The percentage covers extracted known rules only. Unknown/unsupported behavior remains a blocker; passing synthetic tests is not proof of full legacy parity.')
    cs=coverage['summary']
    slide('Complete source accountability',[
        ['Frozen export files / physical lines',f"{cs['source_files']} / {cs['source_lines']}"],
        ['Applicable source lines: verified / total',f"{cs['verified_applicable_lines']} / {cs['applicable_lines']}"],
        ['Blocked / mapped but unverified lines',f"{m['source_blocked_lines']} / {m['source_unverified_mapped_lines']}"],
        ['Non-executable / explicitly out-of-scope lines',f"{cs['non_executable_lines']} / {cs['out_of_scope_lines']}"],
        ['Semantic units: verified / applicable',f"{cs['semantic_units']['verified']} / {cs['semantic_units']['applicable']}"],
        ['Verified platform replacement lines',m['source_platform_replaced_verified_lines']]],
        'coverage.json/CSV/XLSX/HTML preserve every original file and line with target spans, versions, tests and reasons. Line, semantic-unit and extracted-rule percentages use separate denominators.')
    pf=model['portfolio']
    slide('Portfolio progress',[['Production processes',pf['processes']],['Unique selected program versions',pf['unique_program_versions']],['Program memberships across processes',pf['program_memberships']],['Unique selected copybook versions',pf['unique_copybook_versions']],['Completed without blockers',pf['completed_processes']]],'Selected scope excludes unrelated exports; metrics.json lists discovered versions separately. Current completion is counted upon atomic report acceptance. Demonstrations are excluded; memberships preserve reuse.')
    slide('Evidence and decisions',[['Source snapshot',(doc.get('analysis') or {}).get('source_snapshot','Unavailable')[:20]],['SME review','One packet / one return per process'],['Expected vs actual','Frozen source IR vs executed Python'],['Open decisions',m['unresolved_blockers']],['Completion','WITH BLOCKERS' if m['report_final_status']=='COMPLETED_WITH_BLOCKERS' else 'Supported POC boundary verified']],'See metrics.json, source-analysis.json and each synthetic run for complete hashes, source spans, cases, actual results, differences and unresolved answers.')
    prs.save(root/'management.pptx')
    inspected=Presentation(root/'management.pptx');require(len(inspected.slides)==6,'Deck incomplete')
    for s in inspected.slides:
        for shape in s.shapes:require(shape.left>=0 and shape.top>=0 and shape.left+shape.width<=inspected.slide_width and shape.top+shape.height<=inspected.slide_height,'Deck geometry exceeds canvas')
    require(str(m['source_programs']) in '\n'.join(c.text for s in inspected.slides for sh in s.shapes if sh.has_table for row in sh.table.rows for c in row.cells),'Deck metric missing')
    executive_html=render_executive(model['executive'])
    (root/PRIMARY_REPORT).write_text(executive_html,encoding='utf-8')
    inspect_executive((root/PRIMARY_REPORT).read_text(encoding='utf-8'),model['executive'])
    paths=[root/PRIMARY_REPORT]+[root/f for f in ['metrics.json','metrics.csv','metrics.xlsx','management.pptx']]+coverage_paths+rule_paths+factory_paths+economics_paths
    atomic_json(root/'inspection.json',{'verified':True,'primary_report':PRIMARY_REPORT,'executive_schema_version':1,'executive_html_checked':True,'checks':['Executive metrics, safe links and collapsed disclosure','6 editable slides','full source accountability in JSON/CSV/XLSX/HTML','all shapes within canvas','source program metric present'],'powerpoint_render_checked':False,'sha256':{p.name:sha(p.read_bytes()) for p in paths}})
    return paths+[root/'inspection.json']
