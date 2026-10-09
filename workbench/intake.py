"""Small, explicit Markdown/Excel process intake. Missing facts stay unknown."""
import re
from io import BytesIO
from .domain import ValidationError, identity, require, checked_zip, encode, decode, sha, safe_path

HEADERS = ['Job order', 'Job', 'Step order', 'Step', 'Program or utility', 'Input files/tables', 'Output files/tables', 'Condition or dependency']
LINE_SEPARATORS = '\u0085\u2028\u2029'
MAX_PROCESS_MARKDOWN_BYTES = 1024 * 1024


def single_line(value):
    return not any(ord(c) < 32 or c in LINE_SEPARATORS for c in value)


def from_rows(pid, name, rows):
    identity(pid)
    require(isinstance(name, str) and 0 < len(name.strip()) <= 160, 'Process name is required and limited to 160 characters')
    require(single_line(name), 'Process name must be one line without control characters')
    require(isinstance(rows, (list, tuple)), 'Job steps must be a list of rows')
    require(0 < len(rows) <= 200, 'Supply between 1 and 200 job steps')
    jobs, seen, orders, methods = {}, set(), {}, {}
    for row in rows:
        require(isinstance(row, (list, tuple)) and len(row) == 8, 'Each job/step row must have eight columns')
        require(all(not isinstance(v,str) or (len(v)<=4096 and single_line(v) and '|' not in v) for v in row),
                'Intake cells must be single-line values without table delimiters, limited to 4096 characters')
        require(all(row[i] is None or isinstance(row[i], str) for i in (5, 6, 7)),
                'Input, output and condition cells must be text or blank')
        def order(value):
            require(type(value) is int or (isinstance(value,str) and re.fullmatch(r'[0-9]+',value.strip())),
                    'Job and step order must be positive integers; every populated row needs both orders')
            return int(value)
        jo, so = order(row[0]), order(row[2])
        require(0 < jo <= 1000 and 0 < so <= 1000, 'Job/step order out of range')
        require(all(isinstance(row[i],str) and row[i].strip() for i in (1,3,4)), 'Every row needs a job, step and program or utility name')
        job, step, program = map(lambda v: identity(v.strip()), (row[1], row[3], row[4]))
        method=job.lower().replace('-','_')
        require(method not in methods or methods[method]==job,
                'Job names collide after case/hyphen normalization: '+job)
        methods[method]=job
        require((job, so) not in seen, 'Duplicate step order within a job')
        require(not any(s['name'].upper() == step.upper() for s in jobs.get(job, {}).get('steps', [])), 'Duplicate step name within a job (case-insensitive)')
        seen.add((job, so))
        require(jo not in orders or orders[jo] == job, 'Different jobs cannot share a job order')
        orders[jo] = job
        if job in jobs: require(jobs[job]['order'] == jo, 'One job has conflicting orders')
        jobs.setdefault(job, {'name': job, 'order': jo, 'steps': []})['steps'].append({
            'name': step, 'order': so, 'program': program,
            'inputs': [x.strip() for x in re.split('[;,]', str(row[5] or '')) if x.strip()],
            'outputs': [x.strip() for x in re.split('[;,]', str(row[6] or '')) if x.strip()],
            'condition': (row[7] or '').strip() or 'Unknown'})
    result = sorted(jobs.values(), key=lambda j: j['order'])
    for j in result: j['steps'].sort(key=lambda s: s['order'])
    return {'id': pid, 'name': name.strip(), 'jobs': result}


def parse_manifest(text):
    require(isinstance(text, str), 'Provide a UTF-8 Process.md file')
    try: size = len(text.encode('utf-8'))
    except UnicodeError as exc: raise ValidationError('Process.md must be valid UTF-8 without unpaired surrogates') from exc
    require(size <= MAX_PROCESS_MARKDOWN_BYTES, 'Process.md exceeds 1 MiB; keep source code in Endeavor and provide process instructions here')
    require('\x00' not in text, 'Process.md must be text without NUL bytes')
    text = text.removeprefix('\ufeff')
    # Rendered examples and comments are not operator-supplied process facts.
    visible=[]; fence=None; comment=False
    for line in text.splitlines(keepends=True):
        if fence:
            if re.fullmatch(r'[ \t]{0,3}' + re.escape(fence[0]) + '{' + str(fence[1]) + r',}[ \t]*\r?\n?', line): fence=None
            visible.append('\n'); continue
        remainder=line; active=''
        while remainder:
            if comment:
                end=remainder.find('-->')
                if end < 0: remainder=''; break
                comment=False; remainder=remainder[end+3:]
            else:
                start=remainder.find('<!--')
                if start < 0: active+=remainder; break
                active+=remainder[:start]; remainder=remainder[start+4:]; comment=True
        marker=re.match(r'^[ \t]{0,3}(`{3,}|~{3,})', active)
        if marker: fence=(marker[1][0],len(marker[1])); visible.append('\n')
        else: visible.append(active if active.endswith('\n') else active+'\n')
    require(fence is None and not comment, 'Close Markdown example fences and comments before intake')
    text=''.join(visible)
    # Editors commonly bold labels or use headings. Normalize only recognized
    # identity/header syntax; prose remains data and never creates source facts.
    normalized=[]
    for line in text.splitlines():
        match = re.fullmatch(r'[ \t]*[-*]?[ \t]*(?:\*\*|__)(Process ID|Process name)(:)?(?:\*\*|__)[ \t]*(:)?[ \t]*(.*)', line, re.I)
        if match and bool(match[2]) != bool(match[3]):
            line = match[1]+': '+match[4]
        if line.strip().startswith('|'):
            cells=[v.strip() for v in line.strip().strip('|').split('|')]
            if [v.casefold() for v in cells] == [v.casefold() for v in HEADERS]:
                line='| '+' | '.join(HEADERS)+' |'
            elif [v.casefold() for v in cells] == ['transaction','program','mapset','map']:
                line='| Transaction | Program | Mapset | Map |'
        normalized.append(line)
    text='\n'.join(normalized)
    def attr(label):
        # Horizontal whitespace only: an empty attribute must never absorb
        # the next line (for example the job-table header) as its value.
        matches = re.findall(r'^[ \t]*[-*]?[ \t]*' + label + r'[ \t]*:[ \t]*([^\r\n]+)[ \t]*\r?$', text, re.M | re.I)
        require(len(matches)==1, f'Supply exactly one {label}')
        value=matches[0].strip()
        if '`' in value:
            require(value.startswith('`') and value.endswith('`') and value.count('`')==2,
                    'Use balanced inline-code wrappers for '+label)
            value=value[1:-1].strip()
        return value
    online_headers = ['Transaction', 'Program', 'Mapset', 'Map']
    table_lines = [line for line in text.splitlines() if line.strip().startswith('|')]
    if table_lines and [x.strip() for x in table_lines[0].strip().strip('|').split('|')] == online_headers:
        pid, name = attr('Process ID'), attr('Process name')
        identity(pid)
        require(0 < len(name) <= 160 and single_line(name), 'Invalid process name')
        transactions=[]; seen=set()
        for line in table_lines[1:]:
            cells=[x.strip() for x in line.strip().strip('|').split('|')]
            # Preserve empty trailing cells, unlike strip('|').
            cells=[x.strip() for x in line.strip()[1:].removesuffix('|').split('|')]
            if len(cells)==4 and all(re.fullmatch(r':?-{3,}:?',v) for v in cells):continue
            require(len(cells)==4 and all(single_line(v) and len(v)<=80 for v in cells), 'Online rows need Transaction, Program, Mapset and Map')
            tx, program, mapset, screen = [v.upper() for v in cells]
            require(bool(re.fullmatch(r'[A-Z0-9]{1,4}',tx)), 'Transaction IDs need 1–4 alphanumeric characters')
            identity(program)
            require(tx not in seen, 'Duplicate transaction identity');seen.add(tx)
            require(bool(mapset)==bool(screen), 'Supply both Mapset and Map, or leave both unknown')
            if mapset:identity(mapset);identity(screen)
            transactions.append({'id':tx,'program':program,'mapset':mapset or None,'map':screen or None})
        require(0<len(transactions)<=200, 'Supply 1–200 transactions')
        return {'id':pid,'name':name,'jobs':[],'transactions':transactions,'workload':'online'}
    has_batch_header=any([v.strip() for v in line.strip().strip('|').split('|')]==HEADERS
                         for line in text.splitlines() if line.strip().startswith('|'))
    if not has_batch_header and _has_job_roots(text):
        return _parse_job_roots(text, attr('Process ID'), attr('Process name'))
    if not table_lines:
        attr('Process ID');attr('Process name')
        raise ValidationError('Add a Jobs heading and one entry job per line; use the downloadable Process.md template')
    rows = [];header=False
    for line in text.splitlines():
        if not line.strip().startswith('|'): continue
        card=line.strip()[1:]
        if card.endswith('|'):card=card[:-1]
        cells = [v.strip() for v in card.split('|')]
        if cells==HEADERS:
            require(not header, 'Supply one job/step table; repeated headers are ambiguous')
            header=True;continue
        require(header, 'Provide the exact job/step table headers before every data row')
        if len(cells)==8 and all(re.fullmatch(r':?-{3,}:?',v) for v in cells):continue
        require(any(cells), 'Empty table rows are ambiguous; remove the row or provide its job and step')
        rows.append(cells)
    require(header, 'Provide the exact job/step table headers')
    return from_rows(attr('Process ID'), attr('Process name'), rows)


def _has_job_roots(text):
    return bool(re.search(r'^[ \t]*(?:#{1,6}[ \t]+)?(?:Entry jobs|Jobs)(?:[ \t]*\([^\n]*\))?[ \t]*:?[ \t]*$', text, re.M | re.I)
                or re.search(r'^[ \t]*[-*]?[ \t]*(?:Job|Jobs|Entry jobs)[ \t]*:[ \t]*\S', text, re.M | re.I))


def _root_name(value):
    # A single named member may be followed by an attributed explanatory note.
    # Never mine arbitrary prose or guess a member from a filename/path.
    value=value.strip()
    if value.startswith('`'):
        match=re.fullmatch(r'`([^`]+)`(?:[ \t]*(?:[:\u2014\u2013]| - )[ \t]*.*)?',value)
        require(match is not None, 'Use one job member name per Jobs line, for example 1. REFJOB')
        value=match[1]
    else:
        value=re.split(r'[ \t]*(?:[:\u2014\u2013]| - )[ \t]*',value,maxsplit=1)[0].strip()
    require(value.casefold() not in {'unknown','tbd','n/a','none','replacejob','replacejob2','your_job'}, 'Provide at least one known entry job member; missing names stay unresolved')
    try:return identity(value)
    except ValidationError as exc:raise ValidationError('Use one exact job member name per Jobs line; put explanations under Notes, for example 1. REFJOB') from exc


def _parse_job_roots(text, pid, name):
    identity(pid)
    require(0 < len(name) <= 160 and single_line(name), 'Process name is required and limited to 160 characters')
    roots=[]; active=False; mode=None; seen_section=False
    for line in text.splitlines():
        stripped=line.strip()
        heading=re.fullmatch(r'(?:#{1,6}[ \t]+)?(?:Entry jobs|Jobs)(?:[ \t]*\([^\n]*\))?[ \t]*:?',stripped,re.I)
        inline=re.fullmatch(r'[-*]?[ \t]*(?:Job|Jobs|Entry jobs)[ \t]*:[ \t]*(.+)',stripped,re.I)
        if heading or inline:
            require(not seen_section, 'Use one Jobs list; multiple lists are ambiguous')
            active=True;seen_section=True
            if inline:
                values=re.split(r'[;,]',inline[1]);require(all(v.strip() for v in values), 'Remove an empty job name from the Jobs list')
                roots.extend((i,_root_name(v)) for i,v in enumerate(values,1));active=False
            continue
        if not active or not stripped:continue
        if stripped.startswith('#'):active=False;continue
        numbered=re.fullmatch(r'([0-9]+)[.)][ \t]+(.+)',stripped)
        bullet=re.fullmatch(r'[-*+][ \t]+(.+)',stripped)
        # Plain single member lines also work in a Jobs section. All other
        # prose belongs under Notes, rather than silently becoming an identity.
        kind='numbered' if numbered else 'listed'
        require(mode is None or mode==kind, 'Use all numbered jobs or all unnumbered jobs; mixed order is ambiguous')
        mode=kind
        order=int(numbered[1]) if numbered else len(roots)+1
        value=numbered[2] if numbered else bullet[1] if bullet else stripped
        require(0 < order <= 1000 and (not roots or order > roots[-1][0]), 'Jobs must have distinct ascending positive run orders, up to 1000')
        roots.append((order,_root_name(value)))
        require(len(roots)<=1000,'Supply at most 1000 entry jobs; source steps are discovered from their JCL')
    require(roots, 'Add the entry job member under Jobs; the factory discovers its steps and dependencies')
    seen=set();methods=set()
    for _,job in roots:
        method=job.casefold().replace('-','_')
        require(job.casefold() not in seen and method not in methods,'Jobs list contains a duplicate or case-normalized member collision')
        seen.add(job.casefold());methods.add(method)
    return {'id':pid,'name':name,'jobs':[{'name':job,'order':order,'steps':[]} for order,job in roots],
            'process_intake_version':2}


def verify_job_plan(doc, base, reader=None):
    """Validate the exact frozen derivation separately from immutable root intake.

    Both Coordinator transitions and report coverage use this same contract.
    It verifies a registered source-derived plan; it does not derive new jobs,
    grant semantic support or turn source-derived checks into observed parity.
    """
    require(isinstance(doc,dict), 'Job plan requires a process document')
    declared=doc.get('declared_jobs')
    require(isinstance(declared,list), 'Declared jobs are missing from the frozen Process.md intake')
    relative=doc.get('job_plan_artifact')
    if not relative:
        require(encode(doc.get('jobs'))==encode(declared), 'Jobs changed before source-derived planning')
        require(not doc.get('job_plan_gaps'), 'Job plan gaps exist without their frozen plan')
        return None
    require(isinstance(doc.get('artifacts'),list) and relative in doc['artifacts']
            and isinstance(doc.get('artifact_hashes'),dict), 'Job plan is not registered')
    path=safe_path(base,relative)
    if reader is None:
        from .limits import MAX_JSON_DOCUMENT_BYTES
        require(path.is_file() and path.stat().st_size<=MAX_JSON_DOCUMENT_BYTES, 'Frozen job plan exceeds the bounded reader limit')
        with path.open('rb') as stream: raw=stream.read(MAX_JSON_DOCUMENT_BYTES+1)
        require(len(raw)<=MAX_JSON_DOCUMENT_BYTES, 'Frozen job plan changed or exceeds the bounded reader limit')
    else:raw=reader(path)
    digest=sha(raw)
    require(relative=='analysis/job-plan-'+digest+'.json' and digest==doc.get('artifact_hashes',{}).get(relative), 'Frozen job plan changed or has a noncanonical identity')
    plan=decode(raw)
    require(isinstance(plan,dict) and type(plan.get('schema_version')) is int and plan['schema_version']==1
            and plan.get('applicable') is True and plan.get('basis')=='deterministic_source_jcl', 'Unsupported frozen source job plan')
    require(plan.get('process_id')==doc.get('id') and encode(plan.get('jobs'))==encode(doc.get('jobs')), 'Derived jobs differ from the frozen source job plan')
    sources=plan.get('source_files')
    require(isinstance(doc.get('source_files'),dict), 'Process source baseline is missing from the job plan binding')
    require(isinstance(sources,dict) and all(isinstance(path,str) and isinstance(digest,str) and re.fullmatch(r'[a-f0-9]{64}',digest)
            and doc.get('source_files',{}).get(path)==digest for path,digest in sources.items()), 'Derived job plan references a different source version')
    require(plan.get('source_binding_hash')==sha(encode({'declared_jobs':declared,'source_files':sources})), 'Job plan declared/source binding changed')
    gaps=plan.get('gaps')
    require(isinstance(gaps,list) and all(isinstance(gap,dict) and isinstance(gap.get('reason'),str) for gap in gaps), 'Frozen job plan gaps are malformed')
    projection=[{'kind':'unsupported_jcl','message':gap['reason'],'object':gap,'path':gap.get('path'),'line':gap.get('line')} for gap in gaps]
    require(encode(projection)==encode(doc.get('job_plan_gaps',[])), 'Job plan gaps differ from their frozen source obligations')
    return plan


def parse_intake_xlsx(data):
    from zipfile import BadZipFile
    try:
        return _parse_intake_xlsx(data)
    except ValidationError:
        raise
    except (SyntaxError, KeyError, TypeError, ValueError, BadZipFile, RuntimeError, NotImplementedError) as exc:
        raise ValidationError('Invalid XLSX intake structure; use the supplied Intake template and repair malformed workbook XML') from exc


def _parse_intake_xlsx(data):
    from openpyxl import load_workbook
    from xml.etree.ElementTree import fromstring
    checked_zip(data).close()
    # Validate actual cell locations before materializing a worksheet. Streaming
    # readers trust declared dimensions, which can hide jobs or remote cells.
    declared=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
    try:
        require('Intake' in declared.sheetnames,'Workbook needs an Intake sheet')
        sheet=declared['Intake'];dimensions=(sheet.max_row,sheet.max_column)
        require(all(type(n) is int for n in dimensions) and dimensions[0]<=205 and dimensions[1]<=8,'Intake worksheet dimensions are missing or exceed bounds')
        with checked_zip(data) as archive:
            xml=fromstring(archive.read(sheet._worksheet_path))
        namespace='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        dimension_nodes=xml.findall(namespace+'dimension')
        require(len(dimension_nodes)==1, 'Intake worksheet needs one dimension declaration')
        dimension_ref=dimension_nodes[0].get('ref','')
        require(re.fullmatch(r'A1:[A-H][1-9][0-9]{0,2}',dimension_ref) is not None,
                'Intake worksheet dimensions must include the metadata origin A1')
        require(len(xml.findall(namespace+'sheetData'))==1, 'Intake worksheet needs one sheetData section')
        require(not xml.findall(namespace+'mergeCells/'+namespace+'mergeCell'), 'Merged intake cells are ambiguous')
        seen=set();row_ids=set()
        for row in xml.findall(namespace+'sheetData/'+namespace+'row'):
            position=row.get('r','')
            require(position.isdigit() and 1<=int(position)<=205 and position not in row_ids,'Actual intake row exceeds dimensions/bounds or is duplicated')
            row_ids.add(position)
            for cell in row.findall(namespace+'c'):
                require(all(len(cell.findall(namespace+tag))<=1 for tag in ('v','is','f')),
                        'Intake cells must not contain duplicate value elements')
                require(not (cell.find(namespace+'v') is not None and cell.find(namespace+'is') is not None),
                        'Intake cell has conflicting value representations')
                ref=cell.get('r','');match=re.fullmatch(r'([A-H])([1-9][0-9]{0,2})',ref)
                require(match is not None and int(match[2])<=205 and match[2]==position and ref not in seen,'Actual intake cell exceeds dimensions/bounds or is duplicated')
                seen.add(ref)
                require(int(match[2])<=dimensions[0] and ord(match[1])-ord('A')+1<=dimensions[1],'Actual intake cells exceed the declared worksheet dimensions')
    finally:declared.close()
    book = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=False)
    try:
        ws = book['Intake']
        require((ws.max_row,ws.max_column)==dimensions,'Declared and actual intake worksheet dimensions disagree')
        for row in ws:
            require(all(c.data_type != 'f' for c in row), 'Formulas are not accepted in intake')
            require(all(c.data_type != 'e' for c in row), 'Excel error cells are not accepted in intake')
        require(ws['A1'].value=='Process ID' and ws['A2'].value=='Process name', 'Intake metadata labels changed')
        require(all(ws.cell(r,c).value is None or isinstance(ws.cell(r,c).value,str) and not ws.cell(r,c).value.strip()
                    for r in range(1,4) for c in range(1,9) if r==3 or c>2),
                'Unexpected populated cells outside the intake metadata and job table')
        require([ws.cell(4,c).value for c in range(1,9)] == HEADERS, 'Intake headers changed')
        rows = []
        for r in range(5,ws.max_row+1):
            values=[ws.cell(r,c).value for c in range(1,9)]
            if any(v is not None and (not isinstance(v,str) or v.strip()) for v in values):rows.append(values)
        return from_rows(ws['B1'].value, ws['B2'].value, rows)
    finally: book.close()
