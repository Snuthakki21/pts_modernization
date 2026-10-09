"""One editable, plain-language SME packet with immutable question identities."""
from .domain import path_is_link
from io import BytesIO
from html.parser import HTMLParser
from pathlib import Path
import html
import json
import re
import tempfile
import shutil
from .domain import require, sha, encode, decode, checked_zip, write_new, ValidationError, MAX_UPLOAD

GLOBAL_QUESTIONS = [
 ('G_SCOPE','Does the listed program/job/step inventory cover this process?'),
 ('G_LAYOUT','Are the listed field names, lengths, numeric limits and record groups correct?'),
 ('G_ORDER','Are the listed job/step sequence, conditions, and input/output links correct?'),
 ('G_MATCH','Are the listed matching relationships and intentional missing-match cases interpreted correctly?'),
 ('G_TARGET','Is Python with SQLite/JSON files acceptable for this non-production POC?')]

CHECKLIST_HEADER = ['Item ID','What we understood','Source evidence','Category','Answer','If No, what should it be?','Reviewer']
CONTEXT_CHUNK_SIZE = 16000
CHECKLIST_TEXT_LIMIT = 4000
SME_ITEM_LIMIT = 2000
TECHNICAL_GAP_KINDS = frozenset(('unsupported_source','unsupported_jcl','unsupported_utility','source_classification'))
PROGRAM_LIST_TEXT_LIMIT = 240
# Large portable HTML has no ZIP expansion; legacy workbook limits stay separate.
MAX_HTML_REVIEW_BYTES = 128 * 1024 * 1024

# Review wording, not semantic support claims. These categories are deliberately
# conservative; unknown or compound obligations are never guessed from messages.
CONSTRUCT_ASSUMPTIONS = {
    'iteration': 'Confirm loop bounds, termination, ordering and repeated effects.',
    'embedded_sql': 'Confirm SQL inputs, results, NULL handling, errors and transaction boundaries.',
    'cics': 'Confirm transaction, screen, state, resource and error behavior.',
    'file_io': 'Confirm record layouts, access order, keys, EOF, missing records and I/O errors.',
    'arithmetic': 'Confirm precision, rounding, overflow and arithmetic error behavior.',
    'string_handling': 'Confirm string boundaries, delimiters, padding and overflow behavior.',
    'control_flow': 'Confirm branches, transfers of control, scope and termination.',
    'data_layout': 'Confirm storage, record groups, occurrences, redefinitions and condition names.',
    'numeric_representation': 'Confirm signs, decimals, binary/packed representation and encoding.',
    'initialization': 'Confirm initial values, storage lifetime and reset behavior.',
    'procedure_call': 'Confirm called program bindings, parameters, state and return behavior.',
    'copy_dependency': 'Confirm the exact COPY member, replacements and included declarations or logic.',
    'terminal_io': 'Confirm displayed or accepted values, devices and error behavior.',
    'jcl_io': 'Confirm DD bindings, allocation, disposition, records and dataset lifecycle.',
    'jcl_control': 'Confirm step order, conditions, symbols, procedure expansion and restart behavior.',
    'utility': 'Confirm the invoked utility, control cards, inputs, outputs and return codes.',
    'source_classification': 'Confirm the authoritative source type and resolve classification conflicts.',
    'unlocated_source': 'Provide the missing source or exact source binding before interpreting behavior.',
    'unknown': 'Identify the source construct and required behavior from the complete referenced evidence.',
    'mixed': 'Review every construct in this compound obligation and its interactions.',
}


def _construct_families(source):
    """Collect source constructs without losing later verbs on the same line."""
    from .source import normalized_lines
    kind, line = normalized_lines(source)[0] if source else ('blank', '')
    if kind != 'code':
        return {'unknown'}
    line = re.sub(r'"(?:[^"\n]|"")*"|\'(?:[^\'\n]|\'\')*\'', ' ', line).upper().strip()
    if line.startswith('//*'):
        return {'unknown'}
    if re.match(r'//\S*\s+DD(?![A-Z0-9-])', line):
        return {'jcl_io'}
    if re.match(r'//\S*\s+(?:JOB|EXEC|PROC|PEND|INCLUDE|SET|IF|ELSE|ENDIF)(?![A-Z0-9-])', line):
        return {'jcl_control'}
    # Treat SQL/CICS bodies as their own language, so SQL CALL or CICS READ
    # does not manufacture an additional COBOL verb. Scan after END-EXEC too.
    families = set()
    exec_pattern = r'(?<![A-Z0-9-])EXEC\s+(SQL|CICS)(?![A-Z0-9-])'
    while match := re.search(exec_pattern, line):
        families.add('embedded_sql' if match[1] == 'SQL' else 'cics')
        end = re.search(r'(?<![A-Z0-9-])END-EXEC(?![A-Z0-9-])', line[match.end():])
        stop = match.end()+end.end() if end else len(line)
        line = line[:match.start()]+' '*(stop-match.start())+line[stop:]
    line = re.sub(r'^[\s.]+', '', line).strip()
    if not line:
        return families or {'unknown'}
    patterns = (
        ('iteration', r'(?:PERFORM|END-PERFORM|SEARCH|END-SEARCH)(?![A-Z0-9-])'),
        ('file_io', r'(?:OPEN|CLOSE|READ|WRITE|REWRITE|DELETE|START|SORT|MERGE|RELEASE|RETURN)(?![A-Z0-9-])'),
        ('arithmetic', r'(?:ADD|SUBTRACT|MULTIPLY|DIVIDE|COMPUTE)(?![A-Z0-9-])'),
        ('string_handling', r'(?:STRING|UNSTRING|INSPECT)(?![A-Z0-9-])'),
        ('control_flow', r'(?:IF|ELSE|END-IF|EVALUATE|WHEN|END-EVALUATE|GO\s+TO|ALTER|EXIT|GOBACK|STOP\s+RUN)(?![A-Z0-9-])'),
        ('initialization', r'INITIALIZE(?![A-Z0-9-])'),
        ('procedure_call', r'(?:CALL|CANCEL)(?![A-Z0-9-])'),
        ('copy_dependency', r'COPY(?![A-Z0-9-])'),
        ('terminal_io', r'(?:DISPLAY|ACCEPT)(?![A-Z0-9-])'),
    )
    if re.match(r'\d{1,2}\s+[A-Z][A-Z0-9-]*\b', line):
        if re.search(r'\b(?:COMP(?:-[1-5])?|COMPUTATIONAL(?:-[1-5])?|BINARY|PACKED-DECIMAL|SIGN)\b|\bPIC(?:TURE)?\s+[SX9()0-9]*[SV]', line):
            return families | {'numeric_representation'}
        return families | {'data_layout'}
    if not any(re.match(pattern, line) for _, pattern in patterns):
        families.add('unknown')
    for family, pattern in patterns:
        if re.search(r'(?<![A-Z0-9-])'+pattern, line):
            families.add(family)
    return families


def _classify_construct(source):
    """Return the sole family, or explicitly mark a compound/unknown line."""
    families = _construct_families(source)
    if 'unknown' in families:
        return 'unknown'
    return next(iter(families)) if len(families) == 1 else 'mixed'


def _program_list(programs):
    shown = []
    for program in programs:
        if len(', '.join(shown + [program])) > PROGRAM_LIST_TEXT_LIMIT:
            break
        shown.append(program)
    result = ', '.join(shown) or 'See Context'
    if len(shown) < len(programs):
        result += f' (+{len(programs)-len(shown)} more; {len(programs)} programs total)'
    return result



def _celltext(text):
    return "'"+text if text.lstrip().startswith(('=','+','-','@')) else text


def _validate_cell_capacity(text, label):
    require(not isinstance(text,str) or len(text.encode('utf-16-le'))//2<=32767,
            label+' exceeds Excel cell capacity; shorten the identifier/metadata before issuing the review')


def _display_text(item, field):
    return item.get('display_'+field,item[field])


def _metadata_rows(packet):
    return [['Process ID',packet['process_id']],['Packet hash',packet['packet_hash']],
            ['Source snapshot',packet['source_snapshot']],
            ['Evidence class','SOURCE_DERIVED_EXPECTED; no mainframe execution']]


def _context_rows(packet):
    context = packet['context']
    return [['Process information']] + [[context[i:i+CONTEXT_CHUNK_SIZE]]
                                         for i in range(0,len(context),CONTEXT_CHUNK_SIZE)]


def _number_spans(numbers):
    """Lossless inclusive spans for source lines and zero-based blocker indices."""
    spans=[]
    for number in sorted(set(numbers)):
        if spans and number==spans[-1][1]+1:spans[-1][1]=number
        else:spans.append([number,number])
    return spans


def _analysis_reference(analysis):
    return {
        'path':'analysis/source-analysis.json','sha256':sha(encode(analysis)),
        'source_snapshot':analysis['source_snapshot'],
        'evidence_basis':analysis.get('evidence_basis','SOURCE_DERIVED_EXPECTED'),
        'raw_blocker_count':len(analysis['blockers']),'blockers_json_pointer':'/blockers',
        'blocker_index_base':0,'source_line_index_base':1,
        'files':{asset['path']:{'source_sha256':asset['source_hash'],
                               'physical_lines':asset.get('loc',{}).get('physical')}
                 for asset in analysis.get('assets',[])},
        'programs':{name:{'path':program['path'],'source_sha256':program['source_hash'],
                          'coverage_json_pointer':'/programs/'+name+'/coverage',
                          'coverage_sha256':sha(encode(program.get('coverage',[]))),
                          'coverage_line_count':len(program.get('coverage',[]))}
                    for name,program in analysis['programs'].items()},
    }


def _compact_technical_gaps(analysis, context_data):
    # The coordinator freezes and registers this complete artifact before issuing
    # the packet. Referencing its canonical hash avoids duplicating millions of
    # source/coverage rows into Excel while retaining every original blocker.
    context_data['source_analysis']=_analysis_reference(analysis)
    program_by_path={program['path']:name for name,program in analysis['programs'].items()}
    groups={};items=[]
    for index,blocker in enumerate(analysis['blockers']):
        if blocker['kind'] not in TECHNICAL_GAP_KINDS:
            items.append({'id':f'B_{index:03d}','question':'Is this unresolved item described correctly? If no, explain what it should do. '+blocker['message'],
                          'evidence':blocker.get('path','Source/intake evidence'),'kind':'unresolved_item'})
            continue
        path=blocker.get('path','Source/intake evidence')
        program=program_by_path.get(path) or blocker.get('program') or blocker.get('utility_id') or path
        key=(program,path,blocker['kind'])
        group=groups.setdefault(key,{'program':program,'path':path,'kind':blocker['kind'],
                                     'indices':[],'source_lines':{}})
        group['indices'].append(index)
        if blocker.get('lines'):
            group['source_lines'].setdefault(path,[]).extend(blocker['lines'])
        for reference in blocker.get('source_refs',[]):
            if isinstance(reference,dict) and isinstance(reference.get('line'),int):
                group['source_lines'].setdefault(reference.get('path',path),[]).append(reference['line'])
    context_data['technical_gaps']={}
    for key,group in sorted(groups.items()):
        item_id='TG_'+sha(encode(list(key)))[:16]
        indices=group.pop('indices');source_lines=group.pop('source_lines')
        group.update({'blocker_count':len(indices),'blocker_index_spans':_number_spans(indices),
                      'blockers_sha256':sha(encode([analysis['blockers'][i] for i in indices])),
                      'source_line_spans':{path:_number_spans(lines) for path,lines in sorted(source_lines.items())}})
        context_data['technical_gaps'][item_id]=group
        items.append({'id':item_id,
                      'question':f"Are these technical gaps for {group['program']} described correctly? "
                                 f"{group['blocker_count']} {group['kind']} blocker records need a reviewed semantic adapter. "
                                 'Confirm or correct the required source behavior. These gaps remain unsupported until the adapter is implemented and verified; Yes confirms the description only.',
                      'evidence':group['path']+'; see Context > technical_gaps > '+item_id+
                                 ' and source_analysis for exact blocker records, source spans, hashes and complete coverage.',
                      'kind':'technical_gap'})
    return items


def _blocker_lines(blocker):
    """Index valid source coordinates; the hash-bound raw record retains all data."""
    path = blocker.get('path')
    refs = {}
    if path:
        refs[path] = [line for line in blocker.get('lines', []) if type(line) is int and line > 0]
    for reference in blocker.get('source_refs', []):
        if isinstance(reference, dict):
            target = reference.get('path', path)
            line = reference.get('line')
            if target:
                refs.setdefault(target, [])
                if type(line) is int and line > 0:
                    refs[target].append(line)
    return refs


def _blocker_construct(blocker, sources):
    refs = _blocker_lines(blocker)
    if blocker['kind'] == 'missing_source' or not refs or any(path not in sources for path in refs):
        return 'unlocated_source', [], 'Source binding is missing.'
    if blocker['kind'] == 'source_classification':
        return 'source_classification', [], 'Source classification needs human resolution.'
    families = set()
    for path, lines in refs.items():
        source = sources[path]
        if not source or not lines or any(line not in source for line in lines):
            return 'unknown', [], 'Referenced source span is unavailable or not trustworthy.'
        for line in lines:
            families.update(source[line])
    if 'unknown' in families or not families:
        return 'unknown', sorted(families), 'At least one referenced line has no unambiguous construct classification.'
    if blocker['kind'] == 'unsupported_utility' and families == {'jcl_control'}:
        return 'utility', sorted(families), 'Utility finding references JCL invocation evidence.'
    if len(families) > 1:
        return 'mixed', sorted(families), 'One blocker spans multiple construct families.'
    return next(iter(families)), sorted(families), 'Classified from the referenced original source lines.'


def _compact_construct_gaps(analysis, context_data):
    """v3 only: share questions across programs, never merge away raw evidence."""
    context_data['source_analysis'] = _analysis_reference(analysis)
    program_by_path = {program['path']: name for name, program in analysis['programs'].items()}
    needed = {}
    for blocker in analysis['blockers']:
        if blocker['kind'] in TECHNICAL_GAP_KINDS | {'missing_source'}:
            for path, lines in _blocker_lines(blocker).items():
                needed.setdefault(path, set()).update(lines)
    sources = {}
    assets = {asset['path']: asset for asset in [*analysis.get('assets', []), *analysis['programs'].values()]}
    for path in needed.keys() & assets.keys():
        asset = assets[path]; text = asset.get('source_text')
        # File extensions, blocker prose and utility names are not source proof.
        trusted = (isinstance(text, str) and sha(text) == asset['source_hash']
                   and asset.get('kind') in ('cobol_program', 'copybook', 'jcl_job'))
        sources[path] = {i: _construct_families(line) for i, line in enumerate(text.splitlines(), 1)
                         if i in needed[path]} if trusted else None
    groups = {}; items = []
    for index, blocker in enumerate(analysis['blockers']):
        if blocker['kind'] not in TECHNICAL_GAP_KINDS | {'missing_source'}:
            items.append({'id': f'B_{index:03d}',
                          'question': 'Is this unresolved item described correctly? If no, explain what it should do. '+blocker['message'],
                          'evidence': blocker.get('path', 'Source/intake evidence'), 'kind': 'unresolved_item'})
            continue
        family, observed, reason = _blocker_construct(blocker, sources)
        group = groups.setdefault(family, {'family': family, 'indices': [], 'programs': set(),
                                          'source_lines': {}, 'kinds': {}, 'observed_constructs': set(),
                                          'classification_reasons': {}})
        group['indices'].append(index)
        group['observed_constructs'].update(observed)
        group['kinds'][blocker['kind']] = group['kinds'].get(blocker['kind'], 0) + 1
        group['classification_reasons'][reason] = group['classification_reasons'].get(reason, 0) + 1
        if blocker.get('program'):
            group['programs'].add(blocker['program'])
        for path, lines in _blocker_lines(blocker).items():
            group['source_lines'].setdefault(path, []).extend(lines)
            if path in program_by_path:
                group['programs'].add(program_by_path[path])
    context_data['technical_gaps'] = {}
    for family, group in sorted(groups.items()):
        item_id = 'TG_'+sha(encode(['construct', family]))[:16]
        indices = group.pop('indices'); source_lines = group.pop('source_lines')
        group['programs'] = sorted(group['programs'])
        group['observed_constructs'] = sorted(group['observed_constructs'])
        group.update({'program_count': len(group['programs']), 'paths': sorted(source_lines),
                      'blocker_count': len(indices), 'blocker_index_spans': _number_spans(indices),
                      'blockers_sha256': sha(encode([analysis['blockers'][i] for i in indices])),
                      'source_line_spans': {path: _number_spans(lines) for path, lines in sorted(source_lines.items())}})
        context_data['technical_gaps'][item_id] = group
        items.append({'id': item_id,
                      'question': f"ASSUMPTION - {family}: {CONSTRUCT_ASSUMPTIONS[family]} "
                                  f"Programs: {_program_list(group['programs'])}. "
                                  f"{group['blocker_count']} blocker records remain an unverified obligation. "
                                  'Yes confirms the description only; the behavior remains unsupported until a semantic adapter is implemented and verified.',
                      'evidence': 'See Context > technical_gaps > '+item_id+
                                  ' and source_analysis for every program, exact blocker records, source spans, hashes and complete coverage.',
                      'kind': 'technical_gap'})
    return items


def rule_classifications(process):
    """Explicit program defaults and per-rule exceptions, never inferred categories."""
    submitted=process.get('llm',{}).get('analysis',{})
    defaults=submitted.get('rule_classification_defaults',{})
    result={r['id']:defaults[name] for name,p in process['analysis']['programs'].items() if name in defaults for r in p['rules']}
    result.update(submitted.get('rule_classifications',{}))
    return result


def _group_business_items(process, items, context):
    """Group proven repeated decisions; hash input effects incrementally in source order."""
    def predicate_fields(value):
        if not isinstance(value,dict):return set()
        return ({value['field']} if 'field' in value else set()) | predicate_fields(value.get('left')) | predicate_fields(value.get('right'))
    fingerprints={}
    for name,program in process['analysis']['programs'].items():
        invocations=[(j['name'],s['program'],s.get('condition'),s.get('inputs'),s.get('outputs'))
                     for j in process['jobs'] for s in j['steps'] if s['program']==name]
        invocations += [{'transaction':t['id'],'mapset':t.get('mapset'),'map':t.get('map')} for t in process.get('transactions',[]) if t['program']==name]
        semantics=[{k:r.get(k) for k in ('predicate','then','else')} for r in program['rules']]
        base=sha(encode({'program':name,'source':program['path'],'version':program['source_hash'],
            'fields':program['fields'],'dependencies':program.get('dependencies',[]),
            'invocations':invocations,'program_semantics':semantics}))
        effects={}
        for rule,semantic in zip(program['rules'],semantics):
            inputs={field:effects.get(field,'initial') for field in sorted(predicate_fields(semantic['predicate']))}
            fingerprints[rule['id']]=sha(encode([base,inputs,semantic,rule['id'] if program.get('blockers') else None]))
            for field in {a['field'] for a in (semantic['then'] or [])+(semantic['else'] or [])}:
                effects[field]=sha(encode([effects.get(field,'initial'),semantic]))
    buckets={}
    for item in items:
        fingerprint=sha(encode([fingerprints[item['id']],item['kind'],item['question']]))
        buckets.setdefault(fingerprint,[]).append(item)
    grouped=[];context['rule_groups']={}
    for fingerprint,members in buckets.items():
        if len(members)==1:
            grouped.extend(members);continue
        ident='GROUP_'+fingerprint[:24]
        context['rule_groups'][ident]={'members':[m['id'] for m in members],'equivalence_sha256':fingerprint,
            'basis':'Same source, program semantics, preceding input effects, layouts, dependencies and invocation context'}
        grouped.append({'id':ident,'kind':members[0]['kind'],
            'question':members[0]['question']+' Applies to every listed occurrence; corrections leave the group unresolved.',
            'evidence':'; '.join(m['evidence'] for m in members)})
    return grouped


def packet_document(process):
    a=process['analysis']
    # Marker-absent historical documents must reproduce the original v1 packet
    # bytes/hash, including their original size gate. New intake opts into v3.
    packet_version=process.get('sme_packet_version',1)
    require(type(packet_version) is int and packet_version in (1,2,3,4),'Unsupported SME packet version')
    if packet_version in (2,3):
        require(len(a['rules'])<=SME_ITEM_LIMIT,
                'SME packet exceeds supported size; business rules require individual human items; narrow the process scope before issuing')
    items=[{'id':r['id'],'question':r['plain'],'evidence':'; '.join(r['source_refs']),'kind':'business_rule'} for r in a['rules']]
    context_data={'jobs':process['jobs'],'layouts':{n:p['fields'] for n,p in a['programs'].items()},'relationships':a['relationships'],'knowledge_input':process.get('knowledge_context')}
    if process.get('transactions'):context_data['transactions']=process['transactions']
    if packet_version == 4:
        categories=rule_classifications(process)
        for item in items:
            category=categories.get(item['id'],{})
            item['kind']=category.get('category','unclassified')
            item['question']='['+item['kind']+'] '+item['question']
            if category.get('reason'):item['evidence']+='; classification: '+category['reason']
        items=_group_business_items(process,items,context_data)
    items += [{'id':k,'question':q,'evidence':'See process inventory/layout/relationship context','kind':'process_assumption'} for k,q in GLOBAL_QUESTIONS]
    if process.get('knowledge_context'):items.append({'id':'G_KNOWLEDGE','question':'Is the supplied background knowledge correct and applicable to this process? If No, describe the correction.','evidence':'See knowledge_input context and its SHA256','kind':'process_assumption'})
    if process.get('mainframe_knowledge'):
        context_data['file_classifications']=a.get('classifications',{})
        context_data['utility_findings']=a.get('utility_findings',[])
        context_data['application_knowledge']={'status':'UNVERIFIED_INPUT',**process['mainframe_knowledge']['application']}
        context_data['mainframe_knowledge_snapshot']='analysis/mainframe-knowledge.json'
        context_data['mainframe_knowledge_sha256']=sha(encode(process['mainframe_knowledge']))
        checks=[
            ('G_FILE_TYPES','Are the listed file types and dependencies correct, including source, copybooks, JCL/PROCs, control cards, screens and data definitions? If No, name the missing or misclassified file.'),
            ('G_UTILITIES','Does the utility list include every standard, vendor and application utility this process calls? If No, give its name, purpose, inputs, outputs and return codes.'),
            ('G_NATIVE_DATA','Do the provided layouts fully describe record formats, encodings, keys, signs, decimal precision and matching rules? If No, describe the missing format or rule, including empty, duplicate and missing records.'),
            ('G_NATIVE_RUNTIME','Do the listed jobs fully describe runtime behavior, including control cards, DD allocation/deletion, conditional steps, restart, transactions, scheduling and external calls? If No, describe the missing behavior.'),
        ]
        items += [{'id':key,'question':question,'evidence':'See file_classifications, utility_findings, jobs and layouts in Context; recognition alone is not conversion proof.','kind':'mainframe_assumption'} for key,question in checks]
    if process.get('transactions'):
        for item in items:
            if item['id']=='G_ORDER':item['question']='Are the transaction entry programs, mapsets/maps and dependency boundaries in Context correct? Unknown bindings remain unresolved.'
            if item['id']=='G_NATIVE_RUNTIME':item['question']='Does Context describe online runtime behavior, including AID/navigation, session state, COMMAREA/channels, authorization, storage, commits/rollback, recovery and downstream calls? If No, describe missing behavior.'
        items.append({'id':'G_ONLINE_TARGET','question':'Is the stated local record-API/form candidate scope understood? Native CICS/BMS, identity and enterprise transaction behavior require independent verified replacements; approval does not clear those gaps.','evidence':'See transactions and source-specific unresolved obligations in Context','kind':'process_assumption'})
    suggestions=process.get('llm',{}).get('analysis',{})
    for n,text in enumerate(suggestions.get('questions',[])+suggestions.get('assumptions',[])):
        items.append({'id':f'LLM_{n:03d}','question':text,'evidence':'Unverified LLM suggestion; validate against source/context','kind':'provider_suggestion'})
    if packet_version in (3,4):
        items += _compact_construct_gaps(a,context_data)
    elif packet_version==2 and len(items)+len(a['blockers'])>SME_ITEM_LIMIT:
        items += _compact_technical_gaps(a,context_data)
    else:
        for i,b in enumerate(a['blockers']):items.append({'id':f'B_{i:03d}','question':'Is this unresolved item described correctly? If no, explain what it should do. '+b['message'],'evidence':b.get('path','Source/intake evidence'),'kind':'unresolved_item'})
    require(len(items)<=SME_ITEM_LIMIT,'SME packet exceeds supported size; retain a scope blocker before issuing')
    require(len({item['id'] for item in items})==len(items),'SME packet contains duplicate question identities')
    details={}
    for item in items:
        for field in ('question','evidence'):
            if len(item[field])>CHECKLIST_TEXT_LIMIT:
                details.setdefault(item['id'],{})[field]=item[field]
                item['display_'+field]=(item[field][:1500]+'\n[Read the complete '+field+' for '+item['id']+
                                       ' in Context > checklist_details before answering.]')
        for field in ('id','kind'):
            _validate_cell_capacity(item[field],'Checklist '+field)
        for field in ('question','evidence'):
            _validate_cell_capacity(_celltext(_display_text(item,field)),'Checklist '+field)
    if details:context_data['checklist_details']=details
    context=json.dumps(context_data,ensure_ascii=False,indent=2,sort_keys=True)
    doc={'version':packet_version,'process_id':process['id'],'source_snapshot':a['source_snapshot'],'items':items,'context':context}
    html_version=process.get('review_html_contract_version')
    require(html_version is None or type(html_version) is int and html_version in (1,2), 'Unsupported frozen HTML review contract')
    if html_version == 2:doc['html_contract_version']=2
    doc['packet_hash']=sha(encode(doc))
    for row in _metadata_rows(doc):
        for value in row:_validate_cell_capacity(value,'Review metadata')
    require(len(encode(doc))<=MAX_UPLOAD,'SME packet exceeds the supported document size; narrow the process scope before issuing')
    return doc


# The browser renders questions from this hash-bound JSON using textContent.
# Save captures the untouched document shell and changes only the answer payload;
# imports compare that shell without executing HTML or JavaScript.
HTML_RETURN_HEADER = b'<!doctype html>\n<html lang="en" data-workbench-sme="1">'
_HTML_PACKET_ID = 'workbench-sme-packet'
_HTML_RETURN_ID = 'workbench-sme-return'
_HTML_STYLE = '''
:root{font-family:Segoe UI,Arial,sans-serif;color:#202329;background:#f5f4f1;font-size:16px;line-height:1.5;--accent:#aa1728;--border:#d9d9d6}*{box-sizing:border-box}body{margin:0}main{max-width:1100px;margin:auto;padding:28px 24px 70px}h1{font-size:32px;line-height:1.2;margin:0 0 8px}h2{font-size:19px;margin:0}p{margin:8px 0}.eyebrow{color:var(--accent);font-size:12px;font-weight:700;letter-spacing:.08em;text-transform:uppercase}.muted{color:#5d6169}.bar,.card,.context{background:#fff;border:1px solid var(--border);border-radius:12px;padding:20px;margin-top:18px}.bar{position:sticky;top:0;z-index:2;box-shadow:0 3px 12px #20232909}.toolbar{display:flex;gap:12px;align-items:end;flex-wrap:wrap}.field{display:flex;flex-direction:column;gap:6px;flex:1;min-width:190px}label{font-size:14px;font-weight:600}input,select,textarea,button{font:inherit}input,select,textarea{width:100%;border:1px solid #9d9fa4;border-radius:7px;padding:9px 11px;background:#fff;color:inherit}textarea{min-height:80px;resize:vertical}button{border:1px solid #a6a8ad;border-radius:7px;background:#fff;color:inherit;padding:10px 16px;cursor:pointer;font-weight:600}button.primary{background:var(--accent);color:white;border-color:var(--accent)}button:disabled{opacity:.45;cursor:default}:focus-visible{outline:3px solid #d69923;outline-offset:3px}.summary{font-size:14px}.tag{display:inline-block;font-size:12px;background:#f3f2ef;border-radius:5px;padding:3px 7px;margin:8px 5px 8px 0}.question{white-space:pre-wrap;overflow-wrap:anywhere;font-size:17px}.evidence{white-space:pre-wrap;overflow-wrap:anywhere;font-size:13px;color:#5d6169;margin:12px 0}.answer-row{display:grid;grid-template-columns:180px 1fr;gap:18px;align-items:start}.pagination{display:flex;align-items:center;justify-content:space-between;gap:12px;margin-top:18px}.status{min-height:24px;white-space:pre-wrap;overflow-wrap:anywhere;color:#642333;font-size:14px}.metadata{font-size:13px;overflow-wrap:anywhere}.context pre{font-size:12px;white-space:pre-wrap;overflow-wrap:anywhere;max-height:400px;overflow:auto}summary{cursor:pointer;font-weight:600}.empty{padding:30px;text-align:center;color:#5d6169}.footer{margin-top:18px;font-size:13px;color:#5d6169}@media(max-width:600px){main{padding:18px 12px 40px}h1{font-size:26px}.bar,.card,.context{padding:15px}.bar{position:static}.field{min-width:100%}.answer-row{grid-template-columns:1fr}.pagination{flex-wrap:wrap}button.primary{width:100%}}
'''
_HTML_SCRIPT = r'''
'use strict';
(() => {
  const shell = document.documentElement.cloneNode(true);
  const packet = JSON.parse(document.getElementById('workbench-sme-packet').textContent);
  const model = JSON.parse(document.getElementById('workbench-sme-return').textContent);
  const app = document.getElementById('sme-app');
  const answers = new Map(model.items.map(item => [item.id, item]));
  let page = 0;
  const pageSize = 20;
  const make = (tag, text, className) => {
    const node = document.createElement(tag);
    if (text !== undefined) node.textContent = text;
    if (className) node.className = className;
    return node;
  };
  const field = (title, control) => {
    const box = make('div', undefined, 'field');
    const label = make('label', title); label.htmlFor = control.id;
    box.append(label, control); return box;
  };
  app.append(make('p', 'ONE PROCESS / ONE REVIEW PACKET', 'eyebrow'));
  app.append(make('h1', 'Review the process understanding'));
  app.append(make('p', 'Choose Yes, No or Not sure for each statement. Add commentary or corrections where needed. No answer is approved by default.', 'muted'));
  const metadata = make('p', 'Process: ' + packet.process_id + '\nPacket: ' + packet.packet_hash + '\nSource snapshot: ' + packet.source_snapshot, 'metadata');
  metadata.style.whiteSpace = 'pre-wrap'; app.append(metadata);
  const bar = make('div', undefined, 'bar');
  const toolbar = make('div', undefined, 'toolbar');
  const reviewer = make('input'); reviewer.id = 'sme-reviewer'; reviewer.type = 'text'; reviewer.maxLength = 160; reviewer.autocomplete = 'name'; reviewer.value = model.reviewer;
  reviewer.addEventListener('input', () => { model.reviewer = reviewer.value; status.textContent = 'Unsaved changes. Save the portable review file when ready.'; });
  const save = make('button', 'Save review file', 'primary'); save.type = 'button'; save.id = 'sme-save';
  toolbar.append(field('Reviewer name (required to save)', reviewer), save);
  const filters = make('div', undefined, 'toolbar'); filters.style.marginTop = '14px';
  const search = make('input'); search.id = 'sme-search'; search.type = 'search'; search.placeholder = 'Question, program, evidence or ID';
  const filter = make('select'); filter.id = 'sme-filter';
  for (const [value, label] of [['all','All statements'],['','Unanswered'],['No','No'],['Not sure','Not sure'],['Yes','Yes']]) {
    const option = make('option', label); option.value = value; filter.append(option);
  }
  filters.append(field('Find a statement', search), field('Answer filter', filter));
  const counts = make('p', undefined, 'summary'); counts.id = 'sme-counts';
  const status = make('p', 'This file works offline. Save downloads one HTML file; return that saved file to the workbench.', 'status'); status.id = 'sme-status'; status.setAttribute('role','status'); status.setAttribute('aria-live','polite');
  bar.append(toolbar, filters, counts, status); app.append(bar);
  const cards = make('div'); cards.id = 'sme-items'; app.append(cards);
  const navigation = make('div', undefined, 'pagination');
  const previous = make('button', 'Previous'); previous.type = 'button'; previous.id = 'sme-previous';
  const pageLabel = make('span'); pageLabel.id = 'sme-page';
  const next = make('button', 'Next'); next.type = 'button'; next.id = 'sme-next';
  navigation.append(previous, pageLabel, next); app.append(navigation);
  const context = make('details', undefined, 'context'); context.append(make('summary', 'Complete process context and source references'), make('pre', packet.context)); app.append(context);
  app.append(make('p', 'One review round. Blank answers, Not sure, No and corrections remain unresolved. These statements are source-derived expectations, not observed mainframe parity. Your browser may ask where to save the single HTML file.', 'footer'));
  const updateCounts = () => {
    const answered = model.items.filter(item => item.answer !== '').length;
    counts.textContent = answered + ' of ' + packet.items.length + ' answered · ' + (packet.items.length - answered) + ' unresolved blanks';
  };
  const selected = () => {
    const query = search.value.trim().toLocaleLowerCase();
    return packet.items.filter(item => (filter.value === 'all' || answers.get(item.id).answer === filter.value) && (!query || (item.id + ' ' + item.question + ' ' + item.evidence + ' ' + item.kind).toLocaleLowerCase().includes(query)));
  };
  const render = () => {
    updateCounts(); const items = selected(); const pages = Math.max(1, Math.ceil(items.length/pageSize)); page = Math.min(page, pages-1);
    cards.replaceChildren();
    for (const item of items.slice(page*pageSize,(page+1)*pageSize)) {
      const answer = answers.get(item.id); const card = make('section', undefined, 'card'); card.dataset.itemId = item.id;
      card.append(make('h2', item.id), make('span', item.kind.replaceAll('_',' '), 'tag'), make('p', item.question, 'question'), make('p', 'Source evidence: ' + item.evidence, 'evidence'));
      const control = make('select'); control.id = 'answer-' + item.id; control.dataset.answerId = item.id;
      for (const [value,label] of [['','Choose an answer'],['Yes','Yes'],['No','No'],['Not sure','Not sure']]) { const option = make('option',label); option.value = value; control.append(option); }
      control.value = answer.answer;
      control.addEventListener('change', () => { answer.answer = control.value; updateCounts(); status.textContent = 'Unsaved changes. Save the portable review file when ready.'; });
      const commentary = make('textarea'); commentary.id = 'comment-' + item.id; commentary.maxLength = 4000; commentary.value = answer.correction; commentary.dataset.commentId = item.id;
      commentary.addEventListener('input', () => { answer.correction = commentary.value; status.textContent = 'Unsaved changes. Save the portable review file when ready.'; });
      const row = make('div', undefined, 'answer-row'); row.append(field('Answer',control), field('Commentary / correction',commentary)); card.append(row); cards.append(card);
    }
    if (!items.length) cards.append(make('p','No statements match the current filters.','empty'));
    pageLabel.textContent = 'Page ' + (page+1) + ' of ' + pages + ' · ' + items.length + ' statements shown by filter'; previous.disabled = page === 0; next.disabled = page >= pages-1;
  };
  search.addEventListener('input', () => { page = 0; render(); }); filter.addEventListener('change', () => { page = 0; render(); });
  previous.addEventListener('click', () => { page--; render(); }); next.addEventListener('click', () => { page++; render(); });
  const safeJSON = value => JSON.stringify(value).replace(/[<>&\u2028\u2029]/g, character => '\\u' + character.charCodeAt(0).toString(16).padStart(4,'0'));
  save.addEventListener('click', () => {
    model.reviewer = reviewer.value.trim();
    if (!model.reviewer || model.reviewer.length > 160) { status.textContent = 'Enter the actual reviewer name before saving.'; reviewer.focus(); return; }
    if (model.items.some(item => !['','Yes','No','Not sure'].includes(item.answer) || typeof item.correction !== 'string' || item.correction.length > 4000)) { status.textContent = 'Check the answer choices and commentary limits before saving.'; return; }
    const output = shell.cloneNode(true);
    output.querySelector('#workbench-sme-packet').textContent = safeJSON(packet);
    output.querySelector('#workbench-sme-return').textContent = safeJSON(model);
    const blob = new Blob(['<!doctype html>\n',output.outerHTML], {type:'text/html;charset=utf-8'});
    if (blob.size > 8388608) { status.textContent = 'The review exceeds the 8 MiB import limit. Shorten commentary before saving; no file was downloaded.'; return; }
    const link = document.createElement('a'); const url = URL.createObjectURL(blob); link.href = url; link.download = packet.process_id + '-sme-return.html'; link.style.display = 'none'; document.body.append(link); link.click(); link.remove(); setTimeout(() => URL.revokeObjectURL(url), 1000);
    status.textContent = 'Saved one portable HTML review file. ' + model.items.filter(item => item.answer === '').length + ' blank answers remain unresolved. Return the downloaded file, not this original.';
  });
  render();
})();
'''


def _html_json(value):
    # Script raw-text parsing ignores HTML entities, so escape JSON characters.
    text = encode(value).decode('utf-8')
    for character, escaped in (('<', '\\u003c'), ('>', '\\u003e'), ('&', '\\u0026'), ('\u2028', '\\u2028'), ('\u2029', '\\u2029')):
        text = text.replace(character, escaped)
    return text


def _html_response(packet):
    return {'version': 1, 'kind': 'SME_HTML_RETURN', 'process_id': packet['process_id'],
            'packet_hash': packet['packet_hash'], 'source_snapshot': packet['source_snapshot'],
            'reviewer': '', 'items': [{'id': item['id'], 'answer': '', 'correction': ''} for item in packet['items']]}



# Keep the original shell byte-for-byte for already issued immutable packets.
_HTML_SCRIPT_LARGE = _HTML_SCRIPT.replace(
    "if (blob.size > 8388608) { status.textContent = 'The review exceeds the 8 MiB import limit. Shorten commentary before saving; no file was downloaded.'; return; }",
    "if (blob.size > 134217728) { status.textContent = 'This review exceeds 128 MiB. Keep the file local and ask the operator to inspect its packet scope; no file was downloaded.'; return; }")


def review_return_limit(prefix):
    """Bound local transport by detected format before reading the whole file."""
    return MAX_HTML_REVIEW_BYTES if is_html_return(prefix) else MAX_UPLOAD


def _html_contract(packet):
    version=packet.get('html_contract_version')
    require(version is None or type(version) is int and version == 2, 'Unsupported frozen HTML review contract')
    return (_HTML_SCRIPT_LARGE, MAX_HTML_REVIEW_BYTES) if version == 2 else (_HTML_SCRIPT, MAX_UPLOAD)


def render_html_packet(packet, response=None):
    """Render the one issued packet; no external assets or inferred answers."""
    import base64
    import hashlib
    script, limit = _html_contract(packet)
    script_hash = base64.b64encode(hashlib.sha256(script.encode()).digest()).decode('ascii')
    style_hash = base64.b64encode(hashlib.sha256(_HTML_STYLE.encode()).digest()).decode('ascii')
    policy = "default-src 'none'; script-src 'sha256-"+script_hash+"'; style-src 'sha256-"+style_hash+"'; connect-src 'none'; form-action 'none'; base-uri 'none'"
    text = HTML_RETURN_HEADER.decode() + '\n<head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><meta http-equiv="Content-Security-Policy" content="'+policy+'"><title>Process review</title><style>'+_HTML_STYLE+'</style></head>\n<body><main id="sme-app"></main><noscript>This review needs JavaScript enabled in a normal browser. No network connection is needed.</noscript>\n<script type="application/json" id="'+_HTML_PACKET_ID+'">'+_html_json(packet)+'</script>\n<script type="application/json" id="'+_HTML_RETURN_ID+'">'+_html_json(_html_response(packet) if response is None else response)+'</script>\n<script id="workbench-sme-ui">'+script+'</script>\n</body></html>'
    data = text.encode('utf-8')
    require(len(data) <= limit, 'SME HTML document exceeds the supported upload size; retain the complete local file and inspect its scope')
    return data


class _HTMLReviewShape(HTMLParser):
    """Normalize browser serialization while rejecting any altered active shell."""
    def __init__(self, text):
        super().__init__(convert_charrefs=False)
        self.events, self.payloads, self.current = [], {}, None
        self.feed(text); self.close()
        require(self.current is None, 'Unclosed HTML review payload')

    def handle_starttag(self, tag, attrs):
        require(len({key for key, _ in attrs}) == len(attrs), 'Duplicate HTML attribute')
        self.events.append(('start', tag, tuple(sorted(attrs))))
        if tag == 'script' and dict(attrs).get('id') in (_HTML_PACKET_ID, _HTML_RETURN_ID):
            identity = dict(attrs)['id']
            require(self.current is None and identity not in self.payloads, 'Duplicate HTML review payload')
            self.payloads[identity] = ''; self.current = identity

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in ('meta', 'input', 'br', 'hr', 'link', 'img'):
            self.handle_endtag(tag)

    def handle_endtag(self, tag):
        self.events.append(('end', tag))
        if tag == 'script': self.current = None

    def handle_data(self, data):
        if self.current is not None:
            self.payloads[self.current] += data
        elif data.strip():
            self.events.append(('data', data))

    def handle_entityref(self, name): self.events.append(('data', html.unescape('&'+name+';')))
    def handle_charref(self, name): self.events.append(('data', html.unescape('&#'+name+';')))
    def handle_comment(self, data): self.events.append(('comment', data))
    def handle_decl(self, decl): self.events.append(('declaration', decl.lower()))
    def unknown_decl(self, data): self.events.append(('unknown_declaration', data))
    def handle_pi(self, data): self.events.append(('processing_instruction', data))


def is_html_return(data):
    """Typed guard only; read_html_return still validates the complete document."""
    return isinstance(data, bytes) and data.startswith(HTML_RETURN_HEADER)


def _returned_document(data, packet, reviewer, answers):
    if packet.get('version') == 4:
        groups = json.loads(packet['context']).get('rule_groups', {})
        for group_id, group in groups.items():
            answer = answers[group_id]
            for member in group['members']:
                require(member not in answers, 'Duplicate grouped rule identity')
                answers[member] = {**answer, 'group_id': group_id}
    return {'packet_hash': packet['packet_hash'], 'source_snapshot': packet['source_snapshot'],
            'reviewer': reviewer.strip(), 'items': answers, 'return_hash': sha(data)}


def read_html_return(data, packet, reviewer):
    """Read bounded JSON evidence only; never execute uploaded HTML/scripts."""
    require(isinstance(reviewer, str) and 0 < len(reviewer.strip()) <= 160,
            'Name the reviewer responsible for this returned file')
    require(packet.get('packet_hash') == sha(encode({k: v for k, v in packet.items() if k != 'packet_hash'})),
            'Frozen packet content/hash changed')
    _, limit = _html_contract(packet)
    require(is_html_return(data) and len(data) <= limit, 'Invalid or oversized SME HTML return')
    try:
        shape = _HTMLReviewShape(data.decode('utf-8'))
    except ValidationError:
        raise
    except (UnicodeError, RecursionError, ValueError) as exc:
        raise ValidationError('Invalid UTF-8 HTML review document') from exc
    frozen = _HTMLReviewShape(render_html_packet(packet).decode('utf-8'))
    require(shape.events == frozen.events, 'HTML review shell, instructions or script changed')
    require(set(shape.payloads) == {_HTML_PACKET_ID, _HTML_RETURN_ID}, 'Missing HTML review payload')
    embedded = decode(shape.payloads[_HTML_PACKET_ID], limit=limit)
    require(isinstance(embedded, dict) and embedded.get('packet_hash') ==
            sha(encode({key: value for key, value in embedded.items() if key != 'packet_hash'}))
            and encode(embedded, limit=limit) == encode(packet, limit=limit),
            'Original HTML packet changed')
    response = decode(shape.payloads[_HTML_RETURN_ID], limit=limit)
    require(isinstance(response, dict) and set(response) == {'version', 'kind', 'process_id', 'packet_hash', 'source_snapshot', 'reviewer', 'items'}, 'HTML return fields changed')
    require(type(response['version']) is int and response['version'] == 1 and response['kind'] == 'SME_HTML_RETURN', 'Invalid HTML return version or kind')
    require(all(response[key] == packet[key] for key in ('process_id', 'packet_hash', 'source_snapshot')), 'HTML return belongs to a different packet or process')
    require(isinstance(response['reviewer'], str) and response['reviewer'].strip() == reviewer.strip(), 'HTML reviewer attribution must match the named reviewer')
    expected = {item['id']: item for item in packet['items']}; answers = {}
    require(isinstance(response['items'], list) and len(response['items']) == len(expected), 'Missing or extra HTML checklist items')
    for returned in response['items']:
        require(isinstance(returned, dict) and set(returned) == {'id', 'answer', 'correction'}, 'HTML answer fields changed')
        rid, answer, correction = returned['id'], returned['answer'], returned['correction']
        require(isinstance(rid, str) and rid in expected and rid not in answers, 'Unknown or duplicate question ID')
        require(isinstance(answer, str) and answer in ('Yes', 'No', 'Not sure', ''), 'Answer must be Yes, No, Not sure or blank')
        require(isinstance(correction, str) and len(correction) <= 4000, 'Returned commentary exceeds supported bounds')
        original = expected[rid]
        answers[rid] = {'answer': answer or 'Unanswered', 'correction': correction.strip(),
                        'reviewer': reviewer.strip(), 'question': original['question'], 'kind': original['kind']}
    require(set(answers) == set(expected), 'Missing checklist items')
    return {**_returned_document(data, packet, reviewer, answers), 'return_format': 'html'}


def read_return(data, packet, reviewer):
    """One import contract for the HTML file and preserved legacy workbook."""
    return read_html_return(data, packet, reviewer) if is_html_return(data) else read_answers(data, packet, reviewer)


def export_packet(process, directory):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from docx import Document
    directory=Path(directory)
    # A changed default must not rewrite/reissue an existing packet shell.
    candidate=process
    if directory.exists():
        require(not any(path_is_link(path) for path in [directory,*directory.parents]), 'Existing packet path is unsafe')
        require(all((directory/name).is_file() and not path_is_link(directory/name)
                    for name in ('packet.json','sme-checklist.xlsx','sme-checklist.docx','sme-checklist.html')),
                'Existing packet is incomplete; recover its preserved snapshot')
        prior_path=directory/'packet.json'
        require(prior_path.stat().st_size <= MAX_UPLOAD, 'Existing packet exceeds the frozen packet size bound')
        with prior_path.open('rb') as stream: prior_raw=stream.read(MAX_UPLOAD+1)
        prior=decode(prior_raw,MAX_UPLOAD)
        candidate={**process,'review_html_contract_version':prior.get('html_contract_version',1)}
    document=packet_document(candidate)
    if directory.exists():
        require(all((directory/name).is_file() and not path_is_link(directory/name)
                    for name in ('packet.json','sme-checklist.xlsx','sme-checklist.docx','sme-checklist.html')),
                'Existing packet is incomplete; recover its preserved snapshot')
        existing=decode((directory/'packet.json').read_bytes())
        require(existing==document,'Existing packet differs; cannot issue a second packet')
        original=read_answers((directory/'sme-checklist.xlsx').read_bytes(),existing,'Packet export validation')
        require(all(item['answer']=='Unanswered' and not item['correction']
                    and item['reviewer']=='Packet export validation' for item in original['items'].values()),
                'Existing issued workbook contains answers; preserve the original blank packet')
        return existing
    directory.parent.mkdir(parents=True,exist_ok=True)
    # Process directories contain only registered evidence. A crash may leave a
    # draft behind, so stage on the same filesystem outside that strict layout.
    workspace=directory.parents[2] if directory.parent.parent.name=='processes' else directory.parent
    staging=workspace/'.implementation/tmp';staging.mkdir(parents=True,exist_ok=True)
    temp=Path(tempfile.mkdtemp(prefix='review-draft-',dir=staging))
    try:
        book=Workbook();sheet=book.active;sheet.title='Checklist'
        sheet.append(CHECKLIST_HEADER)
        for item in document['items']:sheet.append([item['id'],_celltext(_display_text(item,'question')),_celltext(_display_text(item,'evidence')),item['kind'],'','',''])
        for c in sheet[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='18314F')
        for row in sheet.iter_rows(min_row=2):
            for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
        for col,width in [('A',26),('B',80),('C',42),('D',22),('E',18),('F',60),('G',24)]:sheet.column_dimensions[col].width=width
        sheet.freeze_panes='E2';sheet.auto_filter.ref=sheet.dimensions
        dv=DataValidation(type='list',formula1='"Yes,No,Not sure"');dv.errorTitle='Choose an answer';dv.error='Use Yes, No or Not sure';dv.showErrorMessage=True;sheet.add_data_validation(dv);dv.add(f'E2:E{sheet.max_row}')
        meta=book.create_sheet('Metadata')
        for row in _metadata_rows(document):meta.append(row)
        context=book.create_sheet('Context');context.column_dimensions['A'].width=120
        for row in _context_rows(document):
            context.append(row);context.cell(context.max_row,1).data_type='s';context.cell(context.max_row,1).alignment=Alignment(wrap_text=True)
        book.save(temp/'sme-checklist.xlsx');book.close()
        # Exercise the actual import contract before publication/quota issuance.
        # This catches writer truncation and archive/cell limits without SME answers.
        read_answers((temp/'sme-checklist.xlsx').read_bytes(),document,'Packet export validation')
        word=Document();word.add_heading(process['name']+' — review checklist',0)
        word.add_paragraph('Check each statement. Answer Yes, No or Not sure in the Excel file. For No, write the correction. This is the only review round for this process. Missing or uncertain answers remain unresolved. Source-derived expectations are predictions, not observed mainframe results.')
        word.add_heading('Process context',1);word.add_paragraph(document['context'])
        for item in document['items']:
            word.add_heading(item['id'],2);word.add_paragraph(item['question']);word.add_paragraph('Evidence: '+item['evidence']);word.add_paragraph('Yes / No / Not sure. If No: __________________')
        word.save(temp/'sme-checklist.docx')
        rendered=render_html_packet(document)
        write_new(temp/'sme-checklist.html',rendered);write_new(temp/'packet.json',encode(document));temp.rename(directory)
        return document
    finally:
        shutil.rmtree(temp,ignore_errors=True)


def read_answers(data, packet, reviewer):
    from openpyxl import load_workbook
    require(isinstance(reviewer,str) and 0<len(reviewer.strip())<=160,'Name the reviewer responsible for this returned file')
    require(packet.get('packet_hash')==sha(encode({k:v for k,v in packet.items() if k!='packet_hash'})),
            'Frozen packet content/hash changed')
    checked_zip(data).close()
    frozen = {'Metadata':_metadata_rows(packet),'Context':_context_rows(packet),
              'Checklist':[CHECKLIST_HEADER]+[[None]*7 for _ in packet['items']]}
    def check_dimensions(book):
        require(set(book.sheetnames)==set(frozen),'Review sheet set changed')
        for name,rows in frozen.items():
            sheet=book[name]
            require(sheet.max_row==len(rows) and sheet.max_column==len(rows[0]),
                    name+' dimensions changed')
    # Check declared dimensions as well as actual cells. A forged XML dimension must
    # neither hide extra cells from streaming reads nor add blank rows/columns.
    declared=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
    try:check_dimensions(declared)
    finally:declared.close()
    book=load_workbook(BytesIO(data),read_only=False,data_only=False,keep_links=False)
    try:
        check_dimensions(book)
        require(not any(sheet.merged_cells.ranges for sheet in book),'Merged cells are not accepted in returned review evidence')
        for name in ('Metadata','Context'):
            for cells,row in zip(book[name].iter_rows(),frozen[name]):
                require(all(c.data_type!='f' for c in cells),'Formulas are not accepted in '+name)
                require([c.value for c in cells]==row,'Frozen '+name+' changed')
        sheet=book['Checklist']
        require(all(c.data_type!='f' for c in sheet[1]) and [c.value for c in sheet[1]]==CHECKLIST_HEADER,
                'Checklist headings changed')
        expected={x['id']:x for x in packet['items']};answers={}
        for cells in sheet.iter_rows(min_row=2):
            require(all(c.data_type!='f' for c in cells),'Formulas are not accepted in returned answers')
            values=[c.value for c in cells];rid,q,ref,kind,answer,correction,actor=values
            require(rid in expected and rid not in answers,'Unknown or duplicate question ID')
            original=expected[rid]
            require(q==_celltext(_display_text(original,'question')) and ref==_celltext(_display_text(original,'evidence')) and kind==original['kind'],'Original question/evidence changed')
            require(answer is None or isinstance(answer,str),'Answer must be Yes, No, Not sure or blank')
            require(actor is None or isinstance(actor,str),'Reviewer attribution must be text')
            answer=('' if answer is None else answer).strip()
            correction=('' if correction is None else str(correction)).strip()
            actor=(reviewer if actor is None or actor=='' else actor).strip()
            require(answer in ('Yes','No','Not sure',''),'Answer must be Yes, No, Not sure or blank')
            require(len(correction)<=4000 and 0<len(actor)<=160,'Returned text exceeds supported bounds')
            answers[rid]={'answer':answer or 'Unanswered','correction':correction,'reviewer':actor,'question':original['question'],'kind':kind}
        require(set(answers)==set(expected),'Missing checklist items')
        return _returned_document(data,packet,reviewer,answers)
    finally:book.close()
