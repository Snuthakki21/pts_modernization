"""Operator scope selections are not fabricated SME answers or parity claims."""
import tempfile
import unittest
from pathlib import Path
from workbench.domain import ValidationError
from test_source import COBOL
from test_workflow import MANIFEST

class RequirementsSelectionTests(unittest.TestCase):
    def test_start_waits_for_explicit_save_with_all_requirements_yes(self):
        from workbench.coordinator import Coordinator
        with tempfile.TemporaryDirectory() as temp:
            c=Coordinator(temp)
            try:
                c.create(MANIFEST,{'ELIGIBLE.cbl':COBOL},requirements_selection=True)
                c.start('process-a');c.advance('process-a')
                self.assertEqual(c.ledger.get('process-a')['status'],'WAITING_REQUIREMENTS')
                model=c.requirements_view('process-a')
                self.assertTrue(model['items']);self.assertTrue(all(r['selected'] for r in model['items']))
                self.assertFalse(c.ledger.get('process-a')['packet_issued'])
            finally:c.close()

    def test_no_rule_is_omitted_from_target_but_preserved_in_requirements(self):
        from workbench.requirements import catalog, project, render_markdown
        from workbench.source import analyze_sources
        from workbench.intake import parse_manifest
        from workbench.domain import sha,encode
        doc=parse_manifest(MANIFEST);analysis=analyze_sources({'ELIGIBLE.cbl':COBOL},doc)
        items=catalog(analysis)['items'];rule=next(r for r in items if r['kind']=='rule')
        selection={'schema_version':1,'process_id':doc['id'],'source_snapshot':analysis['source_snapshot'],'catalog_hash':catalog(analysis)['hash'],
                   'revision':1,'excluded_ids':[rule['id']],'saved_at':'2026-10-06T00:00:00+00:00','saved_by':'Fictional requirements operator'}
        chosen=project(analysis,selection)
        self.assertNotIn(rule['rule_id'],[r['id'] for r in chosen['rules']])
        self.assertIn('Not converted because selected No in requirements',render_markdown(selection,catalog(analysis)))
        self.assertEqual(len(analysis['rules']),2)

class SavedRequirementsTests(unittest.TestCase):
    def setUp(self):
        from workbench.coordinator import Coordinator
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.c=Coordinator(self.tmp.name);self.addCleanup(lambda:self.c.close());self.root=Path(self.tmp.name)

    def ready(self,source=COBOL,extra=None):
        self.c.create(MANIFEST,{'ELIGIBLE.cbl':source,**(extra or {})},requirements_selection=True)
        self.c.start('process-a');self.c.advance('process-a')
        return self.c.requirements_view('process-a')

    def save(self,model,exclude=None):
        body={'catalog_hash':model['catalog_hash'],'revision':model['revision'],'excluded_ids':exclude or [],'saved_by':'Fictional scope operator'}
        return self.c.save_requirements('process-a',body),body

    def complete(self):
        from io import BytesIO
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(self.c.artifact('process-a','review/sme-checklist.xlsx').read_bytes()))
        for row in range(2,book['Checklist'].max_row+1):
            book['Checklist'].cell(row,5,'Yes');book['Checklist'].cell(row,7,'Fictional test reviewer')
        out=BytesIO();book.save(out);book.close()
        self.c.import_answers('process-a',out.getvalue(),'Fictional test reviewer')
        self.c.advance('process-a');self.c.advance('process-a')
        return self.c.ledger.get('process-a')

    def test_save_generates_selected_logic_and_reports_every_omission(self):
        from workbench.target import run_generated
        from workbench.domain import decode
        model=self.ready(extra={'JOBA.jcl':'//JOBA JOB\n//S010 EXEC PGM=ELIGIBLE\n'})
        rule=[r for r in model['items'] if r['kind']=='rule'][1]
        saved,_=self.save(model,[rule['id']])
        self.assertTrue((self.root/'processes/process-a/analysis/requirements.md').is_file())
        self.assertEqual(self.c.artifact('process-a',saved['requirements_artifact']).read_bytes(),(self.root/'processes/process-a/analysis/requirements.md').read_bytes())
        self.c.advance('process-a');doc=self.c.ledger.get('process-a')
        code=(self.root/'shared/target/python'/(doc['program_versions']['ELIGIBLE']+'.py')).read_text()
        result=run_generated(code,{'AGE':20,'ACTIVE':'N','DECISION':' '})
        self.assertEqual(result['record']['DECISION'],'Y');self.assertEqual(len(result['trace']),1)
        self.assertNotIn(rule['rule_id'],code)
        doc=self.complete();self.assertEqual(doc['status'],'COMPLETED',doc['blockers'])
        coverage=decode(self.c.artifact('process-a','reports/report-0001/coverage.json').read_bytes())
        rows=[r for r in coverage['rows'] if r.get('requirements_excluded')]
        self.assertEqual(len(rows),rule['end_line']-rule['start_line']+1)
        self.assertEqual(len(coverage['rows']),len(COBOL.splitlines())+2)
        self.assertTrue(all(r['reason']=='Not converted because selected No in requirements.' and not r['target_mappings'] for r in rows))
        self.assertTrue(all(r['evidence'][0]['path']==saved['requirements_artifact'] for r in rows))
        inventory=decode(self.c.artifact('process-a','reports/report-0001/rules.json').read_bytes())
        self.assertEqual(inventory['summary']['business_rule']['excluded_by_requirements'],1)
        self.assertEqual(inventory['summary']['business_rule']['converted_verified'],1)
        comparison=inventory['requirements_comparison'];self.assertEqual(len(comparison['items']),model['total'])
        self.assertTrue(all(r['status']!='unverified' for r in comparison['items']))
        report=self.c.artifact('process-a','reports/report-0001/rules.html').read_text()
        self.assertIn('Not converted because selected No in requirements.',report)
        self.assertIn('Complete source breakdown and replacements',report)
        self.assertTrue(doc['runs'][0]['programs']['ELIGIBLE']['adversarial']['passed'])
        self.assertTrue(doc['logic_validation']['complete'])
        self.assertEqual(doc['accepted_inventory']['verified_version_ids']['batch_cobol'],[])
        estimate=self.c.economics('process-a',{'cohorts':[{'id':'programs','unit':'batch_cobol','total':100}]})
        self.assertEqual(estimate['forecast']['cohorts'][0]['remaining_units'],100)

    def test_changed_draft_does_not_replace_pinned_conversion_input(self):
        model=self.ready();saved,_=self.save(model)
        (self.root/'processes/process-a/analysis/requirements.md').write_text('Changed draft without Save')
        self.c.advance('process-a');self.assertEqual(len(self.c.ledger.get('process-a')['analysis']['rules']),2)
        self.assertIn(b'Fictional scope operator',self.c.artifact('process-a',saved['requirements_artifact']).read_bytes())

    def test_tampered_immutable_requirements_cannot_run(self):
        model=self.ready();saved,_=self.save(model)
        path=self.c.artifact('process-a',saved['requirements_artifact']);path.write_bytes(path.read_bytes()+b'changed')
        with self.assertRaises(ValidationError):self.c.sources(self.c.ledger.get('process-a'))

    def test_save_retry_is_idempotent_and_scope_locks_with_packet(self):
        model=self.ready();saved,body=self.save(model)
        retry=self.c.save_requirements('process-a',body)
        self.assertEqual(retry['requirements_artifact'],saved['requirements_artifact'])
        self.c.advance('process-a');view=self.c.requirements_view('process-a');self.assertFalse(view['editable'])
        with self.assertRaisesRegex(ValidationError,'locked'):self.save(view)

    def test_stale_unknown_and_duplicate_choices_fail_without_conversion(self):
        model=self.ready()
        for body in [dict(catalog_hash='0'*64,revision=0,excluded_ids=[],saved_by='Tester'),dict(catalog_hash=model['catalog_hash'],revision=True,excluded_ids=[],saved_by='Tester'),dict(catalog_hash=model['catalog_hash'],revision=0,excluded_ids=['missing'],saved_by='Tester'),dict(catalog_hash=model['catalog_hash'],revision=0,excluded_ids=[model['items'][0]['id']]*2,saved_by='Tester')]:
            with self.assertRaises(ValidationError):self.c.save_requirements('process-a',body)
        self.assertFalse(self.c.ledger.get('process-a')['packet_issued'])
        self.assertFalse(list((self.root/'shared/target').rglob('*.py')))

    def test_excluding_written_input_dependency_blocks_retained_consumer(self):
        model=self.ready(COBOL.replace('IF ACTIVE = "N"','IF DECISION = "Y"'))
        omitted=next(r for r in model['items'] if r['kind']=='rule');self.save(model,[omitted['id']]);self.c.advance('process-a')
        doc=self.c.ledger.get('process-a');self.assertFalse(doc['program_versions'])
        self.assertTrue(any('reads fields written' in b['message'] for b in doc['blockers']))

    def test_no_required_layout_never_runs_hidden_validation(self):
        model=self.ready();omitted=next(r for r in model['items'] if r['kind']=='data_layout');self.save(model,[omitted['id']]);self.c.advance('process-a')
        doc=self.c.ledger.get('process-a');self.assertFalse(doc['program_versions']);self.assertTrue(any(b['kind']=='requirements_dependency' for b in doc['blockers']))

    def test_no_jcl_prevents_job_execution_even_with_supported_program(self):
        model=self.ready(extra={'JOBA.jcl':'//JOBA JOB\n//S010 EXEC PGM=ELIGIBLE\n'})
        omitted=next(r for r in model['items'] if r['source_path']=='JOBA.jcl');self.save(model,[omitted['id']]);self.c.advance('process-a');doc=self.complete()
        self.assertEqual(doc['status'],'COMPLETED_WITH_BLOCKERS');self.assertFalse(any(p.endswith('jobs.py') for p in doc['artifacts']))

    def test_missing_and_unknown_files_are_visible_without_selectable_false_scope(self):
        model=self.ready(extra={'OTHER.cbl':COBOL.replace('PROGRAM-ID. ELIGIBLE.','PROGRAM-ID. OTHER.')})
        self.assertEqual(len(model['files']),2)
        extra=next(f for f in model['files'] if f['path']=='OTHER.cbl');self.assertFalse(extra['in_process_scope'])
        self.assertFalse(self.c.requirements_view('process-a',path='OTHER.cbl')['items'])
        with self.assertRaises(ValidationError):self.c.requirements_view('process-a',path='absent')

    def test_selection_catalog_reproduces_after_filtering_and_paging(self):
        from workbench.requirements import catalog
        model=self.ready();omitted=next(r for r in model['items'] if r['kind']=='rule');self.save(model,[omitted['id']]);self.c.advance('process-a')
        doc=self.c.ledger.get('process-a');self.assertEqual(catalog(doc['analysis'])['hash'],model['catalog_hash'])
        self.assertFalse(next(r for r in self.c.requirements_view('process-a')['items'] if r['id']==omitted['id'])['selected'])

    def test_canonical_markdown_rejects_table_tampering_or_malformed_encoding(self):
        from workbench.requirements import parse_markdown,catalog
        model=self.ready();analysis=self.c.ledger.get('process-a')['analysis'];saved,_=self.save(model)
        raw=self.c.artifact('process-a',saved['requirements_artifact']).read_bytes()
        for changed in (raw.replace(b'| Yes |',b'| No |',1),b'\xff',raw.replace(b'"schema_version": 1',b'"schema_version": true')):
            if changed!=raw:
                with self.assertRaises(ValidationError):parse_markdown(changed,catalog(analysis))

    def test_changed_adapter_catalog_requires_new_scope_save_without_losing_no(self):
        from unittest.mock import patch
        from copy import deepcopy
        from workbench.source import analyze_sources
        model=self.ready();omitted=next(r for r in model['items'] if r['kind']=='rule');saved,_=self.save(model,[omitted['id']])
        original=analyze_sources(self.c.sources(saved),{k:v for k,v in saved.items() if k!='requirements'})
        changed=deepcopy(original);changed['programs']['ELIGIBLE']['rules'][1]['plain']+=' Updated adapter explanation.'
        with patch('workbench.coordinator.analyze_sources',return_value=changed):self.c.advance('process-a')
        view=self.c.requirements_view('process-a');self.assertEqual(view['status'],'WAITING_REQUIREMENTS')
        self.assertEqual(view['revision'],1);self.assertIn(omitted['id'],view['excluded_ids']);self.assertFalse(self.c.ledger.get('process-a')['packet_issued'])
        self.assertTrue(self.c.artifact('process-a',saved['requirements_artifact']).exists())

    def test_selected_downstream_rule_cannot_hide_omitted_upstream_write(self):
        files={'ELIGIBLE.cbl':COBOL,'FOLLOW.cbl':COBOL.replace('PROGRAM-ID. ELIGIBLE.','PROGRAM-ID. FOLLOW.').replace('IF AGE >= 18','IF DECISION = "Y"')}
        manifest=MANIFEST+'| 1 | JOBA | 2 | S020 | FOLLOW | OUTPUT | FINAL | RC=0 |\n'
        self.c.create(manifest,files,requirements_selection=True);self.c.start('process-a');self.c.advance('process-a')
        model=self.c.requirements_view('process-a');omitted=next(r for r in model['items'] if r['kind']=='rule' and r['program']=='ELIGIBLE')
        self.save(model,[omitted['id']]);self.c.advance('process-a');doc=self.c.ledger.get('process-a')
        self.assertNotIn('FOLLOW',doc['program_versions'])
        self.assertTrue(any(b['kind']=='requirements_dependency' and 'upstream' in b['message'] for b in doc['blockers']))

    def test_saved_markdown_survives_restart_and_failed_convenience_copy(self):
        from unittest.mock import patch
        from workbench.coordinator import Coordinator
        model=self.ready()
        body={'catalog_hash':model['catalog_hash'],'revision':model['revision'],'excluded_ids':[],'saved_by':'Fictional scope operator'}
        with patch('workbench.domain.atomic_bytes',side_effect=OSError('disk copy failed')):
            with self.assertRaises(OSError):self.c.save_requirements('process-a',body)
        saved=self.c.ledger.get('process-a');self.assertEqual(saved['requirements']['revision'],1)
        self.c.close();self.c=Coordinator(self.root)
        retry=self.c.save_requirements('process-a',body)
        self.assertEqual(retry['requirements_artifact'],saved['requirements_artifact'])
        self.assertEqual(len(list((self.root/'processes/process-a/analysis/requirements').glob('*.md'))),1)

    def test_http_intake_waits_for_ui_save_and_copilot_cannot_read_requirements(self):
        from fastapi.testclient import TestClient
        from workbench.api import create_app
        from tools.workbench_mcp import WorkflowBridge
        from unittest.mock import patch
        self.c.close();app=create_app(self.root);self.c=app.state.coordinator
        client=TestClient(app,base_url='http://127.0.0.1:8765');token=client.get('/api/state').json()['token']
        headers={'Origin':'http://127.0.0.1:8765','X-Workbench-Token':token}
        response=client.post('/api/intake',json={'manifest':MANIFEST,'sources':{'ELIGIBLE.cbl':COBOL,'顧客.cbl':COBOL.replace('PROGRAM-ID. ELIGIBLE.','PROGRAM-ID. OTHER.')},'assistant_mode':'deterministic'},headers=headers)
        self.assertEqual(response.status_code,200,response.text)
        self.c.start('process-a');self.c.advance('process-a')
        full_source=client.get('/api/process/process-a/requirements/source?path=ELIGIBLE.cbl')
        unicode_source=client.get('/api/process/process-a/requirements/source',params={'path':'顧客.cbl'})
        self.assertEqual(unicode_source.status_code,200,unicode_source.text)
        self.assertIn("filename*=UTF-8''%",unicode_source.headers['content-disposition'])
        self.assertEqual(full_source.text,COBOL);self.assertIn('attachment',full_source.headers['content-disposition'])
        self.assertEqual(client.get('/api/process/process-a/requirements/source?path=../other').status_code,400)
        model=client.get('/api/process/process-a/requirements').json();self.assertEqual(model['status'],'WAITING_REQUIREMENTS')
        body={'catalog_hash':model['catalog_hash'],'revision':0,'excluded_ids':[],'saved_by':'Fictional UI operator'}
        self.assertEqual(client.post('/api/process/process-a/requirements',json=body).status_code,403)
        self.assertEqual(client.post('/api/process/process-a/requirements',json=body,headers=headers).status_code,200)
        self.c.advance('process-a');bridge=WorkflowBridge()
        view=client.get('/api/process/process-a/requirements').json()
        self.assertEqual(view['revision'],1);self.assertFalse(view['editable'])
        with patch.object(bridge,'request') as request:
            with self.assertRaises(ValidationError):bridge.call('workbench_requirements',{'process_id':'process-a'})
            request.assert_not_called()
        state=client.get('/api/state').json()['processes'][0]
        self.assertNotIn('excluded_units',state['analysis']['requirements'])

    def test_pagination_keeps_distinct_identical_rule_occurrences(self):
        source=COBOL.replace('  GOBACK.','')+('\n  IF ACTIVE = "Y"\n    CONTINUE\n  ELSE\n    CONTINUE\n  END-IF.\n'*55)+'  GOBACK.\n'
        model=self.ready(source);items=list(model['items'])
        while model['has_more']:
            model=self.c.requirements_view('process-a',after=model['next_after']);items.extend(model['items'])
        self.assertEqual(len(items),len({r['id'] for r in items}));self.assertEqual(sum(r['kind']=='rule' for r in items),57)
        last=items[-2];self.save(model,[last['id']]);self.c.advance('process-a')
        self.assertFalse(any(r['id']==last['rule_id'] for r in self.c.ledger.get('process-a')['analysis']['rules']))
