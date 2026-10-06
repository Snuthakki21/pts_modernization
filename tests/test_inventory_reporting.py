"""Declared estate counts cannot create assets; large exports keep every row."""
import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from workbench.coverage import COLUMNS, write_coverage
from workbench.domain import ValidationError, encode, sha
from workbench.inventory import (default_inventory, report_inventory, snapshot_inventory,
                                 validate_inventory, validate_snapshot)
import test_workflow
from test_source import COBOL


def process(pid='process-a',source='PROGRAM-ID. SAMPLE.\nGOBACK.\n',demo=False):
    digest=sha(source)
    asset={'id':sha('program:'+digest),'kind':'cobol_program','name':'SAMPLE','path':'SAMPLE.cbl','source_hash':digest,'source_text':source,'selected':True}
    return {'id':pid,'demo':demo,'jobs':[{'name':'RUN','steps':[{'program':'SAMPLE'}]}],
            'analysis':{'assets':[asset],'programs':{'SAMPLE':asset},'graph':[]}}


def verified_coverage(path='SAMPLE.cbl'):
    return {'summary':{'integrity_errors':[]},'rows':[{'source_path':path,'disposition':'mapped_verified'}]}


class InventoryReportingTests(unittest.TestCase):
    def test_declared_total_and_category_sum_remain_distinct(self):
        model=report_inventory(process())
        self.assertEqual(model['declared_total'],1829)
        self.assertEqual(model['category_total'],1826)
        self.assertEqual(model['unreconciled_count'],3)
        self.assertEqual([row['baseline_count'] for row in model['rows']],[166,599,174,245,80,108,450,4])
        self.assertTrue(all(row['baseline_status']=='USER_REPORTED_UNVERIFIED' for row in model['rows']))
        self.assertEqual({location['name'] for location in model['application_locations']},{'WEDELX','TranRepository'})
        self.assertEqual(sum(row['process_converted_count'] for row in model['rows']),0)

    def test_missing_evidence_remains_unknown_instead_of_baseline_count(self):
        rows={row['key']:row for row in report_inventory(process())['rows']}
        self.assertEqual(rows['batch_cobol']['process_observed_count'],1)
        for key in ('jcl_jobs','jcl_procs','copybooks','cics_screens','db2_tables','ca7_schedules','mq'):
            self.assertIsNone(rows[key]['process_observed_count'])
            self.assertEqual(rows[key]['process_converted_count'],0)

    def test_mutable_inventory_is_frozen_and_snapshot_cannot_be_repinned(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);(root/'knowledge').mkdir()
            path=root/'knowledge/inventory-baseline.json';path.write_bytes(encode(default_inventory()))
            frozen=snapshot_inventory(root)
            changed=default_inventory();changed['categories'][0]['count']=167;path.write_bytes(encode(changed))
            self.assertEqual(frozen['document']['categories'][0]['count'],166)
            self.assertEqual(snapshot_inventory(root)['document']['categories'][0]['count'],167)
            forged=copy.deepcopy(frozen);forged['document']['categories'][0]['count']=900
            with self.assertRaisesRegex(ValidationError,'hash'):validate_snapshot(forged)

    def test_inventory_schema_rejects_duplicates_and_false_verification(self):
        duplicate=default_inventory();duplicate['categories'][-1]=copy.deepcopy(duplicate['categories'][0])
        with self.assertRaises(ValidationError):validate_inventory(duplicate)
        false=default_inventory();false['status']='VERIFIED'
        with self.assertRaisesRegex(ValidationError,'unverified'):validate_inventory(false)

    def test_shared_versions_deduplicate_cumulative_credit_and_keep_memberships(self):
        first=process();first_inventory=report_inventory(first,coverage=verified_coverage())
        first.update(report_verified=True,accepted_inventory=first_inventory)
        second=process('process-b')
        row=report_inventory(second,[first],verified_coverage())['rows'][0]
        self.assertEqual(row['cumulative_converted_count'],1)
        self.assertEqual(row['portfolio_unique_versions'],1)
        self.assertEqual(row['portfolio_unique_assets'],1)
        self.assertEqual(row['portfolio_memberships'],2)
        self.assertEqual(row['remaining_vs_baseline'],165)

    def test_changed_source_keeps_versions_separate_and_logical_asset_single(self):
        first=process();second=process('process-b','PROGRAM-ID. SAMPLE.\nSTOP RUN.\n')
        row=report_inventory(second,[first],verified_coverage())['rows'][0]
        self.assertEqual(row['portfolio_unique_versions'],2)
        self.assertEqual(row['portfolio_unique_assets'],1)
        self.assertEqual(row['portfolio_memberships'],2)

    def test_unaccepted_report_and_demo_do_not_add_cumulative_credit(self):
        first=process();first['accepted_inventory']=report_inventory(first,coverage=verified_coverage())
        second=process('process-b')
        row=report_inventory(second,[first],None)['rows'][0]
        self.assertEqual(row['cumulative_converted_count'],0)
        demonstration=process('demo-a',demo=True)
        row=report_inventory(demonstration,[first],verified_coverage())['rows'][0]
        self.assertEqual(row['process_converted_count'],1)
        self.assertEqual(row['cumulative_converted_count'],0)

    def test_screen_declarations_are_not_cics_transactions(self):
        doc=process();doc['analysis']['assets'].append({'kind':'bms_map','name':'MAPSET','path':'screen.bms','source_text':'SCREEN DFHMDI\n','source_hash':sha('SCREEN DFHMDI\n'),'screens':['SCREEN'],'selected':True})
        rows={row['key']:row for row in report_inventory(doc,coverage=verified_coverage())['rows']}
        self.assertEqual(rows['cics_screens']['process_observed_count'],1)
        self.assertEqual(rows['cics_screens']['process_converted_count'],0)

    def test_distinct_db2_references_are_visible_and_not_converted_tables(self):
        doc=process()
        for path,text in [('a.sql','SELECT * FROM APP.CUSTOMER;'),('b.sql','SELECT ID FROM APP.CUSTOMER;')]:
            doc['analysis']['assets'].append({'kind':'sql','name':path,'path':path,'source_text':text,'source_hash':sha(text),'tables':['APP.CUSTOMER'],'selected':True})
        row=report_inventory(doc,coverage=verified_coverage())['rows'][5]
        self.assertEqual(row['process_observed_count'],1)
        self.assertEqual(row['process_unique_versions'],2)
        self.assertEqual(row['process_converted_count'],0)

    def test_lineage_aggregate_unknown_and_cics_screen_count_are_preserved(self):
        doc=process();doc['lineage']={'nodes':[],'process_assets':[{'kind':'cics_screens','observed_count':3,'status':'OBSERVED_STATIC_CLOSURE'}, {'kind':'mq','observed_count':0,'status':'UNVERIFIED'}]}
        rows={row['key']:row for row in report_inventory(doc)['rows']}
        self.assertEqual(rows['cics_screens']['process_observed_count'],3)
        self.assertIsNone(rows['mq']['process_observed_count'])

    def test_cics_bound_program_is_not_counted_as_batch_without_batch_call_evidence(self):
        doc=process();online=copy.deepcopy(doc['analysis']['assets'][0]);online.update(name='ONLINE',path='ONLINE.cbl',source_text='PROGRAM-ID. ONLINE.',source_hash=sha('PROGRAM-ID. ONLINE.'))
        doc['analysis']['assets'].append(online)
        doc['lineage']={'nodes':[{'id':'root','kind':'manifest_job','name':'RUN','selected':True},
                                {'id':'batch','kind':'program','name':'SAMPLE','path':'SAMPLE.cbl','selected':True,'resolution':'local_source'},
                                {'id':'online','kind':'program','name':'ONLINE','path':'ONLINE.cbl','selected':True,'resolution':'local_source'}],
                        'edges':[{'source':'root','target':'batch','kind':'manifest_step','resolution':'resolved'},
                                 {'source':'batch','target':'online','kind':'has_application_binding','resolution':'resolved'}],
                        'process_assets':[{'kind':'batch_cobol','observed_count':2,'status':'OBSERVED_STATIC_CLOSURE'}]}
        row=report_inventory(doc)['rows'][0]
        self.assertEqual(row['process_observed_count'],1)
        self.assertEqual(row['process_unique_versions'],1)

    def test_coverage_uses_write_only_split_sheets_without_losing_text(self):
        original='\U0001f642'*20000
        rows=[]
        for number in range(12):
            row={column:None for column in COLUMNS};row.update(source_path='source.cbl',source_line=number+1,source_text=original if number==7 else 'line '+str(number),disposition='blocked')
            rows.append(row)
        files=[{'path':'source-'+str(number)+'.cbl','source_text':'sample','classification':{'kind':'cobol_program','evidence':['original classification']}} for number in range(7)]
        model={'summary':{'source_lines':12},'denominators':{'source_lines':'All physical lines'},'files':files,'rows':rows}
        from openpyxl import Workbook
        modes=[]
        def workbook(*args,**kwargs):modes.append(kwargs.get('write_only',False));return Workbook(*args,**kwargs)
        with tempfile.TemporaryDirectory() as directory,patch('workbench.coverage.EXCEL_MAX_ROWS',5),patch('workbench.coverage.Workbook',workbook):
            write_coverage(model,directory)
            book=load_workbook(Path(directory)/'coverage.xlsx')
            try:
                line_sheets=[sheet for sheet in book if sheet.title.startswith('Source lines')]
                self.assertEqual([sheet.max_row for sheet in line_sheets],[5,5,5])
                self.assertEqual(sum(sheet.max_row-1 for sheet in line_sheets),12)
                self.assertTrue(all(sheet.max_row<=5 for sheet in book))
                chunks=[row for sheet in book if sheet.title.startswith('Text chunks') for row in sheet.iter_rows(min_row=2,values_only=True)]
                relevant=[row for row in chunks if row[0]=='Source lines 2' and row[1]==5 and row[2]=='source_text']
                self.assertEqual(''.join(row[4] for row in relevant),original)
                self.assertTrue(all(row[5]==sha(original) for row in relevant))
                self.assertEqual(json.loads((Path(directory)/'coverage.json').read_text()),model)
                self.assertIn('classification',[cell.value for cell in book['Files'][1]])
            finally:book.close()
        self.assertEqual(modes,[True])


class InventoryArtifactTests(unittest.TestCase):
    setUp=test_workflow.WorkflowTests.setUp
    answer=test_workflow.WorkflowTests.answer

    def verified(self):
        self.c.create(test_workflow.MANIFEST,{'ELIGIBLE.cbl':COBOL})
        self.c.start('process-a');self.c.advance('process-a')
        self.c.import_answers('process-a',self.answer(),'Fictional inventory test reviewer')
        self.c.advance('process-a')
        return self.c.ledger.get('process-a')

    def test_primary_report_deck_and_workbook_use_identical_eight_category_counts(self):
        from pptx import Presentation
        self.verified();self.c.advance('process-a')
        model=json.loads(self.c.artifact('process-a','reports/report-0001/metrics.json').read_text())
        self.assertEqual(model['inventory'],model['executive']['inventory'])
        self.assertEqual(json.loads(model['metrics']['estate_inventory_json']),model['inventory'])
        expected=[[row['label'],str(row['baseline_count']),'Unknown' if row['process_observed_count'] is None else str(row['process_observed_count']),str(row['process_converted_count']),str(row['remaining_vs_baseline'])] for row in model['inventory']['rows']]
        deck=Presentation(self.c.artifact('process-a','reports/report-0001/management.pptx'))
        table=next(shape.table for shape in deck.slides[0].shapes if shape.has_table)
        self.assertEqual([[cell.text for cell in row.cells] for row in list(table.rows)[1:]],expected)
        book=load_workbook(self.c.artifact('process-a','reports/report-0001/metrics.xlsx'),data_only=True)
        try:
            values=list(book['Estate inventory'].iter_rows(min_row=2,max_row=9,values_only=True))
            self.assertEqual([row[0] for row in values],[row['label'] for row in model['inventory']['rows']])
            self.assertEqual([row[1] for row in values],[row['baseline_count'] for row in model['inventory']['rows']])
        finally:book.close()
        text=self.c.artifact('process-a','reports/report-0001/executive-report.html').read_text()
        self.assertIn('1,829',text);self.assertIn('1,826',text);self.assertIn('3 unreconciled',text)
        self.assertIn('WEDELX and TranRepository',text)

    def test_changed_frozen_baseline_revokes_coverage_credit_without_target_replay(self):
        from workbench.coverage import build_coverage
        doc=self.verified()
        self.assertTrue(build_coverage(doc,self.root)['summary']['completion_eligible'])
        path=self.root/'processes/process-a/analysis/inventory-baseline.json'
        path.write_text('{"forged":true}')
        with patch('workbench.fixtures.verify_program') as replay:
            model=build_coverage(doc,self.root)
        replay.assert_not_called()
        self.assertEqual(model['summary']['verified_applicable_lines'],0)
        self.assertTrue(any('inventory baseline' in error for error in model['summary']['integrity_errors']))

    def test_changed_preserved_sme_return_revokes_rows_and_inventory_credit_before_replay(self):
        from workbench.coverage import build_coverage
        from workbench.fixtures import verify_program
        doc=self.verified()
        intact=build_coverage(doc,self.root)
        self.assertTrue(intact['summary']['completion_eligible'])
        path=self.root/'processes/process-a/input/sme-return.xlsx'
        path.write_bytes(path.read_bytes()+b'changed preserved review receipt')
        with self.assertRaisesRegex(ValidationError,'SME return changed'):self.c.review_integrity(doc)
        with patch('workbench.fixtures.verify_program',wraps=verify_program) as replay:
            coverage=build_coverage(doc,self.root)
        self.assertEqual(coverage['summary']['verified_applicable_lines'],0)
        replay.assert_not_called()
        self.assertTrue(any('SME' in error for error in coverage['summary']['integrity_errors']))
        row=report_inventory(doc,coverage=coverage)['rows'][0]
        self.assertEqual(row['process_converted_count'],0)
        self.assertEqual(row['cumulative_converted_count'],0)

    def test_recorded_lineage_location_or_journal_integrity_failure_revokes_credit(self):
        from workbench.coverage import build_coverage
        from workbench.fixtures import verify_program
        doc=self.verified();self.c.review_integrity(doc)
        self.assertTrue(build_coverage(doc,self.root)['summary']['completion_eligible'])
        for failure in ('Frozen lineage scope changed','Frozen input location evidence changed','Frozen discovery provenance changed'):
            with self.subTest(failure=failure):
                changed=copy.deepcopy(doc)
                changed['blockers'].append({'kind':'evidence_integrity','message':failure})
                with patch('workbench.fixtures.verify_program',wraps=verify_program) as replay:
                    coverage=build_coverage(changed,self.root)
                replay.assert_not_called()
                self.assertEqual(coverage['summary']['verified_applicable_lines'],0)
                self.assertTrue(any(failure in error for error in coverage['summary']['integrity_errors']))
                self.assertEqual(report_inventory(changed,coverage=coverage)['rows'][0]['process_converted_count'],0)


if __name__=='__main__':unittest.main()
