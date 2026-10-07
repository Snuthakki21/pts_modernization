"""Fictional CICS layout qualifications through the real process engine."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from workbench.coordinator import Coordinator
from workbench.domain import encode,sha
from workbench.coverage import build_coverage
from workbench.rule_inventory import build_rule_inventory
from workbench.comparison import comparison_page,render_comparison_report
from workbench.online_runtime import create_app
from workbench.online import compare_screen_http,packaged_runtime
from workbench.cics import prepare_screen,emit_screen,plan_screen_cases,selected_screens
from test_cics import BMS,COBOL,fixture
from test_cics_requirements_workflow import MANIFEST,CSD


class CicsDeliveryTests(unittest.TestCase):
    def setUp(self):
        staging=Path('.implementation/tmp');staging.mkdir(parents=True,exist_ok=True)
        self.temp=tempfile.TemporaryDirectory(dir=staging);self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name);self.c=Coordinator(self.root);self.addCleanup(self.c.close)

    def process(self,exclude=False,response='Yes',manifest=MANIFEST,sources=None):
        notes=self.root/'.implementation/tmp/process-note.md';notes.parent.mkdir(parents=True,exist_ok=True);notes.write_text('# Fictional process\nT001 uses PGM, MAPS/SCREEN. Native parity is not asserted.\n')
        self.c.create(manifest,sources or {'PGM.cbl':COBOL,'FORM.bms':BMS,'T001.csd':CSD},requirements_selection=True,
                      process_notes=notes)
        self.c.start('cics-ui');self.c.advance('cics-ui')
        self.assertEqual(self.c.ledger.get('cics-ui')['status'],'WAITING_REQUIREMENTS')
        model=self.c.requirements_view('cics-ui')
        omissions=[r['id'] for r in model['items'] if r['kind']=='screen_field' and r['field'].startswith('STATIC_')] if exclude else []
        self.c.save_requirements('cics-ui',{'catalog_hash':model['catalog_hash'],'revision':model['revision'],'excluded_ids':omissions,'saved_by':'Fictional CICS scope operator'})
        self.c.advance('cics-ui')
        self.assertEqual(self.c.ledger.get('cics-ui')['status'],'WAITING_SME')
        raw=self.c.artifact('cics-ui','review/sme-checklist.xlsx').read_bytes()
        book=load_workbook(BytesIO(raw));sheet=book['Checklist']
        for row in range(2,sheet.max_row+1):sheet.cell(row,5,response);sheet.cell(row,7,'Fictional CICS fixture reviewer')
        out=BytesIO();book.save(out);book.close()
        self.c.import_answers('cics-ui',out.getvalue(),'Fictional CICS fixture reviewer')
        self.c.advance('cics-ui');self.c.advance('cics-ui')
        doc=self.c.ledger.get('cics-ui')
        self.assertEqual(doc['status'],'COMPLETED_WITH_BLOCKERS',doc['blockers'])
        return doc

    def test_actual_engine_issues_source_bound_screen_python_fastapi_and_units_per_process(self):
        doc=self.process();run=doc['runs'][-1];screen=selected_screens(doc['analysis'])[0]
        self.assertEqual(doc['online_delivery']['status'],'HTTP_COMPARISON_PASSED')
        from workbench.factory import factory_view
        view=factory_view(doc);self.assertEqual(view['counts']['implemented_layout_candidates'],1);self.assertEqual(view['counts']['implemented_api_candidates'],0)
        self.assertEqual(len(view['transactions'][0]['screen_apis']),1)
        self.assertEqual(view['validation']['screen_layout_count'],1);self.assertGreaterEqual(view['validation']['unit_tests_run'],30)
        self.assertFalse(doc['online_delivery']['native_cics_verified'])
        result=run['screens'][screen['id']]
        self.assertTrue(result['passed']);self.assertTrue(result['adversarial']['passed']);self.assertTrue(result['unit_tests']['passed'])
        self.assertGreaterEqual(result['coverage']['fields']['ACCT'],20)
        self.assertEqual(result['unit_tests']['tests_run'],result['cases'])
        base=self.root/'processes/cics-ui';package=base/'target'/run['id']/'online'
        spec=json.loads((package/'contract.json').read_text())
        self.assertFalse(spec['T001']['record_supported']);self.assertEqual(spec['T001']['controller_status'],'UNVERIFIED')
        self.assertEqual(spec['T001']['screens'][0]['source_hash'],sha(BMS))
        receipt=json.loads((package/'verification.json').read_text())
        self.assertEqual(receipt['screens']['T001:'+screen['id']]['cases'],result['cases'])
        self.assertTrue(receipt['passed']);self.assertTrue((base/doc['requirements_artifact']).is_file())
        self.assertTrue((base/'input/sources/T001.csd').is_file());self.assertTrue((base/'analysis/process-context.json').is_file())
        self.assertTrue((package/'launch.cmd').is_file());self.assertEqual((package/'runtime.py').read_bytes(),packaged_runtime())

    def test_comparison_credits_layout_and_exposes_exact_native_controller_gaps(self):
        doc=self.process();coverage=build_coverage(doc,self.root);inventory=build_rule_inventory(doc,coverage,self.root)
        fields=[row for row in coverage['rows'] if row['source_path']=='FORM.bms' and 'DFHMDF' in row['source_text']]
        self.assertTrue(fields);self.assertTrue(all(row['disposition']=='platform_replaced_verified' for row in fields),fields)
        controllers=[row for row in coverage['rows'] if row['source_path']=='PGM.cbl' and 'EXEC CICS' in row['source_text']]
        self.assertTrue(all(row['disposition']=='blocked' for row in controllers));self.assertTrue(all('no verified Python/FastAPI replacement' in row['reason'] for row in controllers))
        self.assertFalse(coverage['summary']['completion_eligible'])
        program=inventory['programs'][0];self.assertTrue(any(r['cics_screens']==['MAPS/SCREEN'] for r in inventory['rules'] if r.get('cics_screens')))
        page=comparison_page(inventory,status='gaps')
        self.assertEqual(page['cics_contract_version'],1);self.assertTrue(page['rules']);self.assertTrue(page['counts']['verified']>0)
        self.assertTrue(all(r['id'] in program['rule_ids'] for r in inventory['rules'] if r.get('cics_screens')))
        actions=[rule for rule in inventory['rules'] if rule['description'].startswith(('CICS SEND MAP','CICS RECEIVE MAP'))]
        self.assertEqual(len(actions),2);self.assertTrue(all(rule['cics_screens']==['MAPS/SCREEN'] for rule in actions))
        self.assertIn('separately replayed',page['basis'])
        html=render_comparison_report(inventory)
        self.assertIn('Mainframe and Python / FastAPI comparison',html);self.assertIn('SEND',html);self.assertIn('Gaps only',html)

    def test_shared_map_file_keeps_each_programs_screen_ownership_exact(self):
        second=BMS[BMS.index('SCREEN DFHMDI'):BMS.index(' DFHMSD TYPE=FINAL')].replace('SCREEN','OTHER').replace('ACCT','BACC').replace('MSG','BMSG')
        maps=BMS.replace(' DFHMSD TYPE=FINAL',second+' DFHMSD TYPE=FINAL')
        manifest=MANIFEST+'| T002 | OTHER | MAPS | OTHER |\n'
        sources={'PGM.cbl':COBOL,'OTHER.cbl':COBOL.replace('PGM.','OTHER.').replace('SCREEN','OTHER'),
                 'FORM.bms':maps,'T001.csd':CSD+'DEFINE TRANSACTION(T002) PROGRAM(OTHER)\nDEFINE PROGRAM(OTHER)\n'}
        doc=self.process(manifest=manifest,sources=sources)
        inventory=build_rule_inventory(doc,build_coverage(doc,self.root),self.root)
        by_id={rule['id']:rule for rule in inventory['rules']}
        for program in inventory['programs']:
            for ident in program['rule_ids']:
                rule=by_id[ident]
                if rule.get('cics_screens'):
                    self.assertIn(program['program'],rule['cics_programs'])
                    if len(rule['cics_screens'])==1:self.assertEqual(rule['cics_screens'],['MAPS/'+('SCREEN' if program['program']=='PGM' else 'OTHER')])

    def test_observed_db2_metadata_is_context_and_missing_semantics_stay_named(self):
        from test_lineage_discovery import Db2CatalogReceiptTests
        receipt=Db2CatalogReceiptTests().receipt();receipt['description_complete']=False
        program=COBOL.replace('  EXEC CICS RETURN END-EXEC.', '  EXEC SQL SELECT * FROM APP.CUSTOMER END-EXEC.\n  EXEC CICS RETURN END-EXEC.')
        sources={'PGM.cbl':program,'FORM.bms':BMS,'T001.csd':CSD,'db2/description.json':encode(receipt).decode()}
        doc=self.process(sources=sources);coverage=build_coverage(doc,self.root)
        rows=[row for row in coverage['rows'] if row['source_path']=='db2/description.json']
        self.assertTrue(rows);self.assertTrue(all(row['disposition']=='non_executable' and not row['target_mappings'] for row in rows))
        self.assertTrue(any(blocker.get('obligation')=='COLUMN_DESCRIPTION' for blocker in doc['blockers']))
        self.assertTrue(any('metadata alone is not executable support' in blocker['message'] for blocker in doc['blockers']))
        self.assertFalse(coverage['summary']['completion_eligible'])

    def test_saved_cosmetic_no_omits_target_and_retains_exact_markdown_reason(self):
        doc=self.process(exclude=True);screen=selected_screens(doc['analysis'])[0]
        self.assertEqual(len(screen['fields']),2);self.assertEqual(len(screen['omitted_fields']),1)
        code=(self.root/'shared/target/python'/(doc['screen_versions'][screen['id']]+'.py')).read_text()
        self.assertNotIn("'value': 'Customer'",code)
        coverage=build_coverage(doc,self.root)
        omitted=[r for r in coverage['rows'] if r['requirements_excluded']]
        self.assertEqual(len(omitted),2)
        self.assertTrue(all(r['reason']=='Not converted because selected No in requirements.' and not r['target_mappings'] for r in omitted))
        self.assertTrue(all(r['evidence'][0]['path']==doc['requirements_artifact'] for r in omitted))

    def test_changed_target_http_receipt_and_unit_evidence_cannot_receive_layout_credit(self):
        doc=self.process();base=self.root/'processes/cics-ui';run=doc['runs'][-1];ident=next(iter(run['screens']))
        for path in (base/'target'/run['id']/'online/package.json',base/'target'/run['id']/'online/verification.json',base/f'tests/{run["id"]}/{ident}/unit-results.json',
                     base/f'target/{run["id"]}/{ident}.py',base/'target'/run['id']/'online/runtime.py'):
            with self.subTest(path=path):
                original=path.read_bytes();path.write_bytes(original+b' ')
                coverage=build_coverage(doc,self.root)
                self.assertFalse(any(r['disposition']=='platform_replaced_verified' for r in coverage['rows'] if r['source_path']=='FORM.bms'))
                path.write_bytes(original)

    def test_cancel_and_unresolved_human_answers_withhold_layout_credit(self):
        doc=self.process(response='No');coverage=build_coverage(doc,self.root)
        self.assertEqual(doc['online_delivery']['transactions']['T001']['state'],'LAYOUT_TESTED_CONTROLLER_UNVERIFIED')
        package=self.root/'processes/cics-ui/target'/doc['runs'][-1]['id']/'online/package.json'
        self.assertNotIn('Verified character layout',package.read_text())
        self.assertFalse(any(r['disposition']=='platform_replaced_verified' for r in coverage['rows'] if r['source_path']=='FORM.bms'))
        doc['cancel_requested']=True
        with patch('workbench.screen_delivery.verify_screen',side_effect=AssertionError('Cancelled replay executed')):
            coverage=build_coverage(doc,self.root)
        self.assertFalse(any(r['disposition']=='platform_replaced_verified' for r in coverage['rows'] if r['source_path']=='FORM.bms'))

    def test_state_polling_projects_screen_counts_and_keeps_full_receipts_immutable(self):
        doc=self.process();self.c.close()
        from workbench.api import create_app as workbench_app
        app=workbench_app(self.root);self.addCleanup(app.state.coordinator.close)
        with TestClient(app,base_url='http://127.0.0.1:8765') as client:
            displayed=client.get('/api/state').json()['processes'][0]
        analysis=displayed['analysis']['cics'];self.assertEqual(analysis['screen_count'],1)
        self.assertNotIn('units',analysis);self.assertNotIn('screens',analysis);self.assertNotIn('gaps',analysis)
        result=next(iter(displayed['runs'][-1]['screens'].values()))
        self.assertGreaterEqual(result['case_count'],30);self.assertEqual(result['difference_count'],0)
        self.assertNotIn('actual',result);self.assertNotIn('differences',result);self.assertNotIn('coverage',result)
        self.assertNotIn('screen_versions',displayed);self.assertEqual(displayed['screen_target_count'],1)
        self.assertNotIn('screen_layouts',displayed['logic_validation'])
        ident=next(iter(doc['runs'][-1]['screens']))
        raw=(self.root/f'processes/cics-ui/synthetic/{doc["runs"][-1]["id"]}/{ident}/actual-and-comparison.json').read_bytes()
        self.assertEqual(raw,encode(doc['runs'][-1]['screens'][ident]))

    def test_real_fastapi_rejects_scope_auth_duplicate_json_and_controller_execution(self):
        screen=fixture();screen['owners']=['PGM'];code=emit_screen(screen)
        spec={'T001':{'program':'PGM','record_supported':False,'screens':[screen]}}
        app=create_app({},spec,self.root/'state.sqlite','x'*32,screens={screen['id']:prepare_screen(code)})
        headers={'Authorization':'Bearer '+'x'*32};url='/api/screens/T001/'+screen['id']
        with TestClient(app) as client:
            self.assertEqual(client.post(url,json={'values':{'ACCT':'00000001'}}).status_code,401)
            self.assertEqual(client.post(url,headers={**headers,'Origin':'https://untrusted.invalid'},json={'values':{}}).status_code,403)
            self.assertEqual(client.post(url.replace('T001','OTHER'),headers=headers,json={'values':{}}).status_code,404)
            self.assertEqual(client.post(url,headers=headers,content='{"values":{},"values":{}}').status_code,422)
            self.assertEqual(client.post(url,headers=headers,content=b'x'*65537).status_code,413)
            self.assertEqual(client.post('/api/sessions',headers=headers,json={'transaction':'T001'}).status_code,409)
            self.assertEqual(client.post('/api/transactions/T001',headers=headers,json={}).status_code,409)
            suite=plan_screen_cases(screen,21)
            result=compare_screen_http(client,headers,'T001',screen,suite)
            self.assertTrue(result['passed']);self.assertGreaterEqual(result['cases'],30)
            self.assertTrue(any(row['status']==422 for row in result['actual']))
            self.assertIn(screen['id'],json.dumps(app.openapi()))
