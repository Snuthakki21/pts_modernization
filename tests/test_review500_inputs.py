"""R001–R080: bounded intake, portable paths, and immutable document review.

Only synthetic local inputs are used. These checks do not establish live
mainframe parity, native Windows behavior, or atomicity against hostile OS races.
"""
import copy
from io import BytesIO
import json
from pathlib import Path
import struct
import tempfile
import unittest
from unittest.mock import patch
import warnings
from zipfile import ZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as ET

from openpyxl import Workbook

from workbench.domain import (MAX_DOCUMENT_DEPTH, ValidationError, identity, safe_path, decode, encode,
                              checked_zip, write_new, atomic_json, sha)
from workbench.intake import HEADERS, from_rows, parse_manifest, parse_intake_xlsx
from workbench.layout import output_path, validate_workspace


ROW = [1, 'JOBONE', 1, 'STEPONE', 'PROGRAM', 'INPUT', 'OUTPUT', 'Always']
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


class Review500InputTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)

    def rows(self, *rows, name='Example process'):
        return from_rows('example-process', name, list(rows) or [ROW.copy()])

    def manifest(self, name='Example process'):
        return ('- Process ID: example-process\n- Process name: ' + name + '\n'
                '| ' + ' | '.join(HEADERS) + ' |\n'
                '|---|---|---|---|---|---|---|---|\n'
                '| ' + ' | '.join(map(str, ROW)) + ' |\n')

    def workbook(self, mutate=None):
        book = Workbook()
        sheet = book.active
        sheet.title = 'Intake'
        sheet['A1'] = 'Process ID'; sheet['B1'] = 'example-process'
        sheet['A2'] = 'Process name'; sheet['B2'] = 'Example process'
        for col, text in enumerate(HEADERS, 1): sheet.cell(4, col, text)
        for col, value in enumerate(ROW, 1): sheet.cell(5, col, value)
        if mutate: mutate(book, sheet)
        output = BytesIO()
        book.save(output); book.close()
        return output.getvalue()

    def forge(self, mutate):
        output = BytesIO()
        with ZipFile(BytesIO(self.workbook())) as source, ZipFile(output, 'w') as target:
            for info in source.infolist():
                raw = source.read(info)
                if info.filename == 'xl/worksheets/sheet1.xml':
                    node = ET.fromstring(raw); mutate(node); raw = ET.tostring(node)
                target.writestr(info, raw)
        return output.getvalue()

    def archive(self, members):
        output = BytesIO()
        with warnings.catch_warnings():
            warnings.simplefilter('ignore', UserWarning)
            with ZipFile(output, 'w', compression=ZIP_DEFLATED) as archive:
                for name, body in members: archive.writestr(name, body)
        return output.getvalue()

    def place(self, relative, body='evidence'):
        path = self.root / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(body)
        return path

    def test_r001_identity_nontext_is_validation_error(self):
        """Reject nontext process identities before creating paths."""
        for value in (None, 17, True, b'process', ['process']):
            with self.subTest(value=value), self.assertRaises(ValidationError): identity(value)

    def test_r002_identity_grammar_prevents_path_injection(self):
        """Reject leading digits, whitespace, separators and dotted traversal IDs."""
        for value in ('1process', ' process', 'process\n', '../process', 'a/b', 'a.b'):
            with self.subTest(value=value), self.assertRaises(ValidationError): identity(value)

    def test_r003_identity_length_boundary(self):
        """Accept an 80-character ID and reject an 81-character ID."""
        self.assertEqual(identity('p' * 80), 'p' * 80)
        with self.assertRaises(ValidationError): identity('p' * 81)

    def test_r004_identity_windows_device_names(self):
        """Reject case-insensitive Windows device IDs, while allowing ordinary prefixes."""
        for value in ('con', 'PRN', 'Aux', 'NUL', 'com9', 'LPT1'):
            with self.subTest(value=value), self.assertRaises(ValidationError): identity(value)
        self.assertEqual(identity('CONTRACT'), 'CONTRACT')

    def test_r005_json_duplicate_top_level_key(self):
        """Reject conflicting top-level JSON values instead of last-value wins."""
        with self.assertRaises(ValidationError): decode(b'{"approved":false,"approved":true}')

    def test_r006_json_duplicate_nested_key(self):
        """Reject duplicate keys inside arrays and nested evidence objects."""
        with self.assertRaises(ValidationError): decode(b'{"rows":[{"id":"A","id":"B"}]}')

    def test_r007_json_nonfinite_numbers(self):
        """Reject JSON nonfinite constants and finite-looking overflowing exponents."""
        for raw in (b'NaN', b'Infinity', b'-Infinity', b'1e9999'):
            with self.subTest(raw=raw), self.assertRaises(ValidationError): decode(raw)

    def test_r008_json_invalid_unicode_bytes(self):
        """Malformed UTF bytes must become a named validation error."""
        with self.assertRaises(ValidationError): decode(b'{"x":"\xff"}')

    def test_r009_json_excessive_nesting(self):
        """Bound JSON depth while retaining supported 192-comparison program and analysis evidence."""
        depth = MAX_DOCUMENT_DEPTH + 1
        with self.assertRaises(ValidationError): decode(b'[' * depth + b'0' + b']' * depth)
        from test_source import COBOL
        from workbench.source import analyze_program, analyze_sources
        text = COBOL.replace('IF AGE >= 18', 'IF ' + ' OR '.join(['AGE >= 18'] * 192))
        program = analyze_program('ELIGIBLE.cbl', text, {})
        self.assertFalse(program['blockers'])
        self.assertEqual(encode(decode(encode(program))), encode(program))
        manifest = self.rows()
        manifest['jobs'][0]['steps'][0]['program'] = 'ELIGIBLE'
        analysis = analyze_sources({'ELIGIBLE.cbl': text}, manifest)
        self.assertFalse(analysis['blockers'])
        document = {'process': {'analysis': analysis}}
        self.assertEqual(encode(decode(encode(document))), encode(document))

    def test_r010_json_text_limit_counts_encoded_bytes(self):
        """A multibyte JSON string cannot evade the byte upload limit."""
        with self.assertRaises(ValidationError): decode('"\u00e9\u00e9"', limit=5)

    def test_r011_json_non_document_input(self):
        """Non-document objects must produce ValidationError rather than len/type failures."""
        for value in (None, 7, {}, ['{}']):
            with self.subTest(value=value), self.assertRaises(ValidationError): decode(value)

    def test_r012_json_lone_surrogate(self):
        """Reject escaped lone surrogates that cannot be persisted as UTF-8 evidence."""
        with self.assertRaises(ValidationError): decode(b'{"name":"\\ud800"}')

    def test_r013_atomic_nonfinite_value_preserves_original(self):
        """A nonfinite replacement cannot corrupt previously persisted state."""
        path = self.place('state.json', '{"ok":true}\n')
        before = path.read_bytes()
        with self.assertRaises(ValueError): atomic_json(path, {'value': float('nan')})
        self.assertEqual(path.read_bytes(), before)
        self.assertEqual([p.name for p in self.root.iterdir()], ['state.json'])

    def test_r014_zip_nonbinary_input(self):
        """Reject nonbinary archive objects with a named input error."""
        for value in (None, {}, 'not zip'):
            with self.subTest(value=value), self.assertRaises(ValidationError): checked_zip(value)

    def test_r015_json_exact_byte_limit(self):
        """Reject one byte above the limit while preserving documents exactly at it."""
        self.assertEqual(decode(b'{"x":1}', limit=7), {'x': 1})
        with self.assertRaises(ValidationError): decode(b'{"x":1} ', limit=7)

    def test_r016_json_canonical_evidence_roundtrip(self):
        """Canonical encoding preserves Unicode and nested values independent of key insertion order."""
        first = {'b': ['\u00e9', None, False], 'a': {'x': 3}}
        second = {'a': {'x': 3}, 'b': ['\u00e9', None, False]}
        self.assertEqual(encode(first), encode(second))
        self.assertEqual(decode(encode(first)), first)

    def test_r017_path_parent_traversal(self):
        """Reject traversal even when the final path would reenter the process tree."""
        for relative in ('../outside', 'analysis/../../other/file', 'analysis/../target/file'):
            with self.subTest(relative=relative), self.assertRaises(ValidationError): safe_path(self.root, relative)

    def test_r018_path_absolute_and_drive_paths(self):
        """Reject POSIX absolute paths and Windows drive/alternate-stream paths."""
        for relative in ('/tmp/file', 'C:/file', 'analysis/file:stream'):
            with self.subTest(relative=relative), self.assertRaises(ValidationError): safe_path(self.root, relative)

    def test_r019_path_backslash_and_unc(self):
        """Reject alternate path separators and UNC paths on POSIX hosts too."""
        for relative in ('analysis\\file', '\\\\server\\share\\file'):
            with self.subTest(relative=relative), self.assertRaises(ValidationError): safe_path(self.root, relative)

    def test_r020_path_control_characters(self):
        """Reject NUL/newline/tab path characters before filesystem calls."""
        for relative in ('analysis/a\x00b', 'analysis/a\nb', 'analysis/a\tb'):
            with self.subTest(relative=relative), self.assertRaises(ValidationError): safe_path(self.root, relative)

    def test_r021_path_device_basename_with_extension(self):
        """Reject device basenames in nested paths even with normal-looking extensions."""
        for relative in ('sources/con.cbl', 'sources/Com1.txt', 'NUL/file.json'):
            with self.subTest(relative=relative), self.assertRaises(ValidationError): safe_path(self.root, relative)

    def test_r022_path_ambiguous_suffixes(self):
        """Reject trailing dots and spaces that Windows normalizes to another filename."""
        for relative in ('source.cbl.', 'source.cbl ', 'folder./source.cbl'):
            with self.subTest(relative=relative), self.assertRaises(ValidationError): safe_path(self.root, relative)

    def test_r023_path_symlink_leaf(self):
        """Reject existing and dangling symlinks at the artifact leaf."""
        leaf = self.root / 'link'; leaf.symlink_to(self.root / 'missing')
        with self.assertRaises(ValidationError): safe_path(self.root, 'link')

    def test_r024_path_symlink_parent(self):
        """Reject a symlink inside an otherwise safe process-relative path."""
        real = self.root / 'real'; real.mkdir()
        (self.root / 'link').symlink_to(real, target_is_directory=True)
        with self.assertRaises(ValidationError): safe_path(self.root, 'link/new.json')

    def test_r025_path_symlink_root(self):
        """Reject a symlink root before resolve can hide its identity."""
        real = self.root / 'real'; real.mkdir()
        alias = self.root / 'alias'; alias.symlink_to(real, target_is_directory=True)
        with self.assertRaises(ValidationError): safe_path(alias, 'new.json')

    def test_r026_path_symlink_above_root(self):
        """Reject symlink ancestry above a nonexistent workspace root."""
        real = self.root / 'real'; real.mkdir()
        alias = self.root / 'alias'; alias.symlink_to(real, target_is_directory=True)
        with self.assertRaises(ValidationError): safe_path(alias / 'future', 'new.json')

    def test_r027_path_root_alias_is_not_relative_file(self):
        """Reject empty or dot-only paths instead of treating the root itself as an artifact."""
        for relative in ('', '.', './'):
            with self.subTest(relative=relative), self.assertRaises(ValidationError): safe_path(self.root, relative)

    def test_r028_immutable_write_rejects_overwrite(self):
        """A second evidence write must preserve the original bytes and hash."""
        path = self.root / 'evidence.json'
        first = write_new(path, b'original')
        with self.assertRaises(ValidationError): write_new(path, b'replacement')
        self.assertEqual(path.read_bytes(), b'original'); self.assertEqual(first, sha(b'original'))

    def test_r029_immutable_write_rejects_dangling_link(self):
        """A dangling evidence symlink cannot create or overwrite its target."""
        path = self.root / 'evidence'; target = self.root / 'missing'
        path.symlink_to(target)
        with self.assertRaises(ValidationError): write_new(path, b'new')
        self.assertFalse(target.exists())

    def test_r030_immutable_write_rejects_parent_link(self):
        """Immutable output creation must refuse symlink parents."""
        real = self.root / 'real'; real.mkdir()
        alias = self.root / 'alias'; alias.symlink_to(real, target_is_directory=True)
        with self.assertRaises(ValidationError): write_new(alias / 'file', b'new')
        self.assertFalse((real / 'file').exists())

    def test_r031_invalid_evidence_data_leaves_no_partial_file(self):
        """Invalid immutable-write payloads must fail before reserving an empty evidence filename."""
        path = self.root / 'evidence.json'
        with self.assertRaises(ValidationError): write_new(path, 'not bytes')
        self.assertFalse(path.exists())

    def test_r032_atomic_replace_failure_cleans_temp(self):
        """Failed atomic replacement must preserve old state and remove its temporary artifact."""
        path = self.place('state.json', 'old')
        with patch.object(Path, 'replace', side_effect=OSError('synthetic rename failure')):
            with self.assertRaises(OSError): atomic_json(path, {'new': True})
        self.assertEqual(path.read_text(), 'old')
        self.assertEqual([p.name for p in self.root.iterdir()], ['state.json'])

    def test_r033_layout_unknown_root_file(self):
        """Root-level generated results must be reported as misplaced."""
        self.place('results.json')
        self.assertTrue(any('root file' in x for x in validate_workspace(self.root)))

    def test_r034_layout_tmp_prefix_not_exemption(self):
        """A tmp-prefixed root directory cannot hide process output."""
        self.place('tmp-results/result.json')
        self.assertTrue(any('root directory' in x for x in validate_workspace(self.root)))

    def test_r035_layout_workspace_is_regular_directory(self):
        """A regular file cannot masquerade as the workspace."""
        path = self.place('workspace')
        self.assertEqual(validate_workspace(path), ['Workspace must be a directory'])

    def test_r036_layout_invalid_process_identity(self):
        """Malformed process directory identities are visible in layout diagnostics."""
        (self.root / 'processes/1bad').mkdir(parents=True)
        self.assertTrue(any('invalid process ID' in x for x in validate_workspace(self.root)))

    def test_r037_layout_process_must_be_directory(self):
        """An allowlisted process ID stored as a file is rejected."""
        self.place('processes/example-process')
        self.assertTrue(any('process must be a directory' in x for x in validate_workspace(self.root)))

    def test_r038_layout_category_must_be_directory(self):
        """A file named analysis cannot replace the process analysis category."""
        self.place('processes/example-process/analysis')
        self.assertTrue(any('misplaced process output' in x for x in validate_workspace(self.root)))

    def test_r039_layout_input_has_designated_filenames(self):
        """Input allows frozen sources and designated returns, not arbitrary generated evidence."""
        for relative in ('input/results.json', 'input/sources', 'input/process-input.md/extra'):
            with self.subTest(relative=relative), self.assertRaises(ValidationError): output_path(self.root, 'example-process', relative)
        self.assertEqual(output_path(self.root, 'example-process', 'input/sources/original.md').name, 'original.md')

    def test_r040_layout_generated_code_and_database_placement(self):
        """Executable and database artifacts cannot be placed in analysis/review/reports."""
        for relative in ('analysis/jobs.py', 'review/data.sqlite', 'reports/data.db'):
            with self.subTest(relative=relative), self.assertRaises(ValidationError): output_path(self.root, 'example-process', relative)

    def test_r041_layout_no_markdown_per_rule(self):
        """Generated per-rule Markdown is refused while the frozen manifest stays legal."""
        with self.assertRaises(ValidationError): output_path(self.root, 'example-process', 'analysis/rule.md')
        self.assertEqual(output_path(self.root, 'example-process', 'input/process-input.md').name, 'process-input.md')

    def test_r042_layout_shared_version_filename(self):
        """Shared executable versions require exactly a SHA-256 lowercase filename."""
        bad = self.place('shared/target/python/jobs.py')
        self.assertTrue(validate_workspace(self.root))
        bad.rename(bad.with_name('a' * 64 + '.py'))
        self.assertEqual(validate_workspace(self.root), [])

    def test_r043_layout_knowledge_entry_types(self):
        """The canonical knowledge JSON/index entries must be files and inbox must be a directory."""
        (self.root / 'knowledge/records.json').mkdir(parents=True)
        self.place('knowledge/inbox')
        issues = validate_workspace(self.root)
        self.assertTrue(any('records.json' in x for x in issues))
        self.assertTrue(any('inbox' in x for x in issues))

    def test_r044_layout_nested_symlink_is_reported_without_read(self):
        """Nested source symlinks are diagnosed without opening their targets."""
        folder = self.root / 'Endeavor'; folder.mkdir()
        (folder / 'source.cbl').symlink_to(self.root / 'unreadable-secret')
        with patch.object(Path, 'read_bytes', side_effect=AssertionError('must not read')):
            self.assertTrue(any('symlinks' in x for x in validate_workspace(self.root)))

    def test_r045_intake_row_collection_shape(self):
        """A nonsequence row collection must fail through the intake validation boundary."""
        for value in (None, 1, {'row': ROW}):
            with self.subTest(value=value), self.assertRaises(ValidationError): from_rows('process', 'Name', value)

    def test_r046_intake_individual_row_shape(self):
        """A null/scalar/mapping row must not leak len/index errors or be stringified."""
        for value in (None, 1, {i: x for i, x in enumerate(ROW)}):
            with self.subTest(value=value), self.assertRaises(ValidationError): from_rows('process', 'Name', [value])

    def test_r047_intake_row_count_boundary(self):
        """Intake accepts 200 explicit steps and rejects an additional step."""
        rows = [[1, 'JOB', n, 'S' + str(n), 'PROGRAM', '', '', 'Always'] for n in range(1, 201)]
        self.assertEqual(len(self.rows(*rows)['jobs'][0]['steps']), 200)
        with self.assertRaises(ValidationError): self.rows(*(rows + [[1, 'JOB', 201, 'S201', 'PROGRAM', '', '', 'Always']]))

    def test_r048_intake_order_type(self):
        """Boolean and fractional order cells cannot coerce to integer positions."""
        for value in (True, 1.2, None, '1e2'):
            row = ROW.copy(); row[0] = value
            with self.subTest(value=value), self.assertRaises(ValidationError): self.rows(row)

    def test_r049_intake_order_range(self):
        """Both job and step order remain in the explicit 1-to-1000 range."""
        for column, value in ((0, 0), (0, 1001), (2, -1), (2, 1001)):
            row = ROW.copy(); row[column] = value
            with self.subTest(column=column, value=value), self.assertRaises(ValidationError): self.rows(row)

    def test_r050_intake_duplicate_step_order(self):
        """Two different step names cannot occupy one position in a job."""
        other = ROW.copy(); other[3] = 'STEPTWO'
        with self.assertRaises(ValidationError): self.rows(ROW, other)

    def test_r051_intake_duplicate_step_name(self):
        """Case-only step-name differences cannot disguise a duplicate step."""
        other = ROW.copy(); other[2] = 2; other[3] = 'stepone'
        with self.assertRaises(ValidationError): self.rows(ROW, other)

    def test_r052_intake_job_order_conflicts(self):
        """Job order cannot be shared by two jobs or changed midway through one job."""
        other = ROW.copy(); other[1] = 'JOBTWO'
        with self.assertRaises(ValidationError): self.rows(ROW, other)
        other = ROW.copy(); other[0] = 2; other[2] = 2; other[3] = 'STEPTWO'
        with self.assertRaises(ValidationError): self.rows(ROW, other)

    def test_r053_intake_generated_job_method_collision(self):
        """Case and hyphen normalization must not collapse distinct jobs into one method."""
        first = ROW.copy(); first[1] = 'JOB-A'
        other = ROW.copy(); other[0] = 2; other[1] = 'JOB_A'
        with self.assertRaises(ValidationError): self.rows(first, other)

    def test_r054_intake_optional_cells_are_text(self):
        """Input/output/condition objects cannot silently turn into fabricated text facts."""
        for column, value in ((5, {'dsn': 'INPUT'}), (6, 12), (7, False)):
            row = ROW.copy(); row[column] = value
            with self.subTest(column=column), self.assertRaises(ValidationError): self.rows(row)

    def test_r055_intake_unicode_line_separators(self):
        """Unicode line separators in a cell cannot inject rows during Markdown normalization."""
        for separator in ('\u0085', '\u2028', '\u2029', '\x0b'):
            row = ROW.copy(); row[5] = 'INPUT' + separator + 'HIDDEN'
            with self.subTest(separator=repr(separator)), self.assertRaises(ValidationError): self.rows(row)

    def test_r056_intake_process_name_line_separators(self):
        """Process names cannot contain Unicode line boundaries that split the frozen manifest."""
        with self.assertRaises(ValidationError): self.rows(name='Name\u2028Process ID: injected')

    def test_r057_intake_blank_condition_remains_unknown(self):
        """Whitespace-only conditions must retain an explicit Unknown value."""
        row = ROW.copy(); row[7] = '   '
        self.assertEqual(self.rows(row)['jobs'][0]['steps'][0]['condition'], 'Unknown')

    def test_r058_intake_file_lists_preserve_all_names(self):
        """Semicolon/comma-separated source and target lists preserve each nonblank item."""
        row = ROW.copy(); row[5] = ' A; B, C ;; '; row[6] = ' X,Y; Z'
        step = self.rows(row)['jobs'][0]['steps'][0]
        self.assertEqual(step['inputs'], ['A', 'B', 'C']); self.assertEqual(step['outputs'], ['X', 'Y', 'Z'])

    def test_r059_intake_order_not_physical_row_order(self):
        """Out-of-order intake rows must produce stable explicit job/step order."""
        rows = [[2, 'JOBB', 1, 'S1', 'PROGRAM', '', '', 'Always'],
                [1, 'JOBA', 2, 'S2', 'PROGRAM', '', '', 'Always'],
                [1, 'JOBA', 1, 'S1', 'PROGRAM', '', '', 'RC <= 4']]
        jobs = self.rows(*rows)['jobs']
        self.assertEqual([j['name'] for j in jobs], ['JOBA', 'JOBB'])
        self.assertEqual([s['name'] for s in jobs[0]['steps']], ['S1', 'S2'])
        self.assertEqual(jobs[0]['steps'][0]['condition'], 'RC <= 4')

    def test_r060_manifest_duplicate_attributes(self):
        """A second process name/ID declaration must be treated as ambiguous."""
        for extra in ('- Process ID: other', '- Process name: Other'):
            with self.subTest(extra=extra), self.assertRaises(ValidationError): parse_manifest(self.manifest() + extra + '\n')

    def test_r061_manifest_blank_name_cannot_capture_next_line(self):
        """A blank process name must not consume the table header as its name."""
        with self.assertRaises(ValidationError): parse_manifest(self.manifest(name=''))

    def test_r062_manifest_table_header_contract(self):
        """Missing, reordered, or repeated headers cannot be interpreted as valid intake."""
        header = '| ' + ' | '.join(HEADERS) + ' |'
        variants = [self.manifest().replace(header + '\n', ''),
                    self.manifest().replace('Job order | Job', 'Job | Job order'),
                    self.manifest() + header + '\n']
        for text in variants:
            with self.subTest(text=text), self.assertRaises(ValidationError): parse_manifest(text)

    def test_r063_manifest_no_steps(self):
        """A header-only process must be rejected instead of creating an empty workflow."""
        with self.assertRaises(ValidationError): parse_manifest('\n'.join(self.manifest().splitlines()[:-1]))

    def test_r064_manifest_row_width(self):
        """Missing or extra columns cannot silently drop or shift intake facts."""
        for row in (ROW[:-1], ROW + ['extra']):
            text = '\n'.join(self.manifest().splitlines()[:-1]) + '\n| ' + ' | '.join(map(str, row)) + ' |\n'
            with self.subTest(width=len(row)), self.assertRaises(ValidationError): parse_manifest(text)

    def test_r065_zip_duplicate_member_identity(self):
        """Duplicate ZIP members cannot make different readers see different workbook evidence."""
        raw = self.archive([('sheet.xml', b'<first/>'), ('sheet.xml', b'<second/>')])
        with self.assertRaises(ValidationError): checked_zip(raw)

    def test_r066_zip_case_insensitive_xml_guard(self):
        """Uppercase XML suffixes cannot evade declaration/entity inspection."""
        raw = self.archive([('sheet.XML', b'<!DOCTYPE root><root/>')])
        with self.assertRaises(ValidationError): checked_zip(raw)

    def test_r067_zip_relationship_xml_guard(self):
        """OOXML relationship parts receive the same declaration/entity guard as XML parts."""
        raw = self.archive([('_rels/.rels', b'<!DOCTYPE root [<!ENTITY x "bad">]><root/>')])
        with self.assertRaises(ValidationError): checked_zip(raw)

    def test_r068_zip_member_path_traversal(self):
        """Archive members cannot contain traversal, absolute paths, or Windows separators."""
        for name in ('../file.xml', '/file.xml', 'folder\\file.xml'):
            with self.subTest(name=name), self.assertRaises(ValidationError): checked_zip(self.archive([(name, b'<x/>')]))

    def test_r069_zip_member_count_limit(self):
        """An archive with 301 members is rejected before materializing workbook data."""
        raw = self.archive([(str(n) + '.xml', b'<x/>') for n in range(301)])
        with self.assertRaises(ValidationError): checked_zip(raw)

    def test_r070_zip_total_expansion_limit(self):
        """Total declared decompressed bytes are bounded even for tiny compressed input."""
        raw = self.archive([(str(n) + '.bin', b'0' * (7 * 1024 * 1024)) for n in range(5)])
        with self.assertRaises(ValidationError): checked_zip(raw)

    def test_r071_zip_individual_expansion_limit(self):
        """An individual archive member cannot exceed the eight-MiB member bound."""
        raw = self.archive([('one.bin', b'0' * (8 * 1024 * 1024 + 1))])
        with self.assertRaises(ValidationError): checked_zip(raw)

    def test_r072_zip_unsupported_compression_is_validation_error(self):
        """Unsupported compression or corrupt compressed bytes must fail through the ZIP validation boundary."""
        raw = bytearray(self.archive([('sheet.xml', b'<x/>')]))
        local = raw.index(b'PK\x03\x04'); central = raw.index(b'PK\x01\x02')
        struct.pack_into('<H', raw, local + 8, 99); struct.pack_into('<H', raw, central + 10, 99)
        with self.assertRaises(ValidationError): checked_zip(bytes(raw))
        raw = bytearray(self.archive([('sheet.xml', b'<root/>' * 50)]))
        name_length, extra_length = struct.unpack_from('<HH', raw, 26)
        raw[30 + name_length + extra_length] = 255
        with self.assertRaises(ValidationError): checked_zip(bytes(raw))

    def test_r073_xlsx_formula_intake_rejected(self):
        """An Excel formula cannot stand in for a process input fact."""
        raw = self.workbook(lambda book, sheet: setattr(sheet['F5'], 'value', '=1+1'))
        with self.assertRaises(ValidationError): parse_intake_xlsx(raw)

    def test_r074_xlsx_process_name_requires_text(self):
        """Numeric workbook metadata must not be silently stringified into a process name."""
        raw = self.workbook(lambda book, sheet: setattr(sheet['B2'], 'value', 42))
        with self.assertRaises(ValidationError): parse_intake_xlsx(raw)

    def test_r075_xlsx_named_sheet_required(self):
        """An arbitrary first sheet cannot be substituted for the designated Intake sheet."""
        raw = self.workbook(lambda book, sheet: setattr(sheet, 'title', 'Other'))
        with self.assertRaises(ValidationError): parse_intake_xlsx(raw)

    def test_r076_xlsx_forged_dimensions(self):
        """A declared short dimension cannot hide actual populated intake cells."""
        raw = self.forge(lambda node: node.find(NS + 'dimension').set('ref', 'A1:H4'))
        with self.assertRaises(ValidationError): parse_intake_xlsx(raw)

    def test_r077_xlsx_duplicate_cell_coordinates(self):
        """Repeated worksheet cell coordinates are rejected before reader overwrite behavior."""
        def mutate(node):
            row = node.find(NS + 'sheetData')[-1]; row.append(copy.deepcopy(row[0]))
        with self.assertRaises(ValidationError): parse_intake_xlsx(self.forge(mutate))

    def test_r078_xlsx_cell_row_identity_mismatch(self):
        """A cell cannot claim another row than its containing worksheet row."""
        def mutate(node): node.find(NS + 'sheetData')[-1][0].set('r', 'A6')
        with self.assertRaises(ValidationError): parse_intake_xlsx(self.forge(mutate))

    def test_r079_xlsx_hidden_job_is_accounted(self):
        """Hidden populated job rows remain included instead of disappearing from scope."""
        def mutate(book, sheet):
            sheet.append([2, 'HIDDENJOB', 1, 'S1', 'OTHER', 'IN', 'OUT', 'Always'])
            sheet.row_dimensions[6].hidden = True
        jobs = parse_intake_xlsx(self.workbook(mutate))['jobs']
        self.assertEqual([job['name'] for job in jobs], ['JOBONE', 'HIDDENJOB'])

    def test_r080_xlsx_actual_row_bound(self):
        """An actual row outside 205 is rejected even when the dimension is forged short."""
        def mutate(node):
            row = node.find(NS + 'sheetData')[-1]; row.set('r', '206')
            for cell in row: cell.set('r', cell.get('r').rstrip('0123456789') + '206')
        with self.assertRaises(ValidationError): parse_intake_xlsx(self.forge(mutate))


if __name__ == '__main__': unittest.main()
