"""R401–R480: executable source-accountability/report/runner/job reviews.

The only filled SME workbook is an explicitly fictional test fixture. Tests use
one real completed fixture, cloned per case, so tampering cannot leak between
reviews or modify historical repository examples.
"""
import copy
import csv
from io import BytesIO, StringIO
import json
from pathlib import Path
import shlex
import shutil
import tempfile
import threading
import unittest
from unittest.mock import Mock, patch
import zipfile

from openpyxl import load_workbook
from pptx import Presentation
from workbench.coordinator import Coordinator
from workbench.coverage import COLUMNS, DISPOSITIONS, _cell, _read, build_coverage, write_coverage
from workbench.domain import ValidationError, encode, sha
from workbench.orchestration import verify_jobs
from workbench.reports import generate_reports, metrics, portfolio, report_portfolio
from workbench import runner
from workbench.target import emit_jobs, emit_program
from test_source import COBOL
from test_workflow import MANIFEST


class FixtureLedger:
    def __init__(self, root, doc, documents=None):
        self.root = root
        self.doc = doc
        self.documents = documents if documents is not None else [doc]
        self.lock = threading.RLock()
    def get(self, pid):
        return self.doc
    def save(self, doc):
        self.doc = doc
    def list(self, include_demo=False):
        return [p for p in self.documents if include_demo or not p.get('demo')]
    def history(self):
        return []
    def measurements(self):return []
    def measurement_timing(self):return [],[]
    def work_sessions(self):return []


class Review500ReportingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.seed_tmp = tempfile.TemporaryDirectory()
        cls.seed = Path(cls.seed_tmp.name)
        c = Coordinator(cls.seed)
        try:
            c.create(MANIFEST, {'ELIGIBLE.cbl': COBOL})
            c.start('process-a'); c.advance('process-a')
            book = load_workbook(BytesIO(c.artifact('process-a', 'review/sme-checklist.xlsx').read_bytes()))
            for row in range(2, book['Checklist'].max_row + 1):
                book['Checklist'].cell(row, 5, 'Yes')
                book['Checklist'].cell(row, 7, 'Fictional regression reviewer')
            out = BytesIO(); book.save(out); book.close()
            c.import_answers('process-a', out.getvalue(), 'Fictional regression reviewer')
            c.advance('process-a'); c.advance('process-a')
            cls.seed_doc = c.ledger.get('process-a')
            assert cls.seed_doc['status'] == 'COMPLETED', cls.seed_doc['blockers']
        finally:
            c.close()

    @classmethod
    def tearDownClass(cls):
        cls.seed_tmp.cleanup()

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name) / 'workspace'
        shutil.copytree(self.seed, self.root)
        self.doc = copy.deepcopy(self.seed_doc)
        self.base = self.root / 'processes/process-a'
        self.ledger = FixtureLedger(self.root, self.doc)

    def coverage(self):
        return build_coverage(self.doc, self.root)

    def target_path(self):
        return self.root / 'shared/target/python' / (self.doc['program_versions']['ELIGIBLE'] + '.py')

    def result_path(self, filename):
        return self.base / 'target' / self.doc['runs'][-1]['id'] / filename

    def synthetic_path(self, filename):
        return self.base / 'synthetic' / self.doc['runs'][-1]['id'] / 'ELIGIBLE' / filename

    def coordinator(self):
        c = Mock()
        c.root = self.root; c.lock = threading.RLock(); c.ledger = self.ledger
        c.economics.return_value = {'receipt_count':0,'learning_count':0,'usage':{'credits':{'accounts':[]},'budgets':[]},'processes':[]}
        c.process_root.side_effect = lambda pid: self.base
        def artifact(pid, name):
            path = self.base / name
            self.assertEqual(sha(path.read_bytes()), self.doc['artifact_hashes'][name])
            return path
        c.artifact.side_effect = artifact
        def register(doc, name):
            if name not in doc['artifacts']: doc['artifacts'].append(name)
            doc['artifact_hashes'][name] = sha((self.base / name).read_bytes())
        c.register.side_effect = register
        return c

    def exported(self, text):
        model = self.coverage()
        model['rows'][0]['source_text'] = text
        folder = self.root / '.implementation/tmp/export'
        write_coverage(model, folder)
        return model, folder

    def jobs(self):
        return verify_jobs(self.doc, self.root, emit_jobs(self.doc, self.doc['program_versions']))

    def test_r401_physical_source_line_order_is_exact(self):
        rows = self.coverage()['rows']
        self.assertEqual([(r['source_line'], r['source_text']) for r in rows], list(enumerate(COBOL.splitlines(), 1)))

    def test_r402_each_source_line_hash_covers_unmodified_text(self):
        self.assertTrue(all(r['source_line_hash'] == sha(r['source_text']) for r in self.coverage()['rows']))

    def test_r403_all_source_rows_have_disposition_and_reason(self):
        self.assertTrue(all(r['disposition'] in DISPOSITIONS and r['reason'].strip() for r in self.coverage()['rows']))

    def test_r404_rule_span_has_one_semantic_unit(self):
        model = self.coverage(); rule = self.doc['analysis']['programs']['ELIGIBLE']['rules'][0]
        rows = [r for r in model['rows'] if rule['source_start'] <= r['source_line'] <= rule['source_end']]
        self.assertEqual(len({r['unit_id'] for r in rows}), 1)
        self.assertTrue(all((r['source_start'], r['source_end']) == (rule['source_start'], rule['source_end']) for r in rows))

    def test_r405_target_spans_are_bounded_by_actual_file(self):
        for row in self.coverage()['rows']:
            for mapping in row['target_mappings']:
                path = self.root / mapping['file']
                self.assertEqual(sha(path.read_bytes()), mapping['version'])
                self.assertTrue(1 <= mapping['start'] <= mapping['end'] <= len(path.read_text().splitlines()))

    def test_r406_terminal_replacement_names_bounded_adapter(self):
        row = self.coverage()['rows'][-1]
        self.assertEqual(row['disposition'], 'platform_replaced_verified')
        self.assertIn('return', row['replacement']); self.assertTrue(row['tests'] and row['evidence'])

    def test_r407_applicable_denominator_excludes_nonexecutables(self):
        s = self.coverage()['summary']
        self.assertEqual(s['applicable_lines'] + s['non_executable_lines'] + s['out_of_scope_lines'], s['source_lines'])
        self.assertLess(s['semantic_units']['applicable'], s['applicable_lines'])

    def test_r408_missing_manifest_revokes_all_verified_credit(self):
        (self.base / 'input/process-input.md').unlink()
        s = self.coverage()['summary']
        self.assertFalse(s['completion_eligible']); self.assertEqual(s['verified_applicable_lines'], 0)

    def test_r409_source_tamper_preserves_frozen_original_rows(self):
        (self.base / 'input/sources/ELIGIBLE.cbl').write_text('ALTERED\n')
        m = self.coverage()
        self.assertEqual([r['source_text'] for r in m['rows']], COBOL.splitlines())
        self.assertFalse(m['summary']['fully_accounted'])

    def test_r410_missing_source_preserves_frozen_original_rows(self):
        (self.base / 'input/sources/ELIGIBLE.cbl').unlink()
        self.assertEqual([r['source_text'] for r in self.coverage()['rows']], COBOL.splitlines())

    def test_r411_analysis_tamper_blocks_replay(self):
        (self.base / 'analysis/source-analysis.json').write_bytes(b'{}')
        with patch('workbench.fixtures.verify_program') as replay:
            s = self.coverage()['summary']
        replay.assert_not_called(); self.assertEqual(s['verified_applicable_lines'], 0)

    def test_r412_expected_evidence_tamper_revokes_program(self):
        self.synthetic_path('expected.json').write_bytes(b'{}')
        s = self.coverage()['summary']
        self.assertFalse(s['completion_eligible']); self.assertTrue(any('expected' in e.lower() for e in s['integrity_errors']))

    def test_r413_actual_evidence_tamper_revokes_program(self):
        self.synthetic_path('actual-and-comparison.json').write_bytes(b'{}')
        self.assertEqual(self.coverage()['summary']['verified_applicable_lines'], 0)

    def test_r414_missing_target_copy_revokes_program(self):
        self.result_path('ELIGIBLE.py').unlink()
        self.assertFalse(self.coverage()['summary']['completion_eligible'])

    def test_r415_database_hash_tamper_blocks_target_replay(self):
        with self.result_path('target.sqlite').open('ab') as out: out.write(b'forged trailer')
        with patch('workbench.fixtures.verify_program') as replay:
            self.assertEqual(self.coverage()['summary']['verified_applicable_lines'], 0)
        replay.assert_not_called()

    def test_r416_cancelled_evidence_is_not_reexecuted(self):
        self.doc['cancel_requested'] = True
        with patch('workbench.fixtures.verify_program') as replay, patch('workbench.orchestration.verify_jobs') as jobs:
            self.assertFalse(self.coverage()['summary']['completion_eligible'])
        replay.assert_not_called(); jobs.assert_not_called()

    def test_r417_correction_prevents_credit_despite_yes(self):
        rid = self.doc['analysis']['programs']['ELIGIBLE']['rules'][0]['id']
        self.doc['answers']['items'][rid]['correction'] = 'Fictional unresolved correction'
        self.assertEqual(self.coverage()['summary']['verified_applicable_lines'], 0)

    def test_r418_missing_job_comparison_blocks_completion_without_jcl(self):
        self.result_path('job-comparison.json').unlink()
        self.assertFalse(self.coverage()['summary']['completion_eligible'])

    def test_r419_forged_job_comparison_blocks_completion_without_jcl(self):
        self.result_path('job-comparison.json').write_bytes(encode({'matched': True}))
        self.assertFalse(self.coverage()['summary']['completion_eligible'])

    def test_r420_job_order_unconfirmed_blocks_completion_without_jcl(self):
        self.doc['answers']['items']['G_ORDER']['answer'] = 'Not sure'
        self.assertFalse(self.coverage()['summary']['completion_eligible'])

    def test_r421_source_symlink_is_rejected_even_with_matching_bytes(self):
        path = self.base / 'input/sources/ELIGIBLE.cbl'; raw = path.read_bytes()
        other = self.root / '.implementation/tmp/matching'; other.parent.mkdir(parents=True, exist_ok=True); other.write_bytes(raw)
        path.unlink(); path.symlink_to(other)
        self.assertFalse(self.coverage()['summary']['fully_accounted'])

    def test_r422_formula_escape_does_not_truncate_excel_limit(self):
        original = '=' + 'x' * 32766
        _, folder = self.exported(original)
        book = load_workbook(folder / 'coverage.xlsx')
        try:
            value = book['Source lines'].cell(2, COLUMNS.index('source_text') + 1).value
            if value.startswith('[Full text'):
                value = ''.join(r[4] for r in book['Text chunks'].iter_rows(min_row=2, values_only=True) if r[0] == 'Source lines' and r[2] == 'source_text')
            elif value.startswith("'"): value = value[1:]
            self.assertEqual(value, original)
        finally: book.close()

    def test_r423_xml_control_text_chunks_are_reversible(self):
        original = 'first\x00last'
        _, folder = self.exported(original)
        book = load_workbook(folder / 'coverage.xlsx')
        try:
            chunks = [r for r in book['Text chunks'].iter_rows(min_row=2, values_only=True) if r[0] == 'Source lines' and r[2] == 'source_text']
            self.assertEqual(json.loads(''.join(r[4] for r in chunks)), original)
            self.assertTrue(all(r[5] == sha(original) and r[6] == 'json_string' for r in chunks))
        finally: book.close()

    def test_r424_html_escapes_original_markup(self):
        _, folder = self.exported('<script>alert("fixture")</script>')
        text = (folder / 'coverage.html').read_text()
        self.assertNotIn('<script>', text); self.assertIn('&lt;script&gt;', text)

    def test_r425_csv_formula_safety_and_exact_canonical_json(self):
        original = '\t=HYPERLINK("fixture")'
        model, folder = self.exported(original)
        row = next(csv.DictReader(StringIO((folder / 'coverage.csv').read_text())))
        self.assertEqual(row['source_text'], "'" + original)
        self.assertEqual(json.loads((folder / 'coverage.json').read_text()), model)

    def test_r426_known_rule_metric_uses_extracted_rule_denominator(self):
        m = metrics(self.ledger, self.doc)
        self.assertEqual(m['known_rule_verification_percent'], 100 * m['rules_verified'] / m['rules_documented'])
        self.assertIn('known rules', m['verification_percent_denominator'])

    def test_r427_unknown_interface_metrics_are_none(self):
        m = metrics(self.ledger, self.doc)
        for key in ('source_cics_transactions', 'source_vsam_files', 'inbound_interfaces', 'outbound_interfaces'):
            self.assertIsNone(m[key])

    def test_r428_workbench_ui_is_excluded_from_business_replacements(self):
        m = metrics(self.ledger, self.doc)
        self.assertEqual((m['target_react_business_screens'], m['target_business_rest_apis']), (0, 0))

    def test_r429_matching_cases_revoked_with_target_corruption(self):
        self.target_path().write_text('def forged(record): return record\n')
        m = metrics(self.ledger, self.doc)
        self.assertEqual((m['matching_cases'], m['rules_verified']), (0, 0))
        self.assertGreater(m['synthetic_cases'], 0)

    def test_r430_shared_target_loc_is_deduplicated(self):
        baseline = metrics(self.ledger, self.doc)['target_program_code_loc']
        self.doc['program_versions']['ALIAS'] = self.doc['program_versions']['ELIGIBLE']
        self.assertEqual(metrics(self.ledger, self.doc)['target_program_code_loc'], baseline)

    def test_r431_portfolio_demo_memberships_are_excluded(self):
        self.doc['demo'] = True
        p = portfolio(self.ledger)
        self.assertEqual((p['processes'], p['program_memberships'], p['completed_processes']), (0, 0, 0))

    def test_r432_portfolio_shared_versions_keep_two_memberships(self):
        other = copy.deepcopy(self.doc); other['id'] = 'process-b'
        self.ledger.documents.append(other)
        p = portfolio(self.ledger)
        self.assertEqual((p['unique_program_versions'], p['program_memberships']), (1, 2))

    def test_r433_portfolio_changed_hash_is_distinct_version(self):
        other = copy.deepcopy(self.doc); other['id'] = 'process-b'
        other['analysis']['assets'][0]['id'] = sha('different-source')
        self.ledger.documents.append(other)
        self.assertEqual(portfolio(self.ledger)['unique_program_versions'], 2)

    def test_r434_scope_union_selects_version_if_any_membership_selects(self):
        other = copy.deepcopy(self.doc); other['id'] = 'process-b'
        self.doc['analysis']['assets'][0]['selected'] = False
        self.ledger.documents.append(other)
        p = portfolio(self.ledger)
        self.assertEqual((p['unique_program_versions'], p['discovered_program_versions']), (1, 1))

    def test_r435_discovered_out_of_scope_version_is_not_selected(self):
        self.doc['analysis']['assets'][0]['selected'] = False
        p = portfolio(self.ledger)
        self.assertEqual((p['unique_program_versions'], p['discovered_program_versions']), (0, 1))

    def test_r436_report_projection_does_not_mutate_live_status(self):
        self.doc['status'] = 'REPORTING'
        self.assertEqual(report_portfolio(self.ledger, self.doc, 'COMPLETED')['completed_processes'], 1)
        self.assertEqual(self.doc['status'], 'REPORTING')
        self.assertEqual(portfolio(self.ledger)['completed_processes'], 0)

    def test_r437_report_projection_does_not_double_count_completed(self):
        self.assertEqual(report_portfolio(self.ledger, self.doc, 'COMPLETED')['completed_processes'], 1)

    def test_r438_report_projection_removes_revoked_completion(self):
        self.assertEqual(report_portfolio(self.ledger, self.doc, 'COMPLETED_WITH_BLOCKERS')['completed_processes'], 0)

    def test_r439_report_inspection_hashes_all_required_outputs(self):
        folder = self.base / 'reports/report-0001'
        inspection = json.loads((folder / 'inspection.json').read_text())
        self.assertEqual(set(inspection['sha256']), {'economics.json','economics.html','economics.csv','executive-report.html', 'metrics.json', 'metrics.csv', 'metrics.xlsx', 'management.pptx', 'coverage.json', 'coverage.csv', 'coverage.xlsx', 'coverage.html', 'rules.html', 'rules.json', 'rules.csv'})
        self.assertTrue(all(sha((folder / name).read_bytes()) == h for name, h in inspection['sha256'].items()))
        self.assertFalse(inspection['powerpoint_render_checked'])

    def test_r440_management_deck_has_six_editable_bounded_slides(self):
        prs = Presentation(self.base / 'reports/report-0001/management.pptx')
        self.assertEqual(len(prs.slides), 6)
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        diagram=prs.slides[0]
        self.assertEqual(sum(s.has_table for s in diagram.shapes),0)
        self.assertEqual(sum(s.shape_type==MSO_SHAPE_TYPE.AUTO_SHAPE for s in diagram.shapes),6)
        self.assertEqual(sum(s.shape_type==MSO_SHAPE_TYPE.LINE for s in diagram.shapes),5)
        diagram_text='\n'.join(s.text for s in diagram.shapes if s.has_text_frame)
        for label in ('Mainframe + Db2','Claude retrieves','Process folder','Claude develops','Coordinator verifies','Target + report'):
            self.assertIn(label,diagram_text)
        for index,slide in enumerate(prs.slides):
            self.assertEqual(sum(s.has_table for s in slide.shapes),0 if index==0 else 1)
            self.assertTrue(all(s.left >= 0 and s.top >= 0 and s.left + s.width <= prs.slide_width and s.top + s.height <= prs.slide_height for s in slide.shapes))

    def test_r441_management_deck_does_not_render_none_percent(self):
        self.doc['analysis']['rules'] = []
        self.doc['analysis']['programs']['ELIGIBLE']['rules'] = []
        folder = self.root / '.implementation/tmp/no-rules'
        generate_reports(self.ledger, self.doc, folder)
        prs = Presentation(folder / 'management.pptx')
        text = '\n'.join(c.text for s in prs.slides for sh in s.shapes if sh.has_table for row in sh.table.rows for c in row.cells)
        self.assertNotIn('None%', text); self.assertIn('Unknown', text)

    def test_r442_report_output_refuses_existing_frozen_artifacts(self):
        folder = self.base / 'reports/report-0001'; before = {p.name: p.read_bytes() for p in folder.iterdir()}
        with self.assertRaises(ValidationError): generate_reports(self.ledger, self.doc, folder)
        self.assertEqual({p.name: p.read_bytes() for p in folder.iterdir()}, before)

    def test_r443_coverage_output_refuses_existing_artifacts(self):
        folder = self.base / 'reports/report-0001'; before = (folder / 'coverage.json').read_bytes()
        with self.assertRaises(ValidationError): write_coverage(self.coverage(), folder)
        self.assertEqual((folder / 'coverage.json').read_bytes(), before)

    def test_r444_metrics_csv_matches_json_values(self):
        folder = self.base / 'reports/report-0001'
        m = json.loads((folder / 'metrics.json').read_text())['metrics']
        actual = dict(list(csv.reader(StringIO((folder / 'metrics.csv').read_text())))[1:])
        self.assertEqual(actual, {k: 'Unknown' if v is None else str(v) for k, v in m.items()})

    def test_r445_metrics_workbook_matches_json_values(self):
        folder = self.base / 'reports/report-0001'
        m = json.loads((folder / 'metrics.json').read_text())['metrics']
        book = load_workbook(folder / 'metrics.xlsx', data_only=True)
        try: self.assertEqual(dict(book['Metrics'].iter_rows(min_row=2, values_only=True)), {k: 'Unknown' if v is None else v for k, v in m.items()})
        finally: book.close()

    def test_r446_manifest_baseline_is_never_inferred(self):
        self.doc.pop('manifest_hash')
        with self.assertRaisesRegex(ValidationError, 'missing'): runner.manifest_integrity(self.coordinator(), self.doc)
        self.assertNotIn('manifest_hash', self.doc)

    def test_r447_start_requires_exact_manifest_bytes(self):
        with self.assertRaisesRegex(ValidationError, 'different manifest'): runner.manifest_integrity(self.coordinator(), self.doc, MANIFEST.replace('Always', 'ALWAYS').encode())

    def test_r448_continuation_shell_quotes_sensitive_workspace(self):
        c = self.coordinator(); c.root = Path('/tmp/space "quoted" $(touch BAD)')
        command = runner.summary(c, self.doc)['continuation']
        words = shlex.split(command)
        self.assertEqual(words[words.index('--workspace') + 1], str(c.root))
        self.assertIn(shlex.quote(str(c.root)), command)

    def test_r449_summary_separates_packet_from_reports(self):
        result = runner.summary(self.coordinator(), self.doc)
        self.assertTrue(all('/review/' in p for p in result['packet']))
        self.assertTrue(all('/reports/' in p for p in result['reports']))
        self.assertIn('input/sme-return-inbox.xlsx', result['sme_return_inbox'])

    def test_r450_waiting_summary_requires_actual_human_return(self):
        self.doc['status'] = 'WAITING_SME'
        self.assertIn('Never generate SME answers', runner.summary(self.coordinator(), self.doc)['message'])

    def test_r451_wait_timeout_bounds_reject_nonfinite_values(self):
        for timeout in (0, -1, 3600.1, float('nan'), float('inf')):
            with self.subTest(timeout=timeout), self.assertRaises(ValidationError): runner.wait_for_process(self.coordinator(), 'process-a', timeout)

    def test_r452_watch_requires_reviewer_before_worker_launch(self):
        c = self.coordinator()
        with self.assertRaises(ValidationError): runner.wait_for_process(c, 'process-a', watch=True)
        c.launch_worker.assert_not_called()

    def test_r453_wait_stops_at_waiting_sme_without_answers(self):
        self.doc['status'] = 'WAITING_SME'; c = self.coordinator()
        doc, timed = runner.wait_for_process(c, 'process-a')
        self.assertFalse(timed); self.assertEqual(doc['status'], 'WAITING_SME'); c.import_answers.assert_not_called()

    def test_r454_wait_preserves_terminal_states(self):
        for status in runner.TERMINAL | runner.STOPPED:
            self.doc['status'] = status
            doc, timed = runner.wait_for_process(self.coordinator(), 'process-a')
            self.assertEqual(doc['status'], status); self.assertFalse(timed)

    def test_r455_bounded_wait_returns_timeout_without_state_mutation(self):
        self.doc['status'] = 'VERIFYING'; c = self.coordinator()
        with patch('workbench.runner.time.monotonic', side_effect=[0, 2]):
            doc, timed = runner.wait_for_process(c, 'process-a', timeout=1)
        self.assertTrue(timed); self.assertEqual(doc['status'], 'VERIFYING')

    def test_r456_import_idempotency_requires_same_return_and_reviewer(self):
        c = self.coordinator()
        doc = runner.import_return(c, 'process-a', self.base / 'input/sme-return.xlsx', 'Fictional regression reviewer')
        self.assertIs(doc, self.doc); c.import_answers.assert_not_called()

    def test_r457_import_refuses_second_workbook(self):
        path = self.base / 'input/sme-return-inbox.xlsx'; path.write_bytes(b'different')
        with self.assertRaisesRegex(ValidationError, 'already consumed'): runner.import_return(self.coordinator(), 'process-a', path, 'Fictional regression reviewer')

    def test_r458_import_refuses_changed_reviewer(self):
        with self.assertRaisesRegex(ValidationError, 'attribution differs'): runner.import_return(self.coordinator(), 'process-a', self.base / 'input/sme-return.xlsx', 'Other reviewer')

    def test_r459_bundle_is_reused_byte_for_byte(self):
        c = self.coordinator(); first = Path(runner.bundle_process(c, 'process-a')); raw = first.read_bytes()
        second = Path(runner.bundle_process(c, 'process-a'))
        self.assertEqual(first, second); self.assertEqual(raw, second.read_bytes())

    def test_r460_bundle_entries_match_frozen_sources_and_report(self):
        path = runner.bundle_process(self.coordinator(), 'process-a')
        with zipfile.ZipFile(path) as z:
            self.assertEqual(z.read('input/sources/ELIGIBLE.cbl').decode(), COBOL)
            self.assertIn('reports/report-0001/management.pptx', z.namelist())
            self.assertIn('input/sme-return.xlsx', z.namelist())
            self.assertTrue(all(info.date_time == (1980, 1, 1, 0, 0, 0) for info in z.infolist()))

    def test_r461_bundle_refuses_nonterminal_process(self):
        self.doc['status'] = 'WAITING_SME'
        with self.assertRaisesRegex(ValidationError, 'completed or cancelled'): runner.bundle_process(self.coordinator(), 'process-a')

    def test_r462_bundle_requires_inspected_report(self):
        self.doc['report_verified'] = False
        with self.assertRaisesRegex(ValidationError, 'inspected'): runner.bundle_process(self.coordinator(), 'process-a')

    def test_r463_bundle_refuses_missing_registered_inspection_hashes(self):
        self.doc['report_hashes'] = {}
        with self.assertRaisesRegex(ValidationError, 'report.*hash|inspected'): runner.bundle_process(self.coordinator(), 'process-a')

    def test_r464_bundle_refuses_orphaned_report_hash_entry(self):
        self.doc['report_hashes']['reports/report-0001/orphan.csv'] = sha('missing')
        with self.assertRaisesRegex(ValidationError, 'report|inspected'): runner.bundle_process(self.coordinator(), 'process-a')

    def test_r465_bundle_rejects_modified_existing_zip(self):
        c = self.coordinator(); path = Path(runner.bundle_process(c, 'process-a')); path.write_bytes(b'altered')
        with self.assertRaisesRegex(ValidationError, 'Bundle version integrity'): runner.bundle_process(c, 'process-a')

    def test_r466_job_baseline_matches_independent_reference(self):
        result = self.jobs()
        self.assertTrue(result['matched']); self.assertEqual(result['actual'], result['expected'])
        self.assertGreaterEqual(result['integration_cases'],20)
        self.assertGreaterEqual(result['randomized_cases'],20)
        self.assertEqual(result['integration_cases'],len(result['cases']))
        self.assertTrue(all(case['matched'] for case in result['cases']))

    def test_r467_job_result_explicitly_excludes_dataset_and_mainframe_parity(self):
        result = self.jobs()
        self.assertFalse(result['dataset_io_verified']); self.assertFalse(result['observed_mainframe_parity'])
        self.assertEqual(result['io_adapter'], 'in_memory_record_only')

    def test_r468_untrusted_job_code_is_rejected_before_exec(self):
        with self.assertRaisesRegex(ValidationError, 'template differs'): verify_jobs(self.doc, self.root, 'raise Exception("untrusted")')

    def test_r469_target_bytes_are_hashed_without_newline_normalization(self):
        path = self.target_path(); path.write_bytes(path.read_bytes().replace(b'\n', b'\r\n'))
        with self.assertRaisesRegex(ValidationError, 'changed'): self.jobs()

    def test_r470_target_symlink_is_not_accepted_as_frozen_evidence(self):
        path = self.target_path(); alternate = path.parent / 'alternate.py'; alternate.write_bytes(path.read_bytes())
        path.unlink(); path.symlink_to(alternate)
        with self.assertRaisesRegex(ValidationError, '[Ss]ymlink'): self.jobs()

    def test_r471_unresolved_program_blocks_integration_credit(self):
        self.doc['analysis']['programs']['ELIGIBLE']['blockers'].append({'message': 'unsupported fixture'})
        self.assertFalse(self.jobs()['matched'])

    def test_r472_different_field_sets_require_explicit_mapping(self):
        p = copy.deepcopy(self.doc['analysis']['programs']['ELIGIBLE']); p['name'] = 'OTHER'; p['fields'].pop('AGE')
        self.doc['analysis']['programs']['OTHER'] = p
        self.doc['program_versions']['OTHER'] = self.doc['program_versions']['ELIGIBLE']
        self.assertIn('different record layouts', self.jobs()['reason'])

    def test_r473_different_field_widths_require_mapping(self):
        p = copy.deepcopy(self.doc['analysis']['programs']['ELIGIBLE']); p['name'] = 'OTHER'; p['fields']['AGE']['width'] += 1
        self.doc['analysis']['programs']['OTHER'] = p
        self.doc['program_versions']['OTHER'] = self.doc['program_versions']['ELIGIBLE']
        self.assertIn('types or widths differ', self.jobs()['reason'])

    def test_r474_target_version_must_be_lowercase_sha256(self):
        self.doc['program_versions']['ELIGIBLE'] = 'X' * 64
        with self.assertRaisesRegex(ValidationError, 'Invalid pinned'): self.jobs()

    def test_r475_rc_equal_executes_at_zero(self):
        self.doc['jobs'][0]['steps'][0]['condition'] = 'RC=0'
        result = self.jobs()
        self.assertTrue(result['matched']); self.assertIn('record', result['actual']['JOBA'][0])

    def test_r476_rc_greater_skips_at_zero(self):
        self.doc['jobs'][0]['steps'][0]['condition'] = 'RC>0'
        result = self.jobs()
        self.assertFalse(result['matched']); self.assertEqual(result['actual']['JOBA'][0]['status'], 'SKIPPED')
        self.assertTrue(all(case['matched'] for case in result['cases']))
        self.assertFalse(result['validation']['complete'])

    def test_r477_condition_whitespace_is_supported_consistently(self):
        self.doc['jobs'][0]['steps'][0]['condition'] = 'RC\t=\t0'
        result = self.jobs()
        self.assertTrue(result['matched']); self.assertIn('record', result['actual']['JOBA'][0])

    def test_r478_empty_job_sequence_never_receives_integration_credit(self):
        self.doc['jobs'] = []
        self.assertFalse(self.jobs()['matched'])

    def test_r479_orphan_target_version_is_refused(self):
        self.doc['program_versions']['ORPHAN'] = self.doc['program_versions']['ELIGIBLE']
        with self.assertRaisesRegex(ValidationError, '[Vv]ersion|[Pp]rogram'): self.jobs()

    def test_r480_job_steps_preserve_order_and_record_state(self):
        step = copy.deepcopy(self.doc['jobs'][0]['steps'][0]); step['name'] = 'S020'; step['condition'] = 'RC<=0'
        self.doc['jobs'][0]['steps'].append(step)
        result = self.jobs(); actual = result['actual']['JOBA']
        self.assertTrue(result['matched']); self.assertEqual([r['step'] for r in actual], ['S010', 'S020'])
        self.assertEqual(actual[0]['record'], actual[1]['record'])
