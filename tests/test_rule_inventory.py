"""Rule reporting preserves source evidence, counts and conservative credit."""
import copy
import unittest
from workbench.rule_inventory import build_rule_inventory


class RuleInventoryTests(unittest.TestCase):
    def model(self):
        rule={'id':'R1','plain':'Age must be at least 18','source_start':2,'source_end':3}
        doc={'id':'p','jobs':[{'name':'J1','steps':[{'name':'S1','program':'P'},{'name':'S2','program':'P'}]},
                            {'name':'J2','steps':[{'name':'S1','program':'P'}]}],
             'analysis':{'programs':{'P':{'path':'P.cbl','source_hash':'a','rules':[rule],'coverage':[{'line':4,'disposition':'data_layout'}]}},'assets':[]}}
        rows=[]
        for number,unit,disp in [(2,'r','mapped_verified'),(3,'r','mapped_verified'),(4,'layout','mapped_verified'),(5,'u','blocked'),(6,'v','blocked')]:
            rows.append({'source_path':'P.cbl','source_hash':'a','source_kind':'cobol_program','source_line':number,
                         'unit_id':unit,'source_text':'source '+str(number),'disposition':disp,
                         'target_mappings':[{'file':'shared/target/python/a.py','version':'a','start':1,'end':2}] if disp=='mapped_verified' else [],
                         'tests':['boundary'],'evidence':[{'file':'test.json'}],'reason':'verified' if disp=='mapped_verified' else 'unsupported','replacement':None})
        return doc,{'rows':rows,'summary':{'integrity_errors':[]}}

    def test_unique_rules_and_job_memberships_do_not_inflate_reuse(self):
        doc,cov=self.model(); result=build_rule_inventory(doc,cov)
        self.assertEqual(result['summary']['business_rule']['total'],1)
        self.assertEqual(result['summary']['business_rule']['converted_verified'],1)
        self.assertEqual(result['summary']['unclassified']['total'],1)
        self.assertEqual(len(result['jobs']),2)
        self.assertEqual(result['jobs'][0]['counts']['business_rule']['total'],1)
        rule=next(r for r in result['rules'] if r['category']=='business_rule')
        self.assertEqual(len(rule['memberships']),3)
        self.assertEqual(len(rule['source']),2)
        self.assertEqual(len(result['programs']),1)
        self.assertEqual([r['invocation_memberships'] for r in result['job_programs']],[6,3])

    def test_blocked_line_revokes_rule_credit_and_no_target_is_explicit(self):
        doc,cov=self.model();cov['rows'][1]['disposition']='blocked'
        result=build_rule_inventory(doc,cov)
        self.assertEqual(result['summary']['business_rule']['converted_verified'],0)
        self.assertEqual(result['summary']['business_rule']['blocked'],1)
        gap=next(r for r in result['rules'] if r['category']=='unclassified')
        self.assertEqual(gap['modernized_behavior'],'Not implemented')

    def test_unknown_and_zero_denominators_are_not_success(self):
        doc,cov=self.model();cov['rows']=cov['rows'][-2:]
        result=build_rule_inventory(doc,cov)
        self.assertIsNone(result['summary']['business_rule']['verification_percent'])
        self.assertEqual(result['summary']['unclassified']['blocked'],1)

    def test_inventory_is_pure_and_integrity_errors_revoke_credit(self):
        doc,cov=self.model();before=copy.deepcopy((doc,cov))
        build_rule_inventory(doc,cov);self.assertEqual((doc,cov),before)
        cov['summary']['integrity_errors']=['tampered']
        result=build_rule_inventory(doc,cov)
        self.assertEqual(result['summary']['business_rule']['converted_verified'],0)

    def test_historical_external_adapter_mapping_is_preserved(self):
        import tempfile
        from pathlib import Path
        from workbench import reference
        from workbench.domain import sha
        doc,cov=self.model();cov['rows']=[cov['rows'][2]]
        target=Path(reference.__file__).resolve()
        cov['rows'][0]['target_mappings']=[{'file':str(target),'version':sha(target.read_bytes()),'start':1,'end':2}]
        with tempfile.TemporaryDirectory() as folder:
            report=build_rule_inventory(doc,cov,Path(folder))
        self.assertIn('text',report['rules'][0]['targets'][0])
        self.assertEqual(report['summary']['technical_logic']['converted_verified'],1)

    def test_multi_job_member_assigns_cards_to_own_job_only(self):
        doc,cov=self.model()
        doc['lineage']={'nodes':[{'id':'a','name':'J1','kind':'job','path':'both.jcl'},
                                  {'id':'b','name':'J2','kind':'job','path':'both.jcl'}],'edges':[]}
        template=cov['rows'][2];cov['rows']=[]
        for n,(job,text) in enumerate([('J1','//J1 JOB'),('J1','//S1 EXEC PGM=P'),('J2','//J2 JOB'),('J2','//S1 EXEC PGM=P')],1):
            cov['rows'].append({**template,'source_path':'both.jcl','source_kind':'jcl_job','source_line':n,'source_text':text,'unit_id':str(n)})
        report=build_rule_inventory(doc,cov)
        self.assertEqual([j['counts']['technical_logic']['total'] for j in report['jobs']],[2,2])
        self.assertEqual(report['summary']['technical_logic']['total'],4)

    def test_dd_control_member_is_owned_by_job(self):
        doc,cov=self.model()
        doc['lineage']={'nodes':[{'id':'a','name':'J1','kind':'job','path':'job.jcl'},
            {'id':'step','name':'S1','kind':'jcl_step','path':'job.jcl'},
            {'id':'control','name':'SORTIN','kind':'control_member','path':'SORTIN.ctl'}],
            'edges':[{'source':'a','target':'step','kind':'contains_step','resolution':'resolved'},
                     {'source':'step','target':'control','kind':'dd_member','resolution':'resolved'}]}
        cov['rows']=[{**cov['rows'][-1],'source_path':'SORTIN.ctl','source_kind':'utility_control'}]
        report=build_rule_inventory(doc,cov)
        self.assertEqual(report['jobs'][0]['counts']['unclassified']['total'],1)
        self.assertEqual(report['jobs'][1]['counts']['unclassified']['total'],0)
        self.assertEqual(report['unassigned_rule_ids'],[])

    def shared_copybook_model(self, called=False):
        from workbench.source import analyze_sources
        from workbench.domain import sha
        source = ('IDENTIFICATION DIVISION.\nPROGRAM-ID. {name}.\nDATA DIVISION.\n'
                  'LINKAGE SECTION.\nCOPY COMMON.\nPROCEDURE DIVISION USING REC.\n'
                  'IF N > 5\nMOVE "Y" TO RESULT\nELSE\nMOVE "N" TO RESULT\nEND-IF.\nGOBACK.\n')
        files = {name+'.cbl':source.format(name=name) for name in ('P','Q')}
        files['COMMON.cpy']='01 REC.\n05 N PIC 9(2).\n05 RESULT PIC X.\n'
        jobs=[{'name':job,'steps':[{'name':'S1','program':name,'inputs':[],
               'outputs':[],'condition':'ALWAYS'}]} for job,name in [('J1','P'),('J2','Q')]]
        analysis=analyze_sources(files,{'jobs':jobs})
        self.assertFalse(analysis['blockers'])
        original={p['path']:{r['line']:r['disposition'] for r in p['coverage']}
                  for p in analysis['programs'].values()}
        rows=[]
        for path,text in files.items():
            for line,value in enumerate(text.splitlines(),1):
                nonexecuting=original.get(path,{}).get(line) in ('blank','comment','structure','paragraph')
                nonexecuting |= path=='COMMON.cpy' and line==1
                rows.append({'source_path':path,'source_hash':sha(text),
                    'source_kind':'copybook' if path.endswith('.cpy') else 'cobol_program',
                    'source_line':line,'source_text':value,'unit_id':path+':'+str(line),
                    'disposition':'non_executable' if nonexecuting else 'mapped_unverified',
                    'target_mappings':[],'tests':[],'evidence':[],
                    'reason':'Fictional unverified rule projection','replacement':None})
        doc={'id':'fixture','jobs':jobs,'analysis':analysis}
        if called:
            doc['lineage']={'nodes':[{'id':name,'name':name,'kind':'program','path':name+'.cbl'}
                                     for name in ('P','Q')],
                            'edges':[{'source':'P','target':'Q','kind':'calls','resolution':'resolved'}]}
        return doc,{'rows':rows,'summary':{'integrity_errors':[]}}

    def test_shared_copybook_does_not_assign_unreachable_program_to_job(self):
        doc,cov=self.shared_copybook_model();report=build_rule_inventory(doc,cov)
        self.assertEqual({(r['job'],r['program']) for r in report['job_programs']},
                         {('J1','P'),('J2','Q')})
        common=next(r for r in report['rules'] if r['source_path']=='COMMON.cpy')
        self.assertEqual({(m['job'],m['program']) for m in common['memberships']},
                         {('J1','P'),('J2','Q')})
        self.assertEqual([r['invocation_memberships'] for r in report['programs']],
                         [len(r['rule_ids']) for r in report['programs']])

    def test_shared_copybook_preserves_transitively_called_program_membership(self):
        doc,cov=self.shared_copybook_model(called=True);report=build_rule_inventory(doc,cov)
        self.assertEqual({(r['job'],r['program']) for r in report['job_programs']},
                         {('J1','P'),('J1','Q'),('J2','Q')})
        for row in report['job_programs']:
            self.assertEqual(row['invocation_memberships'],len(row['rule_ids']))

    def test_shared_copybook_invocations_within_one_job_remain_program_specific(self):
        doc,cov=self.shared_copybook_model()
        doc['jobs'][0]['steps'].append({**doc['jobs'][1]['steps'][0],'name':'S2'})
        doc['jobs']=doc['jobs'][:1]
        report=build_rule_inventory(doc,cov)
        self.assertEqual({(r['job'],r['program']) for r in report['job_programs']},
                         {('J1','P'),('J1','Q')})
        for row in report['job_programs']:
            self.assertEqual(row['invocation_memberships'],len(row['rule_ids']))

    def test_shared_copybook_rollup_agrees_in_json_csv_html_and_workbook(self):
        import csv,json,tempfile
        from pathlib import Path
        from openpyxl import Workbook
        from workbench.rule_inventory import write_inventory,render_rule_summary,add_workbook
        doc,cov=self.shared_copybook_model();report=build_rule_inventory(doc,cov)
        expected={('J1','P'),('J2','Q')}
        with tempfile.TemporaryDirectory() as folder:
            paths=write_inventory(report,Path(folder))
            frozen=json.loads(paths[0].read_text())
            self.assertEqual({(r['job'],r['program']) for r in frozen['job_programs']},expected)
            with paths[1].open(newline='') as source:
                rows=list(csv.DictReader(source))
            common=next(r for r in rows if r['source_path']=='COMMON.cpy')
            self.assertEqual({(m['job'],m['program']) for m in json.loads(common['memberships'])},expected)
        rendered=render_rule_summary(report)
        for job,program in expected:self.assertIn('<td>'+job+' / '+program+'</td>',rendered)
        for job,program in {('J1','Q'),('J2','P')}:self.assertNotIn('<td>'+job+' / '+program+'</td>',rendered)
        book=Workbook();add_workbook(book,report)
        self.assertEqual({tuple(row[:2]) for row in book['Job Program Summary'].iter_rows(min_row=2,values_only=True)},expected)


class ProgramComparisonTests(unittest.TestCase):
    """Reporting infrastructure fixtures; no new source semantic adapter."""
    def model(self):
        doc, coverage = RuleInventoryTests().model()
        model = build_rule_inventory(doc, coverage)
        model['rules'][0]['targets'][0]['text'] = 'if record["AGE"] >= 18:\n    result = "Y"'
        return model

    def test_comparison_displays_original_and_actual_target_text(self):
        from workbench.comparison import comparison_page
        model = self.model()
        page = comparison_page(model)
        row = page['rules'][0]
        self.assertIn('source 2', row['source_excerpt'])
        self.assertIn('record["AGE"] >= 18', row['targets'][0]['text'])
        self.assertEqual(row['source_start'], 2)
        self.assertEqual(row['source_end'], 3)
        self.assertEqual(page['counts']['total'], 3)
        self.assertEqual(page['matching_total'], 3)

    def test_gaps_only_keeps_selected_no_separate(self):
        from workbench.comparison import comparison_page
        doc, cov = RuleInventoryTests().model()
        for row in cov['rows'][:2]:
            row.update(requirements_excluded=True, disposition='out_of_scope')
        model = build_rule_inventory(doc, cov)
        page = comparison_page(model, status='gaps')
        self.assertEqual(len(page['rules']), 1)
        self.assertEqual(page['counts']['excluded'], 1)
        no = comparison_page(model, status='excluded')['rules'][0]
        self.assertIn('selected No in requirements', no['modernized_behavior'])
        self.assertFalse(no['gaps'])

    def test_shared_copybook_has_both_program_owners_without_total_inflation(self):
        from workbench.comparison import comparison_page
        doc, cov = RuleInventoryTests().shared_copybook_model()
        model = build_rule_inventory(doc, cov)
        all_rows = comparison_page(model)
        keys = {p['program']: p['key'] for p in all_rows['programs']}
        p = comparison_page(model, program=keys['P'])
        q = comparison_page(model, program=keys['Q'])
        self.assertTrue(set(r['id'] for r in p['rules']) & set(r['id'] for r in q['rules']))
        self.assertLess(all_rows['counts']['total'], p['counts']['total'] + q['counts']['total'])

    def test_missing_job_memberships_do_not_make_owned_program_unassigned(self):
        from workbench.comparison import comparison_page
        doc, cov = RuleInventoryTests().model(); doc['jobs'] = []
        model = build_rule_inventory(doc, cov)
        self.assertTrue(model['unassigned_rule_ids'])
        self.assertEqual(comparison_page(model, program='unassigned')['matching_total'], 0)

    def test_invalid_counts_and_query_are_rejected_without_credit(self):
        from workbench.comparison import comparison_page
        from workbench.domain import ValidationError
        model = self.model()
        for args in ({'status':'success'}, {'program':'P'}, {'after':-1}, {'limit':0}, {'limit':51}):
            with self.subTest(args=args), self.assertRaises(ValidationError): comparison_page(model, **args)
        model['summary']['business_rule']['converted_verified'] = 99
        with self.assertRaises(ValidationError): comparison_page(model)

    def test_bounded_excerpts_and_pagination_are_explicit(self):
        from workbench.comparison import comparison_page
        model = self.model(); model['rules'][0]['source'][0]['text'] = 'X' * 5000
        page = comparison_page(model, limit=1)
        self.assertEqual(page['next_after'], 1)
        self.assertFalse(page['rules'][0]['source_excerpt_complete'])
        self.assertLessEqual(len(page['rules'][0]['source_excerpt']), 2400)
        self.assertEqual(comparison_page(model, after=1, limit=1)['rules'][0]['id'], model['rules'][1]['id'])

    def test_frozen_gap_details_name_rule_deficit_and_program_failure_separately(self):
        from workbench.comparison import comparison_page
        doc, cov = RuleInventoryTests().model()
        for row in cov['rows'][:2]: row['disposition'] = 'mapped_unverified'; row['reason'] = 'Program fixture coverage is incomplete'
        doc['runs'] = [{'id':'run-0001', 'programs':{'P':{
            'coverage':{'complete':False, 'record_count_gaps':[{'rule_id':'R1','required':20,'observed':7,'reason':'Only seven distinct states were observed','status':'bounded_generation_exhausted'}]},
            'differences':[{'case_id':'C9','expected':{'R':'Y'},'actual':{'R':'N'}}],
            'adversarial':{'passed':True}, 'unit_tests':{'passed':True}}}}]
        model = build_rule_inventory(doc, cov)
        rule = comparison_page(model)['rules'][0]
        issues = rule['program_gates']
        deficit = next(g for g in issues if g['kind'] == 'record_count_gaps')
        self.assertEqual(deficit['scope'], 'rule')
        self.assertEqual(deficit['facts']['observed'], 7)
        self.assertEqual(deficit['source_rule_id'], 'R1')
        failure = next(g for g in issues if g['kind'] == 'target_difference')
        self.assertEqual(failure['scope'], 'program')
        self.assertEqual(failure['facts']['case_ids'], ['C9'])
        self.assertTrue(rule['gaps'][0]['resolution'])


class AcceptedComparisonTests(unittest.TestCase):
    import test_workflow as _workflow
    setUp = _workflow.WorkflowTests.setUp
    create = _workflow.WorkflowTests.create
    answer = _workflow.WorkflowTests.answer

    def accepted(self):
        self.create(); self.c.start('process-a'); self.c.advance('process-a')
        self.c.import_answers('process-a', self.answer(), 'Example reviewer')
        self.c.advance('process-a'); self.c.advance('process-a')
        return self.c.ledger.get('process-a')

    def test_only_accepted_report_supports_comparison(self):
        from workbench.domain import ValidationError
        self.create()
        with self.assertRaisesRegex(ValidationError, 'accepted report'):
            self.c.program_comparison('process-a')

    def test_accepted_comparison_is_read_only_and_does_not_replay_targets(self):
        from unittest.mock import patch
        from workbench.domain import encode
        doc = self.accepted(); before = encode(doc)
        with patch('workbench.coverage.build_coverage', side_effect=AssertionError('No filter-time replay')):
            page = self.c.program_comparison('process-a')
            again = self.c.program_comparison('process-a', status='gaps')
        self.assertEqual(again['matching_total'], 0)
        self.assertTrue(page['snapshot']['inventory_sha256'])
        self.assertTrue(any(t['text'] for r in page['rules'] for t in r['targets']))
        self.assertEqual(encode(self.c.ledger.get('process-a')), before)

    def test_tampered_inventory_fails_without_repinning(self):
        from workbench.domain import ValidationError, encode
        doc = self.accepted(); before = encode(doc)
        path = self.c.artifact('process-a', 'reports/report-0001/rules.json')
        path.write_bytes(path.read_bytes() + b' ')
        with self.assertRaisesRegex(ValidationError, 'changed after freezing'):
            self.c.program_comparison('process-a')
        self.assertEqual(encode(self.c.ledger.get('process-a')), before)

    def test_comparison_http_filters_and_duplicate_parameters(self):
        from fastapi.testclient import TestClient
        from workbench.api import create_app
        self.accepted(); self.c.close()
        app = create_app(self.root); self.addCleanup(app.state.coordinator.close)
        client = TestClient(app, base_url='http://127.0.0.1:8765'); self.addCleanup(client.close)
        response = client.get('/api/process/process-a/comparison?status=gaps')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['matching_total'], 0)
        for query in ('status=gaps&status=all', 'limit=0', 'program=unknown', 'unexpected=1'):
            self.assertIn(client.get('/api/process/process-a/comparison?' + query).status_code, (400,422))

    def test_exact_read_binding_rejects_inventory_changed_between_reads(self):
        from unittest.mock import patch
        from workbench.domain import ValidationError
        self.accepted(); original = self.c.artifact
        def changed(pid, relative):
            path = original(pid, relative)
            if relative.endswith('/rules.json'): path.write_bytes(path.read_bytes() + b' ')
            return path
        with patch.object(self.c, 'artifact', side_effect=changed):
            with self.assertRaisesRegex(ValidationError, 'changed while reading'):
                self.c.program_comparison('process-a')

    def test_tampered_primary_and_metrics_reject_inventory_display(self):
        from workbench.domain import ValidationError
        self.accepted()
        for name in ('executive-report.html','metrics.json'):
            with self.subTest(name=name):
                path = self.c.artifact('process-a','reports/report-0001/'+name); raw = path.read_bytes()
                try:
                    path.write_bytes(raw+b' ')
                    with self.assertRaises(ValidationError): self.c.program_comparison('process-a')
                finally: path.write_bytes(raw)


class ComparisonReportTests(unittest.TestCase):
    model = ProgramComparisonTests.model
    def test_html_filtering_preserves_full_source_and_escapes_untrusted_text(self):
        from workbench.comparison import render_comparison_report
        model=self.model(); model['rules'][0]['source'][0]['text']='</pre><script>alert(1)</script>'
        page=render_comparison_report(model)
        self.assertIn('Gaps only',page)
        self.assertIn('Selected No',page)
        self.assertIn('data-owners="PROGRAM_',page)
        self.assertIn('&lt;script&gt;alert(1)&lt;/script&gt;',page)
        self.assertNotIn('<script>alert(1)',page)
        self.assertEqual(page.count('<article class="comparison"'),len(model['rules']))
        self.assertIn('Full counts and requirements accounting',page)

    def test_historical_missing_diagnostics_are_explicit(self):
        from workbench.comparison import comparison_page
        model=self.model();model.pop('comparison_contract_version');model.pop('program_gates')
        gap=next(r for r in comparison_page(model)['rules'] if r['gaps'])
        self.assertFalse(gap['diagnostics_recorded'])
        self.assertEqual(gap['program_gates'],[])
        self.assertIn('unsupported',gap['gaps'][0]['reason'])

    def test_selected_rule_shortfall_survives_many_unrelated_gate_previews(self):
        from workbench.comparison import comparison_page, program_key
        model=self.model();key=program_key(model['programs'][0]);rule=model['rules'][0]
        rule['status']='implemented_unverified'
        from workbench.rule_inventory import _counts
        model['summary']=_counts(model['rules']);model['programs'][0]['counts']=_counts(model['rules'])
        issues=[{'id':'other'+str(i),'kind':'record_count_gaps','scope':'rule','program':'P','source_rule_id':'OTHER'+str(i),
                 'reason':'Other source obligation','facts':{'required':20,'observed':0},'resolution':'Provide witnesses.','evidence':[]} for i in range(30)]
        issues.append({'id':'own','kind':'record_count_gaps','scope':'rule','program':'P','source_rule_id':'R1',
                       'reason':'Selected rule needs distinct states','facts':{'required':20,'observed':7},'resolution':'Provide witnesses.','evidence':[]})
        issues.append({'id':'whole','kind':'program_coverage','scope':'program','program':'P','source_rule_id':None,
                       'reason':'Whole-program coverage incomplete','facts':{},'resolution':'Resolve every obligation.','evidence':[]})
        model['program_gates']=[{'key':key,'program':'P','issues':issues}]
        result=comparison_page(model)['rules'][0]
        self.assertEqual([g['id'] for g in result['program_gates'][:2]],['own','whole'])
        self.assertEqual(result['program_gate_count'],32)
        self.assertFalse(result['program_gates_complete'])

    def test_process_integrity_gap_cannot_display_a_positive_line_reason(self):
        from workbench.comparison import comparison_page
        doc,cov=RuleInventoryTests().model();cov['summary']['integrity_errors']=['Q: frozen expectation receipt changed']
        model=build_rule_inventory(doc,cov)
        rule=comparison_page(model)['rules'][0]
        self.assertTrue(any(g['scope']=='process' and 'Q: frozen expectation receipt changed' in g['reason'] for g in rule['gaps']))
        self.assertFalse(any(g['reason']=='verified' for g in rule['gaps']))

    def test_actual_sme_item_id_retains_exact_correction_and_single_packet_limit(self):
        from workbench.comparison import comparison_page
        doc,cov=RuleInventoryTests().model()
        for row in cov['rows'][:2]:row['disposition']='mapped_unverified';row['reason']='SME confirmation is unresolved or has a correction'
        doc['answers']={'items':{'R1':{'answer':'No','correction':'Cutoff must be 21','reviewer':'Fictional reviewer'}},'packet_hash':'packet','return_hash':'return'}
        doc['blockers']=[{'kind':'sme_unresolved','item_id':'R1','message':'SME item R1: No. Cutoff must be 21','correction_preserved':True}]
        rule=comparison_page(build_rule_inventory(doc,cov))['rules'][0]
        gate=next(g for g in rule['program_gates'] if g['kind']=='sme_unresolved')
        self.assertEqual(gate['source_rule_id'],'R1')
        self.assertEqual(gate['facts']['correction'],'Cutoff must be 21')
        self.assertEqual(gate['facts']['reviewer'],'Fictional reviewer')
        self.assertIn('one-packet',gate['resolution'])
        self.assertIn('input/sme-return.xlsx',gate['evidence'])

    def test_gap_reason_and_process_preview_have_visible_completeness_counts(self):
        from workbench.comparison import comparison_page
        doc,cov=RuleInventoryTests().model()
        for i in range(10):
            cov['rows'].append({**cov['rows'][-1],'source_line':7+i,'source_text':'unknown '+str(i),'reason':'Unparsed clause '+str(i)})
        cov['summary']['integrity_errors']=['Changed file '+str(i) for i in range(14)]
        page=comparison_page(build_rule_inventory(doc,cov));rule=page['rules'][-1]
        self.assertGreater(rule['gap_count'],len(rule['gaps']))
        self.assertFalse(rule['gaps_complete'])
        self.assertEqual(page['process_gate_count'],14)
        self.assertEqual(len(page['process_gates']),12)
        self.assertFalse(page['process_gates_complete'])

    def test_shared_source_own_gates_are_prioritized_without_definition_owner(self):
        from workbench.comparison import comparison_page,program_key
        doc,cov=RuleInventoryTests().shared_copybook_model();model=build_rule_inventory(doc,cov)
        rule=next(r for r in model['rules'] if r['source_path']=='COMMON.cpy')
        self.assertIsNone(rule['program']);entry=model['programs'][0]
        issues=[{'id':'g'+str(i),'kind':'record_count_gaps','scope':'rule','program':entry['program'],'source_rule_id':'OTHER'+str(i),
                 'reason':'Unrelated rule deficit','facts':{},'resolution':'Provide witnesses.','evidence':[]} for i in range(20)]
        issues.append({**issues[0],'id':'shared-own','source_rule_id':rule['source_rule_id'],'reason':'Actual shared source obligation'})
        model['program_gates']=[{'key':program_key(entry),'program':entry['program'],'issues':issues}]
        result=next(r for r in comparison_page(model)['rules'] if r['id']==rule['id'])
        self.assertEqual(result['program_gates'][0]['id'],'shared-own')

    def test_process_sme_gate_remains_visible_in_html_with_no_rule_gaps(self):
        from workbench.comparison import comparison_page,render_comparison_report
        doc,cov=RuleInventoryTests().model();cov['rows']=cov['rows'][:3]
        doc['blockers']=[{'kind':'sme_unresolved','item_id':'G_SCOPE','message':'SME item G_SCOPE: No. Selected scope excludes downstream settlement.'}]
        doc['answers']={'items':{'G_SCOPE':{'answer':'No','correction':'Selected scope excludes downstream settlement.','reviewer':'Fictional reviewer','kind':'scope'}}}
        model=build_rule_inventory(doc,cov);page=comparison_page(model,status='gaps')
        self.assertEqual(page['matching_total'],0)
        self.assertEqual(page['process_gate_count'],1)
        rendered=render_comparison_report(model)
        self.assertIn('Process-wide gates (1)',rendered)
        self.assertIn('Selected scope excludes downstream settlement.',rendered)
        self.assertIn('Fictional reviewer',rendered)


class ComparisonReceiptLinkTests(unittest.TestCase):
    setUp=AcceptedComparisonTests.setUp
    create=AcceptedComparisonTests.create
    answer=AcceptedComparisonTests.answer
    def test_preserved_sme_return_is_a_reference_not_an_unregistered_download(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.create();self.c.start('process-a');self.c.advance('process-a')
        book=load_workbook(BytesIO(self.answer()));book['Checklist'].cell(2,6,'Fictional correction: cutoff must be 21')
        out=BytesIO();book.save(out);book.close();self.c.import_answers('process-a',out.getvalue(),'Example reviewer')
        self.c.advance('process-a');self.c.advance('process-a');page=self.c.program_comparison('process-a')
        refs={ref for rule in page['rules'] for gate in rule['program_gates'] for ref in gate.get('evidence',[])}
        self.assertIn('input/sme-return.xlsx',refs)
        self.assertNotIn('input/sme-return.xlsx',page['downloadable_evidence'])
        self.assertIn('review/packet.json',page['downloadable_evidence'])


class ComparisonScaleTests(unittest.TestCase):
    def model(self,count):
        from workbench.rule_inventory import _counts
        from workbench.comparison import program_key
        model=ProgramComparisonTests().model();template=copy.deepcopy(model['rules'][0]);rules=[];issues=[]
        for index in range(count):
            rule=copy.deepcopy(template);rule.update(id=f'RULE_{index:06}',source_rule_id=f'R_{index:06}',status='implemented_unverified',gap_reasons=['Distinct-state evidence is incomplete'])
            rules.append(rule)
            issues.append({'id':f'GATE_{index:06}','kind':'record_count_gaps','scope':'rule','program':'P','source_rule_id':rule['source_rule_id'],
                           'reason':'Insufficient distinct states','resolution':'Execute the missing source states and compare actual outputs.',
                           'facts':{'required':20,'observed':7,'witness':f'CASE-GAP-{index:06}'},'evidence':['synthetic/run-0001/P/actual-and-comparison.json']})
        program=model['programs'][0];program.update(rule_ids=[r['id'] for r in rules],counts=_counts(rules))
        issues.append({'id':'PROGRAM_GATE','kind':'unit_tests','scope':'program','program':'P','source_rule_id':None,
                       'reason':'PROGRAM_WIDE_RECEIPT_ONCE','resolution':'Correct the unit-test failure.','facts':{},'evidence':['tests/run-0001/P/unit-results.json']})
        model.update(rules=rules,summary=_counts(rules),programs=[program],jobs=[],job_programs=[],unassigned_rule_ids=[],
                     program_gates=[{'key':program_key(program),'program':'P','issues':issues}],
                     process_gates=[{'id':'PROCESS_GATE','kind':'integrity_gate','scope':'process','program':None,'source_rule_id':None,
                                     'reason':'PROCESS_RECEIPT_ONCE','resolution':'Recover original evidence.','facts':{},'evidence':[]}])
        return model

    def test_large_gap_report_retains_specific_facts_without_quadratic_duplication(self):
        from workbench.comparison import render_comparison_report
        small=render_comparison_report(self.model(110));large=render_comparison_report(self.model(220))
        self.assertLess(len(large),len(small)*2.6)
        self.assertEqual(large.count('PROGRAM_WIDE_RECEIPT_ONCE'),1)
        self.assertEqual(large.count('PROCESS_RECEIPT_ONCE'),1)
        for index in range(220):self.assertEqual(large.count(f'CASE-GAP-{index:06}'),2)
        self.assertEqual(large.count('<article class="comparison"'),220)
