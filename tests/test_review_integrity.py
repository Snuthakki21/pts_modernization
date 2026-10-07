"""Returned SME evidence is bound to the issued, frozen packet."""
import copy
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
import os
import shutil
import subprocess
import sys
from zipfile import ZipFile
import xml.etree.ElementTree as ET
from unittest.mock import patch

from openpyxl import load_workbook
from workbench.coordinator import Coordinator
from workbench.domain import ValidationError
from workbench.review import export_packet, packet_document, read_answers
from test_source import COBOL
from test_workflow import MANIFEST


class ReviewIntegrityTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.c = Coordinator(self.root)
        self.addCleanup(self.c.close)
        self.c.create(MANIFEST, {'ELIGIBLE.cbl': COBOL})
        self.c.start('process-a')
        self.c.advance('process-a')
        self.original = self.c.artifact('process-a', 'review/sme-checklist.xlsx').read_bytes()

    def returned(self, edit=None):
        book = load_workbook(BytesIO(self.original))
        for row in book['Checklist'].iter_rows(min_row=2):
            row[4].value = 'Yes'
        if edit:
            edit(book)
        output = BytesIO()
        book.save(output)
        book.close()
        return output.getvalue()

    def test_unchanged_yes_return_remains_usable_and_one_round_only(self):
        self.c.import_answers('process-a', self.returned(), 'Reviewer')
        self.assertTrue(self.c.ledger.get('process-a')['packet_imported'])
        with self.assertRaises(ValidationError):
            self.c.import_answers('process-a', self.returned(), 'Reviewer')

    def test_packet_replay_canonicalizes_mapping_order(self):
        from workbench.domain import decode, encode

        def reverse_mappings(value):
            if isinstance(value, dict):
                return {key: reverse_mappings(item) for key, item in reversed(list(value.items()))}
            if isinstance(value, list):
                return [reverse_mappings(item) for item in value]
            return value

        process = self.c.ledger.get('process-a')
        process['knowledge_context'] = {'text': 'Supplied background', 'evidence': {'path': 'context.md', 'sha256': 'a' * 64}}
        reordered = reverse_mappings(process)
        self.assertEqual(process, reordered)
        folder = self.root / 'replay-review'
        issued = export_packet(reordered, folder)
        # The ledger's canonical JSON round trip can change insertion order after
        # a crash/restart without changing any program, job or review evidence.
        restored = decode(encode(reordered))
        self.assertEqual(packet_document(restored), issued)
        self.assertEqual(export_packet(restored, folder), issued)
        book = load_workbook(folder / 'sme-checklist.xlsx')
        for row in book['Checklist'].iter_rows(min_row=2):
            row[4].value = 'Yes'
        output = BytesIO(); book.save(output); book.close()
        answers = read_answers(output.getvalue(), packet_document(restored), 'Reviewer')
        self.assertEqual(answers['packet_hash'], issued['packet_hash'])
        self.assertTrue(all(item['answer'] == 'Yes' for item in answers['items'].values()))

    def test_supported_long_question_is_returnable_with_full_context(self):
        from workbench.domain import decode
        condition = ' AND '.join("FLAG = '" + 'A' * 256 + "'" for _ in range(130))
        source = '\n'.join([
            'IDENTIFICATION DIVISION.', 'PROGRAM-ID. ELIGIBLE.', 'DATA DIVISION.',
            'LINKAGE SECTION.', '01 INPUT.', '05 FLAG PIC X(256).', '05 RESULT PIC X.',
            'PROCEDURE DIVISION USING INPUT.', 'IF ' + condition,
            "MOVE 'Y' TO RESULT", 'ELSE', "MOVE 'N' TO RESULT", 'END-IF.', 'GOBACK.',
        ])
        coordinator = Coordinator(self.root / 'long-workspace')
        self.addCleanup(coordinator.close)
        coordinator.create(MANIFEST, {'ELIGIBLE.cbl': source})
        coordinator.start('process-a'); coordinator.advance('process-a')
        doc = coordinator.ledger.get('process-a')
        self.assertEqual(doc['status'], 'WAITING_SME')
        self.assertFalse(doc['analysis']['blockers'])
        packet = decode(coordinator.artifact('process-a', 'review/packet.json').read_bytes())
        question = packet['items'][0]['question']
        self.assertGreater(len(question), 32767)
        book = load_workbook(coordinator.artifact('process-a', 'review/sme-checklist.xlsx'))
        for row in book['Checklist'].iter_rows(min_row=2):
            row[4].value = 'Yes'
        output = BytesIO(); book.save(output); book.close()
        # An untouched issued question must be importable before any quota is spent.
        coordinator.import_answers('process-a', output.getvalue(), 'Reviewer')
        self.assertTrue(coordinator.ledger.get('process-a')['packet_imported'])
        import json
        self.assertEqual(json.loads(packet['context'])['checklist_details'][packet['items'][0]['id']]['question'], question)
        returned = load_workbook(BytesIO(output.getvalue()))
        try:
            self.assertLessEqual(len(returned['Checklist']['B2'].value), 4000)
            self.assertIn('Context', returned['Checklist']['B2'].value)
            chunks = [cell[0].value for cell in returned['Context'].iter_rows(min_row=2)]
            self.assertTrue(all(len(chunk) <= 16000 for chunk in chunks))
            self.assertEqual(json.loads(''.join(chunks))['checklist_details'][packet['items'][0]['id']]['question'], question)
        finally:
            returned.close()

    def test_packet_draft_stays_outside_process_and_cleans_up_on_failure(self):
        process = self.c.ledger.get('process-a')
        destination = self.root / 'processes/another-process/review'
        saved_paths = []

        def fail_save(book, filename):
            saved_paths.append(Path(filename))
            raise OSError('simulated save failure')

        with patch('openpyxl.workbook.workbook.Workbook.save', fail_save):
            with self.assertRaises(OSError):
                export_packet(process, destination)
        self.assertEqual(len(saved_paths), 1)
        self.assertTrue(saved_paths[0].is_relative_to(self.root / '.implementation/tmp'))
        self.assertFalse(saved_paths[0].parent.exists())
        self.assertFalse(destination.exists())

    def test_long_evidence_and_formula_like_context_chunks_are_returnable(self):
        import json
        process = self.c.ledger.get('process-a')
        evidence = 'source-evidence ' * 2500
        process['analysis']['rules'][0]['source_refs'] = [evidence]
        process['knowledge_context'] = {'text': 'PLACEHOLDER'}
        base = packet_document(process)['context']
        position = base.index('PLACEHOLDER')
        boundary = ((position // 16000) + 1) * 16000
        process['knowledge_context']['text'] = 'x' * (boundary - position) + '=literal context fragment'
        packet = export_packet(process, self.root / 'long-evidence-review')
        self.assertEqual(packet['context'][boundary], '=')
        book = load_workbook(self.root / 'long-evidence-review/sme-checklist.xlsx')
        output = BytesIO()
        try:
            self.assertLessEqual(len(book['Checklist']['C2'].value), 4000)
            for row in book['Checklist'].iter_rows(min_row=2):
                row[4].value = 'Yes'
            self.assertTrue(all(row[0].data_type != 'f' for row in book['Context'].iter_rows()))
            self.assertEqual(json.loads(''.join(row[0].value for row in book['Context'].iter_rows(min_row=2)))['checklist_details'][packet['items'][0]['id']]['evidence'], evidence)
            book.save(output)
        finally:
            book.close()
        self.assertTrue(read_answers(output.getvalue(), packet, 'Reviewer')['items'])

    def test_oversized_metadata_is_rejected_before_packet_is_published(self):
        process = self.c.ledger.get('process-a')
        process['analysis']['source_snapshot'] = 'x' * 32768
        destination = self.root / 'oversized-metadata-review'
        with self.assertRaisesRegex(ValidationError, 'metadata.*Excel cell capacity'):
            export_packet(process, destination)
        self.assertFalse(destination.exists())

    def test_hard_crash_leftover_draft_does_not_block_workspace_restart(self):
        workspace = self.root / 'crash-workspace'
        package = Path(__file__).resolve().parents[1]
        program = '''
import os, sys
from unittest.mock import patch
from workbench.coordinator import Coordinator
from test_source import COBOL
from test_workflow import MANIFEST
c = Coordinator(sys.argv[1])
c.create(MANIFEST, {'ELIGIBLE.cbl': COBOL})
c.start('process-a')
def crash_save(*args, **kwargs):
    os._exit(19)
with patch('openpyxl.workbook.workbook.Workbook.save', crash_save):
    c.advance('process-a')
'''
        result = subprocess.run([sys.executable, '-c', program, str(workspace)],
                                env=dict(os.environ, PYTHONPATH=str(package)+os.pathsep+str(package/'tests')),
                                capture_output=True, text=True, timeout=20)
        self.assertEqual(result.returncode, 19, result.stderr)
        leftovers = list((workspace / '.implementation/tmp').glob('review-draft-*'))
        self.assertEqual(len(leftovers), 1)
        self.assertFalse(list((workspace / 'processes/process-a').glob('*review-draft*')))
        coordinator = Coordinator(workspace)
        self.addCleanup(coordinator.close)
        for _ in range(5):
            coordinator.advance('process-a')
            if coordinator.ledger.get('process-a')['status'] == 'WAITING_SME':
                break
        self.assertEqual(coordinator.ledger.get('process-a')['status'], 'WAITING_SME')
        self.assertTrue(coordinator.artifact('process-a', 'review/sme-checklist.xlsx').is_file())

    def test_every_context_or_metadata_change_is_rejected_without_consuming_round(self):
        mutations = {
            'context content': lambda b: setattr(b['Context']['A2'], 'value', 'Altered layouts'),
            'context title': lambda b: setattr(b['Context']['A1'], 'value', 'Changed title'),
            'context missing': lambda b: b.remove(b['Context']),
            'context extra row': lambda b: b['Context'].append(['extra data']),
            'context extra column': lambda b: setattr(b['Context']['B1'], 'value', 'extra data'),
            'context formula': lambda b: setattr(b['Context']['A2'], 'value', '=1+1'),
            'context dimension': lambda b: setattr(b['Context']['B5'], 'value', None),
            'metadata identity': lambda b: setattr(b['Metadata']['B1'], 'value', 'other-process'),
            'metadata label': lambda b: setattr(b['Metadata']['A1'], 'value', 'Ignored label'),
            'metadata evidence': lambda b: setattr(b['Metadata']['B4'], 'value', 'MAINFRAME_OBSERVED'),
            'metadata extra row': lambda b: b['Metadata'].append(['extra data']),
            'metadata formula': lambda b: setattr(b['Metadata']['A4'], 'value', '=1+1'),
            'checklist header': lambda b: setattr(b['Checklist']['B1'], 'value', 'Changed question heading'),
            'extra sheet': lambda b: b.create_sheet('Altered evidence'),
        }
        for label, mutation in mutations.items():
            with self.subTest(label=label):
                with self.assertRaises(ValidationError):
                    self.c.import_answers('process-a', self.returned(mutation), 'Reviewer')
                self.assertFalse(self.c.ledger.get('process-a')['packet_imported'])
        self.c.import_answers('process-a', self.returned(), 'Reviewer')
        self.assertTrue(self.c.ledger.get('process-a')['packet_imported'])

    def test_context_chunks_are_bound_individually(self):
        process = self.c.ledger.get('process-a')
        process['knowledge_context'] = {'text': 'x' * 32000}
        packet = export_packet(process, self.root / 'large-review')
        data = (self.root / 'large-review/sme-checklist.xlsx').read_bytes()
        book = load_workbook(BytesIO(data))
        # Repartition the identical joined text: comparing only joined strings is insufficient.
        one, two = book['Context']['A2'], book['Context']['A3']
        one.value, two.value = one.value[:-1], one.value[-1:] + two.value
        output = BytesIO(); book.save(output); book.close()
        with self.assertRaises(ValidationError):
            read_answers(output.getvalue(), packet, 'Reviewer')

    def test_original_formula_like_literal_questions_remain_usable(self):
        for index, question in enumerate(('=Literal', '+Literal', '-Literal', '@Literal', "'=Literal")):
            with self.subTest(question=question):
                process = self.c.ledger.get('process-a')
                process['analysis']['rules'][0]['plain'] = question
                folder = self.root / f'literal-review-{index}'
                packet = export_packet(process, folder)
                book = load_workbook(folder / 'sme-checklist.xlsx')
                for row in book['Checklist'].iter_rows(min_row=2):
                    row[4].value = 'Yes'
                data = BytesIO(); book.save(data); book.close()
                self.assertEqual(read_answers(data.getvalue(), packet, 'Reviewer')['items'][packet['items'][0]['id']]['answer'], 'Yes')

    def test_frozen_packet_content_must_match_its_hash(self):
        from workbench.domain import decode
        packet = decode(self.c.artifact('process-a', 'review/packet.json').read_bytes())
        packet = copy.deepcopy(packet)
        packet['context'] = 'changed frozen content'
        with self.assertRaises(ValidationError):
            read_answers(self.returned(), packet, 'Reviewer')

    def test_forged_sheet_dimension_cannot_hide_extra_data(self):
        from workbench.domain import decode
        packet = decode(self.c.artifact('process-a', 'review/packet.json').read_bytes())
        mutated = self.returned(lambda b: setattr(b['Context']['B2'], 'value', 'hidden changed context'))
        result = BytesIO()
        with ZipFile(BytesIO(mutated)) as source, ZipFile(result, 'w') as destination:
            for member in source.infolist():
                contents = source.read(member.filename)
                if member.filename == 'xl/worksheets/sheet3.xml':
                    xml = ET.fromstring(contents)
                    xml.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}dimension').set('ref', 'A1:A2')
                    contents = ET.tostring(xml)
                destination.writestr(member, contents)
        with self.assertRaises(ValidationError):
            read_answers(result.getvalue(), packet, 'Reviewer')


class SetupScriptTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name) / 'workspace with spaces'
        (self.root / 'scripts').mkdir(parents=True)
        original = Path(__file__).resolve().parents[1] / 'scripts/setup.sh'
        shutil.copy(original, self.root / 'scripts/setup.sh')
        (self.root / 'requirements.lock').write_text('fixture lock')
        self.bin = Path(self.tmp.name) / 'bin'
        self.bin.mkdir()
        self.log = Path(self.tmp.name) / 'calls'
        self.interpreter = self.bin / 'python3.12'
        self.interpreter.write_text("""#!/bin/bash
if [[ \"$1\" == '-c' ]]; then exit \"${FAKE_BAD_VERSION:-0}\"; fi
if [[ \"$1 $2\" == '-m venv' ]]; then
    echo venv >> \"$FAKE_LOG\"
    mkdir -p \"$3/bin\"
    cp \"$FAKE_TEMPLATE\" \"$3/bin/python\"
    exit \"${FAKE_VENV_EXIT:-0}\"
fi
if [[ \"$1 $2\" == '-m pip' ]]; then
    echo \"pip:$*\" >> \"$FAKE_LOG\"
    if [[ \"$3\" == '--version' ]]; then exit 0; fi
    exit \"${FAKE_PIP_EXIT:-0}\"
fi
if [[ \"$1 $2\" == '-m workbench.preflight' ]]; then
    echo \"preflight:$*\" >> \"$FAKE_LOG\"
    exit \"${FAKE_PREFLIGHT_EXIT:-0}\"
fi
exit 9
""")
        self.interpreter.chmod(0o755)
        (self.bin / 'python3').symlink_to(self.interpreter)
        self.environment = dict(os.environ, PATH=str(self.bin) + ':' + os.environ['PATH'],
                                FAKE_LOG=str(self.log), FAKE_TEMPLATE=str(self.interpreter))

    def run_setup(self, **environment):
        return subprocess.run(['/bin/bash', str(self.root / 'scripts/setup.sh')],
                              cwd=self.tmp.name, env=dict(self.environment, **environment),
                              text=True, capture_output=True)

    def test_root_with_spaces_and_idempotent_environment_use_hashed_wheels(self):
        first = self.run_setup()
        self.assertEqual(first.returncode, 0, first.stderr)
        second = self.run_setup()
        self.assertEqual(second.returncode, 0, second.stderr)
        calls = self.log.read_text().splitlines()
        self.assertEqual(calls.count('venv'), 1)
        pip_calls = [line for line in calls if line.startswith('pip:')]
        self.assertEqual(len(pip_calls), 4)
        probes = [line for line in pip_calls if '--version' in line]
        installs = [line for line in pip_calls if ' install ' in line]
        self.assertEqual(len(probes), 2)
        self.assertEqual(len(installs), 2)
        self.assertTrue(all('--require-hashes' in line and '--only-binary=:all:' in line for line in installs))
        preflight_calls=[line for line in calls if line.startswith('preflight:')]
        self.assertEqual(len(preflight_calls),2)
        self.assertTrue(all('--initialize-knowledge' in line and '--workspace' in line for line in preflight_calls))
        self.assertTrue((self.root / 'Endeavor').is_dir())
        self.assertTrue((self.root / 'knowledge/inbox').is_dir())

    def test_failed_install_stops_before_reporting_success_or_creating_input_folders(self):
        result = self.run_setup(FAKE_PIP_EXIT='23')
        self.assertNotEqual(result.returncode, 0)
        self.assertNotIn('complete', result.stdout.lower())
        self.assertFalse((self.root / 'Endeavor').exists())

    def test_wrong_python_version_has_actionable_preflight(self):
        result = self.run_setup(FAKE_BAD_VERSION='4')
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('Python 3.12', result.stderr)
        self.assertFalse(self.log.exists())

    def test_missing_python_has_actionable_preflight(self):
        self.interpreter.unlink()
        (self.bin / 'python3').unlink()
        (self.bin / 'dirname').symlink_to('/usr/bin/dirname')
        result = self.run_setup(PATH=str(self.bin))
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('Python 3.12', result.stderr)

    def test_environment_creation_failure_stops_before_install(self):
        result = self.run_setup(FAKE_VENV_EXIT='27')
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(any(line.startswith('pip:') for line in self.log.read_text().splitlines()))

    def test_failed_preflight_stops_before_reporting_success(self):
        result=self.run_setup(FAKE_PREFLIGHT_EXIT='2')
        self.assertNotEqual(result.returncode,0)
        self.assertNotIn('Setup complete',result.stdout)
        self.assertIn('Setup checks found blockers',result.stderr)
        self.assertFalse((self.root/'Endeavor').exists())


if __name__ == '__main__':
    unittest.main()
