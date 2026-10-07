"""Source intake, frozen notes, rule grouping, and host-neutral modernization."""
import copy
import json
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from workbench.coordinator import Coordinator
from workbench.domain import ValidationError, sha
from workbench.intake import parse_manifest
from workbench.source import analyze_sources
from workbench.review import packet_document, export_packet, read_answers
from workbench.process_context import freeze_context, excerpt
from workbench.copilot import build_task, validate_submission
from workbench.zowe_setup import inspect_project_config
from test_source import COBOL
from test_workflow import MANIFEST


class ModernizationUpgradeTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name)

    def process(self):
        doc=parse_manifest(MANIFEST);doc['sme_packet_version']=4
        doc['analysis']=analyze_sources({'ELIGIBLE.cbl':COBOL},doc)
        doc['llm']={'analysis':{'rule_classifications':{r['id']:{'category':'business_rule','reason':'Eligibility outcome is a business decision'} for r in doc['analysis']['rules']}}}
        return doc

    def test_external_export_and_original_notes_are_frozen_without_modifying_source(self):
        folder=self.root/'external';folder.mkdir();raw=COBOL.replace('\n','\r\n').encode();(folder/'ELIGIBLE.cbl').write_bytes(raw)
        notes=self.root/'notes.md';notes.write_text('# Process\nApplication facts\n# Job\nJOBA uses ELIGIBLE\n')
        workspace=self.root/'workspace'
        c=Coordinator(workspace);self.addCleanup(c.close)
        doc=c.create(MANIFEST,source_folder=folder,process_notes=notes,assistant_mode='agent')
        self.assertEqual(doc['sme_packet_version'],4)
        self.assertEqual((workspace/'processes/process-a/input/sources/ELIGIBLE.cbl').read_bytes(),raw)
        self.assertEqual((folder/'ELIGIBLE.cbl').read_bytes(),raw)
        context=json.loads(c.artifact(doc['id'],'analysis/process-context.json').read_text())
        self.assertEqual(context['documents'][0]['sha256'],sha(notes.read_bytes()))
        notes.write_text('changed')
        self.assertIn('Application facts',excerpt(context,'NOTE_1',1,2)['text'])

    def test_context_index_reads_large_notes_and_rejects_symlinks(self):
        folder=self.root/'knowledge/inbox';folder.mkdir(parents=True)
        (folder/'process.md').write_text('# Process\n'+('Facts\n'*4000))
        context=freeze_context(self.root)
        self.assertGreater(len(context['documents'][0]['text']),16000)
        self.assertEqual(excerpt(context,'NOTE_1',2,3)['text'],'Facts\nFacts')
        (folder/'alias.md').symlink_to(folder/'process.md')
        with self.assertRaises(ValidationError):freeze_context(self.root)

    def test_v4_equivalent_repeated_rules_share_one_authentic_answer(self):
        doc=self.process();program=doc['analysis']['programs']['ELIGIBLE']
        repeated=copy.deepcopy(program['rules'][0]);repeated['id']='REPEATED';repeated['source_refs']=['ELIGIBLE.cbl:90-95']
        program['rules'].append(repeated);doc['analysis']['rules'].append(repeated)
        doc['llm']['analysis']['rule_classifications']['REPEATED']=copy.deepcopy(doc['llm']['analysis']['rule_classifications'][program['rules'][0]['id']])
        packet=packet_document(doc);groups=json.loads(packet['context'])['rule_groups']
        self.assertEqual(len(groups),1)
        folder=self.root/'review';export_packet(doc,folder)
        book=load_workbook(folder/'sme-checklist.xlsx');group_id=next(iter(groups))
        for row in book['Checklist'].iter_rows(min_row=2):
            row[4].value='Yes'
            if row[0].value==group_id:row[5].value='Threshold is wrong'
        out=BytesIO();book.save(out);book.close()
        answers=read_answers(out.getvalue(),packet,'Fixture reviewer')
        for member in groups[group_id]['members']:
            self.assertEqual(answers['items'][member]['group_id'],group_id)
            self.assertEqual(answers['items'][member]['correction'],'Threshold is wrong')
        self.assertEqual(len(doc['analysis']['rules']),3)

    def test_preceding_input_mutation_prevents_grouping_even_with_identical_prose(self):
        doc=self.process();program=doc['analysis']['programs']['ELIGIBLE']
        first=program['rules'][0]; repeated=copy.deepcopy(first);repeated['id']='REPEATED'
        # Altering predicate input invalidates equivalence of the later decision.
        field=first['predicate']['left']['field']
        program['rules'][1]['then']=[{'field':field,'value':99}]
        program['rules'].append(repeated);doc['analysis']['rules'].append(repeated)
        packet=packet_document(doc)
        self.assertFalse(json.loads(packet['context'])['rule_groups'])

    def test_unknown_classification_is_explicit_and_v3_still_has_individual_rules(self):
        doc=self.process();doc.pop('llm')
        packet=packet_document(doc)
        self.assertTrue(any(i['kind']=='unclassified' for i in packet['items']))
        doc['sme_packet_version']=3
        packet=packet_document(doc)
        self.assertEqual(len([i for i in packet['items'] if i['kind']=='business_rule']),len(doc['analysis']['rules']))

    def test_zowe_normalization_is_semantic_noop_with_backup_and_no_network(self):
        (self.root/'zowe.config.json').write_text(json.dumps({'profiles':{'base':{'type':'base','secure':['user','password','user']}},'defaults':{'base':'base'}}))
        (self.root/'zowe.schema.json').write_text('{"type":"object"}')
        original=(self.root/'zowe.config.json').read_bytes()
        self.assertFalse(inspect_project_config(self.root)['changed'])
        result=inspect_project_config(self.root,True)
        self.assertTrue(result['changed']);self.assertEqual(result['connectivity'],'UNVERIFIED')
        self.assertEqual(next((self.root/'.implementation/tmp').glob('zowe-config-*.json')).read_bytes(),original)
        self.assertFalse(inspect_project_config(self.root,True)['changed'])

    def test_i_z_candidates_never_resolve_missing_environment_identity(self):
        from workbench.lineage import map_lineage
        doc=parse_manifest(MANIFEST.replace('ELIGIBLE','IPROG'));doc['sme_packet_version']=4
        files={'ZPROG.cbl':COBOL.replace('ELIGIBLE','ZPROG'),'JOBA.jcl':'//JOBA JOB\n//S010 EXEC PGM=IPROG\n'}
        lineage=map_lineage(files,doc)
        candidates=[g for g in lineage['closure']['gaps'] if g.get('environment_candidates')]
        self.assertTrue(candidates);self.assertFalse(lineage['closure']['complete'])
        self.assertEqual(candidates[0]['environment_candidates'][0]['status'],'UNVERIFIED_ENVIRONMENT_RELATIONSHIP')

    def test_large_repeated_inventory_uses_program_defaults_and_keeps_all_rule_ids(self):
        from workbench.domain import encode
        doc=self.process();program=doc['analysis']['programs']['ELIGIBLE']
        base=copy.deepcopy(program['rules'][0])
        program['rules']=[{**copy.deepcopy(base),'id':'ELIGIBLE_R'+str(i),'source_start':2*i+1,'source_end':2*i+2} for i in range(19000)]
        doc['analysis']['rules']=program['rules']
        doc['llm']={'analysis':{'rule_classification_defaults':{'ELIGIBLE':{'category':'business_rule','reason':'Repeated identical eligibility decision'}}}}
        packet=packet_document(doc);groups=json.loads(packet['context'])['rule_groups']
        self.assertEqual(len(groups),1);self.assertEqual(len(next(iter(groups.values()))['members']),19000)
        self.assertEqual(len(doc['analysis']['rules']),19000)
        doc['source_files']={'ELIGIBLE.cbl':sha(COBOL)};doc['manifest_hash']=sha(MANIFEST)
        task=build_task(doc,{'ELIGIBLE.cbl':COBOL},doc['analysis'])
        self.assertEqual(task['rule_catalog_total'],19000)
        self.assertEqual(len(task['rule_catalog']),50)
        submitted={k:task[k] for k in ('process_id','task_hash','lineage_hash')}
        submitted.update(summary='Repeated eligibility decision',assumptions=[],questions=[],source_refs=[{'path':'ELIGIBLE.cbl','start_line':1,'end_line':1,'source_hash':sha(COBOL)}],adapter_tasks=[],rule_classification_defaults=doc['llm']['analysis']['rule_classification_defaults'])
        returned=validate_submission(task,submitted,doc['analysis'])
        self.assertIn('ELIGIBLE',returned['analysis']['rule_classification_defaults'])

    def test_classification_cannot_reference_another_program_or_rule(self):
        doc=self.process();doc['source_files']={'ELIGIBLE.cbl':sha(COBOL)};doc['manifest_hash']=sha(MANIFEST)
        task=build_task(doc,{'ELIGIBLE.cbl':COBOL},doc['analysis'])
        submitted={k:task[k] for k in ('process_id','task_hash','lineage_hash')}
        submitted.update(summary='Evidence',assumptions=[],questions=[],source_refs=[{'path':'ELIGIBLE.cbl','start_line':1,'end_line':1,'source_hash':sha(COBOL)}],adapter_tasks=[],rule_classification_defaults={'OTHER':{'category':'business_rule','reason':'Wrong program'}})
        with self.assertRaises(ValidationError):validate_submission(task,submitted,doc['analysis'])

    def test_real_agent_workflow_report_reconciles_all_formats(self):
        from workbench.executive import accepted_executive
        c=Coordinator(self.root);self.addCleanup(c.close)
        files={'ELIGIBLE.cbl':COBOL,'JOBA.jcl':'//JOBA JOB\n//S010 EXEC PGM=ELIGIBLE\n'}
        c.create(MANIFEST,files,assistant_mode='agent');c.start('process-a');c.advance('process-a')
        self.assertEqual(c.ledger.get('process-a')['status'],'WAITING_COPILOT')
        task=c.agent_task('process-a')
        submitted={k:task[k] for k in ('process_id','task_hash','lineage_hash')}
        submitted.update(summary='Eligibility decisions with job orchestration',assumptions=[],questions=[],source_refs=[{'path':'ELIGIBLE.cbl','start_line':10,'end_line':19,'source_hash':sha(COBOL)}],adapter_tasks=[],rule_classification_defaults={'ELIGIBLE':{'category':'business_rule','reason':'Both decisions set the eligibility output'}})
        c.submit_agent_analysis('process-a',submitted);c.advance('process-a')
        book=load_workbook(c.artifact('process-a','review/sme-checklist.xlsx'))
        for row in book['Checklist'].iter_rows(min_row=2):row[4].value='Yes'
        out=BytesIO();book.save(out);book.close()
        c.import_answers('process-a',out.getvalue(),'Fictional test reviewer');c.advance('process-a');c.advance('process-a')
        doc=c.ledger.get('process-a');self.assertIn(doc['status'],('COMPLETED','COMPLETED_WITH_BLOCKERS'))
        result=json.loads(c.artifact('process-a','reports/report-0001/rules.json').read_text())
        executive=accepted_executive(c,doc)['executive']
        self.assertEqual(result['summary'],executive['rule_inventory']['summary'])
        self.assertEqual(result['summary']['business_rule']['total'],2)
        book=load_workbook(c.artifact('process-a','reports/report-0001/metrics.xlsx'))
        try:
            self.assertEqual(book['Business Rules'].max_row,3)
            self.assertIn('Job Program Summary',book.sheetnames)
        finally:book.close()
        detail=c.artifact('process-a','reports/report-0001/rules.html').read_text()
        self.assertIn('IF AGE &gt;= 18',detail)
        self.assertIn('Modernized implementation',detail)
        import csv
        with c.artifact('process-a','reports/report-0001/rules.csv').open() as stream:
            rows=list(csv.DictReader(stream))
        self.assertEqual(len(rows),len(result['rules']))

    def test_unsupported_intervening_behavior_cannot_establish_equivalence(self):
        doc=self.process();program=doc['analysis']['programs']['ELIGIBLE']
        repeated=copy.deepcopy(program['rules'][0]);repeated['id']='REPEATED'
        program['rules'].append(repeated);doc['analysis']['rules'].append(repeated)
        doc['llm']['analysis']['rule_classification_defaults']={'ELIGIBLE':{'category':'business_rule','reason':'Candidate decision'}}
        program['blockers']=[{'kind':'unsupported_source','message':'Unknown input effect'}]
        self.assertFalse(json.loads(packet_document(doc)['context'])['rule_groups'])

    def test_refresh_refuses_stale_loaded_adapters_and_records_unchanged_analysis(self):
        from unittest.mock import patch
        c=Coordinator(self.root);self.addCleanup(c.close)
        c.create(MANIFEST,{'ELIGIBLE.cbl':COBOL,'JOBA.jcl':'//JOBA JOB\n//S010 EXEC PGM=ELIGIBLE\n'},assistant_mode='agent')
        c.start('process-a');c.advance('process-a')
        with patch('workbench.coordinator.adapter_fingerprint',return_value='changed'):
            with self.assertRaisesRegex(ValidationError,'Restart'):c.refresh_analysis('process-a')
        self.assertEqual(c.ledger.get('process-a')['status'],'WAITING_COPILOT')
        c.refresh_analysis('process-a');c.advance('process-a')
        self.assertEqual(c.agent_task('process-a')['continuation']['progress'],'UNCHANGED')
