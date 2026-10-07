import base64
import json
import tempfile
import socket
import subprocess
import sys
import time
import unittest
import urllib.request
import urllib.error
from io import BytesIO
from openpyxl import load_workbook
from workbench.coordinator import Coordinator

class HttpTests(unittest.TestCase):
    def test_one_click_review_return_and_automatic_report(self):
        with tempfile.TemporaryDirectory() as root:
            with socket.socket() as listener:
                listener.bind(('127.0.0.1',0));port=listener.getsockname()[1]
            url=f'http://127.0.0.1:{port}'
            # Drain service logs to disk: unread PIPE buffers can deadlock a chatty
            # server while this test repeatedly polls the real HTTP workflow.
            logs=tempfile.TemporaryFile();self.addCleanup(logs.close)
            service=subprocess.Popen([sys.executable,'-m','workbench','--root',root,'--port',str(port)],stdout=logs,stderr=subprocess.STDOUT)
            try:
                def get(path):return urllib.request.urlopen(url+path,timeout=10).read()
                until=time.monotonic()+15
                while True:
                    try:state=json.loads(get('/api/state'));break
                    except (urllib.error.URLError,ConnectionError):
                        if service.poll() is not None or time.monotonic()>until:
                            logs.seek(0);self.fail('Uvicorn startup failed: '+logs.read().decode(errors='replace')[-4000:])
                        time.sleep(.05)
                token=state['token']
                def post(path,payload):
                    q=urllib.request.Request(url+path,data=json.dumps(payload).encode(),headers={'Content-Type':'application/json','Origin':url,'X-Workbench-Token':token});return json.load(urllib.request.urlopen(q,timeout=10))
                def wait_for(pid,statuses):
                    until=time.monotonic()+45
                    while time.monotonic()<until:
                        p=next(x for x in json.loads(get('/api/state'))['processes'] if x['id']==pid)
                        if p['status'] in statuses:return p
                        time.sleep(.05)
                    self.fail('Automatic worker did not reach '+str(statuses)+'; last status: '+str(p.get('status'))+'; blockers: '+str(p.get('blockers')))
                p=post('/api/demo',{});pid=p['id'];p=wait_for(pid,['WAITING_SME','FAILED']);self.assertEqual(p['status'],'WAITING_SME')
                raw=get('/api/process/'+pid+'/artifact?path=review/sme-checklist.xlsx');book=load_workbook(BytesIO(raw))
                for row in book['Checklist'].iter_rows(min_row=2):row[4].value='Yes';row[6].value='Fictional reviewer'
                out=BytesIO();book.save(out);book.close();post('/api/process/'+pid+'/answers',{'xlsx':base64.b64encode(out.getvalue()).decode(),'reviewer':'Fictional reviewer'})
                p=wait_for(pid,['COMPLETED','COMPLETED_WITH_BLOCKERS','FAILED','REPORTING_FAILED']);self.assertEqual(p['status'],'COMPLETED',p.get('blockers'))
                self.assertTrue(get('/api/process/'+pid+'/artifact?path=reports/report-0001/management.pptx').startswith(b'PK'))
                self.assertEqual(json.loads(get('/api/state'))['portfolio']['processes'],0)
                with self.assertRaises(urllib.error.HTTPError):post('/api/process/'+pid+'/answers',{'xlsx':base64.b64encode(out.getvalue()).decode(),'reviewer':'Fictional reviewer'})
            finally:
                service.terminate()
                try:service.communicate(timeout=10)
                except subprocess.TimeoutExpired:service.kill();service.communicate(timeout=5)
                # Lifespan shutdown releases the same-workspace single-writer lock.
                recovered=Coordinator(root);recovered.close()

if __name__=='__main__':unittest.main()
