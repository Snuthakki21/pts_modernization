"""One portable, offline SME HTML return shares the existing packet identity."""
import copy
from io import BytesIO
from pathlib import Path
import random
import tempfile
import unittest

from openpyxl import load_workbook
from workbench.domain import MAX_UPLOAD, ValidationError, decode, encode, sha
from workbench.intake import parse_manifest
from workbench.review import (HTML_RETURN_HEADER, _HTMLReviewShape, _html_response,
                              export_packet, is_html_return, packet_document,
                              read_answers, read_html_return, read_return,
                              render_html_packet)
from workbench.source import analyze_sources
from test_source import COBOL
from test_workflow import MANIFEST


class HTMLReviewTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.folder = Path(self.tmp.name)/'packet'
        self.process = parse_manifest(MANIFEST.replace('Referral eligibility', 'Fictional offline SME fixture'))
        self.process['analysis'] = analyze_sources({'ELIGIBLE.cbl': COBOL}, self.process)
        self.packet = export_packet(self.process, self.folder)
        self.blank = (self.folder/'sme-checklist.html').read_bytes()
        self.reviewer = 'Fictional fixture reviewer'

    def response(self):
        response = _html_response(self.packet)
        response['reviewer'] = self.reviewer
        return response

    def returned(self, response=None):
        return render_html_packet(self.packet, self.response() if response is None else response)

    def test_export_keeps_one_packet_four_formats_and_no_default_approval(self):
        self.assertEqual({p.name for p in self.folder.iterdir()},
                         {'packet.json', 'sme-checklist.xlsx', 'sme-checklist.docx', 'sme-checklist.html'})
        self.assertTrue(is_html_return(self.blank))
        self.assertTrue(self.blank.startswith(HTML_RETURN_HEADER))
        self.assertNotIn(b'Return the Excel workbook.', self.blank)
        shape = _HTMLReviewShape(self.blank.decode())
        response = decode(shape.payloads['workbench-sme-return'])
        self.assertEqual(response['reviewer'], '')
        self.assertTrue(all(item['answer'] == '' for item in response['items']))
        with self.assertRaisesRegex(ValidationError, 'reviewer attribution'):
            read_html_return(self.blank, self.packet, self.reviewer)
        self.assertEqual(export_packet(self.process, self.folder), self.packet)
        self.assertEqual((self.folder/'sme-checklist.html').read_bytes(), self.blank)

    def test_html_and_legacy_workbook_have_same_answer_meaning_and_group_expansion(self):
        response = self.response()
        for index, item in enumerate(response['items']):
            item.update(answer=['Yes', 'No', 'Not sure', ''][index % 4], correction='Fictional correction' if index % 2 else '')
        book = load_workbook(self.folder/'sme-checklist.xlsx')
        for row, item in zip(book['Checklist'].iter_rows(min_row=2), response['items']):
            row[4].value, row[5].value = item['answer'], item['correction']
        out = BytesIO(); book.save(out); book.close()
        workbook = read_answers(out.getvalue(), self.packet, self.reviewer)
        returned = read_return(self.returned(response), self.packet, self.reviewer)
        self.assertEqual(returned['return_format'], 'html')
        for key in ('packet_hash', 'source_snapshot', 'reviewer', 'items'):
            self.assertEqual(returned[key], workbook[key])
        self.assertEqual(read_return(out.getvalue(), self.packet, self.reviewer), workbook)
        self.assertNotIn('return_format', workbook)
        self.assertEqual(returned['return_hash'], sha(self.returned(response)))
        grouped = copy.deepcopy(self.packet)
        grouped['version'] = 4
        context = decode(grouped['context']); context['rule_groups'] = {grouped['items'][0]['id']: {'members': ['GROUP-MEMBER-A', 'GROUP-MEMBER-B']}}
        grouped['context'] = encode(context).decode(); grouped['packet_hash'] = sha(encode({k:v for k,v in grouped.items() if k != 'packet_hash'}))
        grouped_response = _html_response(grouped); grouped_response['reviewer'] = self.reviewer
        grouped_response['items'][0]['answer'] = 'No'
        answer = read_html_return(render_html_packet(grouped, grouped_response), grouped, self.reviewer)
        self.assertEqual(answer['items']['GROUP-MEMBER-A']['answer'], 'No')
        self.assertEqual(answer['items']['GROUP-MEMBER-B']['group_id'], grouped['items'][0]['id'])

    def test_20_runtime_randomized_answer_states_replay_and_match_actual_workbook_target(self):
        seed = random.SystemRandom().randrange(2**32)
        generator = random.Random(seed)
        states = set()
        for index in range(20):
            response = self.response()
            for item in response['items']:
                item['answer'] = generator.choice(['Yes', 'No', 'Not sure', ''])
                item['correction'] = f'Fictional state {index}; identifier 00{generator.randrange(10000):04d}'
            data = self.returned(response)
            states.add(sha(data))
            actual = read_html_return(data, self.packet, self.reviewer)
            book = load_workbook(self.folder/'sme-checklist.xlsx')
            for row, item in zip(book['Checklist'].iter_rows(min_row=2), response['items']):
                row[4].value, row[5].value = item['answer'], item['correction']
            out = BytesIO(); book.save(out); book.close()
            expected = read_answers(out.getvalue(), self.packet, self.reviewer)
            self.assertEqual(actual['items'], expected['items'], f'seed={seed}, state={index}')
            self.assertEqual(read_html_return(data, self.packet, self.reviewer), actual)
        self.assertEqual(len(states), 20)

    def test_browser_serialization_and_crlf_are_accepted_without_shell_changes(self):
        data = self.returned()
        browser_serialized = data.replace(b'<!doctype html>', b'<!DOCTYPE html>', 1).replace(b'<meta charset="utf-8">', b'<meta charset="utf-8"/>')
        # The typed guard stays exact; the browser's Save explicitly emits lower-case doctype.
        self.assertFalse(is_html_return(browser_serialized))
        self.assertEqual(_HTMLReviewShape(browser_serialized.decode()).events,
                         _HTMLReviewShape(data.decode()).events)
        browser_serialized = browser_serialized.replace(b'<!DOCTYPE html>', b'<!doctype html>')
        self.assertEqual(read_html_return(browser_serialized, self.packet, self.reviewer)['items'],
                         read_html_return(data, self.packet, self.reviewer)['items'])
        # Source/packet JSON line endings may vary, but fixed trusted script bytes must not.
        response = encode(self.response()).decode().replace('\n', '\r\n')
        text = self.returned().decode(); shape = _HTMLReviewShape(text)
        old = shape.payloads['workbench-sme-return']
        changed = text.replace(old, response).encode()
        self.assertEqual(read_html_return(changed, self.packet, self.reviewer)['items'],
                         read_html_return(data, self.packet, self.reviewer)['items'])

    def test_dangerous_source_and_commentary_are_json_escaped_and_retained_as_text(self):
        dangerous = '</script><img src=x onerror="alert(1)">&\u2028\u2029'
        packet = copy.deepcopy(self.packet)
        packet['items'][0]['question'] = dangerous
        packet['items'][0]['evidence'] = dangerous
        packet['context'] += dangerous
        packet['packet_hash'] = sha(encode({k:v for k,v in packet.items() if k != 'packet_hash'}))
        response = _html_response(packet); response['reviewer'] = self.reviewer
        response['items'][0]['correction'] = dangerous
        data = render_html_packet(packet, response)
        self.assertNotIn(dangerous.encode(), data)
        self.assertIn(b'\\u003c/script\\u003e', data)
        self.assertIn(b'\\u2028\\u2029', data)
        result = read_html_return(data, packet, self.reviewer)
        self.assertEqual(result['items'][packet['items'][0]['id']]['question'], dangerous)
        self.assertEqual(result['items'][packet['items'][0]['id']]['correction'], dangerous.strip())
        self.assertIn(b'node.textContent = text', data)
        self.assertNotIn(b'innerHTML', data)
        self.assertNotIn(b'fetch(', data)
        self.assertIn(b"connect-src 'none'", data)

    def test_altered_packet_process_snapshot_and_hash_are_rejected(self):
        for field, value in [('process_id','other-process'), ('packet_hash','f'*64), ('source_snapshot','e'*64)]:
            response = self.response(); response[field] = value
            with self.subTest(field=field), self.assertRaisesRegex(ValidationError, 'different packet'):
                read_html_return(self.returned(response), self.packet, self.reviewer)
        text = self.returned().decode(); shape = _HTMLReviewShape(text)
        packet = copy.deepcopy(self.packet); packet['items'][0]['question'] += ' Altered'
        text = text.replace(shape.payloads['workbench-sme-packet'], encode(packet).decode())
        with self.assertRaisesRegex(ValidationError, 'Original HTML packet'):
            read_html_return(text.encode(), self.packet, self.reviewer)
        packet['packet_hash'] = 'f'*64
        with self.assertRaisesRegex(ValidationError, 'Frozen packet'):
            read_html_return(self.returned(), packet, self.reviewer)

    def test_embedded_packet_types_are_hash_checked_not_python_numeric_equal(self):
        for version, changed in ((3, 3.0), (1, True), (4, 4.0)):
            packet = copy.deepcopy(self.packet); packet['version'] = version
            packet['packet_hash'] = sha(encode({k:v for k,v in packet.items() if k != 'packet_hash'}))
            response = _html_response(packet); response['reviewer'] = self.reviewer
            data = render_html_packet(packet, response); shape = _HTMLReviewShape(data.decode())
            embedded = copy.deepcopy(packet); embedded['version'] = changed
            self.assertEqual(embedded, packet)  # Python equality is insufficient for issued evidence.
            forged = data.decode().replace(shape.payloads['workbench-sme-packet'], encode(embedded).decode()).encode()
            with self.subTest(version=version, changed_type=type(changed).__name__), self.assertRaisesRegex(ValidationError, 'Original HTML packet'):
                read_html_return(forged, packet, self.reviewer)
            embedded['packet_hash'] = sha(encode({k:v for k,v in embedded.items() if k != 'packet_hash'}))
            forged = data.decode().replace(shape.payloads['workbench-sme-packet'], encode(embedded).decode()).encode()
            with self.assertRaisesRegex(ValidationError, 'Original HTML packet'):
                read_html_return(forged, packet, self.reviewer)

    def test_missing_duplicate_unknown_items_and_extra_fields_are_rejected(self):
        changes = [lambda r:r['items'].pop(), lambda r:r['items'].append(copy.deepcopy(r['items'][0])),
                   lambda r:r['items'][0].update(id='UNKNOWN'), lambda r:r['items'][1].update(id=r['items'][0]['id']),
                   lambda r:r['items'][0].update(reviewer='Unexpected actor'), lambda r:r.update(extra=True)]
        for index, change in enumerate(changes):
            response = self.response(); change(response)
            with self.subTest(index=index), self.assertRaises(ValidationError):
                read_html_return(self.returned(response), self.packet, self.reviewer)

    def test_answer_type_enum_commentary_and_explicit_attribution_bounds(self):
        for key, value in [('answer',True), ('answer','Approved'), ('answer',' Yes '),
                           ('correction',12), ('correction','x'*4001)]:
            response = self.response(); response['items'][0][key] = value
            with self.subTest(key=key,value_type=type(value).__name__), self.assertRaises(ValidationError):
                read_html_return(self.returned(response), self.packet, self.reviewer)
        for actor in ('', 'Other reviewer', 12, 'x'*161):
            response = self.response(); response['reviewer'] = actor
            with self.subTest(actor=actor), self.assertRaises(ValidationError):
                read_html_return(self.returned(response), self.packet, self.reviewer)
        for actor in ('', 'x'*161, None):
            with self.subTest(import_actor=actor), self.assertRaises(ValidationError):
                read_html_return(self.returned(), self.packet, actor)

    def test_active_html_shell_script_style_and_duplicate_payload_mutations_fail_closed(self):
        text = self.returned().decode()
        mutants = [text.replace('node.textContent = text', 'node.innerHTML = text'),
                   text.replace('background:#f5f4f1', 'background:red'),
                   text.replace('<main id="sme-app"></main>', '<main id="sme-app"><p>Altered questions</p></main>'),
                   text.replace('</body>', '<script src="https://evil.invalid/x"></script></body>'),
                   text.replace('id="sme-app"', 'id="sme-app" onclick="alert(1)"'),
                   text.replace('<head>', '<head><!--altered-->'),
                   text.replace('</body>', '<script type="application/json" id="workbench-sme-return">{}</script></body>'),
                   text.replace('id="sme-app"', 'id="sme-app" id="duplicate"')]
        for index, changed in enumerate(mutants):
            with self.subTest(index=index), self.assertRaises(ValidationError):
                read_html_return(changed.encode(), self.packet, self.reviewer)
        shape = _HTMLReviewShape(text)
        payload = shape.payloads['workbench-sme-return']
        duplicate_key = payload.replace('"version": 1', '"version": 1,"version": 1')
        with self.assertRaisesRegex(ValidationError, 'Invalid JSON document'):
            read_html_return(text.replace(payload,duplicate_key).encode(), self.packet, self.reviewer)

    def test_oversized_html_is_denied_before_parsing_or_consuming_a_return(self):
        data = HTML_RETURN_HEADER + b' ' * MAX_UPLOAD
        with self.assertRaisesRegex(ValidationError, 'oversized'):
            read_html_return(data, self.packet, self.reviewer)
        response = self.response(); response['items'][0]['correction'] = 'x' * MAX_UPLOAD
        with self.assertRaisesRegex(ValidationError, 'upload size'):
            self.returned(response)

    def test_invalid_utf8_truncated_or_wrong_type_returns_are_not_executed(self):
        for data in (self.returned()+b'\xff', self.returned()[:-50], b'<html>untrusted</html>', 'not bytes'):
            with self.subTest(kind=type(data).__name__), self.assertRaises(ValidationError):
                read_html_return(data, self.packet, self.reviewer)
        self.assertFalse(is_html_return(b'PK workbook'))
        self.assertFalse(is_html_return('not bytes'))


if __name__ == '__main__':
    unittest.main()
