"""SME technical questions group by behavior without granting conversion credit."""
import copy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from workbench.domain import ValidationError, encode, sha
from workbench.review import packet_document, export_packet, read_answers
from test_review_compaction import process_with_unsupported_lines, expand_spans


class ConstructReviewTests(unittest.TestCase):
    def process(self, records):
        process = process_with_unsupported_lines(0, version=3)
        analysis = process['analysis']; analysis['blockers'] = []
        for name, lines in records.items():
            path = name + '.cbl'; text = '\n'.join(lines)
            program = copy.deepcopy(analysis['programs']['ELIGIBLE'])
            program.update(name=name, path=path, source_text=text, source_hash=sha(text))
            analysis['programs'][name] = program
            analysis['assets'].append(dict(program))
            for line in range(1, len(lines) + 1):
                analysis['blockers'].append({'kind':'unsupported_source', 'path':path,
                                            'lines':[line], 'message':'Adapter remains unverified'})
        return process

    def test_construct_catalog_covers_nineteen_families_and_source_keywords(self):
        from workbench.review import CONSTRUCT_ASSUMPTIONS, _classify_construct
        samples = {'iteration':'PERFORM VARYING I FROM 1 BY 1 UNTIL I > 9',
                   'embedded_sql':'EXEC SQL SELECT X FROM A END-EXEC',
                   'cics':'EXEC CICS SEND MAP("SCREEN") END-EXEC',
                   'file_io':'REWRITE ACCOUNT-RECORD INVALID KEY CONTINUE',
                   'sorting':'SORT WORK-FILE ON ASCENDING KEY ACCOUNT',
                   'arithmetic':'COMPUTE AMOUNT ROUNDED = PRICE * RATE',
                   'conditional':'EVALUATE PRODUCT', 'assignment':'MOVE SPACES TO OUTPUT',
                   'calling':'CALL "OTHER" USING INPUT', 'control_flow':'GO TO CLEANUP',
                   'data_layout':'05 AMOUNT PIC S9(7)V99 COMP-3.',
                   'string_handling':'UNSTRING FULL-NAME DELIMITED BY SPACE INTO FIRST LAST',
                   'table_handling':'SEARCH ALL LOOKUP WHEN ID(I) = VALUE',
                   'jcl_execution':'//RUN EXEC PGM=APP,COND=(4,LT)',
                   'jcl_data':'//INPUT DD DSN=APP.INPUT,DISP=SHR',
                   'scheduling':'JOB,JOB=APP,SCHID=002', 'messaging':'CALL "MQPUT" USING HANDLE',
                   'unclassified_source':'APPLICATION-SPECIFIC-VERB X'}
        self.assertEqual(len(CONSTRUCT_ASSUMPTIONS), 19)
        self.assertEqual(set(samples) | {'unlocated_source'}, set(CONSTRUCT_ASSUMPTIONS))
        for family, source in samples.items():
            with self.subTest(family=family):
                self.assertEqual(_classify_construct(source), family)

    def test_literals_comments_and_field_names_do_not_manufacture_behavior(self):
        from workbench.review import _classify_construct
        for source in ('DISPLAY "PERFORM EXEC SQL MQPUT"', "DISPLAY 'CICS''PERFORM'", 'DISPLAY X *> EXEC SQL'):
            self.assertEqual(_classify_construct(source), 'file_io')
        for source in ('000100*PERFORM UNTIL DONE', '*> EXEC CICS SEND', '// *not a valid card'):
            self.assertIsNone(_classify_construct(source))
        self.assertEqual(_classify_construct('000100 PERFORM WORK.'), 'iteration')
        self.assertEqual(_classify_construct('05 PERFORM-COUNT PIC 9(3).'), 'data_layout')
        self.assertEqual(_classify_construct('MOVE SQL-RESULT TO CICS-FLAG'), 'assignment')

    def test_v3_compacts_across_programs_even_for_a_small_packet(self):
        process = self.process({'FIRST':['PERFORM WORK'], 'SECOND':['PERFORM OTHER']})
        original = copy.deepcopy(process['analysis']); packet = packet_document(process)
        gaps = [i for i in packet['items'] if i['kind'] == 'technical_gap']
        self.assertEqual(packet['version'], 3); self.assertEqual(len(gaps), 1)
        self.assertTrue(gaps[0]['question'].startswith('ASSUMPTION - iteration:'))
        self.assertIn('unverified obligation', gaps[0]['question'])
        self.assertIn('Yes confirms', gaps[0]['question'])
        context = json.loads(packet['context']); group = context['technical_gaps'][gaps[0]['id']]
        self.assertEqual(group['programs'], ['FIRST', 'SECOND'])
        self.assertEqual(expand_spans(group['blocker_index_spans']), [0, 1])
        self.assertEqual(group['source_line_spans'], {'FIRST.cbl':[[1,1]], 'SECOND.cbl':[[1,1]]})
        self.assertEqual(context['source_analysis']['sha256'], sha(encode(original)))
        self.assertEqual(process['analysis'], original)

    def test_one_blocker_with_multiple_behaviors_keeps_each_obligation(self):
        process = self.process({'FIRST':['IF PRODUCT = 1', 'PERFORM WORK', 'MOVE 0 TO RESULT']})
        process['analysis']['blockers'] = [{'kind':'unsupported_source', 'path':'FIRST.cbl',
                                         'lines':[1,2,3], 'message':'Unsupported nested source'}]
        context = json.loads(packet_document(process)['context'])
        self.assertEqual({g['construct'] for g in context['technical_gaps'].values()},
                         {'conditional','iteration','assignment'})
        self.assertTrue(all(g['blocker_count'] == 1 for g in context['technical_gaps'].values()))
        self.assertEqual(context['source_analysis']['raw_blocker_count'], 1)

    def test_multiple_statements_on_one_physical_line_keep_all_families(self):
        process=self.process({'FIRST':['PERFORM WORK MOVE 1 TO RESULT',
            'EXEC SQL SELECT ID FROM T END-EXEC PERFORM WORK',
            'IF READY = 1 PERFORM WORK ELSE MOVE 0 TO RESULT END-IF']})
        groups=json.loads(packet_document(process)['context'])['technical_gaps'].values()
        by_family={g['construct']:expand_spans(g['blocker_index_spans']) for g in groups}
        self.assertEqual(by_family,{'iteration':[0,1,2], 'assignment':[0,2], 'embedded_sql':[1], 'conditional':[2]})

    def test_sql_comments_and_continued_quotes_cannot_close_or_open_exec_blocks(self):
        process=self.process({'FIRST':['EXEC SQL', '-- END-EXEC is a comment',
            'SELECT "END-EXEC"', '/* END-EXEC */ FROM A', 'END-EXEC', 'PERFORM WORK',
            'DISPLAY "text continued', 'EXEC SQL END-EXEC"', 'MOVE 1 TO RESULT']})
        groups=json.loads(packet_document(process)['context'])['technical_gaps'].values()
        by_family={g['construct']:expand_spans(g['blocker_index_spans']) for g in groups}
        self.assertEqual(by_family['embedded_sql'],[0,2,3,4])
        self.assertEqual(by_family['iteration'],[5])
        self.assertEqual(by_family['assignment'],[8])
        self.assertNotIn('data_layout',by_family)

    def test_jcl_prefix_inside_continued_literals_and_sql_comments_is_not_a_card(self):
        for lines in ([ '000100 DISPLAY "PREFIX', '000200-    "//INPUT DD DSN=APP.DATA,DISP=SHR"', '000300 PERFORM WORK'],
                      ['DISPLAY "PREFIX', '//INPUT DD DSN=APP.DATA,DISP=SHR"', 'PERFORM WORK'],
                      ['EXEC SQL /* comment', '//INPUT DD DSN=APP.DATA,DISP=SHR */ SELECT X', 'END-EXEC', 'PERFORM WORK']):
            with self.subTest(lines=lines):
                groups=json.loads(packet_document(self.process({'FIRST':lines}))['context'])['technical_gaps'].values()
                by_family={g['construct']:expand_spans(g['blocker_index_spans']) for g in groups}
                self.assertNotIn('jcl_data',by_family)
                self.assertEqual(by_family['iteration'],[len(lines)-1])

    def test_split_exec_language_does_not_leave_a_provisional_unknown_question(self):
        groups=json.loads(packet_document(self.process({'FIRST':['EXEC','SQL SELECT X FROM A','END-EXEC']}))['context'])['technical_gaps'].values()
        self.assertEqual({g['construct']:expand_spans(g['blocker_index_spans']) for g in groups},
                         {'embedded_sql':[0,1,2]})

    def test_identifier_boundaries_and_quoted_call_names_do_not_create_extra_families(self):
        process=self.process({'FIRST':['MOVE PERFORM-COUNT TO SQL-RESULT','MOVE RETURN-CODE TO X',
            'CALL "EXEC SQL"','DISPLAY "END-EXEC"']})
        groups=json.loads(packet_document(process)['context'])['technical_gaps'].values()
        by_family={g['construct']:expand_spans(g['blocker_index_spans']) for g in groups}
        self.assertEqual(by_family,{'assignment':[0,1], 'calling':[2], 'file_io':[3]})

    def test_each_family_points_only_to_its_evidenced_program_and_lines(self):
        process=self.process({'FIRST':['IF READY = 1'], 'SECOND':['PERFORM WORK']})
        process['analysis']['blockers']=[{'kind':'unsupported_source',
            'source_refs':[{'path':'FIRST.cbl','line':1},{'path':'SECOND.cbl','line':1}], 'message':'Cross-source blocker'}]
        groups={g['construct']:g for g in json.loads(packet_document(process)['context'])['technical_gaps'].values()}
        self.assertEqual(groups['conditional']['programs'],['FIRST'])
        self.assertEqual(groups['iteration']['programs'],['SECOND'])
        self.assertEqual(groups['conditional']['source_line_spans'],{'FIRST.cbl':[[1,1]]})
        self.assertEqual(groups['iteration']['source_line_spans'],{'SECOND.cbl':[[1,1]]})
        self.assertTrue(all(g['blocker_index_spans']==[[0,0]] for g in groups.values()))

    def test_missing_source_summary_keeps_actual_program_names(self):
        from workbench.source import analyze_sources
        from test_source import COBOL
        process=process_with_unsupported_lines(0,version=3)
        process['jobs'][0]['steps'][0]['program']='MISSING1'
        second=copy.deepcopy(process['jobs'][0]['steps'][0]);second.update(name='S2',program='MISSING2',order=2)
        process['jobs'][0]['steps'].append(second)
        process['analysis']=analyze_sources({'ELIGIBLE.cbl':COBOL},process)
        blockers=[b for b in process['analysis']['blockers'] if b['kind']=='missing_source']
        self.assertEqual({b['program'] for b in blockers},{'MISSING1','MISSING2'})
        packet=packet_document(process)
        group=next(g for g in json.loads(packet['context'])['technical_gaps'].values() if g['construct']=='unlocated_source')
        self.assertEqual(group['programs'],['MISSING1','MISSING2'])
        self.assertIn('(2 programs total)',next(i['question'] for i in packet['items'] if i['kind']=='technical_gap'))
        for blocker in blockers:blocker.pop('program')
        legacy=packet_document(process)
        legacy_group=next(g for g in json.loads(legacy['context'])['technical_gaps'].values() if g['construct']=='unlocated_source')
        self.assertEqual(legacy_group['programs'],['MISSING1','MISSING2'])

    def test_multiline_sql_and_cics_bodies_remain_in_their_construct(self):
        process = self.process({'FIRST':['EXEC SQL', 'SELECT X', 'FROM A', 'END-EXEC',
                                        'EXEC CICS', 'SEND MAP("SCREEN")', 'END-EXEC']})
        groups = json.loads(packet_document(process)['context'])['technical_gaps'].values()
        self.assertEqual({g['construct'] for g in groups}, {'embedded_sql', 'cics'})
        self.assertEqual({g['construct']:expand_spans(g['blocker_index_spans']) for g in groups},
                         {'embedded_sql':[0,1,2,3], 'cics':[4,5,6]})

    def test_missing_or_invalid_source_evidence_is_an_unlocated_obligation(self):
        process = self.process({'FIRST':['PERFORM WORK']})
        process['analysis']['blockers'].extend([
            {'kind':'missing_source', 'program':'MISSING', 'message':'Source unavailable'},
            {'kind':'unsupported_source', 'path':'FIRST.cbl', 'lines':[900], 'message':'Not evidenced'},
            {'kind':'unsupported_source', 'path':'NOFILE.cbl', 'lines':[1], 'message':'Missing export'}])
        groups = json.loads(packet_document(process)['context'])['technical_gaps'].values()
        missing = next(g for g in groups if g['construct'] == 'unlocated_source')
        self.assertEqual(expand_spans(missing['blocker_index_spans']), [1,2,3])
        self.assertIn('MISSING', missing['programs'])
        self.assertIn('FIRST.cbl', missing['source_line_spans'])

    def test_located_opaque_source_is_not_misrepresented_as_missing(self):
        process = self.process({'FIRST':['APPLICATION-SPECIFIC-VERB X']})
        gaps = json.loads(packet_document(process)['context'])['technical_gaps'].values()
        self.assertEqual([g['construct'] for g in gaps], ['unclassified_source'])

    def test_jcl_cards_are_grounded_to_the_exact_source_line(self):
        from workbench.source import analyze_sources
        from test_source import COBOL
        process=process_with_unsupported_lines(0,version=3)
        process['analysis']=analyze_sources({'ELIGIBLE.cbl':COBOL,'JOB.jcl':
            '//NIGHT JOB\n//S1 EXEC PGM=ELIGIBLE,PARM=RUN\n//IN DD DSN=APP.INPUT,DISP=SHR'},process)
        blockers=[b for b in process['analysis']['blockers'] if b['kind']=='unsupported_jcl']
        self.assertEqual([b['lines'] for b in blockers],[[2],[3]])
        groups=json.loads(packet_document(process)['context'])['technical_gaps'].values()
        self.assertEqual({g['construct'] for g in groups},{'jcl_execution','jcl_data'})

    def test_real_ca7_cics_and_dclgen_classifications_route_to_their_families(self):
        from workbench.mainframe import load_knowledge
        from workbench.source import analyze_sources
        from test_source import COBOL
        process=process_with_unsupported_lines(0,version=3)
        with tempfile.TemporaryDirectory() as folder:
            process['mainframe_knowledge']=load_knowledge(Path(folder))
            process['analysis']=analyze_sources({'ELIGIBLE.cbl':COBOL,
                'run.ca7':'JOB,JOB=NIGHT,SCHID=001',
                'resources.csd':'DEFINE TRANSACTION(ABCD) GROUP(APP) PROGRAM(ELIGIBLE)',
                'account.dclgen':'EXEC SQL DECLARE APP.ACCOUNT TABLE (ID INTEGER) END-EXEC.'},process)
        classifications=process['analysis']['classifications']
        self.assertEqual(classifications['run.ca7']['kind'],'scheduler_definition')
        self.assertEqual(classifications['resources.csd']['kind'],'cics_definition')
        self.assertEqual(classifications['account.dclgen']['kind'],'dclgen')
        groups=json.loads(packet_document(process)['context'])['technical_gaps'].values()
        paths={g['construct']:g['paths'] for g in groups}
        self.assertEqual(paths['scheduling'],['run.ca7'])
        self.assertEqual(paths['cics'],['resources.csd'])
        self.assertEqual(paths['data_layout'],['account.dclgen'])
        self.assertTrue(all(b['kind']=='unsupported_source' for b in process['analysis']['blockers']))

    def test_source_references_to_multiple_files_keep_all_object_identities(self):
        process=self.process({'FIRST':['PERFORM WORK'],'SECOND':['PERFORM OTHER']})
        process['analysis']['blockers']=[{'kind':'unsupported_source','path':'FIRST.cbl',
            'source_refs':[{'path':'FIRST.cbl','line':1},{'path':'SECOND.cbl','line':1}], 'message':'Multi-object obligation'}]
        group=next(iter(json.loads(packet_document(process)['context'])['technical_gaps'].values()))
        self.assertEqual(group['programs'],['FIRST','SECOND'])
        self.assertEqual(group['paths'],['FIRST.cbl','SECOND.cbl'])

    def test_program_list_truncates_at_name_boundaries_and_retains_every_name(self):
        records = {'PROGRAM_%03d_LONG_NAME'%i:['PERFORM WORK'] for i in range(80)}
        packet = packet_document(self.process(records))
        item = next(i for i in packet['items'] if i['kind'] == 'technical_gap')
        names = json.loads(packet['context'])['technical_gaps'][item['id']]['programs']
        self.assertEqual(names, sorted(records))
        shown = [name for name in names if name in item['question']]
        self.assertGreater(len(shown), 0); self.assertLess(len(shown), 80)
        self.assertIn('(+%d more; 80 programs total)'%(80-len(shown)), item['question'])
        self.assertNotIn('display_question', item)

    def test_all_individual_business_rules_and_nontechnical_questions_are_preserved(self):
        process = self.process({'FIRST':['PERFORM WORK']}); process['llm']={'analysis':{'questions':['Which business date applies?']}}
        process['analysis']['blockers'].append({'kind':'unresolved_condition','message':'Unknown job condition'})
        packet = packet_document(process)
        self.assertEqual([i['id'] for i in packet['items'] if i['kind']=='business_rule'],
                         [r['id'] for r in process['analysis']['rules']])
        self.assertEqual(next(i for i in packet['items'] if i['kind']=='unresolved_item')['id'], 'B_001')
        self.assertEqual(next(i for i in packet['items'] if i['kind']=='provider_suggestion')['id'], 'LLM_000')

    def test_same_family_from_thousands_of_files_keeps_full_index_and_hash(self):
        process = self.process({'FILE%04d'%i:['PERFORM WORK'] for i in range(2100)})
        packet = packet_document(process)
        group = next(iter(json.loads(packet['context'])['technical_gaps'].values()))
        self.assertEqual(len(group['programs']), 2100)
        self.assertEqual(expand_spans(group['blocker_index_spans']), list(range(2100)))
        self.assertEqual(group['blockers_sha256'], sha(encode(process['analysis']['blockers'])))
        self.assertLess(len(packet['items']), 20)

    def test_v1_and_v2_packet_fingerprints_stay_unchanged(self):
        self.assertEqual(packet_document(process_with_unsupported_lines(0, version=None))['packet_hash'],
                         'b61caeb11f1910c14fac6cc3ee5689b0975d22dbce4327e68df0f6d7db858c89')
        self.assertEqual(packet_document(process_with_unsupported_lines(2001))['packet_hash'],
                         '295cb551c750dd30d603b2b52b48edc99f07fc5a6b15aec0af241a4354189937')

    def test_v3_packet_replays_and_full_context_is_checked_on_return(self):
        process = self.process({'FIRST':['PERFORM WORK'], 'SECOND':['PERFORM OTHER']})
        with tempfile.TemporaryDirectory() as folder:
            packet = export_packet(process, Path(folder)/'review')
            from workbench.domain import decode
            self.assertEqual(packet_document(decode(encode(process))), packet)
            book = load_workbook(Path(folder)/'review/sme-checklist.xlsx')
            for row in book['Checklist'].iter_rows(min_row=2): row[4].value = 'Yes'
            output = BytesIO(); book.save(output)
            self.assertEqual(len(read_answers(output.getvalue(), packet, 'Fixture reviewer')['items']), len(packet['items']))
            book['Context']['A2'] = 'Changed construct evidence'
            altered = BytesIO(); book.save(altered); book.close()
            with self.assertRaisesRegex(ValidationError, 'Frozen Context changed'):
                read_answers(altered.getvalue(), packet, 'Fixture reviewer')

    def test_new_intake_uses_v3_and_yes_answers_cannot_clear_the_adapter_gap(self):
        from workbench.coordinator import Coordinator
        from workbench.domain import decode
        from test_source import COBOL
        from test_workflow import MANIFEST
        with tempfile.TemporaryDirectory() as folder:
            coordinator = Coordinator(Path(folder)); self.addCleanup(coordinator.close)
            process = coordinator.create(MANIFEST, {'ELIGIBLE.cbl':COBOL.replace('  GOBACK.', '  PERFORM WORK.\n  GOBACK.')})
            self.assertEqual(process['sme_packet_version'], 3)
            coordinator.start('process-a'); coordinator.advance('process-a')
            packet = decode(coordinator.artifact('process-a', 'review/packet.json').read_bytes())
            self.assertEqual(packet['version'], 3)
            before = coordinator.artifact('process-a', 'analysis/source-analysis.json').read_bytes()
            book = load_workbook(coordinator.artifact('process-a', 'review/sme-checklist.xlsx'))
            for row in book['Checklist'].iter_rows(min_row=2): row[4].value = 'Yes'
            returned = BytesIO(); book.save(returned); book.close()
            coordinator.import_answers('process-a', returned.getvalue(), 'Fixture reviewer'); coordinator.advance('process-a')
            state = coordinator.ledger.get('process-a')
            self.assertTrue(any(b['kind']=='unsupported_source' for b in state['blockers']))
            self.assertEqual(state['program_versions'], {})
            self.assertEqual(coordinator.artifact('process-a', 'analysis/source-analysis.json').read_bytes(), before)
