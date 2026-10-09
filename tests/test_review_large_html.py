"""Versioned portable review accepts large returns without relaxing ZIP gates."""
import copy
from io import BytesIO
from pathlib import Path
import random
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from workbench.domain import MAX_UPLOAD, ValidationError, encode, sha
from workbench.intake import parse_manifest
from workbench.review import (HTML_RETURN_HEADER, MAX_HTML_REVIEW_BYTES,
                              _html_response, export_packet, read_answers,
                              read_html_return, render_html_packet,
                              review_return_limit)
from workbench.source import analyze_sources
from test_source import COBOL
from test_workflow import MANIFEST


class LargeHTMLReviewTests(unittest.TestCase):
    def process(self):
        process=parse_manifest(MANIFEST)
        process['analysis']=analyze_sources({'ELIGIBLE.cbl':COBOL},process)
        process['review_html_contract_version']=2
        return process

    def packet(self,context='Synthetic',count=1):
        packet={'version':3,'html_contract_version':2,'process_id':'sample-a','source_snapshot':'a'*64,
                'items':[{'id':f'SYN_{i}','question':'Synthetic requirement','evidence':'Synthetic source witness','kind':'business'} for i in range(count)],
                'context':context}
        packet['packet_hash']=sha(encode(packet));return packet

    def returned(self,packet):
        response=_html_response(packet);response['reviewer']='Synthetic reviewer'
        for item in response['items']:item['answer']='Yes'
        return response

    def test_valid_review_above_8_mib_keeps_every_question_answer_and_comment(self):
        packet=self.packet('x'*(2*1024*1024),2000);response=self.returned(packet)
        for index,item in enumerate(response['items']):
            item['answer']=('Yes','No','Not sure','')[index%4]
            item['correction']=f'Synthetic {index}: '+('x'*3500)
        data=render_html_packet(packet,response)
        self.assertGreater(len(data),MAX_UPLOAD);self.assertLess(len(data),MAX_HTML_REVIEW_BYTES)
        returned=read_html_return(data,packet,'Synthetic reviewer')
        self.assertEqual(len(returned['items']),2000)
        for item in response['items']:
            actual=returned['items'][item['id']]
            self.assertEqual(actual['answer'],item['answer'] or 'Unanswered')
            self.assertEqual(actual['correction'],item['correction'])
        self.assertEqual(returned['return_hash'],sha(data))
        self.assertIn(b'134217728',data);self.assertNotIn(b'8 MiB import limit',data)

    def test_64_distinct_random_returns_match_executed_workbook_import(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp)/'review';packet=export_packet(self.process(),folder)
            self.assertEqual(packet['html_contract_version'],2)
            seed=random.SystemRandom().randrange(2**32);rng=random.Random(seed);states=set()
            for state in range(64):
                response=self.returned(packet)
                for item in response['items']:
                    item['answer']=rng.choice(['Yes','No','Not sure',''])
                    item['correction']=f'Synthetic {state}; {rng.randrange(10**12)}'
                data=render_html_packet(packet,response);states.add(sha(data))
                actual=read_html_return(data,packet,'Synthetic reviewer')
                book=load_workbook(folder/'sme-checklist.xlsx')
                for row,item in zip(book['Checklist'].iter_rows(min_row=2),response['items']):
                    row[4].value=item['answer'];row[5].value=item['correction']
                buffer=BytesIO();book.save(buffer);book.close()
                expected=read_answers(buffer.getvalue(),packet,'Synthetic reviewer')
                self.assertEqual(actual['items'],expected['items'],f'seed={seed},state={state}')
            self.assertEqual(len(states),64)

    def test_already_issued_old_packet_remains_identical_when_new_default_changes(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp)/'review';process=self.process();process.pop('review_html_contract_version')
            original=export_packet(process,folder);old_html=(folder/'sme-checklist.html').read_bytes()
            self.assertNotIn('html_contract_version',original);self.assertIn(b'8388608',old_html)
            process['review_html_contract_version']=2
            self.assertEqual(export_packet(process,folder),original)
            self.assertEqual((folder/'sme-checklist.html').read_bytes(),old_html)
            response=self.returned(original)
            self.assertEqual(read_html_return(render_html_packet(original,response),original,'Synthetic reviewer')['packet_hash'],original['packet_hash'])

    def test_format_limit_does_not_expand_workbook_archives(self):
        self.assertEqual(review_return_limit(HTML_RETURN_HEADER),128*1024*1024)
        self.assertEqual(review_return_limit(b'PK workbook'),MAX_UPLOAD)
        self.assertEqual(review_return_limit(b'<html>untrusted'),MAX_UPLOAD)

    def test_excess_size_and_shell_mutations_fail_before_quota_consumption(self):
        packet=self.packet();response=self.returned(packet);data=render_html_packet(packet,response)
        with patch('workbench.review.MAX_HTML_REVIEW_BYTES',len(data)-1):
            with self.assertRaisesRegex(ValidationError,'oversized'):
                read_html_return(data,packet,'Synthetic reviewer')
            with self.assertRaisesRegex(ValidationError,'upload size'):
                render_html_packet(packet,response)
        for mutated in (data.replace(b'134217728',b'8388608'),data.replace(b'node.textContent = text',b'node.innerHTML = text'),data.replace(b'</body>',b'<script src="https://untrusted.invalid/x"></script></body>')):
            with self.assertRaisesRegex(ValidationError,'shell'):
                read_html_return(mutated,packet,'Synthetic reviewer')
        for changed in (True,1,2.0,3):
            bad=copy.deepcopy(packet);bad['html_contract_version']=changed
            bad['packet_hash']=sha(encode({k:v for k,v in bad.items() if k!='packet_hash'}))
            with self.assertRaisesRegex(ValidationError,'Unsupported frozen HTML'):
                render_html_packet(bad,response)
        response['items'][0]['correction']='x'*4001
        with self.assertRaisesRegex(ValidationError,'commentary'):
            read_html_return(render_html_packet(packet,response),packet,'Synthetic reviewer')


if __name__=='__main__':unittest.main()
