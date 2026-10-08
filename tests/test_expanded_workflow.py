"""R701–R800: new stage, transaction and provenance interactions.

All workbook responses are explicit fictional test fixtures, never operational SME answers.
"""
import copy
import json
import shutil
import sqlite3
import tempfile
import threading
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from workbench.coordinator import Coordinator, StageInterrupted
from workbench.domain import ValidationError, encode, decode, sha
from workbench.knowledge import update_knowledge
from workbench.ledger import Ledger
from workbench.review import export_packet, packet_document, read_answers
from test_source import COBOL
from test_workflow import MANIFEST

REVIEWER = "Fictional expanded workflow fixture reviewer"
_TEMPLATES = None


def templates():
    global _TEMPLATES
    if _TEMPLATES is None:
        temporary = tempfile.TemporaryDirectory()
        root = Path(temporary.name)
        c = Coordinator(root / 'ready')
        c.create(MANIFEST, {'ELIGIBLE.cbl': COBOL})
        c.close()
        shutil.copytree(root / 'ready', root / 'waiting')
        c = Coordinator(root / 'waiting')
        c.start('process-a'); c.advance('process-a'); c.close()
        shutil.copytree(root / 'waiting', root / 'accepted')
        c = Coordinator(root / 'accepted')
        book = load_workbook(BytesIO(c.artifact('process-a', 'review/sme-checklist.xlsx').read_bytes()))
        for row in book['Checklist'].iter_rows(min_row=2):
            row[4].value = 'Yes'
        output = BytesIO(); book.save(output); book.close()
        c.import_answers('process-a', output.getvalue(), REVIEWER); c.close()
        _TEMPLATES = temporary, root
    return _TEMPLATES[1]


class ExpandedFixture(unittest.TestCase):
    state = 'waiting'

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name) / 'work'
        shutil.copytree(templates() / self.state, self.root)
        self.c = Coordinator(self.root)
        self.addCleanup(lambda: self.c.close())

    def doc(self): return self.c.ledger.get('process-a')
    def save(self, status=None, **updates):
        doc = self.doc(); doc.update(updates); return self.c.ledger.save(doc, status)
    def restart(self):
        self.c.close(); self.c = Coordinator(self.root); return self.doc()
    def raw(self): return (self.c.process_root('process-a') / 'review/sme-checklist.xlsx').read_bytes()
    def packet(self): return decode((self.c.process_root('process-a') / 'review/packet.json').read_bytes())
    def edit(self, fn):
        book = load_workbook(BytesIO(self.raw()))
        try:
            fn(book); stream = BytesIO(); book.save(stream); return stream.getvalue()
        finally: book.close()
    def accept(self, edit=None):
        data = self.raw() if edit is None else self.edit(edit)
        return self.c.import_answers('process-a', data, REVIEWER)
    def record_count(self): return self.c.ledger.db.execute('SELECT COUNT(*) FROM knowledge').fetchone()[0]
    def events(self): return self.c.ledger.events('process-a')
    def failing_trigger(self, name, table, timing='BEFORE INSERT', condition=''):
        self.c.ledger.db.executescript(f"CREATE TRIGGER {name} {timing} ON {table} {condition} BEGIN SELECT RAISE(ABORT,'fixture rejection'); END;")
    def completion(self):
        return self.save('REPORTING', report_verified=True,
            report_hashes={'reports/report-0001/metrics.json': 'h'},
            artifact_hashes={**self.doc()['artifact_hashes'], 'reports/report-0001/metrics.json': 'h'})


class ExpandedStageTests(ExpandedFixture):

    def test_r701_interrupted_reporting_restarts_as_reporting(self):
        self.save('REPORTING', verification_finished=False)
        self.assertEqual(self.restart()['status'], 'QUEUED_REPORT')

    def test_r702_interrupted_verification_without_return_stays_waiting(self):
        self.save('VERIFYING')
        self.assertEqual(self.restart()['status'], 'WAITING_SME')

    def test_r703_accepted_return_restart_queues_verification(self):
        before = self.accept(); self.save('VERIFYING')
        after = self.restart()
        self.assertEqual(after['status'], 'QUEUED_VERIFY'); self.assertEqual(after['answers'], before['answers'])

    def test_r704_finished_verification_restart_skips_target_reexecution(self):
        self.accept(); self.save('VERIFYING', verification_finished=True)
        self.assertEqual(self.restart()['status'], 'QUEUED_REPORT')

    def test_r705_cancelled_interrupted_analysis_restarts_report(self):
        self.save('ANALYZING', cancel_requested=True, blockers=[{'kind':'cancelled','message':'fixture'}])
        after = self.restart(); self.assertEqual(after['status'], 'QUEUED_REPORT'); self.assertTrue(after['cancel_requested'])

    def test_r706_pause_waiting_resume_does_not_issue_new_packet(self):
        before = self.doc(); self.c.control('process-a','pause'); after = self.c.control('process-a','resume')
        self.assertEqual(after['status'],'WAITING_SME'); self.assertEqual(after['packet_hash'],before['packet_hash'])

    def test_r707_resume_failed_analysis_with_packet_waits_for_human(self):
        self.save('FAILED', resume_status='QUEUED_ANALYSIS')
        after = self.c.control('process-a','resume'); self.assertEqual(after['status'],'WAITING_SME')

    def test_r708_resume_failed_analysis_with_return_routes_verify(self):
        self.accept(); self.save('FAILED',resume_status='QUEUED_ANALYSIS')
        self.assertEqual(self.c.control('process-a','resume')['status'],'QUEUED_VERIFY')

    def test_r709_resume_finished_verification_routes_report(self):
        self.accept(); self.save('FAILED',resume_status='QUEUED_VERIFY',verification_finished=True)
        self.assertEqual(self.c.control('process-a','resume')['status'],'QUEUED_REPORT')

    def test_r710_resume_retry_counter_resets_only_failed_stage(self):
        self.save('FAILED',resume_status='QUEUED_VERIFY',stage_attempts={'QUEUED_VERIFY':3,'QUEUED_ANALYSIS':2})
        after=self.c.control('process-a','resume'); self.assertEqual(after['stage_attempts'],{'QUEUED_VERIFY':0,'QUEUED_ANALYSIS':2})

    def test_r711_resume_pause_preserves_attempt_budget(self):
        self.save('QUEUED_VERIFY',stage_attempts={'QUEUED_VERIFY':2}); self.c.control('process-a','pause')
        after=self.c.control('process-a','resume'); self.assertEqual(after['stage_attempts']['QUEUED_VERIFY'],2)

    def test_r712_future_retry_has_no_stage_or_event_side_effects(self):
        self.save('QUEUED_ANALYSIS',next_retry_at=10**20); before=self.doc(); events=self.events()
        with patch.object(self.c,'analyze') as execute: self.c.advance('process-a'); execute.assert_not_called()
        self.assertEqual(self.doc(),before); self.assertEqual(self.events(),events)

    def test_r713_active_stage_cannot_reenter(self):
        self.save('QUEUED_ANALYSIS'); before=self.doc(); self.c.active.add('process-a')
        try:
            with patch.object(self.c,'analyze') as execute:
                self.c.advance('process-a'); execute.assert_not_called()
            self.assertEqual(self.doc(),before)
        finally: self.c.active.discard('process-a')


    def test_r714_validation_failure_is_not_transient_retry(self):
        self.save('QUEUED_ANALYSIS')
        with patch.object(self.c,'analyze',side_effect=ValidationError('fixture gate')):
            with self.assertRaises(ValidationError): self.c.advance('process-a')
        self.assertEqual(self.doc()['status'],'FAILED'); self.assertEqual(self.doc()['next_retry_at'],0)

    def test_r715_missing_evidence_is_not_transient_retry(self):
        self.save('QUEUED_VERIFY')
        with patch.object(self.c,'verify',side_effect=FileNotFoundError('fixture evidence')):
            with self.assertRaises(FileNotFoundError): self.c.advance('process-a')
        self.assertEqual(self.doc()['status'],'FAILED'); self.assertEqual(self.doc()['resume_status'],'QUEUED_VERIFY')

    def test_r716_report_failure_revokes_prior_inspection(self):
        self.save('QUEUED_REPORT',report_verified=True)
        with patch.object(self.c,'report',side_effect=ValidationError('fixture inspection')):
            with self.assertRaises(ValidationError): self.c.advance('process-a')
        self.assertEqual(self.doc()['status'],'REPORTING_FAILED'); self.assertFalse(self.doc()['report_verified'])

    def test_r717_failure_message_is_bounded_but_stage_is_retained(self):
        self.save('QUEUED_ANALYSIS')
        with patch.object(self.c,'analyze',side_effect=RuntimeError('x'*2000)):
            with self.assertRaises(RuntimeError): self.c.advance('process-a')
        f=self.doc()['last_failure']; self.assertEqual(len(f['message']),1000); self.assertEqual(f['stage'],'QUEUED_ANALYSIS')

    def test_r718_exception_clears_active_process_guard(self):
        self.save('QUEUED_ANALYSIS')
        with patch.object(self.c,'analyze',side_effect=RuntimeError('fixture')):
            with self.assertRaises(RuntimeError): self.c.advance('process-a')
        self.assertNotIn('process-a',self.c.active)

    def test_r719_pause_during_exception_preserves_operator_decision(self):
        self.save('QUEUED_ANALYSIS')
        def fail(doc):
            self.c.control('process-a','pause'); raise OSError('fixture')
        with patch.object(self.c,'analyze',side_effect=fail): self.c.advance('process-a')
        self.assertEqual(self.doc()['status'],'PAUSED'); self.assertFalse(any(b['kind']=='stage_failure' for b in self.doc()['blockers']))

    def test_r720_cancel_during_exception_does_not_requeue_analysis(self):
        self.save('QUEUED_ANALYSIS')
        def fail(doc):
            self.c.control('process-a','cancel'); raise OSError('fixture')
        with patch.object(self.c,'analyze',side_effect=fail): self.c.advance('process-a')
        self.assertEqual(self.doc()['status'],'QUEUED_REPORT'); self.assertTrue(self.doc()['cancel_requested'])

    def test_r721_control_event_failure_rolls_back_pause(self):
        before=self.doc(); self.failing_trigger('reject_pause','events',condition="WHEN NEW.stage='pause'")
        with self.assertRaises(sqlite3.IntegrityError): self.c.control('process-a','pause')
        self.assertEqual(self.doc(),before)

    def test_r722_control_save_failure_does_not_emit_phantom_event(self):
        before=self.events(); self.failing_trigger('reject_state','processes','BEFORE UPDATE',"WHEN NEW.status='PAUSED'")
        with self.assertRaises(sqlite3.IntegrityError): self.c.control('process-a','pause')
        self.assertEqual(self.events(),before); self.assertEqual(self.doc()['status'],'WAITING_SME')

    def test_r723_cancelled_report_checkpoint_allows_accountability(self):
        doc=self.c.control('process-a','cancel'); doc['_active_stage']='QUEUED_REPORT'
        self.c.checkpoint(doc,'REPORTING'); self.assertEqual(self.doc()['status'],'REPORTING'); self.assertTrue(self.doc()['cancel_requested'])

    def test_r724_cancelled_report_target_checkpoint_denies_execution(self):
        doc=self.c.control('process-a','cancel'); doc['_active_stage']='QUEUED_REPORT'
        with self.assertRaises(ValidationError): self.c.target_checkpoint(doc)
        self.assertTrue(self.doc()['cancel_requested'])

    def test_r725_nonpersisting_checkpoint_reads_durable_control_revision(self):
        doc=self.doc(); doc['_active_stage']='QUEUED_REPORT'; self.save(control_revision=12)
        with patch.object(self.c.ledger,'get',side_effect=AssertionError('full document read')): self.c.checkpoint(doc,persist=False)
        self.assertEqual(doc['control_revision'],12)


class ExpandedLedgerTests(ExpandedFixture):
    state = 'waiting'

    def test_r726_return_event_failure_rolls_back_quota(self):
        self.failing_trigger('reject_return_event','events',condition="WHEN NEW.stage='review'")
        with self.assertRaises(sqlite3.IntegrityError): self.accept()
        self.assertFalse(self.doc()['packet_imported']); self.assertIsNone(self.doc()['answers'])

    def test_r727_return_event_failure_keeps_replayable_bytes(self):
        raw=self.raw(); self.failing_trigger('reject_return_event','events',condition="WHEN NEW.stage='review'")
        with self.assertRaises(sqlite3.IntegrityError): self.accept()
        self.assertEqual((self.c.process_root('process-a')/'input/sme-return.xlsx').read_bytes(),raw)
        self.c.ledger.db.execute('DROP TRIGGER reject_return_event'); self.c.ledger.db.commit(); self.accept(); self.assertTrue(self.doc()['packet_imported'])

    def test_r728_return_commit_emits_one_audit_event(self):
        before=len(self.events()); self.accept()
        events=self.events()[before:]; self.assertEqual(len(events),1); self.assertEqual(events[0]['stage'],'review')

    def test_r729_stale_completion_cannot_discard_cancel_control(self):
        doc=self.completion(); latest=self.doc(); latest.update(control_revision=1,cancel_requested=True); latest['blockers'].append({'kind':'cancelled','message':'fixture'}); self.c.ledger.save(latest,'REPORTING')
        with self.assertRaises(ValidationError): self.c.ledger.complete_report(doc,{})
        self.assertTrue(self.doc()['cancel_requested']); self.assertEqual(self.c.ledger.history('process-a'),[])

    def test_r730_stale_completion_cannot_discard_pause_resume_revision(self):
        doc=self.completion(); self.save('REPORTING',control_revision=2)
        with self.assertRaises(ValidationError): self.c.ledger.complete_report(doc,{})
        self.assertEqual(self.doc()['control_revision'],2)

    def test_r731_completion_current_cancelled_candidate_is_blocked(self):
        doc=self.completion(); doc.update(cancel_requested=True,control_revision=1); doc['blockers']=[]; self.c.ledger.save(doc,'REPORTING')
        result=self.c.ledger.complete_report(doc,{}); self.assertEqual(result['status'],'COMPLETED_WITH_BLOCKERS'); self.assertTrue(result['cancel_requested'])

    def test_r732_completion_process_update_failure_rolls_back_snapshot(self):
        doc=self.completion(); self.failing_trigger('reject_complete','processes','BEFORE UPDATE',"WHEN NEW.status='COMPLETED'")
        with self.assertRaises(sqlite3.IntegrityError): self.c.ledger.complete_report(doc,{'new':1})
        self.assertEqual(self.c.ledger.history('process-a'),[]); self.assertEqual(self.doc()['status'],'REPORTING')

    def test_r733_snapshot_failure_prevents_terminal_event(self):
        doc=self.completion(); before=self.events(); self.failing_trigger('reject_snapshot','snapshots')
        with self.assertRaises(sqlite3.IntegrityError): self.c.ledger.complete_report(doc,{'new':1})
        self.assertEqual(self.events(),before); self.assertEqual(self.doc()['status'],'REPORTING')

    def test_r734_identical_metrics_across_processes_keep_separate_history(self):
        other=self.doc(); other['id']='process-b'; self.c.ledger.create(other)
        self.c.ledger.snapshot('process-a',{'n':1}); self.c.ledger.snapshot('process-b',{'n':1})
        self.assertEqual(len(self.c.ledger.history()),2)

    def test_r735_event_pagination_uses_global_cursor_without_crossprocess_leak(self):
        other=self.doc(); other['id']='process-b'; self.c.ledger.create(other)
        self.c.ledger.event('process-a','fixture','first'); cursor=self.events()[-1]['seq']; self.c.ledger.event('process-b','fixture','other'); self.c.ledger.event('process-a','fixture','last')
        self.assertEqual([e['message'] for e in self.c.ledger.events('process-a',cursor)],['last'])

    def test_r736_event_foreign_key_failure_does_not_poison_next_write(self):
        with self.assertRaises(sqlite3.IntegrityError): self.c.ledger.event('missing','fixture','bad')
        self.c.ledger.event('process-a','fixture','good'); self.assertEqual(self.events()[-1]['message'],'good')

    def test_r737_asset_membership_failure_rolls_back_global_insert(self):
        asset={'id':'novel','kind':'cobol','name':'N','source_hash':'n'}
        with self.assertRaises(sqlite3.IntegrityError): self.c.ledger.register_assets('missing',[asset])
        self.assertIsNone(self.c.ledger.db.execute("SELECT id FROM assets WHERE id='novel'").fetchone())

    def test_r738_shared_asset_multiple_memberships_do_not_duplicate_portfolio(self):
        other=self.doc(); other['id']='process-b'; self.c.ledger.create(other); assets=self.c.ledger.assets('process-a')
        self.c.ledger.register_assets('process-b',assets); self.assertEqual(len(self.c.ledger.assets()),len(assets)); self.assertEqual(len(self.c.ledger.assets('process-b')),len(assets))

    def test_r739_mutating_get_result_does_not_mutate_durable_answers(self):
        self.accept(); value=self.doc(); value['answers']['reviewer']='mutated'
        self.assertEqual(self.doc()['answers']['reviewer'],REVIEWER)

    def test_r740_two_ledger_handles_race_for_single_packet(self):
        other=Ledger(self.root); self.addCleanup(other.close)
        new=self.doc(); new['id']='process-b'; self.c.ledger.create(new)
        barrier=threading.Barrier(2); results=[]
        def issue(ledger,fingerprint):
            try:
                barrier.wait(timeout=5); ledger.issue_packet('process-b',fingerprint); results.append(('accepted',fingerprint))
            except ValidationError: results.append(('rejected',fingerprint))
        threads=[threading.Thread(target=issue,args=(ledger,fingerprint)) for ledger,fingerprint in [(self.c.ledger,'first'),(other,'second')]]
        for thread in threads: thread.start()
        for thread in threads: thread.join(timeout=5); self.assertFalse(thread.is_alive())
        self.assertEqual(sorted(r[0] for r in results),['accepted','rejected'])
        winner=next(r[1] for r in results if r[0]=='accepted'); self.assertEqual(self.c.ledger.get('process-b')['packet_hash'],winner)


    def test_r741_two_ledger_handles_race_for_single_return(self):
        other=Ledger(self.root); self.addCleanup(other.close); answers=read_answers(self.raw(),self.packet(),REVIEWER)
        barrier=threading.Barrier(2); results=[]; before=len(self.events())
        def consume(ledger):
            try:
                barrier.wait(timeout=5); ledger.consume_return('process-a',answers); results.append('accepted')
            except ValidationError: results.append('rejected')
        threads=[threading.Thread(target=consume,args=(ledger,)) for ledger in [self.c.ledger,other]]
        for thread in threads: thread.start()
        for thread in threads: thread.join(timeout=5); self.assertFalse(thread.is_alive())
        self.assertEqual(sorted(results),['accepted','rejected']); self.assertEqual(self.doc()['answers'],answers); self.assertEqual(len(self.events())-before,1)


    def test_r742_list_returns_demo_only_when_requested(self):
        other=self.doc(); other['id']='fixture-demo'; self.c.ledger.create(other,demo=True)
        self.assertEqual([p['id'] for p in self.c.ledger.list()],['process-a']); self.assertEqual(len(self.c.ledger.list(True)),2)

    def test_r743_controls_missing_process_fails_without_synthetic_state(self):
        with self.assertRaises(ValidationError): self.c.ledger.controls('missing')
        self.assertEqual(len(self.c.ledger.list()),1)

    def test_r744_completion_deduplicates_preexisting_matching_snapshot(self):
        doc=self.completion(); self.c.ledger.snapshot('process-a',{'n':7})
        self.c.ledger.complete_report(doc,{'n':7}); self.assertEqual(len(self.c.ledger.history('process-a')),1); self.assertEqual(self.events()[-1]['stage'],'complete')

    def test_r745_start_state_failure_rolls_back_authorization_event(self):
        self.c.close(); shutil.rmtree(self.root); shutil.copytree(templates()/'ready',self.root); self.c=Coordinator(self.root)
        before=self.events(); self.failing_trigger('reject_start','processes','BEFORE UPDATE',"WHEN NEW.status='QUEUED_ANALYSIS'")
        with self.assertRaises(sqlite3.IntegrityError): self.c.start('process-a')
        self.assertEqual(self.events(),before); self.assertNotIn('authorization',self.doc())


class ExpandedPacketTests(ExpandedFixture):
    state = 'waiting'

    def test_r746_provider_suggestions_and_background_share_one_packet(self):
        doc=self.doc(); doc['knowledge_context']={'text':'fixture','sha256':'a','status':'UNVERIFIED_INPUT'}; doc['llm']={'analysis':{'questions':['Question fixture'],'assumptions':['Assumption fixture']}}
        packet=packet_document(doc); ids=[i['id'] for i in packet['items']]; self.assertIn('G_KNOWLEDGE',ids); self.assertIn('LLM_001',ids); self.assertEqual(len(ids),len(set(ids)))

    def test_r747_long_rule_details_remain_complete_in_context(self):
        doc=self.doc(); text='Long rule '+('detail '*1000); doc['analysis']['rules'][0]['plain']=text
        packet=packet_document(doc); item=packet['items'][0]; context=json.loads(packet['context']); self.assertEqual(context['checklist_details'][item['id']]['question'],text); self.assertLess(len(item['display_question']),4000)

    def test_r748_long_evidence_reference_remains_bound_to_full_record(self):
        doc=self.doc(); refs=['fixture.cbl:'+str(i) for i in range(600)]; doc['analysis']['rules'][0]['source_refs']=refs
        packet=packet_document(doc); item=packet['items'][0]; self.assertEqual(json.loads(packet['context'])['checklist_details'][item['id']]['evidence'],'; '.join(refs))

    def test_r749_packet_input_is_not_mutated_by_display_shortening(self):
        doc=self.doc(); doc['analysis']['rules'][0]['plain']='x'*6000; before=copy.deepcopy(doc)
        packet_document(doc); self.assertEqual(doc,before)

    def test_r750_formula_looking_source_question_is_plain_text(self):
        doc=self.doc(); doc['analysis']['rules'][0]['plain']='=SUM(1,2) source wording'; folder=self.root/'.implementation/tmp/formula-packet'; packet=export_packet(doc,folder)
        book=load_workbook(folder/'sme-checklist.xlsx',data_only=False); self.assertEqual(book['Checklist']['B2'].data_type,'s'); book.close()
        answers=read_answers((folder/'sme-checklist.xlsx').read_bytes(),packet,REVIEWER); self.assertEqual(next(iter(answers['items'].values()))['question'],'=SUM(1,2) source wording')

    def test_r751_html_packet_escapes_source_markup(self):
        doc=self.doc(); doc['analysis']['rules'][0]['plain']='<script>fixture()</script>'; folder=self.root/'.implementation/tmp/html-packet'; export_packet(doc,folder)
        from workbench.review import _HTMLReviewShape
        html=(folder/'sme-checklist.html').read_text()
        self.assertNotIn('<script>fixture()</script>',html)
        shape=_HTMLReviewShape(html)
        embedded=decode(shape.payloads['workbench-sme-packet'])
        self.assertEqual(embedded['items'][0]['question'],'<script>fixture()</script>')
        self.assertIn(r'\u003cscript\u003e',shape.payloads['workbench-sme-packet'])
        self.assertIn('node.textContent = text',html)
        self.assertIn("connect-src 'none'",html)

    def test_r752_export_failure_leaves_no_partial_published_packet(self):
        folder=self.root/'.implementation/tmp/failure-packet'
        with patch('docx.document.Document.save',side_effect=OSError('fixture disk')):
            with self.assertRaises(OSError): export_packet(self.doc(),folder)
        self.assertFalse(folder.exists())

    def test_r753_export_failure_cleans_draft_directory(self):
        folder=self.root/'.implementation/tmp/failure-packet'
        with patch('docx.document.Document.save',side_effect=OSError('fixture disk')):
            with self.assertRaises(OSError): export_packet(self.doc(),folder)
        self.assertEqual(list((folder.parent/'.implementation/tmp').glob('review-draft-*')),[])

    def test_r754_packet_recovery_preserves_all_companion_bytes(self):
        folder=self.c.process_root('process-a')/'review'; before={p.name:p.read_bytes() for p in folder.iterdir()}
        export_packet(self.doc(),folder); self.assertEqual({p.name:p.read_bytes() for p in folder.iterdir()},before)

    def test_r755_changed_context_cannot_reissue_into_existing_directory(self):
        doc=self.doc(); doc['jobs']=copy.deepcopy(doc['jobs']); doc['jobs'][0]['name']='changed-job'
        with self.assertRaises(ValidationError): export_packet(doc,self.c.process_root('process-a')/'review')
        self.assertEqual(self.packet()['packet_hash'],self.doc()['packet_hash'])

    def test_r756_row_reviewer_overrides_only_own_attribution(self):
        data=self.edit(lambda b:setattr(b['Checklist']['G2'],'value','Fictional delegated reviewer'))
        answers=read_answers(data,self.packet(),REVIEWER); first=self.packet()['items'][0]['id']; second=self.packet()['items'][1]['id']
        self.assertEqual(answers['items'][first]['reviewer'],'Fictional delegated reviewer'); self.assertEqual(answers['items'][second]['reviewer'],REVIEWER)

    def test_r757_workbook_sheet_reordering_preserves_answer_binding(self):
        def edit(b): b._sheets=[b['Context'],b['Metadata'],b['Checklist']]
        answers=read_answers(self.edit(edit),self.packet(),REVIEWER); self.assertEqual(set(answers['items']),{i['id'] for i in self.packet()['items']})

    def test_r758_answer_style_edit_does_not_change_semantics(self):
        from openpyxl.styles import Font
        def edit(b): b['Checklist']['E2'].font=Font(bold=True,color='FF0000')
        answers=read_answers(self.edit(edit),self.packet(),REVIEWER); self.assertTrue(all(a['answer']=='Unanswered' for a in answers['items'].values()))

    def test_r759_return_hash_binds_formatting_changes(self):
        from openpyxl.styles import Font
        def edit(b): b['Checklist']['E2'].font=Font(italic=True)
        raw=self.edit(edit); answers=read_answers(raw,self.packet(),REVIEWER); self.assertEqual(answers['return_hash'],sha(raw)); self.assertNotEqual(sha(raw),sha(self.raw()))

    def test_r760_hidden_existing_answer_row_is_still_consumed(self):
        def edit(b): b['Checklist'].row_dimensions[2].hidden=True; b['Checklist']['E2']='Not sure'
        answers=read_answers(self.edit(edit),self.packet(),REVIEWER); self.assertEqual(answers['items'][self.packet()['items'][0]['id']]['answer'],'Not sure')

    def test_r761_long_checklist_question_roundtrips_full_semantics(self):
        doc=self.doc(); question='Complete explanation '+('condition detail '*700); doc['analysis']['rules'][0]['plain']=question
        folder=self.root/'.implementation/tmp/long-question'; packet=export_packet(doc,folder)
        answers=read_answers((folder/'sme-checklist.xlsx').read_bytes(),packet,REVIEWER)
        self.assertEqual(answers['items'][doc['analysis']['rules'][0]['id']]['question'],question)
        self.assertEqual(answers['items'][doc['analysis']['rules'][0]['id']]['answer'],'Unanswered')


    def test_r762_unicode_context_chunks_roundtrip_without_truncation(self):
        doc=self.doc(); background='Context \U0001f9ea ' * 4200; doc['knowledge_context']={'text':background,'sha256':sha(background),'status':'UNVERIFIED_INPUT'}
        folder=self.root/'.implementation/tmp/large-context'; packet=export_packet(doc,folder)
        book=load_workbook(folder/'sme-checklist.xlsx'); pieces=[r[0] for r in book['Context'].iter_rows(min_row=2,values_only=True)]; book.close()
        self.assertGreater(len(pieces),2); self.assertEqual(''.join(pieces),packet['context']); self.assertEqual(json.loads(''.join(pieces))['knowledge_input']['text'],background)
        answers=read_answers((folder/'sme-checklist.xlsx').read_bytes(),packet,REVIEWER); self.assertIn('G_KNOWLEDGE',answers['items'])


    def test_r763_word_companion_preserves_full_rule_and_context(self):
        from docx import Document
        doc=self.doc(); statement='Complete rule '+('fixture meaning '*400); doc['analysis']['rules'][0]['plain']=statement
        folder=self.root/'.implementation/tmp/word-context'; packet=export_packet(doc,folder)
        paragraphs=[p.text for p in Document(folder/'sme-checklist.docx').paragraphs]
        self.assertIn(statement,paragraphs); self.assertIn(packet['context'],paragraphs)
        self.assertIn(doc['analysis']['rules'][0]['id'],paragraphs)


    def test_r764_packet_hash_changes_when_provider_assumption_changes(self):
        doc=self.doc(); doc['llm']={'analysis':{'questions':[],'assumptions':['first']}}; first=packet_document(doc)
        doc['llm']['analysis']['assumptions']=['second']; second=packet_document(doc); self.assertNotEqual(first['packet_hash'],second['packet_hash'])

    def test_r765_crossprocess_workbook_rejected_even_with_same_source(self):
        doc=self.doc(); doc['id']='process-b'; packet=packet_document(doc)
        with self.assertRaises(ValidationError): read_answers(self.raw(),packet,REVIEWER)


class ExpandedKnowledgeTests(ExpandedFixture):
    state = 'accepted'

    def test_r766_projection_rebuild_after_index_deletion_uses_same_records(self):
        count=update_knowledge(self.c.ledger,self.doc()); before=(self.root/'knowledge/records.json').read_bytes(); (self.root/'knowledge/INDEX.md').unlink()
        self.assertEqual(update_knowledge(self.c.ledger,self.doc()),count); self.assertEqual((self.root/'knowledge/records.json').read_bytes(),before); self.assertTrue((self.root/'knowledge/INDEX.md').is_file())

    def test_r767_records_projection_rebuild_after_deletion(self):
        count=update_knowledge(self.c.ledger,self.doc()); (self.root/'knowledge/records.json').unlink()
        self.assertEqual(update_knowledge(self.c.ledger,self.doc()),count); self.assertEqual(len(decode((self.root/'knowledge/records.json').read_bytes())),count)

    def test_r768_canonical_record_timestamp_survives_replay(self):
        update_knowledge(self.c.ledger,self.doc()); before=list(self.c.ledger.db.execute('SELECT id,updated FROM knowledge ORDER BY id'))
        update_knowledge(self.c.ledger,self.doc()); self.assertEqual([tuple(r) for r in self.c.ledger.db.execute('SELECT id,updated FROM knowledge ORDER BY id')],[tuple(r) for r in before])

    def test_r769_records_projection_symlink_blocks_before_database_write(self):
        victim=self.root/'.implementation/tmp/victim'; victim.parent.mkdir(parents=True,exist_ok=True); victim.write_text('keep'); (self.root/'knowledge').mkdir(exist_ok=True); (self.root/'knowledge/records.json').symlink_to(victim)
        with self.assertRaises(ValidationError): update_knowledge(self.c.ledger,self.doc())
        self.assertEqual(self.record_count(),0); self.assertEqual(victim.read_text(),'keep')

    def test_r770_missing_accepted_return_prevents_knowledge_commit(self):
        (self.c.process_root('process-a')/'input/sme-return.xlsx').unlink()
        with self.assertRaises(FileNotFoundError): update_knowledge(self.c.ledger,self.doc())
        self.assertEqual(self.record_count(),0)

    def test_r771_analysis_file_loss_prevents_knowledge_commit(self):
        (self.c.process_root('process-a')/'analysis/source-analysis.json').unlink()
        with self.assertRaises(FileNotFoundError): update_knowledge(self.c.ledger,self.doc())
        self.assertEqual(self.record_count(),0)

    def test_r772_packet_file_loss_prevents_knowledge_commit(self):
        (self.c.process_root('process-a')/'review/packet.json').unlink()
        with self.assertRaises(FileNotFoundError): update_knowledge(self.c.ledger,self.doc())
        self.assertEqual(self.record_count(),0)

    def test_r773_analysis_hash_failure_blocks_even_equivalent_json(self):
        path=self.c.process_root('process-a')/'analysis/source-analysis.json'; path.write_text(json.dumps(self.doc()['analysis']))
        with self.assertRaises(ValidationError): update_knowledge(self.c.ledger,self.doc())
        self.assertEqual(self.record_count(),0)

    def test_r774_packet_hash_failure_blocks_equivalent_json(self):
        path=self.c.process_root('process-a')/'review/packet.json'; path.write_text(json.dumps(self.packet()))
        with self.assertRaises(ValidationError): update_knowledge(self.c.ledger,self.doc())
        self.assertEqual(self.record_count(),0)

    def test_r775_changed_ledger_answer_requires_return_reparse(self):
        doc=self.doc(); doc['answers']['items'][next(iter(doc['answers']['items']))]['reviewer']='Forged fixture reviewer'; self.c.ledger.save(doc)
        with self.assertRaises(ValidationError): update_knowledge(self.c.ledger,self.doc())
        self.assertEqual(self.record_count(),0)

    def test_r776_wrong_source_snapshot_with_rebound_answer_is_rejected(self):
        doc=self.doc(); doc['analysis']['source_snapshot']='f'*64; doc['answers']['source_snapshot']='f'*64; self.c.ledger.save(doc)
        with self.assertRaises(ValidationError): update_knowledge(self.c.ledger,self.doc())
        self.assertEqual(self.record_count(),0)

    def test_r777_knowledge_second_insert_failure_rolls_back_first_rule(self):
        rules=self.doc()['analysis']['rules']; self.assertGreaterEqual(len(rules),2)
        self.failing_trigger('reject_knowledge','knowledge',condition="WHEN (SELECT COUNT(*) FROM knowledge)=1")
        with self.assertRaises(sqlite3.IntegrityError): update_knowledge(self.c.ledger,self.doc())
        self.assertEqual(self.record_count(),0); self.assertFalse((self.root/'knowledge/records.json').exists())


    def test_r778_index_replace_failure_retains_canonical_records_for_replay(self):
        original=Path.replace
        def replace(path,target):
            if Path(target).name=='INDEX.md': raise OSError('fixture index replace')
            return original(path,target)
        with patch.object(Path,'replace',replace):
            with self.assertRaises(OSError): update_knowledge(self.c.ledger,self.doc())
        count=self.record_count(); self.assertGreater(count,0); self.assertEqual(update_knowledge(self.c.ledger,self.doc()),count)

    def test_r779_index_replace_failure_cleans_temporary_projection(self):
        original=Path.replace
        def replace(path,target):
            if Path(target).name=='INDEX.md': raise OSError('fixture index replace')
            return original(path,target)
        with patch.object(Path,'replace',replace):
            with self.assertRaises(OSError): update_knowledge(self.c.ledger,self.doc())
        self.assertEqual(list((self.root/'.implementation/tmp').glob('knowledge-index-*')),[])

    def test_r780_knowledge_record_ids_are_sorted_in_projection(self):
        update_knowledge(self.c.ledger,self.doc()); records=decode((self.root/'knowledge/records.json').read_bytes())
        self.assertEqual([r['id'] for r in records],sorted(r['id'] for r in records))

    def test_r781_global_assumptions_do_not_create_rule_knowledge(self):
        update_knowledge(self.c.ledger,self.doc()); rows=[json.loads(r[0]) for r in self.c.ledger.db.execute('SELECT document FROM knowledge')]
        self.assertEqual({r['rule_id'] for r in rows},{r['id'] for r in self.doc()['analysis']['rules']}); self.assertFalse(any(r['rule_id'].startswith('G_') for r in rows))

    def test_r782_knowledge_publication_does_not_grant_verification_credit(self):
        update_knowledge(self.c.ledger,self.doc()); doc=self.doc()
        self.assertEqual(doc['status'],'QUEUED_VERIFY'); self.assertFalse(doc.get('verification_finished')); self.assertFalse(doc.get('report_verified')); self.assertEqual(doc['runs'],[])

    def test_r783_knowledge_replay_preserves_unrelated_process_audit(self):
        before=self.events(); update_knowledge(self.c.ledger,self.doc()); update_knowledge(self.c.ledger,self.doc())
        self.assertEqual(self.events(),before); self.assertTrue(self.doc()['packet_imported'])

    def test_r784_knowledge_return_symlink_cannot_rebind_provenance(self):
        path=self.c.process_root('process-a')/'input/sme-return.xlsx'; raw=path.read_bytes(); path.unlink(); victim=self.root/'.implementation/tmp/return.xlsx'; victim.parent.mkdir(parents=True,exist_ok=True); victim.write_bytes(raw); path.symlink_to(victim)
        with self.assertRaises(ValidationError): update_knowledge(self.c.ledger,self.doc())
        self.assertEqual(self.record_count(),0)

    def test_r785_historical_conflict_rolls_back_new_rule_insertion(self):
        update_knowledge(self.c.ledger,self.doc()); rules=self.doc()['analysis']['rules']; self.assertEqual(len(rules),2)
        records=[dict(r) for r in self.c.ledger.db.execute('SELECT id,document FROM knowledge')]
        first=next(r for r in records if json.loads(r['document'])['rule_id']==rules[0]['id']); second=next(r for r in records if json.loads(r['document'])['rule_id']==rules[1]['id'])
        self.c.ledger.db.execute('DELETE FROM knowledge WHERE id=?',(first['id'],)); self.c.ledger.db.execute('UPDATE knowledge SET document=? WHERE id=?',(second['document']+' ',second['id'])); self.c.ledger.db.commit()
        with self.assertRaises(ValidationError): update_knowledge(self.c.ledger,self.doc())
        self.assertEqual(self.record_count(),1); self.assertIsNone(self.c.ledger.db.execute('SELECT id FROM knowledge WHERE id=?',(first['id'],)).fetchone())
        self.assertEqual(self.c.ledger.db.execute('SELECT document FROM knowledge WHERE id=?',(second['id'],)).fetchone()[0],second['document']+' ')



class ExpandedEvidenceTests(ExpandedFixture):
    state = 'waiting'

    def test_r786_return_import_event_failure_does_not_start_verification(self):
        self.failing_trigger('reject_return_event','events',condition="WHEN NEW.stage='review'")
        with self.assertRaises(sqlite3.IntegrityError): self.accept()
        with patch.object(self.c,'verify') as verify: self.c.advance('process-a'); verify.assert_not_called()
        self.assertEqual(self.doc()['status'],'WAITING_SME')

    def test_r787_accepted_return_integrity_reparse_rejects_reviewer_drift(self):
        doc=self.accept(); doc['answers']['reviewer']='Another fictional reviewer'; self.c.ledger.save(doc)
        with self.assertRaises(ValidationError): self.c.review_integrity(self.doc())

    def test_r788_review_integrity_rejects_symlink_accepted_return(self):
        self.accept(); path=self.c.process_root('process-a')/'input/sme-return.xlsx'; raw=path.read_bytes(); path.unlink(); alternate=self.root/'.implementation/tmp/return.xlsx'; alternate.parent.mkdir(parents=True,exist_ok=True); alternate.write_bytes(raw); path.symlink_to(alternate)
        with self.assertRaises(ValidationError): self.c.review_integrity(self.doc())

    def test_r789_source_integrity_failure_routes_verify_to_blocked_report(self):
        self.accept(); path=self.c.process_root('process-a')/'input/sources/ELIGIBLE.cbl'; path.write_text(COBOL+'\nmodified')
        with patch('workbench.coordinator.verify_program') as execute: self.c.advance('process-a'); execute.assert_not_called()
        self.assertEqual(self.doc()['status'],'QUEUED_REPORT'); self.assertTrue(any(b['kind']=='evidence_integrity' for b in self.doc()['blockers']))

    def test_r790_finished_verification_still_checks_source_before_knowledge(self):
        self.accept(); self.save(verification_finished=True); path=self.c.process_root('process-a')/'input/sources/ELIGIBLE.cbl'; path.write_text('changed')
        with patch('workbench.coordinator.update_knowledge') as knowledge: self.c.advance('process-a'); knowledge.assert_not_called()
        self.assertTrue(any(b['kind']=='evidence_integrity' for b in self.doc()['blockers']))

    def test_r791_registered_artifact_hash_without_listing_is_not_downloadable(self):
        doc=self.doc(); doc['artifacts'].remove('review/sme-checklist.html'); self.c.ledger.save(doc)
        with self.assertRaises(ValidationError): self.c.artifact('process-a','review/sme-checklist.html')

    def test_r792_artifact_registration_does_not_reorder_existing_inventory(self):
        doc=self.doc(); before=list(doc['artifacts']); self.c.register(doc,'review/packet.json')
        self.assertEqual(doc['artifacts'],before)

    def test_r793_report_acceptance_requires_executive_inspection_confirmation(self):
        from workbench.reports import generate_reports
        self.save('QUEUED_REPORT')
        def unconfirmed(*args, **kwargs):
            paths=generate_reports(*args, **kwargs)
            inspection=next(path for path in paths if path.name=='inspection.json')
            model=decode(inspection.read_bytes()); model['executive_html_checked']=False; inspection.write_bytes(encode(model))
            return paths
        with patch('workbench.reports.generate_reports',side_effect=unconfirmed):
            with self.assertRaises(ValidationError): self.c.advance('process-a')
        self.assertEqual(self.doc()['status'],'REPORTING_FAILED'); self.assertFalse(self.doc().get('report_verified'))
        self.assertEqual(self.c.ledger.history('process-a'),[])


    def test_r794_closing_idle_coordinator_releases_database_and_lock(self):
        self.c.close(); self.assertIsNone(self.c.ledger); self.c=Coordinator(self.root)
        self.assertEqual(self.doc()['status'],'WAITING_SME')

    def test_r795_closing_coordinator_twice_is_idempotent(self):
        self.c.close(); self.c.close(); self.assertIsNone(self.c.ledger)
        self.c=Coordinator(self.root); self.assertTrue(self.doc()['packet_issued'])

    def test_r796_closed_active_checkpoint_retains_resume_stage(self):
        doc=self.doc(); doc['_active_stage']='QUEUED_ANALYSIS'; self.c.closed=True
        with self.assertRaises(StageInterrupted): self.c.checkpoint(doc)
        self.assertEqual(self.doc()['status'],'QUEUED_ANALYSIS'); self.c.closed=False

    def test_r797_process_b_control_cannot_pause_process_a(self):
        self.c.create(MANIFEST.replace('process-a','process-b'),{'ELIGIBLE.cbl':COBOL}); self.c.start('process-b'); self.c.control('process-b','pause')
        doc=self.doc(); doc['_active_stage']='QUEUED_ANALYSIS'; self.c.checkpoint(doc)
        self.assertEqual(self.doc()['status'],'WAITING_SME'); self.assertEqual(self.c.ledger.get('process-b')['status'],'PAUSED')

    def test_r798_process_b_cancellation_does_not_block_process_a_target_gate(self):
        self.c.create(MANIFEST.replace('process-a','process-b'),{'ELIGIBLE.cbl':COBOL}); self.c.control('process-b','cancel')
        doc=self.doc(); doc['_active_stage']='QUEUED_VERIFY'; self.c.target_checkpoint(doc)
        self.assertFalse(doc.get('cancel_requested')); self.assertTrue(self.c.ledger.get('process-b')['cancel_requested'])

    def test_r799_packet_integrity_failure_during_restart_releases_writer_lock(self):
        path=self.c.process_root('process-a')/'review/sme-checklist.html'; original=path.read_bytes(); path.write_bytes(b'changed'); self.c.close()
        with self.assertRaises(ValidationError): Coordinator(self.root)
        path.write_bytes(original); self.c=Coordinator(self.root); self.assertTrue(self.doc()['packet_issued'])

    def test_r800_transient_retry_recovery_clears_only_own_failure_metadata(self):
        self.save('QUEUED_VERIFY',blockers=[{'kind':'stage_failure','stage':'QUEUED_VERIFY'},{'kind':'sme_unresolved','message':'retained'}],last_error='OSError',last_failure={'stage':'QUEUED_VERIFY'},next_retry_at=9)
        doc=self.doc(); self.c.stage_success(doc,'QUEUED_VERIFY'); self.c.ledger.save(doc,'QUEUED_REPORT')
        after=self.doc(); self.assertNotIn('last_error',after); self.assertNotIn('last_failure',after); self.assertEqual(after['next_retry_at'],0); self.assertEqual(after['blockers'],[{'kind':'sme_unresolved','message':'retained'}])
