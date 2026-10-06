import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from workbench.domain import ValidationError
from workbench.intake import from_rows, parse_manifest, parse_intake_xlsx


EXAMPLES=Path(__file__).resolve().parent.parent/'examples'


class IntakeHardeningTests(unittest.TestCase):
    def manifest(self):return (EXAMPLES/'process-input.md').read_text()

    def test_populated_markdown_rows_cannot_silently_disappear(self):
        for order in ('', '2.5', '-1', 'two', '1e2'):
            with self.subTest(order=order),self.assertRaisesRegex(ValidationError,'order|integers'):
                parse_manifest(self.manifest()+f'\n| {order} | LOSTJOB | 1 | STEP1 | SECRET | IN | OUT | ALWAYS |\n')

    def test_missing_final_cell_is_not_stripped_into_an_accepted_row(self):
        with self.assertRaisesRegex(ValidationError,'eight columns'):
            parse_manifest(self.manifest()+'\n| 2 | LOSTJOB | 1 | STEP1 | SECRET | IN | OUT |\n')

    def test_populated_excel_rows_require_order_even_after_blank_rows(self):
        book=load_workbook(EXAMPLES/'intake-template.xlsx')
        try:
            sheet=book['Intake']; row=sheet.max_row+2
            for col,value in enumerate([None,'LOSTJOB',1,'STEP1','SECRET','IN','OUT','ALWAYS'],1):sheet.cell(row,col).value=value
            data=BytesIO();book.save(data)
            with self.assertRaisesRegex(ValidationError,'every populated row'):
                parse_intake_xlsx(data.getvalue())
        finally:book.close()

    def test_empty_excel_rows_remain_harmless(self):
        book=load_workbook(EXAMPLES/'intake-template.xlsx')
        try:
            book['Intake'].cell(10,2).value='   '
            data=BytesIO();book.save(data)
            self.assertEqual(len(parse_intake_xlsx(data.getvalue())['jobs']),1)
        finally:book.close()

    def test_non_integer_orders_do_not_truncate_or_accept_booleans(self):
        for value in (1.5, True, False, None, float('nan')):
            with self.subTest(value=value),self.assertRaises(ValidationError):
                from_rows('process-a','A',[[value,'JOBA',1,'S1','PROGRAM','','','Always']])

    def test_job_method_collisions_fail_before_packet_or_target_creation(self):
        for first,second in (('JOBA','joba'),('JOB-A','JOB_A')):
            with self.subTest(first=first,second=second),self.assertRaisesRegex(ValidationError,'collide'):
                from_rows('process-a','A',[[1,first,1,'S1','PROGRAM','','','Always'],[2,second,1,'S1','PROGRAM','','','Always']])

    def test_step_names_are_case_insensitive(self):
        with self.assertRaisesRegex(ValidationError,'Duplicate step'):
            from_rows('process-a','A',[[1,'JOBA',1,'S1','PROGRAM','','','Always'],[1,'JOBA',2,'s1','PROGRAM','','','Always']])

    def test_missing_identity_cells_are_not_stringified_as_none(self):
        for column in (1,3,4):
            row=[1,'JOBA',1,'S1','PROGRAM','','','Always'];row[column]=None
            with self.subTest(column=column),self.assertRaisesRegex(ValidationError,'needs a job'):
                from_rows('process-a','A',[row])

    def test_duplicate_process_attributes_are_rejected(self):
        for extra in ('- Process ID: other-process','- Process name: Other'):
            with self.subTest(extra=extra),self.assertRaisesRegex(ValidationError,'exactly one'):
                parse_manifest(self.manifest()+'\n'+extra+'\n')

    def test_repeated_or_missing_headers_fail_explicitly(self):
        with self.assertRaisesRegex(ValidationError,'repeated headers'):
            parse_manifest(self.manifest()+'\n'+next(x for x in self.manifest().splitlines() if x.startswith('| Job order')))
        with self.assertRaisesRegex(ValidationError,'headers'):
            parse_manifest('\n'.join(x for x in self.manifest().splitlines() if not x.startswith('| Job order')))

    def test_excel_cells_cannot_inject_markdown_rows_during_normalization(self):
        for value in ('IN|OUT','IN\n| 2 | INJECTED |'):
            with self.subTest(value=value),self.assertRaisesRegex(ValidationError,'single-line'):
                from_rows('process-a','A',[[1,'JOBA',1,'S1','PROGRAM',value,'','Always']])
        with self.assertRaisesRegex(ValidationError,'one line'):
            from_rows('process-a','A\n- Process ID: other',[[1,'JOBA',1,'S1','PROGRAM','','','Always']])


if __name__=='__main__':unittest.main()
