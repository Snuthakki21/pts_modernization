"""Factory contracts and explicitly bounded online delivery regressions."""
import json
from pathlib import Path
import tempfile
import unittest
from workbench.intake import parse_manifest
from workbench.domain import ValidationError

ONLINE = '''Process ID: online-pilot
Process name: Fictional eligibility transaction
| Transaction | Program | Mapset | Map |
|---|---|---|---|
| ELIG | ELIGIBLE | | |
'''

class FactoryTests(unittest.TestCase):
    def test_online_intake_has_no_invented_batch_job(self):
        doc = parse_manifest(ONLINE)
        self.assertEqual(doc['jobs'], [])
        self.assertEqual(doc['transactions'][0]['program'], 'ELIGIBLE')

    def test_knowledge_guidance_matches_ingestion_limits(self):
        text = Path('knowledge/README.md').read_text()
        self.assertNotIn('maximum 16 KB for automatic ingestion', text)
        self.assertIn('1 MiB', text)

    def test_factory_unknowns_do_not_become_supported(self):
        from workbench.factory import factory_view
        view = factory_view({'id':'p','status':'READY','jobs':[], 'blockers':[], 'artifacts':[]})
        self.assertEqual(view['counts']['verified_capabilities'], 0)
        self.assertTrue(all(c['state']=='NOT_ASSESSED' for c in view['capabilities']))
        guidance=' '.join(view['investigation_strategy'])
        self.assertIn('behavior-preserving replacement',guidance)
        self.assertIn('never silently repair legacy logic',guidance)
        self.assertNotIn('target design change',guidance)

    def test_transaction_duplicate_and_partial_map_binding_rejected(self):
        with self.assertRaises(ValidationError):parse_manifest(ONLINE+'| elig | ELIGIBLE | | |\n')
        with self.assertRaises(ValidationError):parse_manifest(ONLINE.replace('| | |','| MAPSET | |'))

    def test_online_manifest_tampering_is_rejected(self):
        from workbench.coordinator import Coordinator
        from test_source import COBOL
        with tempfile.TemporaryDirectory() as temp:
            c=Coordinator(temp)
            try:
                doc=c.create(ONLINE,{'ELIGIBLE.cbl':COBOL})
                doc['transactions'][0]['program']='OTHER'
                with self.assertRaises(ValidationError):c.manifest_integrity(doc)
            finally:c.close()

    def test_online_lineage_resolves_program_without_job(self):
        from workbench.lineage import map_lineage
        from test_source import COBOL
        doc=parse_manifest(ONLINE)
        result=map_lineage({'ELIGIBLE.cbl':COBOL},doc)
        self.assertTrue(result['closure']['complete'])
        self.assertEqual(result['scope']['root_transactions'],['ELIG'])
        self.assertIn('ELIGIBLE.cbl',result['scope']['selected_files'])

    def test_online_runtime_boundaries_recovery_and_real_source_comparisons(self):
        from workbench.online_runtime import create_app
        from workbench.source import analyze_program
        from workbench.target import emit_program,run_generated
        from workbench.reference import run_reference
        from fastapi.testclient import TestClient
        from test_source import COBOL
        p=analyze_program('ELIGIBLE.cbl',COBOL,{})
        p['target_contract_version']=2;code=emit_program(p)
        now=[100.0]; fail=[False]
        def execute(record):
            if fail[0]:raise ValueError('private error')
            return run_generated(code,record)
        with tempfile.TemporaryDirectory() as temp:
            db=Path(temp)/'sessions.sqlite';token='x'*32
            spec={'ELIG':{'program':'ELIGIBLE','fields':p['fields']},'OTHR':{'program':'ELIGIBLE','fields':p['fields']}}
            app=create_app({'ELIGIBLE':execute},spec,db,token,clock=lambda:now[0],ttl=30)
            published=app.openapi()
            self.assertEqual(published['components']['schemas']['ELIGRecord']['properties']['AGE']['maximum'],999)
            self.assertIn('requestBody',published['paths']['/api/transactions/{transaction}']['post'])
            with TestClient(app) as client:
                auth={'Authorization':'Bearer '+token}
                self.assertEqual(client.post('/api/sessions',json={'transaction':'ELIG'}).status_code,401)
                self.assertEqual(client.post('/api/sessions',json={'transaction':[]},headers=auth).status_code,422)
                self.assertEqual(client.post('/api/sessions',json={'transaction':'ELIG'},headers={**auth,'Origin':'https://evil.invalid'}).status_code,403)
                sid=client.post('/api/sessions',json={'transaction':'ELIG'},headers=auth).json()['session_id']
                headers={**auth,'X-Session-ID':sid,'Idempotency-Key':'key'}
                record={'AGE':17,'ACTIVE':'Y','DECISION':' '}
                data={'record':record,'revision':0}
                first=client.post('/api/transactions/ELIG',json=data,headers=headers)
                self.assertEqual(first.status_code,200);self.assertEqual(first.json()['result'],run_reference(p,record))
                self.assertEqual(client.post('/api/transactions/ELIG',json=data,headers=headers).json(),first.json())
                self.assertEqual(client.post('/api/transactions/ELIG',json={**data,'record':{**record,'AGE':18}},headers=headers).status_code,409)
                self.assertEqual(client.post('/api/transactions/ELIG',json=data,headers={**headers,'Idempotency-Key':'new'}).status_code,409)
                self.assertEqual(client.post('/api/transactions/OTHR',json=data,headers=headers).status_code,409)
                # 20 distinct source states through the exported HTTP target.
                for i in range(20):
                    record={'AGE':i+10,'ACTIVE':'N' if i%2 else 'Y','DECISION':' '}
                    response=client.post('/api/transactions/ELIG',json={'record':record,'revision':i+1},headers={**headers,'Idempotency-Key':str(i)})
                    self.assertEqual(response.status_code,200);self.assertEqual(response.json()['result'],run_reference(p,record))
                invalid=client.post('/api/transactions/ELIG',json={'record':{'AGE':-1},'revision':21},headers={**headers,'Idempotency-Key':'invalid'})
                self.assertEqual(invalid.status_code,422)
                fail[0]=True
                failed=client.post('/api/transactions/ELIG',json={'record':record,'revision':21},headers={**headers,'Idempotency-Key':'retry'})
                self.assertEqual(failed.status_code,500);self.assertNotIn('private error',failed.text)
                fail[0]=False
                retried=client.post('/api/transactions/ELIG',json={'record':record,'revision':21},headers={**headers,'Idempotency-Key':'retry'})
                self.assertEqual(retried.status_code,200)
                self.assertEqual(client.post('/api/transactions/ELIG',content=b'{"revision":0,"revision":1}',headers=headers).status_code,422)
                self.assertEqual(client.post('/api/transactions/ELIG',content=b'x'*65537,headers=headers).status_code,413)
                now[0]=131
                self.assertEqual(client.post('/api/transactions/ELIG',json=data,headers=headers).status_code,410)
            # Restart retains idempotency state and expiry.
            with TestClient(create_app({'ELIGIBLE':execute},spec,db,token,clock=lambda:100)) as client:
                self.assertEqual(client.post('/api/transactions/ELIG',json=data,headers=headers).json(),first.json())

    def test_packaged_online_runtime_supports_lf_and_windows_crlf_implementation_checkouts(self):
        from unittest.mock import patch
        from workbench.online import packaged_runtime
        original = Path('workbench/online_runtime.py').read_bytes().replace(b'\r\n', b'\n')
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for mode, data in (('LF', original), ('CRLF', original.replace(b'\n', b'\r\n'))):
                with self.subTest(mode=mode):
                    (root/'online_runtime.py').write_bytes(data)
                    with patch('workbench.online.__file__', str(root/'online.py')):
                        packed = packaged_runtime()
                    scope = {'__name__':'standalone_runtime'}
                    exec(compile(packed, 'runtime.py', 'exec'), scope)
                    app = scope['create_app']({}, {}, root/(mode+'.sqlite'), 'x'*32)
                    self.assertEqual(app.title, 'Modernized business capabilities')
                    self.assertTrue(callable(scope['path_is_link']))

    def test_factory_end_to_end_fictional_online_report(self):
        from workbench.coordinator import Coordinator
        from test_source import COBOL
        from openpyxl import load_workbook
        from io import BytesIO
        with tempfile.TemporaryDirectory() as temp:
            c=Coordinator(temp)
            try:
                c.create(ONLINE.replace('ELIGIBLE','RUNTIME'),{'RUNTIME.cbl':COBOL.replace('ELIGIBLE','RUNTIME')},demo=True);c.start('online-pilot');c.advance('online-pilot')
                packet=json.loads(c.artifact('online-pilot','review/packet.json').read_text())
                self.assertEqual(json.loads(packet['context'])['transactions'][0]['id'],'ELIG')
                self.assertIn('G_ONLINE_TARGET',{item['id'] for item in packet['items']})
                book=load_workbook(c.artifact('online-pilot','review/sme-checklist.xlsx'))
                for row in book['Checklist'].iter_rows(min_row=2):row[4].value='Yes'
                data=BytesIO();book.save(data);book.close()
                c.import_answers('online-pilot',data.getvalue(),'Fictional test reviewer');c.advance('online-pilot');c.advance('online-pilot')
                doc=c.ledger.get('online-pilot')
                self.assertEqual(doc['status'],'COMPLETED_WITH_BLOCKERS')
                self.assertEqual(doc['online_delivery']['status'],'HTTP_COMPARISON_PASSED')
                self.assertFalse(doc['online_delivery']['native_cics_verified'])
                factory=json.loads(c.artifact(doc['id'],'reports/report-0001/factory.json').read_text())
                self.assertEqual(factory['transactions'][0]['api'],'/api/transactions/ELIG')
                self.assertTrue(c.artifact(doc['id'],'target/run-0001/online/openapi.json').is_file())
                package=json.loads(c.artifact(doc['id'],doc['online_delivery']['package']).read_text())
                self.assertEqual(len(package['files']),len({p.casefold() for p in package['files']}))
                self.assertTrue(any(p.startswith('modules/') for p in package['files']))
                import os,subprocess,sys
                folder=c.artifact(doc['id'],doc['online_delivery']['package']).parent
                launched=subprocess.run([sys.executable,'-c',"import application; assert application.app.title == 'Modernized business capabilities'"],cwd=folder,env={**os.environ,'ONLINE_TOKEN':'fixture-'+'x'*32,'ONLINE_STATE':str(Path(temp)/'.implementation/tmp/online-state')},capture_output=True,text=True)
                self.assertEqual(launched.returncode,0,launched.stderr)
                self.assertEqual(sum(b['kind']=='online_semantics_gap' for b in doc['blockers']),1)
                self.assertIn('Factory capabilities',load_workbook(c.artifact(doc['id'],'reports/report-0001/metrics.xlsx')).sheetnames)
            finally:c.close()

    def test_online_bms_binding_requires_correct_map_and_mapset(self):
        from workbench.lineage import map_lineage
        from test_source import COBOL
        files={'ELIGIBLE.cbl':COBOL,'FORM.bms':'FORM DFHMSD TYPE=MAP\nSCREEN DFHMDI SIZE=(24,80)\n DFHMSD TYPE=FINAL\n END\n'}
        doc=parse_manifest(ONLINE.replace('| | |','| FORM | SCREEN |'))
        result=map_lineage(files,doc)
        self.assertTrue(result['closure']['complete'],result['closure']['gaps'])
        doc['transactions'][0]['map']='MISSING'
        self.assertFalse(map_lineage(files,doc)['closure']['complete'])

    def test_concurrent_online_requests_commit_once_and_sessions_are_isolated(self):
        from workbench.online_runtime import create_app
        from fastapi.testclient import TestClient
        from concurrent.futures import ThreadPoolExecutor
        calls=[]
        def execute(record):calls.append(record);return {'input_status':'ACCEPT_INPUT','record':record,'return_code':0}
        with tempfile.TemporaryDirectory() as temp:
            with TestClient(create_app({'P':execute},{'TEST':{'program':'P'}},Path(temp)/'state.sqlite','x'*32)) as client:
                auth={'Authorization':'Bearer '+'x'*32}
                a=client.post('/api/sessions',json={'transaction':'TEST'},headers=auth).json()['session_id']
                b=client.post('/api/sessions',json={'transaction':'TEST'},headers=auth).json()['session_id']
                def request(sid):return client.post('/api/transactions/TEST',json={'record':{'VALUE':1},'revision':0},headers={**auth,'X-Session-ID':sid,'Idempotency-Key':'same'})
                with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(request,[a,a]))
                self.assertEqual([r.status_code for r in results],[200,200]);self.assertEqual(len(calls),1)
                self.assertEqual(request(b).status_code,200);self.assertEqual(len(calls),2)

    def test_consistency_blocks_false_completion_and_unsupported_target(self):
        from workbench.factory import consistency_findings
        doc={'status':'COMPLETED','blockers':[{'kind':'unknown'}],'program_versions':{'X':'hash'},'analysis':{'programs':{},'rules':[]}}
        ids={f['id'] for f in consistency_findings(doc)}
        self.assertIn('INVALID_COMPLETION',ids);self.assertIn('TARGET_WITHOUT_SUPPORTED_SOURCE:X',ids)

    def test_real_http_factory_rules_obligations_and_notes_remain_aligned(self):
        from workbench.api import create_app
        from fastapi.testclient import TestClient
        from test_source import COBOL
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);notes=root/'.implementation/tmp/process.md';notes.parent.mkdir(parents=True);notes.write_text('# Process\nFictional eligibility\n')
            app=create_app(root);c=app.state.coordinator
            try:
                c.create(ONLINE,{'ELIGIBLE.cbl':COBOL},assistant_mode='agent',process_notes=notes)
                c.start('online-pilot');c.advance('online-pilot')
                client=TestClient(app,base_url='http://127.0.0.1:8765')
                try:
                    for suffix in ('factory','agent/rules','agent/obligations','agent/context?document_id=NOTE_1&start_line=1&end_line=2'):
                        response=client.get('/api/process/online-pilot/'+suffix)
                        self.assertEqual(response.status_code,200,(suffix,response.text))
                        self.assertIsInstance(response.json(),dict)
                    self.assertEqual(client.get('/api/process/online-pilot/agent/rules?after=99999').status_code,400)
                    self.assertEqual(client.get('/api/process/online-pilot/factory?after=-1').status_code,400)
                finally:client.close()
            finally:c.close()

    def test_online_startup_closes_the_real_sqlite_connection(self):
        import sqlite3
        from unittest.mock import patch
        from workbench.online_runtime import create_app
        original=sqlite3.connect;connections=[]
        def tracked(*args,**kwargs):
            db=original(*args,**kwargs);connections.append(db);return db
        with tempfile.TemporaryDirectory() as temp:
            try:
                with patch('workbench.online_runtime.sqlite3.connect',side_effect=tracked):
                    create_app({}, {}, Path(temp)/'state.sqlite','x'*32)
                self.assertEqual(len(connections),1)
                with self.assertRaisesRegex(sqlite3.ProgrammingError,'closed'):
                    connections[0].execute('SELECT 1')
            finally:
                for db in connections:db.close()

    def test_online_startup_failures_close_connections_and_preserve_the_error(self):
        import sqlite3
        from unittest.mock import patch
        from workbench.online_runtime import create_app
        original=sqlite3.connect
        for stage in ('pragma','schema'):
            with self.subTest(stage=stage), tempfile.TemporaryDirectory() as temp:
                connections=[]
                class FailingConnection(sqlite3.Connection):
                    def execute(self,sql,*args):
                        if stage=='pragma' and sql=='PRAGMA foreign_keys=ON':
                            raise sqlite3.OperationalError('fixture pragma failure')
                        return super().execute(sql,*args)
                    def executescript(self,sql):
                        raise sqlite3.OperationalError('fixture schema failure')
                def tracked(*args,**kwargs):
                    db=original(*args,**kwargs,factory=FailingConnection);connections.append(db);return db
                try:
                    with patch('workbench.online_runtime.sqlite3.connect',side_effect=tracked):
                        with self.assertRaisesRegex(sqlite3.OperationalError,'fixture '+stage+' failure'):
                            create_app({}, {}, Path(temp)/'state.sqlite','x'*32)
                    self.assertEqual(len(connections),1)
                    with self.assertRaisesRegex(sqlite3.ProgrammingError,'closed'):
                        connections[0].execute('SELECT 1')
                finally:
                    for db in connections:db.close()

    def test_online_session_delete_closes_and_commits_the_cascade(self):
        self._assert_online_delete_connection_lifecycle(False)

    def test_online_session_delete_failure_closes_and_rolls_back(self):
        self._assert_online_delete_connection_lifecycle(True)

    def _assert_online_delete_connection_lifecycle(self, fail):
        import sqlite3
        from contextlib import closing
        from types import SimpleNamespace
        from unittest.mock import patch
        from workbench.online_runtime import create_app
        original=sqlite3.connect;connections=[]
        def tracked(*args,**kwargs):
            db=original(*args,**kwargs);connections.append(db);return db
        with tempfile.TemporaryDirectory() as temp:
            database=Path(temp)/'state.sqlite'
            app=create_app({}, {}, database,'x'*32)
            with closing(original(database)) as oracle, oracle:
                oracle.execute("INSERT INTO sessions VALUES ('session','TEST',900,0)")
                oracle.execute("INSERT INTO requests VALUES ('session','key','hash','{}')")
                if fail:
                    oracle.executescript("CREATE TRIGGER deny_delete BEFORE DELETE ON sessions BEGIN SELECT RAISE(ABORT,'protected session'); END;")
            endpoint=next(r.endpoint for r in app.routes if r.path=='/api/sessions/{sid}')
            request=SimpleNamespace(headers={'authorization':'Bearer '+'x'*32},base_url='http://testserver/')
            try:
                with patch('workbench.online_runtime.sqlite3.connect',side_effect=tracked):
                    if fail:
                        with self.assertRaisesRegex(sqlite3.IntegrityError,'protected session'):
                            endpoint('session',request)
                    else:
                        self.assertEqual(endpoint('session',request),{'closed':True})
                self.assertEqual(len(connections),1)
                with self.assertRaisesRegex(sqlite3.ProgrammingError,'closed'):
                    connections[0].execute('SELECT 1')
                with closing(original(database)) as oracle:
                    expected=1 if fail else 0
                    self.assertEqual(oracle.execute('SELECT COUNT(*) FROM sessions').fetchone()[0],expected)
                    self.assertEqual(oracle.execute('SELECT COUNT(*) FROM requests').fetchone()[0],expected)
            finally:
                for db in connections:db.close()
