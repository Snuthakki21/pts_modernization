"""Large technical-gap inventories keep one bounded, evidence-complete SME packet."""
import copy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from workbench.domain import ValidationError, encode, sha
from workbench.intake import parse_manifest
from workbench.review import export_packet, packet_document, read_answers
from workbench.source import analyze_sources
from test_source import COBOL
from test_workflow import MANIFEST


def process_with_unsupported_lines(count=21509, version=2):
    process = parse_manifest(MANIFEST)
    if version is not None:
        process['sme_packet_version'] = version
    source = COBOL.replace('  GOBACK.', 'DISPLAY X\n' * count + '  GOBACK.')
    process['analysis'] = analyze_sources({'ELIGIBLE.cbl': source}, process)
    return process


def expand_spans(spans):
    return [number for first, last in spans for number in range(first, last + 1)]


class ReviewCompactionTests(unittest.TestCase):
    def test_21509_unsupported_lines_issue_one_technical_gap_without_losing_evidence(self):
        process = process_with_unsupported_lines()
        original = copy.deepcopy(process['analysis'])
        self.assertEqual(len(original['blockers']), 21509)
        packet = packet_document(process)
        self.assertEqual(packet['version'], 2)
        self.assertLessEqual(len(packet['items']), 2000)
        rules = [item for item in packet['items'] if item['kind'] == 'business_rule']
        self.assertEqual([item['id'] for item in rules], [rule['id'] for rule in original['rules']])
        self.assertEqual([item['question'] for item in rules], [rule['plain'] for rule in original['rules']])
        gaps = [item for item in packet['items'] if item['kind'] == 'technical_gap']
        self.assertEqual(len(gaps), 1)
        self.assertIn('semantic adapter', gaps[0]['question'])
        self.assertIn('unsupported', gaps[0]['question'])
        context = json.loads(packet['context'])
        analysis_ref = context['source_analysis']
        self.assertEqual(analysis_ref['path'], 'analysis/source-analysis.json')
        self.assertEqual(analysis_ref['sha256'], sha(encode(original)))
        self.assertEqual(analysis_ref['raw_blocker_count'], 21509)
        group = context['technical_gaps'][gaps[0]['id']]
        self.assertEqual(group['program'], 'ELIGIBLE')
        self.assertEqual(group['kind'], 'unsupported_source')
        indices = expand_spans(group['blocker_index_spans'])
        self.assertEqual(indices, list(range(21509)))
        self.assertEqual(group['blockers_sha256'], sha(encode([original['blockers'][i] for i in indices])))
        expected_lines = sorted({line for blocker in original['blockers'] for line in blocker['lines']})
        self.assertEqual(expand_spans(group['source_line_spans']['ELIGIBLE.cbl']), expected_lines)
        program_ref = analysis_ref['programs']['ELIGIBLE']
        self.assertEqual(program_ref['source_sha256'], original['programs']['ELIGIBLE']['source_hash'])
        self.assertEqual(program_ref['coverage_sha256'], sha(encode(original['programs']['ELIGIBLE']['coverage'])))
        self.assertEqual(process['analysis'], original)
        blocked_lines = set(expected_lines)
        self.assertTrue(all(line['disposition'] == 'unsupported' for line in original['programs']['ELIGIBLE']['coverage'] if line['line'] in blocked_lines))

    def test_new_packet_compacts_by_program_and_gap_kind_while_preserving_other_human_items(self):
        process = process_with_unsupported_lines(2001)
        analysis = process['analysis']
        analysis['blockers'].extend([
            {'kind': 'unsupported_source', 'path': 'SECOND.cbl', 'lines': [17, 19], 'message': 'Different program needs an adapter'},
            {'kind': 'unsupported_jcl', 'path': 'job.jcl', 'message': 'Dataset allocation needs an I/O adapter'},
            {'kind': 'unresolved_condition', 'message': 'Ask the owner what this job condition means'},
        ])
        second = copy.deepcopy(analysis['programs']['ELIGIBLE'])
        second.update({'name': 'SECOND', 'path': 'SECOND.cbl'})
        analysis['programs']['SECOND'] = second
        process['llm'] = {'analysis': {'questions': ['Which business date should the process use?']}}
        packet = packet_document(process)
        groups = json.loads(packet['context'])['technical_gaps'].values()
        self.assertEqual({(group['program'], group['kind']) for group in groups}, {
            ('ELIGIBLE', 'unsupported_source'), ('SECOND', 'unsupported_source'), ('job.jcl', 'unsupported_jcl')})
        self.assertEqual(len([item for item in packet['items'] if item['kind'] == 'technical_gap']), 3)
        human_item = next(item for item in packet['items'] if item['kind'] == 'unresolved_item')
        self.assertEqual(human_item['id'], 'B_2003')
        self.assertIn('Ask the owner', human_item['question'])
        self.assertEqual(next(item for item in packet['items'] if item['kind'] == 'provider_suggestion')['id'], 'LLM_000')

    def test_business_rule_overflow_still_blocks_instead_of_truncating(self):
        process = process_with_unsupported_lines(2001)
        rule = process['analysis']['rules'][0]
        process['analysis']['rules'] = [dict(rule, id='ACTUAL_R%04d' % i) for i in range(2001)]
        with self.assertRaisesRegex(ValidationError, 'business rules'):
            packet_document(process)
        self.assertEqual(len(process['analysis']['rules']), 2001)

    def test_many_distinct_gap_groups_still_fail_the_actual_human_item_limit(self):
        process = process_with_unsupported_lines(0)
        process['analysis']['blockers'] = [
            {'kind': 'unsupported_source', 'path': 'member-%04d.cbl' % i, 'message': 'Distinct source adapter required'}
            for i in range(2001)]
        with self.assertRaisesRegex(ValidationError, 'SME packet exceeds supported size'):
            packet_document(process)
        self.assertEqual(len(process['analysis']['blockers']), 2001)

    def test_two_thousand_actual_human_items_are_retained_and_one_extra_blocks(self):
        process = process_with_unsupported_lines(2001)
        rule = process['analysis']['rules'][0]
        process['analysis']['rules'] = [dict(rule, id='ACTUAL_R%04d' % i) for i in range(1994)]
        packet = packet_document(process)
        self.assertEqual(len(packet['items']), 2000)
        self.assertEqual(len([item for item in packet['items'] if item['kind'] == 'business_rule']), 1994)
        process['analysis']['rules'].append(dict(rule, id='ONE_EXTRA_REAL_RULE'))
        with self.assertRaisesRegex(ValidationError, 'SME packet exceeds supported size'):
            packet_document(process)
        self.assertEqual(len(process['analysis']['rules']), 1995)

    def test_compacted_packet_replay_survives_canonical_persistence(self):
        from workbench.domain import decode
        process = process_with_unsupported_lines(2001)
        packet = packet_document(process)
        self.assertEqual(packet_document(decode(encode(process))), packet)
        source_blocker = process['analysis']['blockers'][0]
        source_blocker['message'] += '; exact correction evidence'
        corrected = packet_document(process)
        self.assertNotEqual(corrected['packet_hash'], packet['packet_hash'])
        self.assertNotEqual(json.loads(corrected['context'])['source_analysis']['sha256'],
                            json.loads(packet['context'])['source_analysis']['sha256'])

    def test_v1_fingerprint_is_unchanged_and_historical_large_packets_still_fail_closed(self):
        process = process_with_unsupported_lines(0, version=None)
        historical = packet_document(process)
        self.assertEqual(historical['version'], 1)
        self.assertEqual(historical['packet_hash'], 'b61caeb11f1910c14fac6cc3ee5689b0975d22dbce4327e68df0f6d7db858c89')
        process['sme_packet_version'] = 1
        self.assertEqual(packet_document(process), historical)
        old_large = process_with_unsupported_lines(2001, version=None)
        with self.assertRaisesRegex(ValidationError, 'SME packet exceeds supported size'):
            packet_document(old_large)

    def test_compacted_export_round_trips_answers_and_rejects_context_tampering(self):
        process = process_with_unsupported_lines(2001)
        with tempfile.TemporaryDirectory() as folder:
            packet = export_packet(process, Path(folder) / 'review')
            self.assertEqual(export_packet(process, Path(folder) / 'review'), packet)
            book = load_workbook(Path(folder) / 'review/sme-checklist.xlsx')
            for row in book['Checklist'].iter_rows(min_row=2):
                row[4].value = 'Yes'
            output = BytesIO(); book.save(output)
            answers = read_answers(output.getvalue(), packet, 'Actual reviewer')
            self.assertEqual(len(answers['items']), len(packet['items']))
            self.assertEqual(answers['packet_hash'], packet['packet_hash'])
            book['Context']['A2'] = 'changed technical gap evidence'
            altered = BytesIO(); book.save(altered); book.close()
            with self.assertRaisesRegex(ValidationError, 'Frozen Context changed'):
                read_answers(altered.getvalue(), packet, 'Actual reviewer')

    def test_coordinator_preserves_frozen_analysis_and_yes_does_not_support_the_gap(self):
        from workbench.coordinator import Coordinator
        from workbench.domain import decode
        source = COBOL.replace('  GOBACK.', 'DISPLAY X\n' * 2001 + '  GOBACK.')
        with tempfile.TemporaryDirectory() as folder:
            coordinator = Coordinator(Path(folder))
            try:
                process = coordinator.create(MANIFEST, {'ELIGIBLE.cbl': source})
                process['sme_packet_version'] = 2
                coordinator.ledger.save(process)
                coordinator.start('process-a'); coordinator.advance('process-a')
                issued = coordinator.ledger.get('process-a')
                self.assertEqual(issued['status'], 'WAITING_SME')
                packet = decode(coordinator.artifact('process-a', 'review/packet.json').read_bytes())
                analysis_bytes = coordinator.artifact('process-a', 'analysis/source-analysis.json').read_bytes()
                self.assertEqual(json.loads(packet['context'])['source_analysis']['sha256'], sha(analysis_bytes))
                self.assertEqual(len(decode(analysis_bytes)['blockers']), 2001)
                book = load_workbook(coordinator.artifact('process-a', 'review/sme-checklist.xlsx'))
                for row in book['Checklist'].iter_rows(min_row=2):
                    row[4].value = 'Yes'
                returned = BytesIO(); book.save(returned); book.close()
                coordinator.import_answers('process-a', returned.getvalue(), 'Test fixture reviewer')
                coordinator.advance('process-a')
                verified = coordinator.ledger.get('process-a')
                self.assertEqual(len([blocker for blocker in verified['blockers'] if blocker['kind'] == 'unsupported_source']), 2001)
                self.assertEqual(verified['program_versions'], {})
                self.assertEqual(coordinator.artifact('process-a', 'analysis/source-analysis.json').read_bytes(), analysis_bytes)
            finally:
                coordinator.close()


class ConstructFamilyPacketTests(unittest.TestCase):
    def process(self, statements, names=('ELIGIBLE', 'SECOND')):
        process = parse_manifest(MANIFEST)
        process['sme_packet_version'] = 3
        files = {name+'.cbl': COBOL.replace('PROGRAM-ID. ELIGIBLE.', 'PROGRAM-ID. '+name+'.')
                 .replace('  GOBACK.', '\n'.join(statements)+'\n  GOBACK.') for name in names}
        process['lineage_scope'] = list(files)
        process['analysis'] = analyze_sources(files, process)
        return process

    def groups(self, packet):
        return {group['family']: group for group in json.loads(packet['context'])['technical_gaps'].values()}

    def test_v2_hashes_remain_byte_identical_for_small_and_compacted_packets(self):
        self.assertEqual(packet_document(process_with_unsupported_lines(0, 2))['packet_hash'],
                         'cfd719a2498d51d2a4d6db42141bfdc83b0b4a8d75fe1f82e36145d36db2ad67')
        self.assertEqual(packet_document(process_with_unsupported_lines(2001, 2))['packet_hash'],
                         '295cb551c750dd30d603b2b52b48edc99f07fc5a6b15aec0af241a4354189937')

    def test_small_v3_packet_groups_families_across_programs_and_keeps_rules_individual(self):
        process = self.process(['PERFORM WORK UNTIL AGE = 99', 'EXEC SQL SELECT X FROM T END-EXEC',
                                'EXEC CICS SEND MAP("SCREEN") END-EXEC'])
        original = copy.deepcopy(process['analysis'])
        packet = packet_document(process)
        self.assertEqual(packet['version'], 3)
        groups = self.groups(packet)
        self.assertEqual(set(groups), {'iteration', 'embedded_sql', 'cics'})
        for group in groups.values():
            self.assertEqual(group['programs'], ['ELIGIBLE', 'SECOND'])
            self.assertEqual(group['program_count'], 2)
            self.assertEqual(group['blocker_count'], 2)
            indices = expand_spans(group['blocker_index_spans'])
            self.assertEqual(group['blockers_sha256'], sha(encode([original['blockers'][i] for i in indices])))
            self.assertEqual(set(group['source_line_spans']), {'ELIGIBLE.cbl', 'SECOND.cbl'})
        for item in packet['items']:
            if item['kind'] == 'technical_gap':
                self.assertTrue(item['question'].startswith('ASSUMPTION - '))
                self.assertIn('unverified obligation', item['question'])
                self.assertIn('Yes confirms the description only', item['question'])
        self.assertEqual([item['id'] for item in packet['items'] if item['kind'] == 'business_rule'],
                         [rule['id'] for rule in original['rules']])
        self.assertEqual(process['analysis'], original)
        from workbench.domain import decode
        self.assertEqual(packet_document(decode(encode(process))), packet)

    def test_source_classifier_uses_tokens_without_literals_comments_or_name_prefixes(self):
        from workbench.review import _classify_construct
        examples = [
            ('000100 PERFORM CHECK-RECORD', 'iteration'),
            ('SEARCH ITEMS AT END DISPLAY "missing"', 'mixed'),
            ('EXEC SQL SELECT COL FROM T END-EXEC', 'embedded_sql'),
            ('EXEC CICS LINK PROGRAM("P") END-EXEC', 'cics'),
            ('READ INPUT-FILE', 'file_io'), ('ADD 1 TO AGE', 'arithmetic'),
            ('STRING "EXEC SQL" INTO BUFFER', 'string_handling'),
            ('EVALUATE AGE', 'control_flow'), ('INITIALIZE RECORD-AREA', 'initialization'),
            ('CALL "PERFORM"', 'procedure_call'), ('COPY MY-RECORD.', 'copy_dependency'),
            ('DISPLAY "EXEC SQL PERFORM"', 'terminal_io'),
            ('05 VALUE-FIELD PIC 9(3) COMP-3.', 'numeric_representation'),
            ('05 VALUE-FIELD PIC S9(3)V99.', 'numeric_representation'),
            ('05 VALUE-FIELD PIC X VALUE "COMP-3".', 'data_layout'),
            ('05 ITEMS OCCURS 10 TIMES.', 'data_layout'),
            ('//IN DD DSN=APP.INPUT,DISP=SHR', 'jcl_io'),
            ('//STEP EXEC PGM=UTILITY', 'jcl_control'),
            ('*> EXEC SQL SELECT', 'unknown'), ('000100*PERFORM LOOP', 'unknown'),
            ('//* EXEC PGM=PROGRAM', 'unknown'), ('PERFORM-COUNT.', 'unknown'),
            ('MOVE "PERFORM" TO ACTION', 'unknown'), ('', 'unknown'),
        ]
        for source, expected in examples:
            with self.subTest(source=source):
                self.assertEqual(_classify_construct(source), expected)

    def test_compound_lines_keep_each_observed_construct_without_parsing_literals_as_verbs(self):
        examples = [
            ("PERFORM WORK CALL 'EXTERNAL'", ['iteration', 'procedure_call']),
            ('PERFORM 2 TIMES DISPLAY "tick" END-PERFORM', ['iteration', 'terminal_io']),
            ('READ INPUT-FILE AT END DISPLAY "missing" END-READ', ['file_io', 'terminal_io']),
            ('EXEC SQL CALL MYPROC END-EXEC DISPLAY "SQL done"', ['embedded_sql', 'terminal_io']),
            ('EXEC CICS READ FILE("CALL") END-EXEC DISPLAY "done"', ['cics', 'terminal_io']),
        ]
        for source, observed in examples:
            with self.subTest(source=source):
                process = self.process([source], names=('ELIGIBLE',))
                group = self.groups(packet_document(process))['mixed']
                self.assertEqual(group['observed_constructs'], observed)
                self.assertEqual(group['blocker_count'], 1)
                index = expand_spans(group['blocker_index_spans'])[0]
                original = process['analysis']['blockers'][index]
                self.assertEqual(group['blockers_sha256'], sha(encode([original])))
                self.assertEqual(expand_spans(group['source_line_spans']['ELIGIBLE.cbl']), original['lines'])
        from workbench.review import _construct_families
        self.assertEqual(_construct_families('DISPLAY "CALL PERFORM"'), {'terminal_io'})
        self.assertEqual(_construct_families('EXEC SQL CALL MYPROC END-EXEC'), {'embedded_sql'})
        self.assertEqual(_construct_families('EXEC CICS READ FILE("A") END-EXEC'), {'cics'})
        self.assertEqual(_construct_families('PERFORM CALL-PROCESS DISPLAY-VALUE'), {'iteration'})
        process = self.process(['PERFORM LOOP'], names=('ELIGIBLE',))
        process['analysis']['blockers'][0]['source_refs'] = [{'path': 'missing.cbl', 'line': 0}]
        self.assertIn('unlocated_source', self.groups(packet_document(process)))

    def test_missing_unknown_mixed_and_nontechnical_blockers_remain_explicit(self):
        process = self.process(['PERFORM LOOP', 'EXEC SQL SELECT X FROM T END-EXEC'])
        analysis = process['analysis']; first, second = analysis['blockers'][:2]
        analysis['blockers'] = [
            {'kind': 'missing_source', 'message': 'Missing invocation target', 'program': 'MISSING'},
            {'kind': 'unsupported_source', 'path': 'absent.cbl', 'lines': [5], 'message': 'PERFORM is only prose'},
            dict(first, lines=[999999], message='EXEC SQL is only prose'),
            dict(first, lines=first['lines']+second['lines']),
            {'kind': 'unsupported_jcl', 'path': 'ELIGIBLE.cbl', 'message': 'No exact line evidence'},
            {'kind': 'unresolved_condition', 'message': 'Business calendar question'},
        ]
        packet = packet_document(process); groups = self.groups(packet)
        self.assertEqual(set(groups), {'unlocated_source', 'unknown', 'mixed'})
        self.assertEqual(groups['unlocated_source']['blocker_count'], 2)
        self.assertEqual(groups['unknown']['blocker_count'], 2)
        self.assertEqual(groups['mixed']['observed_constructs'], ['embedded_sql', 'iteration'])
        covered = sorted(i for group in groups.values() for i in expand_spans(group['blocker_index_spans']))
        self.assertEqual(covered, list(range(5)))
        self.assertEqual(next(i for i in packet['items'] if i['kind'] == 'unresolved_item')['id'], 'B_005')

    def test_source_refs_preserve_cross_file_lines_hashes_and_unknown_language(self):
        process = self.process(['PERFORM LOOP'])
        analysis = process['analysis']; refs = []
        for blocker in analysis['blockers']:
            refs.append({'path': blocker['path'], 'line': blocker['lines'][0]})
        analysis['blockers'] = [{'kind': 'unsupported_source', 'message': 'Joint obligation', 'source_refs': refs}]
        packet = packet_document(process); group = self.groups(packet)['iteration']
        self.assertEqual(group['programs'], ['ELIGIBLE', 'SECOND'])
        self.assertEqual(group['blocker_index_spans'], [[0, 0]])
        for ref in refs:
            self.assertEqual(expand_spans(group['source_line_spans'][ref['path']]), [ref['line']])
        original_hash = packet['packet_hash']
        analysis['blockers'][0]['source_refs'][0]['line'] += 1
        self.assertNotEqual(packet_document(process)['packet_hash'], original_hash)
        # An extension/content hint or stale source cannot certify the family.
        analysis['programs']['ELIGIBLE']['kind'] = 'other_source'
        self.assertIn('unknown', self.groups(packet_document(process)))
        analysis['programs']['ELIGIBLE']['kind'] = 'cobol_program'
        analysis['programs']['ELIGIBLE']['source_text'] += '\n'
        self.assertIn('unknown', self.groups(packet_document(process)))

    def test_long_program_display_truncates_only_between_names_and_context_keeps_all(self):
        names = tuple('PROGRAM-'+str(i).zfill(3)+'-LONG-NAME' for i in range(25))
        process = self.process(['PERFORM LOOP'], names=names)
        packet = packet_document(process); group = self.groups(packet)['iteration']
        self.assertEqual(group['programs'], sorted(names))
        question = next(i['question'] for i in packet['items'] if i['kind'] == 'technical_gap')
        displayed = question.split('Programs: ', 1)[1].split(' (+', 1)[0].split(', ')
        self.assertTrue(set(displayed) < set(names))
        self.assertIn(f'(+{25-len(displayed)} more; 25 programs total)', question)
        self.assertEqual(group['program_count'], 25)

    def test_v3_business_rules_and_other_questions_still_obey_human_item_cap(self):
        process = self.process(['PERFORM LOOP'], names=('ELIGIBLE',))
        rule = process['analysis']['rules'][0]
        process['analysis']['rules'] = [dict(rule, id='R%04d'%i) for i in range(1994)]
        self.assertEqual(len(packet_document(process)['items']), 2000)
        process['analysis']['blockers'].append({'kind': 'unresolved_condition', 'message': 'One extra real obligation'})
        with self.assertRaisesRegex(ValidationError, 'SME packet exceeds supported size'):
            packet_document(process)

    def test_v3_export_is_idempotent_and_frozen_context_cannot_be_changed(self):
        process = self.process(['PERFORM LOOP'])
        with tempfile.TemporaryDirectory() as folder:
            directory = Path(folder)/'review'
            packet = export_packet(process, directory)
            issued = (directory/'sme-checklist.xlsx').read_bytes()
            self.assertEqual(export_packet(process, directory), packet)
            self.assertEqual((directory/'sme-checklist.xlsx').read_bytes(), issued)
            book = load_workbook(BytesIO(issued))
            book['Context']['A2'] = 'Changed family membership'
            output = BytesIO(); book.save(output); book.close()
            with self.assertRaisesRegex(ValidationError, 'Frozen Context changed'):
                read_answers(output.getvalue(), packet, 'Test fixture reviewer')

    def test_new_coordinator_intake_uses_v3_and_yes_never_clears_unsupported_semantics(self):
        from workbench.coordinator import Coordinator
        from workbench.domain import decode
        source = COBOL.replace('  GOBACK.', 'PERFORM WORK\n  GOBACK.')
        with tempfile.TemporaryDirectory() as folder:
            coordinator = Coordinator(Path(folder))
            try:
                process = coordinator.create(MANIFEST, {'ELIGIBLE.cbl': source})
                self.assertEqual(process['sme_packet_version'], 3)
                coordinator.start('process-a'); coordinator.advance('process-a')
                packet = decode(coordinator.artifact('process-a', 'review/packet.json').read_bytes())
                self.assertEqual(packet['version'], 3)
                self.assertIn('iteration', self.groups(packet))
                analysis_path = coordinator.artifact('process-a', 'analysis/source-analysis.json')
                analysis_bytes = analysis_path.read_bytes()
                self.assertEqual(json.loads(packet['context'])['source_analysis']['sha256'], sha(analysis_bytes))
                book = load_workbook(coordinator.artifact('process-a', 'review/sme-checklist.xlsx'))
                for row in book['Checklist'].iter_rows(min_row=2):
                    row[4].value = 'Yes'
                output = BytesIO(); book.save(output); book.close()
                coordinator.import_answers('process-a', output.getvalue(), 'Test fixture reviewer')
                coordinator.advance('process-a')
                verified = coordinator.ledger.get('process-a')
                self.assertTrue(any(b['kind'] == 'unsupported_source' for b in verified['blockers']))
                self.assertEqual(verified['program_versions'], {})
                self.assertEqual(analysis_path.read_bytes(), analysis_bytes)
                self.assertTrue(verified['packet_issued'])
                self.assertTrue(verified['packet_imported'])
                self.assertEqual(verified['packet_hash'], packet['packet_hash'])
            finally:
                coordinator.close()


if __name__ == '__main__':
    unittest.main()
