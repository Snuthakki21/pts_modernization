"""New process policy and executable, immutable generated unit evidence."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from workbench.coordinator import Coordinator
from test_source import COBOL
from test_workflow import MANIFEST

class RuntimePolicyTests(unittest.TestCase):
    def test_new_process_pins_sixty_four_and_runtime_seed_once(self):
        with tempfile.TemporaryDirectory() as folder:
            with CoordinatorContext(folder) as c:
                doc=c.create(MANIFEST,{'ELIGIBLE.cbl':COBOL})
                self.assertEqual(doc['logic_validation_min_records'],64)
                self.assertEqual(doc['fixture_contract_version'],5)
                with patch('secrets.randbits',return_value=90123) as random_seed:
                    doc=c.start('process-a')
                    self.assertEqual(doc['authorization']['seed'],90123)
                    random_seed.assert_called_once_with(63)
                c.control('process-a','pause');c.control('process-a','resume')
                self.assertEqual(c.ledger.get('process-a')['authorization']['seed'],90123)
            with CoordinatorContext(folder) as c:
                self.assertEqual(c.ledger.get('process-a')['authorization']['seed'],90123)

class CoordinatorContext:
    def __init__(self,root):self.coordinator=Coordinator(root)
    def __enter__(self):return self.coordinator
    def __exit__(self,*args):self.coordinator.close()

class GeneratedUnitTests(unittest.TestCase):
    def test_export_executes_each_frozen_case_and_detects_target_tampering(self):
        from workbench.unit_evidence import generate_unit_tests, run_unit_tests
        from workbench.source import analyze_sources
        from workbench.intake import parse_manifest
        from workbench.fixtures import plan_cases
        from workbench.target import emit_program
        from workbench.domain import encode,sha
        p=analyze_sources({'ELIGIBLE.cbl':COBOL},parse_manifest(MANIFEST))['programs']['ELIGIBLE']
        p['target_contract_version']=2
        suite=plan_cases(p,812,budget=4096,min_records_per_logic=20,fixture_contract_version=4)
        code=emit_program(p)
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder);path=base/'tests/run-0001/ELIGIBLE/test_generated.py'
            path.parent.mkdir(parents=True)
            expected=base/'synthetic/run-0001/ELIGIBLE/expected.json';expected.parent.mkdir(parents=True);expected.write_bytes(encode(suite))
            target=base/'target/run-0001/ELIGIBLE.py';target.parent.mkdir(parents=True);target.write_bytes(code.encode())
            script=generate_unit_tests(p,suite,sha(code));path.write_text(script)
            result=run_unit_tests(script,path)
            self.assertTrue(result['passed'],result)
            self.assertEqual(result['tests_run'],len(suite['cases']))
            self.assertEqual(result['skipped'],0)
            import subprocess,sys
            standalone=subprocess.run([sys.executable,str(path)],capture_output=True,text=True)
            self.assertEqual(standalone.returncode,0,standalone.stderr)
            # Even when bytes are correctly pinned, wrong target behavior must fail.
            mutated=code.replace("row['DECISION'] = 'Y'", "row['DECISION'] = 'N'")
            self.assertNotEqual(code,mutated)
            target.write_bytes(mutated.encode())
            failed=run_unit_tests(generate_unit_tests(p,suite,sha(mutated)),path)
            self.assertFalse(failed['passed']);self.assertGreater(failed['failures'],0)
            wrong_type=code.replace("'return_code': 0", "'return_code': False")
            target.write_bytes(wrong_type.encode())
            self.assertFalse(run_unit_tests(generate_unit_tests(p,suite,sha(wrong_type)),path)['passed'])
            target.write_bytes((code+'# tampered\n').encode())
            with self.assertRaisesRegex(ValueError,'Target hash'):run_unit_tests(script,path)

class RuntimeEvidenceTests(unittest.TestCase):
    def test_missing_unit_receipt_revokes_report_validation_credit(self):
        from io import BytesIO
        import json
        from openpyxl import load_workbook
        from workbench.coverage import build_coverage
        from workbench.reports import metrics
        from workbench.factory import validation_view
        with tempfile.TemporaryDirectory() as folder:
            with CoordinatorContext(folder) as c:
                c.create(MANIFEST,{'ELIGIBLE.cbl':COBOL,'JOBA.jcl':'//JOBA JOB\n//S010 EXEC PGM=ELIGIBLE\n'},demo=True)
                c.start('process-a');c.advance('process-a')
                book=load_workbook(c.artifact('process-a','review/sme-checklist.xlsx'))
                for row in book['Checklist'].iter_rows(min_row=2):row[4].value='Yes'
                out=BytesIO();book.save(out);book.close()
                c.import_answers('process-a',out.getvalue(),'Fictional test reviewer')
                c.advance('process-a');c.advance('process-a')
                doc=c.ledger.get('process-a');self.assertEqual(doc['status'],'COMPLETED',doc['blockers'])
                self.assertEqual(validation_view(doc)['status'],'RECORDED_PASS')
                (Path(folder)/'processes/process-a/tests/run-0001/ELIGIBLE/unit-results.json').unlink()
                coverage=build_coverage(doc,Path(folder))
                self.assertFalse(coverage['summary']['completion_eligible'])
                result=json.loads(metrics(c.ledger,doc,coverage)['logic_validation_json'])
                self.assertFalse(result['complete'])
                self.assertEqual(result['logic_meeting_minimum'],0)
                self.assertEqual(validation_view(doc,coverage)['status'],'GAPS')
