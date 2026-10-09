"""Exercise generated operator surfaces against the current Claude retrieval policy."""
import tempfile
import unittest
from pathlib import Path
from fastapi.testclient import TestClient
from pptx import Presentation
from workbench.api import create_app
from workbench.deck import framework_deck
from test_workflow import MANIFEST

class ClaudePolicyAlignmentTests(unittest.TestCase):
    def test_retrieval_http_instructions_assign_approved_mcp_to_claude(self):
        with tempfile.TemporaryDirectory() as tmp:
            app=create_app(tmp);c=app.state.coordinator
            try:
                process=c.prepare_process(MANIFEST)
                with TestClient(app,base_url='http://127.0.0.1:8765') as client:
                    response=client.get('/api/process/'+process['id']+'/retrieval')
                    self.assertEqual(response.status_code,200)
                    text=response.json()['instruction']
                    self.assertIn('Claude',text)
                    self.assertIn('approved',text)
                    self.assertNotIn('Copilot',text)
                    self.assertNotIn('without MCP',text)
                    self.assertIn('Raw records stay local',text)
            finally:c.close()

    def test_actual_framework_deck_uses_claude_and_private_data_boundary(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'framework.pptx';framework_deck(path)
            prs=Presentation(path)
            self.assertEqual(len(prs.slides),4)
            text='\n'.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
            self.assertIn('Claude retrieves',text)
            self.assertIn('safe context',text)
            self.assertIn('provider credits',text)
            self.assertNotIn('Copilot',text)
            self.assertIn('Raw records stay local',text)
