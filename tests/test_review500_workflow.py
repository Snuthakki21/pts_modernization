"""R241-R320: bounded adversarial workflow reviews using fictional SME fixtures only."""
import copy
import json
import sqlite3
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook

from workbench.coordinator import Coordinator, StageInterrupted
from workbench.domain import ValidationError, encode, decode, sha
from workbench.intake import parse_manifest
from workbench.knowledge import update_knowledge
from workbench.ledger import Ledger
from workbench.review import export_packet, packet_document, read_answers
from test_source import COBOL
from test_workflow import MANIFEST

FICTIONAL_REVIEWER = 'Fictional review500 fixture reviewer'


class Review500PacketTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.temp.name)
        c = Coordinator(cls.root)
        try:
            c.create(MANIFEST, {'ELIGIBLE.cbl': COBOL}); c.start('process-a'); c.advance('process-a')
            cls.process = c.ledger.get('process-a')
            cls.packet = decode(c.artifact('process-a', 'review/packet.json').read_bytes())
            cls.original = c.artifact('process-a', 'review/sme-checklist.xlsx').read_bytes()
        finally: c.close()

    @classmethod
    def tearDownClass(cls): cls.temp.cleanup()

    def returned(self, edit=None):
        book = load_workbook(BytesIO(self.original))
        try:
            if edit: edit(book)
            out = BytesIO(); book.save(out); return out.getvalue()
        finally: book.close()

    def answer(self, edit=None, reviewer=FICTIONAL_REVIEWER):
        return read_answers(self.returned(edit), self.packet, reviewer)

    def reject(self, edit):
        with self.assertRaises(ValidationError): self.answer(edit)

    def first(self, answers): return answers['items'][self.packet['items'][0]['id']]

    def test_r241_blank_answer_stays_unanswered(self):
        self.assertEqual(self.first(self.answer())['answer'], 'Unanswered')

    def test_r242_zero_correction_is_retained(self):
        answer = self.first(self.answer(lambda b: setattr(b['Checklist']['F2'], 'value', 0)))
        self.assertEqual(answer['correction'], '0')

    def test_r243_false_correction_is_retained(self):
        answer = self.first(self.answer(lambda b: setattr(b['Checklist']['F2'], 'value', False)))
        self.assertEqual(answer['correction'], 'False')

    def test_r244_numeric_zero_answer_is_not_blank(self):
        self.reject(lambda b: setattr(b['Checklist']['E2'], 'value', 0))

    def test_r245_boolean_answer_is_not_blank(self):
        self.reject(lambda b: setattr(b['Checklist']['E2'], 'value', False))

    def test_r246_numeric_reviewer_is_not_silently_replaced(self):
        self.reject(lambda b: setattr(b['Checklist']['G2'], 'value', 0))

    def test_r247_empty_global_reviewer_rejected(self):
        with self.assertRaises(ValidationError): self.answer(reviewer='  ')

    def test_r248_oversized_reviewer_rejected(self):
        with self.assertRaises(ValidationError): self.answer(reviewer='x' * 161)

    def test_r249_formula_answer_rejected(self):
        self.reject(lambda b: setattr(b['Checklist']['E2'], 'value', '=1'))

    def test_r250_formula_correction_rejected(self):
        self.reject(lambda b: setattr(b['Checklist']['F2'], 'value', '=WEBSERVICE("https://example.invalid")'))

    def test_r251_frozen_context_edit_rejected(self):
        self.reject(lambda b: setattr(b['Context']['A2'], 'value', 'different source meaning'))

    def test_r252_packet_metadata_process_edit_rejected(self):
        self.reject(lambda b: setattr(b['Metadata']['B1'], 'value', 'process-b'))

    def test_r253_repeated_question_id_rejected(self):
        self.reject(lambda b: setattr(b['Checklist']['A3'], 'value', b['Checklist']['A2'].value))

    def test_r254_missing_question_row_rejected(self):
        self.reject(lambda b: b['Checklist'].delete_rows(2))

    def test_r255_extra_hidden_question_rejected(self):
        def edit(b):
            s = b['Checklist']; s.append(['FORGED', 'invented', 'none', 'business_rule', 'Yes', '', '']); s.row_dimensions[s.max_row].hidden = True
        self.reject(edit)

    def test_r256_extra_hidden_sheet_rejected(self):
        def edit(b): b.create_sheet('Hidden answers').sheet_state = 'hidden'
        self.reject(edit)

    def test_r257_evidence_reference_edit_rejected(self):
        self.reject(lambda b: setattr(b['Checklist']['C2'], 'value', 'different.cbl:1'))

    def test_r258_rule_category_edit_rejected(self):
        self.reject(lambda b: setattr(b['Checklist']['D2'], 'value', 'process_assumption'))

    def test_r259_invalid_answer_vocabulary_rejected(self):
        self.reject(lambda b: setattr(b['Checklist']['E2'], 'value', 'Approved'))

    def test_r260_no_answer_correction_preserved(self):
        def edit(b): b['Checklist']['E2'] = 'No'; b['Checklist']['F2'] = 'Use zero as the threshold'
        answer = self.first(self.answer(edit)); self.assertEqual(answer['answer'], 'No'); self.assertEqual(answer['correction'], 'Use zero as the threshold')

    def test_r261_uncertainty_preserved(self):
        self.assertEqual(self.first(self.answer(lambda b: setattr(b['Checklist']['E2'], 'value', 'Not sure')))['answer'], 'Not sure')

    def test_r262_oversized_correction_rejected(self):
        self.reject(lambda b: setattr(b['Checklist']['F2'], 'value', 'x' * 4001))

    def test_r263_merged_answer_and_correction_rejected(self):
        self.reject(lambda b: b['Checklist'].merge_cells('E2:F2'))

    def test_r264_changed_packet_payload_with_old_hash_rejected(self):
        packet = copy.deepcopy(self.packet); packet['context'] += 'forged'
        with self.assertRaises(ValidationError): read_answers(self.original, packet, FICTIONAL_REVIEWER)

    def test_r265_reordered_rows_keep_identity_binding(self):
        def edit(b):
            s = b['Checklist']; a = [x.value for x in s[2]]; d = [x.value for x in s[3]]
            for i, value in enumerate(d, 1): s.cell(2, i, value)
            for i, value in enumerate(a, 1): s.cell(3, i, value)
        self.assertEqual(self.answer(edit)['items'], self.answer()['items'])

    def test_r266_duplicate_packet_rule_ids_rejected_before_publication(self):
        p = copy.deepcopy(self.process); p['analysis']['rules'].append(copy.deepcopy(p['analysis']['rules'][0]))
        with self.assertRaises(ValidationError): packet_document(p)

    def test_r267_rule_id_colliding_with_global_question_rejected(self):
        p = copy.deepcopy(self.process); p['analysis']['rules'][0]['id'] = 'G_SCOPE'
        with self.assertRaises(ValidationError): packet_document(p)

    def test_r268_recovery_rejects_forged_packet_self_hash(self):
        with tempfile.TemporaryDirectory() as t:
            folder = Path(t) / 'review'; export_packet(self.process, folder)
            doc = decode((folder / 'packet.json').read_bytes()); doc['context'] += 'forged'; (folder / 'packet.json').write_bytes(encode(doc))
            with self.assertRaises(ValidationError): export_packet(self.process, folder)

    def test_r269_recovery_rejects_missing_companion(self):
        with tempfile.TemporaryDirectory() as t:
            folder = Path(t) / 'review'; export_packet(self.process, folder); (folder / 'sme-checklist.docx').unlink()
            with self.assertRaises(ValidationError): export_packet(self.process, folder)

    def test_r270_recovery_rejects_answered_issued_workbook(self):
        with tempfile.TemporaryDirectory() as t:
            folder = Path(t) / 'review'; export_packet(self.process, folder)
            (folder / 'sme-checklist.xlsx').write_bytes(self.returned(lambda b: setattr(b['Checklist']['E2'], 'value', 'Yes')))
            with self.assertRaises(ValidationError): export_packet(self.process, folder)


class Review500LedgerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name); self.ledger = Ledger(self.root); self.addCleanup(lambda: self.ledger.close())
        self.doc = self.ledger.create(parse_manifest(MANIFEST)); self.doc['analysis'] = {'source_snapshot': 's' * 64}; self.ledger.save(self.doc)
        self.answers = {'packet_hash': 'p' * 64, 'source_snapshot': 's' * 64, 'reviewer': FICTIONAL_REVIEWER, 'return_hash': 'r' * 64, 'items': {}}

    def issue(self): self.ledger.issue_packet('process-a', self.answers['packet_hash'])
    def consume(self): self.ledger.consume_return('process-a', self.answers)
    def accepted_doc(self):
        doc = self.ledger.get('process-a'); doc.update(report_verified=True, report_hashes={'reports/report-0001/metrics.json': 'h'}, artifact_hashes={'reports/report-0001/metrics.json': 'h'})
        self.ledger.save(doc, 'REPORTING'); return doc

    def test_r271_double_packet_issue_keeps_original_fingerprint(self):
        self.issue()
        with self.assertRaises(ValidationError): self.ledger.issue_packet('process-a', 'new')
        self.assertEqual(self.ledger.get('process-a')['packet_hash'], 'p' * 64)

    def test_r272_return_requires_issued_packet(self):
        with self.assertRaises(ValidationError): self.consume()
        self.assertFalse(self.ledger.get('process-a')['packet_imported'])

    def test_r273_return_packet_fingerprint_must_match(self):
        self.issue(); self.answers['packet_hash'] = 'wrong'
        with self.assertRaises(ValidationError): self.consume()
        self.assertFalse(self.ledger.get('process-a')['packet_imported'])

    def test_r274_return_source_snapshot_must_match(self):
        self.issue(); self.answers['source_snapshot'] = 'wrong'
        with self.assertRaises(ValidationError): self.consume()

    def test_r275_return_reviewer_required_at_commit(self):
        self.issue(); self.answers['reviewer'] = ' '
        with self.assertRaises(ValidationError): self.consume()

    def test_r276_return_byte_fingerprint_required(self):
        self.issue(); self.answers.pop('return_hash')
        with self.assertRaises(ValidationError): self.consume()

    def test_r277_double_return_keeps_first_answer(self):
        self.issue(); self.consume(); first = self.ledger.get('process-a')['answers']
        with self.assertRaises(ValidationError): self.consume()
        self.assertEqual(self.ledger.get('process-a')['answers'], first)

    def test_r278_return_transaction_rolls_back_failed_process_update(self):
        self.issue(); self.ledger.db.executescript("CREATE TRIGGER reject_return BEFORE UPDATE ON processes WHEN NEW.packet_imported=1 BEGIN SELECT RAISE(ABORT,'fixture crash'); END;")
        with self.assertRaises(sqlite3.IntegrityError): self.consume()
        self.assertFalse(self.ledger.get('process-a')['packet_imported']); self.assertIsNone(self.ledger.get('process-a')['answers'])

    def test_r279_document_save_cannot_reset_quota_columns(self):
        self.issue(); self.consume(); doc = self.ledger.get('process-a'); doc.update(packet_imported=False, packet_issued=False, packet_hash='forged'); self.ledger.save(doc)
        actual = self.ledger.get('process-a'); self.assertTrue(actual['packet_imported']); self.assertEqual(actual['packet_hash'], 'p' * 64)

    def test_r280_identical_snapshots_are_deduplicated(self):
        self.ledger.snapshot('process-a', {'count': 1}); self.ledger.snapshot('process-a', {'count': 1}); self.assertEqual(len(self.ledger.history('process-a')), 1)

    def test_r281_changed_snapshots_preserve_history(self):
        self.ledger.snapshot('process-a', {'count': 1}); self.ledger.snapshot('process-a', {'count': 2}); self.assertEqual([h['document']['count'] for h in self.ledger.history('process-a')], [1, 2])

    def test_r282_completion_rejects_missing_inspection(self):
        doc = self.accepted_doc(); doc['report_verified'] = False
        with self.assertRaises(ValidationError): self.ledger.complete_report(doc, {})
        self.assertEqual(self.ledger.history('process-a'), [])

    def test_r283_completion_rejects_unpinned_report(self):
        doc = self.accepted_doc(); doc['artifact_hashes'] = {}
        with self.assertRaises(ValidationError): self.ledger.complete_report(doc, {})

    def test_r284_completion_rejects_stale_nonreporting_stage(self):
        doc = self.accepted_doc(); self.ledger.save(doc, 'PAUSED')
        with self.assertRaises(ValidationError): self.ledger.complete_report(doc, {})
        self.assertEqual(self.ledger.get('process-a')['status'], 'PAUSED')

    def test_r285_completion_snapshot_and_event_rollback_together(self):
        doc = self.accepted_doc(); self.ledger.db.executescript("CREATE TRIGGER reject_event BEFORE INSERT ON events BEGIN SELECT RAISE(ABORT,'fixture crash'); END;")
        with self.assertRaises(sqlite3.IntegrityError): self.ledger.complete_report(doc, {'count': 1})
        self.assertEqual(self.ledger.history('process-a'), []); self.assertEqual(self.ledger.get('process-a')['status'], 'REPORTING')

    def test_r286_completion_keeps_blocker_status(self):
        doc = self.accepted_doc(); doc['blockers'] = [{'kind': 'unsupported', 'message': 'fixture'}]
        self.assertEqual(self.ledger.complete_report(doc, {})['status'], 'COMPLETED_WITH_BLOCKERS')

    def test_r287_asset_identity_collision_rejected(self):
        asset = {'id': 'asset', 'kind': 'cobol', 'name': 'A', 'source_hash': 'one'}; self.ledger.register_assets('process-a', [asset])
        self.ledger.register_assets('process-a', [{**asset, 'path': 'different/source/location', 'in_scope': False}])
        with self.assertRaises(ValidationError): self.ledger.register_assets('process-a', [{**asset, 'source_hash': 'two'}])
        self.assertEqual(self.ledger.assets('process-a')[0]['source_hash'], 'one')

    def test_r288_asset_batch_rolls_back_prior_insert_on_collision(self):
        asset = {'id': 'asset', 'kind': 'cobol', 'name': 'A', 'source_hash': 'one'}; self.ledger.register_assets('process-a', [asset])
        with self.assertRaises(ValidationError): self.ledger.register_assets('process-a', [{**asset, 'id': 'new'}, {**asset, 'source_hash': 'two'}])
        self.assertEqual(len(self.ledger.assets('process-a')), 1)

    def test_r289_portfolio_excludes_demo_membership(self):
        d = parse_manifest(MANIFEST.replace('process-a', 'demo')); self.ledger.create(d, demo=True)
        self.ledger.register_assets('demo', [{'id': 'demo-asset', 'kind': 'cobol', 'name': 'D', 'source_hash': 'd'}])
        self.assertEqual(self.ledger.assets(), []); self.assertEqual(len(self.ledger.assets(include_demo=True)), 1)

    def test_r290_durable_controls_do_not_drop_cancellation(self):
        doc = self.ledger.get('process-a'); doc.update(cancel_requested=True, control_revision=4); self.ledger.save(doc, 'QUEUED_REPORT')
        controls = self.ledger.controls('process-a'); self.assertTrue(controls['cancel_requested']); self.assertEqual(controls['control_revision'], 4); self.assertEqual(controls['blockers'][0]['kind'], 'cancelled')


class CoordinatorFixture(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name); self.c = Coordinator(self.root); self.addCleanup(lambda: self.c.close())
        self.c.create(MANIFEST, {'ELIGIBLE.cbl': COBOL})

    def prepare(self, response='Yes', correction='', reviewer=FICTIONAL_REVIEWER, pid='process-a'):
        self.c.start(pid); self.c.advance(pid)
        book = load_workbook(BytesIO(self.c.artifact(pid, 'review/sme-checklist.xlsx').read_bytes()))
        for row in book['Checklist'].iter_rows(min_row=2): row[4].value = response; row[5].value = correction
        out = BytesIO(); book.save(out); book.close(); raw = out.getvalue()
        self.c.import_answers(pid, raw, reviewer); return self.c.ledger.get(pid)


class Review500CoordinatorTests(CoordinatorFixture):
    def test_r291_changed_source_before_return_does_not_consume_quota(self):
        self.c.start('process-a'); self.c.advance('process-a'); raw = self.c.artifact('process-a', 'review/sme-checklist.xlsx').read_bytes()
        (self.c.process_root('process-a') / 'input/sources/ELIGIBLE.cbl').write_text(COBOL + '\nchanged')
        with self.assertRaises(ValidationError): self.c.import_answers('process-a', raw, FICTIONAL_REVIEWER)
        self.assertFalse(self.c.ledger.get('process-a')['packet_imported'])

    def test_r292_added_frozen_source_file_cannot_disappear_from_accounting(self):
        (self.c.process_root('process-a') / 'input/sources/UNRECORDED.txt').write_text('unaccounted export')
        with self.assertRaises(ValidationError): self.c.sources(self.c.ledger.get('process-a'))

    def test_r293_deleted_source_file_remains_integrity_failure(self):
        (self.c.process_root('process-a') / 'input/sources/ELIGIBLE.cbl').unlink()
        with self.assertRaises((ValidationError, FileNotFoundError)): self.c.sources(self.c.ledger.get('process-a'))

    def test_r294_foreign_packet_process_identity_cannot_consume_return(self):
        self.c.start('process-a'); self.c.advance('process-a')
        packet_path = self.c.process_root('process-a') / 'review/packet.json'; packet = decode(packet_path.read_bytes()); packet['process_id'] = 'foreign'
        packet['packet_hash'] = sha(encode({k: v for k, v in packet.items() if k != 'packet_hash'})); packet_path.write_bytes(encode(packet))
        doc = self.c.ledger.get('process-a'); doc['artifact_hashes']['review/packet.json'] = sha(packet_path.read_bytes()); self.c.ledger.save(doc)
        self.c.ledger.db.execute('UPDATE processes SET packet_hash=? WHERE id=?', (packet['packet_hash'], 'process-a')); self.c.ledger.db.commit()
        with patch('workbench.coordinator.read_answers', return_value={'packet_hash': packet['packet_hash'], 'source_snapshot': doc['analysis']['source_snapshot'], 'reviewer': FICTIONAL_REVIEWER, 'return_hash': sha(b'fictional'), 'items': {}}):
            with self.assertRaises(ValidationError): self.c.import_answers('process-a', b'fictional', FICTIONAL_REVIEWER)
        self.assertFalse(self.c.ledger.get('process-a')['packet_imported'])

    def test_r295_pause_checkpoint_keeps_durable_pause(self):
        doc = self.c.start('process-a'); doc['_active_stage'] = 'QUEUED_ANALYSIS'; self.c.control('process-a', 'pause')
        with self.assertRaises(StageInterrupted): self.c.checkpoint(doc, 'ANALYZING')
        self.assertEqual(self.c.ledger.get('process-a')['status'], 'PAUSED')

    def test_r296_cancel_checkpoint_preserves_blocker(self):
        doc = self.c.start('process-a'); doc['_active_stage'] = 'QUEUED_ANALYSIS'; self.c.control('process-a', 'cancel')
        with self.assertRaises(StageInterrupted): self.c.checkpoint(doc)
        current = self.c.ledger.get('process-a'); self.assertTrue(current['cancel_requested']); self.assertTrue(any(x['kind'] == 'cancelled' for x in current['blockers']))

    def test_r297_retry_bound_stops_after_three_transient_failures(self):
        self.c.start('process-a')
        with patch.object(self.c, 'analyze', side_effect=OSError('fictional transient')) as stage:
            for _ in range(3):
                doc = self.c.ledger.get('process-a'); doc['next_retry_at'] = 0; self.c.ledger.save(doc)
                with self.assertRaises(OSError): self.c.advance('process-a')
            self.c.advance('process-a'); self.assertEqual(stage.call_count, 3)
        self.assertEqual(self.c.ledger.get('process-a')['status'], 'FAILED')

    def test_r298_permission_error_is_not_retried(self):
        self.c.start('process-a')
        with patch.object(self.c, 'analyze', side_effect=PermissionError('fictional denied')):
            with self.assertRaises(PermissionError): self.c.advance('process-a')
        self.assertEqual(self.c.ledger.get('process-a')['status'], 'FAILED')

    def test_r299_stage_success_only_clears_its_failure(self):
        doc = self.c.ledger.get('process-a'); doc['blockers'] = [{'kind': 'stage_failure', 'stage': 'QUEUED_VERIFY'}, {'kind': 'stage_failure', 'stage': 'QUEUED_REPORT'}, {'kind': 'sme_unresolved'}]
        self.c.stage_success(doc, 'QUEUED_VERIFY'); self.assertEqual(len(doc['blockers']), 2); self.assertEqual(doc['blockers'][0]['stage'], 'QUEUED_REPORT')

    def test_r300_preserved_return_conflict_does_not_consume_quota(self):
        self.c.start('process-a'); self.c.advance('process-a'); raw = self.c.artifact('process-a', 'review/sme-checklist.xlsx').read_bytes()
        (self.c.process_root('process-a') / 'input/sme-return.xlsx').write_bytes(b'first preserved fixture')
        with self.assertRaises(ValidationError): self.c.import_answers('process-a', raw, FICTIONAL_REVIEWER)
        self.assertFalse(self.c.ledger.get('process-a')['packet_imported'])

    def test_r301_return_write_crash_replay_consumes_only_once(self):
        self.c.start('process-a'); self.c.advance('process-a'); raw = self.c.artifact('process-a', 'review/sme-checklist.xlsx').read_bytes()
        with patch.object(self.c.ledger, 'consume_return', side_effect=OSError('fictional crash')):
            with self.assertRaises(OSError): self.c.import_answers('process-a', raw, FICTIONAL_REVIEWER)
        self.assertFalse(self.c.ledger.get('process-a')['packet_imported']); self.assertEqual((self.c.process_root('process-a') / 'input/sme-return.xlsx').read_bytes(), raw)
        self.c.import_answers('process-a', raw, FICTIONAL_REVIEWER); self.assertTrue(self.c.ledger.get('process-a')['packet_imported'])

    def test_r302_registered_artifact_cannot_be_rebaselined(self):
        self.c.start('process-a'); self.c.advance('process-a'); doc = self.c.ledger.get('process-a'); path = self.c.process_root('process-a') / 'review/sme-checklist.html'; path.write_text('forged')
        with self.assertRaises(ValidationError): self.c.register(doc, 'review/sme-checklist.html')

    def test_r303_paused_return_is_not_imported(self):
        self.c.start('process-a'); self.c.advance('process-a'); raw = self.c.artifact('process-a', 'review/sme-checklist.xlsx').read_bytes(); self.c.control('process-a', 'pause')
        with self.assertRaises(ValidationError): self.c.import_answers('process-a', raw, FICTIONAL_REVIEWER)
        self.assertFalse(self.c.ledger.get('process-a')['packet_imported'])

    def test_r304_completed_controls_cannot_rewrite_terminal_evidence(self):
        doc = self.c.ledger.get('process-a'); self.c.ledger.save(doc, 'COMPLETED_WITH_BLOCKERS')
        with self.assertRaises(ValidationError): self.c.control('process-a', 'cancel')
        self.assertEqual(self.c.ledger.get('process-a')['status'], 'COMPLETED_WITH_BLOCKERS')

    def test_r305_restart_interrupted_analysis_preserves_issued_packet(self):
        self.c.start('process-a'); self.c.advance('process-a'); before = self.c.ledger.get('process-a'); self.c.ledger.save(before, 'ANALYZING'); self.c.close(); self.c = Coordinator(self.root)
        after = self.c.ledger.get('process-a'); self.assertEqual(after['status'], 'WAITING_SME'); self.assertEqual(after['packet_hash'], before['packet_hash']); self.assertTrue(after['packet_issued'])


class Review500KnowledgeTests(CoordinatorFixture):
    def records(self):
        return [json.loads(row[0]) for row in self.c.ledger.db.execute('SELECT document FROM knowledge ORDER BY id')]

    def test_r306_unconsumed_hypothetical_answers_not_published(self):
        self.c.start('process-a'); self.c.advance('process-a'); doc = self.c.ledger.get('process-a')
        doc['answers'] = {'items': {r['id']: {'answer': 'Yes', 'correction': '', 'reviewer': FICTIONAL_REVIEWER} for r in doc['analysis']['rules']}}
        with self.assertRaises(ValidationError): update_knowledge(self.c.ledger, doc)
        self.assertEqual(self.records(), [])

    def test_r307_no_answer_never_becomes_approved_knowledge(self):
        doc = self.prepare(response='No'); self.assertEqual(update_knowledge(self.c.ledger, doc), 0)

    def test_r308_uncertain_answer_never_becomes_approved_knowledge(self):
        doc = self.prepare(response='Not sure'); self.assertEqual(update_knowledge(self.c.ledger, doc), 0)

    def test_r309_blank_answer_never_becomes_approved_knowledge(self):
        doc = self.prepare(response=''); self.assertEqual(update_knowledge(self.c.ledger, doc), 0)

    def test_r310_yes_with_correction_never_becomes_approved_knowledge(self):
        doc = self.prepare(correction='Fictional corrected threshold'); self.assertEqual(update_knowledge(self.c.ledger, doc), 0)

    def test_r311_yes_rule_keeps_reviewer_packet_and_source_provenance(self):
        doc = self.prepare(); self.assertGreater(update_knowledge(self.c.ledger, doc), 0)
        record = self.records()[0]; self.assertEqual(record['reviewer'], FICTIONAL_REVIEWER); self.assertEqual(record['packet_hash'], doc['packet_hash']); self.assertEqual(record['applicability']['source_snapshot'], doc['analysis']['source_snapshot'])

    def test_r312_same_source_second_process_keeps_first_review_provenance(self):
        first = self.prepare(); update_knowledge(self.c.ledger, first); before = copy.deepcopy(self.records())
        self.c.create(MANIFEST.replace('process-a', 'process-b'), {'ELIGIBLE.cbl': COBOL}); second = self.prepare(pid='process-b', reviewer='Second fictional fixture reviewer'); update_knowledge(self.c.ledger, second)
        after = self.records(); self.assertTrue(all(r in after for r in before)); self.assertEqual(len(after), 2 * len(before))

    def test_r313_knowledge_replay_does_not_duplicate_records(self):
        doc = self.prepare(); count = update_knowledge(self.c.ledger, doc); self.assertEqual(update_knowledge(self.c.ledger, doc), count)

    def test_r314_changed_inmemory_answer_not_published(self):
        doc = self.prepare(response='No')
        for item in doc['answers']['items'].values(): item['answer'] = 'Yes'
        with self.assertRaises(ValidationError): update_knowledge(self.c.ledger, doc)
        self.assertEqual(self.records(), [])

    def test_r315_changed_rule_statement_not_published(self):
        doc = self.prepare(); doc['analysis']['rules'][0]['plain'] = 'Forged new behavior'
        with self.assertRaises(ValidationError): update_knowledge(self.c.ledger, doc)
        self.assertEqual(self.records(), [])

    def test_r316_changed_accepted_return_not_published(self):
        doc = self.prepare(); (self.c.process_root('process-a') / 'input/sme-return.xlsx').write_bytes(b'tampered')
        with self.assertRaises(ValidationError): update_knowledge(self.c.ledger, doc)
        self.assertEqual(self.records(), [])

    def test_r317_records_projection_matches_canonical_database(self):
        doc = self.prepare(); update_knowledge(self.c.ledger, doc)
        self.assertEqual(decode((self.root / 'knowledge/records.json').read_bytes()), self.records())

    def test_r318_index_symlink_cannot_overwrite_unrelated_file(self):
        doc = self.prepare(); victim = self.root / '.implementation/tmp/victim.txt'; victim.parent.mkdir(parents=True, exist_ok=True); victim.write_text('preserve')
        (self.root / 'knowledge').mkdir(exist_ok=True); (self.root / 'knowledge/INDEX.md').symlink_to(victim)
        with self.assertRaises(ValidationError): update_knowledge(self.c.ledger, doc)
        self.assertEqual(victim.read_text(), 'preserve'); self.assertEqual(self.records(), [])

    def test_r319_knowledge_projection_failure_is_replayable(self):
        doc = self.prepare()
        with patch('workbench.knowledge.atomic_json', side_effect=OSError('fictional projection failure')):
            with self.assertRaises(OSError): update_knowledge(self.c.ledger, doc)
        count = update_knowledge(self.c.ledger, doc); self.assertEqual(count, len(doc['analysis']['rules'])); self.assertEqual(decode((self.root / 'knowledge/records.json').read_bytes()), self.records())

    def test_r320_conflicting_existing_knowledge_is_not_overwritten(self):
        doc = self.prepare(); update_knowledge(self.c.ledger, doc); record = self.records()[0]; record['statement'] = 'Conflicting historical statement'
        self.c.ledger.db.execute('UPDATE knowledge SET document=? WHERE id=?', (encode(record).decode(), record['id'])); self.c.ledger.db.commit()
        with self.assertRaises(ValidationError): update_knowledge(self.c.ledger, doc)
        self.assertIn(record, self.records())


if __name__ == '__main__': unittest.main()
