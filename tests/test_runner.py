import json
import os
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook
from test_source import COBOL
from test_workflow import MANIFEST


class RunnerTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        (self.root / 'Endeavor').mkdir()
        (self.root / 'Endeavor/ELIGIBLE.cbl').write_text(COBOL)
        self.manifest = self.root / 'process-input.md'; self.manifest.write_text(MANIFEST)

    def command(self, *args, expected=0):
        result = subprocess.run([sys.executable, '-m', 'workbench.runner', *args,
                                 '--workspace', str(self.root)], capture_output=True, text=True,
                                timeout=30, env={**os.environ, 'WB_LLM_ENDPOINT':'', 'WB_ZOWE_PROFILE':'', 'WB_DB2_MCP_URL':''})
        self.assertEqual(result.returncode, expected, result.stdout + result.stderr)
        return json.loads(result.stdout)

    def start(self):
        return self.command('run', '--manifest', str(self.manifest), '--timeout', '15')

    def returned_file(self):
        path = self.root / 'processes/process-a/review/sme-checklist.xlsx'
        book = load_workbook(path)
        for row in range(2, book['Checklist'].max_row + 1):
            book['Checklist'].cell(row, 5, 'Yes')
            book['Checklist'].cell(row, 7, 'Example reviewer')
        out = BytesIO(); book.save(out); book.close()
        inbox = self.root / 'processes/process-a/input/sme-return-inbox.xlsx'
        inbox.write_bytes(out.getvalue()); return inbox

    def test_start_is_idempotent_and_never_manufactures_answers(self):
        first = self.start(); second = self.start()
        self.assertEqual(first['status'], 'WAITING_SME')
        self.assertEqual(first['packet_hash'], second['packet_hash'])
        self.assertFalse(second['packet_imported'])
        self.assertFalse((self.root / 'processes/process-a/input/sme-return.xlsx').exists())
        self.assertEqual(len(list((self.root / 'processes/process-a/review').glob('sme-checklist.xlsx'))), 1)
        self.manifest.write_text(MANIFEST.replace('Referral eligibility', 'Changed name'))
        self.assertIn('manifest', self.command('start', '--manifest', str(self.manifest), expected=2)['error'])

    def test_real_return_continues_and_registered_bundle_is_reused(self):
        self.start(); inbox = self.returned_file()
        final = self.command('resume', 'process-a', '--reviewer', 'Example reviewer', '--timeout', '15')
        self.assertEqual(final['status'], 'COMPLETED')
        self.assertTrue(final['packet_imported'])
        again = self.command('import', 'process-a', '--file', str(inbox), '--reviewer', 'Example reviewer', '--timeout', '15')
        self.assertEqual(again['status'], 'COMPLETED')
        self.assertEqual(self.command('bundle', 'process-a')['bundle'], self.command('bundle', 'process-a')['bundle'])
        self.assertTrue(self.command('report', 'process-a')['reports'])

    def test_watch_requires_reviewer_and_does_not_consume_unattributed_return(self):
        self.command('run', '--manifest', str(self.manifest), '--watch', expected=2)
        self.start(); self.returned_file()
        result = self.command('resume', 'process-a', expected=2)
        self.assertIn('reviewer', result['error'])
        self.assertFalse(self.command('status', 'process-a')['packet_imported'])

    def test_layout_failure_prevents_intake(self):
        (self.root / 'untracked-output.md').write_text('misplaced')
        self.command('run', '--manifest', str(self.manifest), expected=2)
        self.assertFalse((self.root / '.migration/ledger.sqlite').exists())

    def test_watch_imports_a_later_actual_return_and_bounded_timeout_preserves_packet(self):
        first = self.start()
        timed = self.command('resume', 'process-a', '--watch', '--reviewer', 'Example reviewer', '--timeout', '.15', expected=3)
        self.assertTrue(timed['timed_out']); self.assertEqual(timed['packet_hash'], first['packet_hash'])
        writer = threading.Thread(target=lambda: (time.sleep(.3), self.returned_file()))
        writer.start(); self.addCleanup(writer.join)
        final = self.command('resume', 'process-a', '--watch', '--reviewer', 'Example reviewer', '--timeout', '15')
        self.assertEqual(final['status'], 'COMPLETED')
        inbox = self.root / 'processes/process-a/input/sme-return-inbox.xlsx'
        inbox.write_bytes(b'different workbook')
        refused = self.command('import', 'process-a', '--file', str(inbox), '--reviewer', 'Example reviewer', expected=2)
        self.assertIn('already consumed', refused['error'])

    def test_tampered_manifest_snapshot_and_matching_supplied_input_are_refused(self):
        self.start()
        changed = MANIFEST.replace('Referral eligibility', 'Tampered name')
        self.manifest.write_text(changed)
        (self.root / 'processes/process-a/input/process-input.md').write_text(changed)
        result = self.command('start', '--manifest', str(self.manifest), expected=2)
        self.assertIn('manifest', result['error'].lower())

    def test_terminal_bundle_refuses_tampered_manifest(self):
        self.start(); self.returned_file()
        self.command('resume', 'process-a', '--reviewer', 'Example reviewer', '--timeout', '15')
        (self.root / 'processes/process-a/input/process-input.md').write_text(MANIFEST.replace('Referral eligibility', 'Tampered name'))
        result = self.command('bundle', 'process-a', expected=2)
        self.assertIn('manifest', result['error'].lower())
        self.assertFalse(list((self.root / 'processes/process-a/reports').glob('bundle-*.zip')))

    def test_historical_missing_manifest_baseline_is_not_derived_on_start(self):
        from workbench.coordinator import Coordinator
        with_baseline = Coordinator(self.root)
        try:
            doc = with_baseline.create(MANIFEST, {'ELIGIBLE.cbl': COBOL})
            original_quota = (doc['packet_issued'], doc['packet_imported'], doc['packet_hash'])
            original_sources = dict(doc['source_files'])
            doc.pop('manifest_hash', None); with_baseline.ledger.save(doc)
        finally: with_baseline.close()
        result = self.command('start', '--manifest', str(self.manifest), expected=2)
        self.assertIn('manifest hash', result['error'].lower())
        self.assertIn('missing', result['error'].lower())
        c = Coordinator(self.root)
        try:
            doc = c.ledger.get('process-a')
            self.assertNotIn('manifest_hash', doc)
            self.assertEqual(doc['status'], 'QUEUED_REPORT')
            self.assertTrue(any(b['kind'] == 'legacy_missing_baseline' for b in doc['blockers']))
            self.assertEqual((doc['packet_issued'], doc['packet_imported'], doc['packet_hash']), original_quota)
            self.assertEqual(doc['source_files'], original_sources)
            self.assertNotIn('authorization', doc)
            self.assertEqual(doc['runs'], [])
        finally: c.close()
