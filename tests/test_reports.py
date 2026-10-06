import tempfile
import unittest
from pathlib import Path
from pptx import Presentation
import test_workflow
from workbench.reports import metrics, portfolio

class ReportTests(unittest.TestCase):
    setUp=test_workflow.WorkflowTests.setUp
    create=test_workflow.WorkflowTests.create
    answer=test_workflow.WorkflowTests.answer
    def test_complete_only_after_real_ppt_and_snapshot(self):
        self.create();self.c.start('process-a');self.c.advance('process-a')
        self.c.import_answers('process-a',self.answer(),'Example reviewer')
        self.c.advance('process-a');self.c.advance('process-a')
        doc=self.c.ledger.get('process-a')
        self.assertEqual(doc['status'],'COMPLETED')
        deck=self.c.artifact('process-a','reports/report-0001/management.pptx')
        self.assertEqual(len(Presentation(deck).slides),6)
        m=metrics(self.c.ledger,doc)
        self.assertEqual(m['source_programs'],1)
        self.assertEqual(m['rules_verified'],m['rules_documented'])
        self.assertFalse(m['observed_mainframe_parity'])
        self.assertEqual(m['source_applicable_line_verification_percent'], 100.0)
        self.assertTrue(m['source_full_accounting'])
        import json
        for name in ('coverage.json', 'coverage.csv', 'coverage.xlsx', 'coverage.html'):
            self.assertTrue(self.c.artifact('process-a','reports/report-0001/'+name).is_file())
        coverage=json.loads(self.c.artifact('process-a','reports/report-0001/coverage.json').read_text())
        self.assertEqual(coverage['summary']['source_lines'], m['source_accounted_lines'])
        text='\n'.join(c.text for slide in Presentation(deck).slides for sh in slide.shapes if sh.has_table for row in sh.table.rows for c in row.cells)
        self.assertIn('All selected source code LOC', text)
        self.assertIn('Applicable source lines', text)
        self.assertEqual(len(self.c.ledger.history('process-a')),1)

    def test_first_real_report_includes_its_accepted_completion_consistently(self):
        import json
        from unittest.mock import patch
        from openpyxl import load_workbook
        from workbench.reports import generate_reports
        self.create();self.c.start('process-a');self.c.advance('process-a')
        self.c.import_answers('process-a',self.answer(),'Example reviewer');self.c.advance('process-a')
        self.assertEqual(portfolio(self.c.ledger)['completed_processes'],0)
        def inspect_pending(*args,**kwargs):
            paths=generate_reports(*args,**kwargs)
            self.assertEqual(portfolio(self.c.ledger)['completed_processes'],0)
            self.assertEqual(self.c.ledger.history('process-a'),[])
            return paths
        with patch('workbench.reports.generate_reports',side_effect=inspect_pending):self.c.advance('process-a')
        model=json.loads(self.c.artifact('process-a','reports/report-0001/metrics.json').read_text())
        self.assertEqual(model['portfolio']['completed_processes'],1)
        self.assertEqual(portfolio(self.c.ledger)['completed_processes'],1)
        self.assertEqual(model['metrics']['portfolio_completed_processes'],1)
        self.assertEqual(model['metrics']['report_final_status'],'COMPLETED')
        self.assertEqual(self.c.ledger.history('process-a')[0]['document'],model['metrics'])
        deck=Presentation(self.c.artifact('process-a','reports/report-0001/management.pptx'))
        completed=[row.cells[1].text for slide in deck.slides for shape in slide.shapes if shape.has_table for row in shape.table.rows if row.cells[0].text=='Completed without blockers']
        self.assertEqual(completed,['1'])
        book=load_workbook(self.c.artifact('process-a','reports/report-0001/metrics.xlsx'))
        try:
            self.assertEqual(dict(book['Portfolio'].iter_rows(min_row=2,values_only=True))['completed_processes'],1)
            history=list(book['History'].iter_rows(min_row=2,values_only=True))
            self.assertEqual(len(history),1);self.assertEqual(history[0][0],'process-a')
            self.assertEqual(history[0][2],model['metrics']['source_programs'])
            self.assertEqual(history[0][3],model['metrics']['rules_verified'])
        finally:book.close()
        regenerate=self.root/'regenerated'
        generate_reports(self.c.ledger,self.c.ledger.get('process-a'),regenerate)
        self.assertEqual(json.loads((regenerate/'metrics.json').read_text())['portfolio']['completed_processes'],1)
        book=load_workbook(regenerate/'metrics.xlsx')
        try:self.assertEqual(book['History'].max_row,2)
        finally:book.close()
        self.assertEqual(len(self.c.ledger.history('process-a')),1)

    def test_demo_report_never_projects_production_completion(self):
        import json
        from test_source import COBOL
        self.c.create(test_workflow.MANIFEST,{'ELIGIBLE.cbl':COBOL},True)
        self.c.start('process-a');self.c.advance('process-a')
        self.c.import_answers('process-a',self.answer(),'Fictional reviewer')
        self.c.advance('process-a');self.c.advance('process-a')
        model=json.loads(self.c.artifact('process-a','reports/report-0001/metrics.json').read_text())
        self.assertEqual(self.c.ledger.get('process-a')['status'],'COMPLETED')
        self.assertEqual(model['portfolio']['processes'],0)
        self.assertEqual(model['portfolio']['completed_processes'],0)
        self.assertEqual(model['metrics']['portfolio_completed_processes'],0)
        self.assertEqual(portfolio(self.c.ledger)['completed_processes'],0)

    def test_blocked_report_does_not_project_completion(self):
        import json
        self.create();self.c.start('process-a');self.c.advance('process-a')
        self.c.import_answers('process-a',self.answer(response='Not sure'),'Example reviewer')
        self.c.advance('process-a');self.c.advance('process-a')
        model=json.loads(self.c.artifact('process-a','reports/report-0001/metrics.json').read_text())
        self.assertEqual(model['portfolio']['completed_processes'],0)
        self.assertEqual(model['metrics']['report_final_status'],'COMPLETED_WITH_BLOCKERS')
        self.assertEqual(self.c.ledger.history('process-a')[0]['document'],model['metrics'])

    def test_metrics_revoke_credit_after_target_evidence_is_changed(self):
        self.create();self.c.start('process-a');self.c.advance('process-a')
        self.c.import_answers('process-a',self.answer(),'Example reviewer');self.c.advance('process-a')
        doc=self.c.ledger.get('process-a')
        self.assertEqual(metrics(self.c.ledger,doc)['rules_verified'], 2)
        target=self.root/'shared/target/python'/(doc['program_versions']['ELIGIBLE']+'.py')
        target.write_text(target.read_text().replace('>= 18','> 18'))
        model=metrics(self.c.ledger,doc)
        self.assertEqual(model['rules_verified'], 0)
        self.assertEqual(model['source_applicable_line_verification_percent'], 0.0)

    def test_uncertainty_does_not_receive_verified_credit(self):
        self.create();self.c.start('process-a');self.c.advance('process-a')
        self.c.import_answers('process-a',self.answer(response='Not sure'),'Example reviewer')
        self.c.advance('process-a');self.c.advance('process-a')
        doc=self.c.ledger.get('process-a')
        self.assertEqual(doc['status'],'COMPLETED_WITH_BLOCKERS')
        self.assertEqual(metrics(self.c.ledger,doc)['rules_verified'],0)

if __name__=='__main__':unittest.main()
