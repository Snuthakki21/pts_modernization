import tempfile
import unittest
from pathlib import Path
from io import BytesIO
import time
from unittest.mock import patch
from openpyxl import load_workbook
from workbench.coordinator import Coordinator
from workbench.domain import ValidationError
from test_source import COBOL

MANIFEST='''# Process
- Process ID: process-a
- Process name: Referral eligibility
| Job order | Job | Step order | Step | Program or utility | Input files/tables | Output files/tables | Condition or dependency |
|---|---|---|---|---|---|---|---|
| 1 | JOBA | 1 | S010 | ELIGIBLE | INPUT | OUTPUT | Always |
'''


class WorkflowTests(unittest.TestCase):
    def test_failed_verification_resumes_verification_and_cannot_certify_empty_run(self):
        self.create();self.c.start('process-a');self.c.advance('process-a')
        self.c.import_answers('process-a',self.answer(),'Example reviewer')
        with patch('workbench.coordinator.plan_cases',side_effect=OSError('fixture failure')):
            self.c.launch_worker()
            until=time.monotonic()+3
            while self.c.ledger.get('process-a')['status']!='FAILED' and time.monotonic()<until:time.sleep(.02)
            failed=self.c.ledger.get('process-a');self.assertEqual(failed['status'],'FAILED')
            self.assertEqual(failed['resume_status'],'QUEUED_VERIFY')
            self.assertTrue(any(b['kind']=='stage_failure' for b in failed['blockers']))
        self.c.stopped.set();self.c.worker.join();self.c.worker=None
        self.c.control('process-a','resume');self.c.advance('process-a');self.c.advance('process-a')
        doc=self.c.ledger.get('process-a');self.assertEqual(doc['status'],'COMPLETED');self.assertEqual(len(doc['runs']),1)
        self.assertFalse(any(b['kind']=='stage_failure' for b in doc['blockers']))

    def test_yes_with_correction_is_preserved_unresolved_and_uncredited(self):
        from workbench.reports import metrics
        self.create();self.c.start('process-a');self.c.advance('process-a')
        book=load_workbook(BytesIO(self.answer()));rid=book['Checklist']['A2'].value
        book['Checklist']['F2']='Wrong threshold; must be 21';out=BytesIO();book.save(out);book.close()
        self.c.import_answers('process-a',out.getvalue(),'Example reviewer');self.c.advance('process-a');self.c.advance('process-a')
        doc=self.c.ledger.get('process-a');self.assertEqual(doc['status'],'COMPLETED_WITH_BLOCKERS')
        self.assertTrue(any(b.get('item_id')==rid for b in doc['blockers']))
        self.assertLess(metrics(self.c.ledger,doc)['rules_verified'],metrics(self.c.ledger,doc)['rules_documented'])
        import json
        records=json.loads((self.root/'knowledge/records.json').read_text())
        self.assertFalse(any(r['rule_id']==rid for r in records))
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.c=Coordinator(self.root)
        self.addCleanup(lambda:self.c.close())

    def create(self):return self.c.create(MANIFEST,{'ELIGIBLE.cbl':COBOL})

    def answer(self,alter=False,response='Yes'):
        doc=self.c.ledger.get('process-a')
        data=self.c.artifact('process-a','review/sme-checklist.xlsx').read_bytes()
        book=load_workbook(BytesIO(data));sheet=book['Checklist']
        for row in range(2,sheet.max_row+1):sheet.cell(row,5,response);sheet.cell(row,7,'Example reviewer')
        if alter:sheet.cell(2,2,'Changed question')
        out=BytesIO();book.save(out);book.close();return out.getvalue()

    def test_one_packet_survives_restart_and_duplicate_start(self):
        self.create();self.c.start('process-a');self.c.advance('process-a')
        first=self.c.ledger.get('process-a')
        self.assertEqual(first['status'],'WAITING_SME');self.assertTrue(first['packet_issued'])
        with self.assertRaises(ValidationError):self.c.start('process-a')
        self.c.close();self.c=Coordinator(self.root)
        self.assertEqual(self.c.ledger.get('process-a')['packet_hash'],first['packet_hash'])
        self.c.import_answers('process-a',self.answer(),'Example reviewer')
        with self.assertRaises(ValidationError):self.c.import_answers('process-a',self.answer(),'Example reviewer')
        self.assertEqual(self.c.ledger.get('process-a')['status'],'QUEUED_VERIFY')

    def test_tampered_question_rejected_without_consuming_round(self):
        self.create();self.c.start('process-a');self.c.advance('process-a')
        with self.assertRaises(ValidationError):self.c.import_answers('process-a',self.answer(alter=True),'Example reviewer')
        self.assertFalse(self.c.ledger.get('process-a')['packet_imported'])

    def test_uncertain_answer_is_consumed_and_remains_blocked(self):
        self.create();self.c.start('process-a');self.c.advance('process-a')
        self.c.import_answers('process-a',self.answer(response='Not sure'),'Example reviewer')
        self.c.advance('process-a')
        doc=self.c.ledger.get('process-a')
        self.assertTrue(any(b['kind']=='sme_unresolved' for b in doc['blockers']))
        self.assertTrue(doc['packet_imported'])

    def test_pause_and_cancel_preserve_authorization_and_quota(self):
        self.create();self.c.start('process-a');self.c.control('process-a','pause')
        self.c.advance('process-a');self.assertEqual(self.c.ledger.get('process-a')['status'],'PAUSED')
        self.c.control('process-a','resume');self.c.advance('process-a')
        self.assertEqual(self.c.ledger.get('process-a')['status'],'WAITING_SME')
        self.c.control('process-a','cancel');self.assertTrue(self.c.ledger.get('process-a')['packet_issued'])


if __name__=='__main__':unittest.main()
