"""R501–R600: new ingress, placement and recognition interactions.

Fixtures are fictional; recognition never establishes native execution parity.
"""
import copy
from io import BytesIO
import json
import os
from pathlib import Path
import stat
import struct
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZipInfo, ZIP_STORED
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.comments import Comment
from workbench.domain import ValidationError, atomic_json, checked_zip, decode, encode, safe_path, sha, write_new, MAX_DOCUMENT_DEPTH
from workbench.intake import HEADERS, from_rows, parse_manifest, parse_intake_xlsx
from workbench.layout import output_path, validate_workspace
from workbench.mainframe import classify_files, load_knowledge, utility_findings, validate_snapshot

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
ROW = [1, 'J', 1, 'S', 'PROGRAM', 'IN', 'OUT', 'Always']


class ExpandedInputTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.snapshot = load_knowledge(self.root)

    def rows(self, rows=None, name='Fictional process'):
        return from_rows('fictional', name, rows or [ROW.copy()])

    def manifest(self, rows=None, name='Fictional process'):
        return '- Process ID: fictional\n- Process name: ' + name + '\n' + '| ' + ' | '.join(HEADERS) + ' |\n|---|---|---|---|---|---|---|---|\n' + ''.join('| ' + ' | '.join(map(str, row)) + ' |\n' for row in (rows or [ROW]))

    def workbook(self, mutate=None):
        book=Workbook(); ws=book.active; ws.title='Intake'
        ws.append(['Process ID', 'fictional']); ws.append(['Process name','Fictional process'])
        for c, v in enumerate(HEADERS,1): ws.cell(4,c,v)
        for c, v in enumerate(ROW,1): ws.cell(5,c,v)
        if mutate: mutate(book,ws)
        buf=BytesIO(); book.save(buf); book.close(); return buf.getvalue()

    def forge(self, mutate, member='xl/worksheets/sheet1.xml'):
        out=BytesIO()
        with ZipFile(BytesIO(self.workbook())) as source, ZipFile(out,'w') as target:
            for info in source.infolist():
                raw=source.read(info)
                if info.filename==member:
                    node=ET.fromstring(raw); mutate(node); raw=ET.tostring(node)
                target.writestr(info,raw)
        return out.getvalue()

    def archive(self, entries):
        out=BytesIO()
        with ZipFile(out,'w',compression=ZIP_STORED) as z:
            for name, value in entries: z.writestr(name,value)
        return out.getvalue()

    def place(self, relative, value='evidence'):
        p=self.root/relative; p.parent.mkdir(parents=True,exist_ok=True); p.write_text(value); return p

    def classify(self, text, path='MEMBER', extra=None):
        return classify_files({path:text, **(extra or {})},{'jobs':[]},self.snapshot)[path]

    def findings(self, text='', manifest=None, files=None):
        return utility_findings(manifest or {'jobs':[]},self.snapshot,files if files is not None else {'JOB':text})

    def test_r501_json_utf16_autodetection_cannot_bypass_utf8_contract(self):
        """Valid UTF-16/32 JSON bytes must fail the UTF-8 wire contract; R008 covered corrupt bytes, not accepted alternate encodings."""
        for encoding in ('utf-16','utf-32'):
            with self.subTest(encoding=encoding), self.assertRaises(ValidationError): decode('{"a":1}'.encode(encoding))
        self.assertEqual(decode(b'\xef\xbb\xbf{"a":1}'), {'a':1})

    def test_r502_encode_rejects_lossy_object_key_coercion(self):
        """Integer and boolean object keys must not silently become text; R016 checked a canonical text-key document only."""
        for value in ({1:'a'},{False:'b'},{None:'c'}):
            with self.subTest(value=value), self.assertRaises(ValidationError): encode(value)

    def test_r503_encode_surrogate_has_named_failure(self):
        """An in-memory surrogate must produce ValidationError before persistence; R012 exercised decoding escaped surrogates only."""
        with self.assertRaises(ValidationError): encode({'value':'\ud800'})

    def test_r504_encoder_and_decoder_share_depth_budget(self):
        """An in-memory tree exceeding the read budget must not be written unreadably; R009 bounded incoming JSON, not encode."""
        value=0
        for _ in range(MAX_DOCUMENT_DEPTH+1): value=[value]
        with self.assertRaises(ValidationError): encode(value)

    def test_r505_cyclic_state_fails_with_document_diagnostic(self):
        """A cyclic in-memory object graph must fail intentionally; prior JSON tests used acyclic wire documents."""
        value=[]; value.append(value)
        with self.assertRaises(ValidationError): encode(value)

    def test_r506_nonserializable_state_is_named_validation(self):
        """A nested Python-only value must not leak TypeError from the state writer; prior tests covered nonfinite JSON numbers."""
        with self.assertRaises(ValidationError): encode({'row':{'members'}})

    def test_r507_bad_state_does_not_create_output_directories(self):
        """Serialization failure at a new state location must leave no directory footprint; R013 protected an existing original file."""
        path=self.root/'new'/'nested'/'state.json'
        with self.assertRaises(ValueError): atomic_json(path,{'number':float('nan')})
        self.assertFalse((self.root/'new').exists())

    def test_r508_atomic_state_directory_destination_is_validation(self):
        """A directory occupying the state filename must be diagnosed before temporary writes; R032 simulated replace failure only."""
        path=self.root/'state.json'; path.mkdir()
        with self.assertRaises(ValidationError): atomic_json(path,{'new':1})
        self.assertEqual(list(path.iterdir()),[])

    def test_r509_atomic_random_collision_preserves_foreign_temp(self):
        """A pre-existing generated temporary filename must survive a failed exclusive create; R032 only checked cleanup of our own temp."""
        self.place('state.json','old'); foreign=self.place('.state.json.fixed','foreign')
        with patch('workbench.domain.uuid.uuid4',return_value=type('UUID',(),{'hex':'fixed'})()):
            with self.assertRaises(FileExistsError): atomic_json(self.root/'state.json',{'new':1})
        self.assertEqual(foreign.read_text(),'foreign'); self.assertEqual((self.root/'state.json').read_text(),'old')

    def test_r510_mutable_evidence_buffer_is_copied_and_hashed(self):
        """Mutating a caller bytearray after writing must not alter the persisted digest or bytes; R031 tested invalid data types."""
        data=bytearray(b'original'); path=self.root/'evidence.bin'; digest=write_new(path,data); data[:]=b'changed!'
        self.assertEqual(path.read_bytes(),b'original'); self.assertEqual(digest,sha(b'original'))

    def test_r511_archive_dot_alias_is_not_distinct_identity(self):
        """Two lexical member paths resolving to one file must be rejected; R065 checked only identical ZIP names."""
        with self.assertRaises(ValidationError): checked_zip(self.archive([('a/item','one'),('a/./item','two')]))

    def test_r512_archive_portable_aliases_do_not_overwrite(self):
        """Case and normalization aliases must not form distinct portable member identities; prior ZIP checks used byte-equal names."""
        for names in (('Data.xml','data.xml'),('caf\u00e9','cafe\u0301')):
            with self.subTest(names=names), self.assertRaises(ValidationError): checked_zip(self.archive([(names[0],'one'),(names[1],'two')]))

    def test_r513_archive_file_parent_collision_is_rejected(self):
        """A member cannot simultaneously be a file and a parent directory; prior traversal checks did not examine member relationships."""
        with self.assertRaises(ValidationError): checked_zip(self.archive([('a','one'),('a/b','two')]))

    def test_r514_archive_symlink_metadata_cannot_become_evidence(self):
        """A Unix ZIP symlink entry must be refused even with a benign name; prior archive checks did not inspect entry mode."""
        info=ZipInfo('source'); info.create_system=3; info.external_attr=(stat.S_IFLNK|0o777)<<16
        with self.assertRaises(ValidationError): checked_zip(self.archive([(info,'../outside')]))

    def test_r515_archive_windows_device_identity_is_rejected(self):
        """An archive path using a device basename must fail before consumers read it; R021 applied the rule only to filesystem output paths."""
        with self.assertRaises(ValidationError): checked_zip(self.archive([('xl/CON.xml','data')]))

    def test_r516_archive_directory_payload_cannot_disappear(self):
        """Nonempty bytes hidden behind a directory entry must be rejected; prior expansion tests treated every entry as a regular file."""
        with self.assertRaises(ValidationError): checked_zip(self.archive([('folder/','hidden bytes')]))

    def test_r517_nonxml_archive_crc_is_checked_before_return(self):
        """Corrupt binary member CRC must fail the archive gate, not wait for a later consumer; prior XML guards read only XML members."""
        raw=bytearray(self.archive([('binary.bin','unique-body')])); offset=raw.find(b'unique-body'); raw[offset]=ord('X')
        with self.assertRaises(ValidationError): checked_zip(bytes(raw))

    def test_r518_archive_explicit_directory_and_child_remain_legal(self):
        """A genuine directory plus its child must remain readable after collision hardening; prior archive tests had no directory hierarchy."""
        with checked_zip(self.archive([('folder/',''),('folder/a','value')])) as archive:
            self.assertEqual(archive.read('folder/a'),b'value')

    def test_r519_fenced_example_does_not_override_live_intake(self):
        """Markdown example fences must not supply process attributes or extra jobs; prior duplicate-header tests used visible tables."""
        example=self.manifest(name='Example only').replace('fictional','sample')
        for fence in ('```','~~~'):
            self.assertEqual(parse_manifest(fence+'markdown\n'+example+fence+'\n'+self.manifest()),self.rows())

    def test_r520_html_comment_intake_is_not_operator_input(self):
        """An HTML-commented sample must not conflict with the active manifest; earlier tests did not distinguish rendered and hidden text."""
        self.assertEqual(parse_manifest('<!--\n'+self.manifest(name='Hidden')+'-->\n'+self.manifest()),self.rows())

    def test_r521_unclosed_fence_does_not_activate_hidden_rows(self):
        """A manifest enclosed in an unclosed code fence must be rejected, not treated as active intake; R063 had no hidden table."""
        with self.assertRaises(ValidationError): parse_manifest('```markdown\n'+self.manifest())

    def test_r522_unbalanced_attribute_code_span_is_rejected(self):
        """An unmatched process attribute backtick must not be silently repaired; prior attribute tests covered duplicate and empty values."""
        with self.assertRaises(ValidationError): parse_manifest(self.manifest().replace('fictional','`fictional'))

    def test_r523_balanced_attribute_code_spans_roundtrip(self):
        """Balanced inline-code wrappers on both attributes retain semantic equivalence with XLSX; prior tests did not compare the two intake routes."""
        text=self.manifest().replace(': fictional',': `fictional`').replace(': Fictional process',': `Fictional process`')
        self.assertEqual(parse_manifest(text),parse_intake_xlsx(self.workbook()))

    def test_r524_windows_manifest_matches_excel_semantics(self):
        """CRLF Markdown and workbook intake must produce identical ordered process values; prior byte-preservation checks concerned source exports."""
        self.assertEqual(parse_manifest(self.manifest().replace('\n','\r\n')),parse_intake_xlsx(self.workbook()))

    def test_r525_unicode_process_name_survives_both_formats(self):
        """A name with combining marks and supplementary Unicode must survive both ingress formats without normalization; prior tests rejected line separators."""
        name='Cafe\u0301 \U0001f680 process'
        data=self.workbook(lambda b,w:setattr(w['B2'],'value',name))
        self.assertEqual(parse_manifest(self.manifest(name=name)),parse_intake_xlsx(data)); self.assertEqual(parse_intake_xlsx(data)['name'],name)

    def test_r526_workbook_xml_escaped_conditions_remain_literal(self):
        """XML-escaped comparison and ampersand text must retain its exact intake condition; prior workbook tests focused on formulas and dimensions."""
        condition='RC < 4 & OWNER'; data=self.workbook(lambda b,w:setattr(w['H5'],'value',condition))
        self.assertEqual(parse_intake_xlsx(data)['jobs'][0]['steps'][0]['condition'],condition)

    def test_r527_duplicate_file_references_are_not_deduplicated(self):
        """Repeated declared file references must remain evidence rather than being set-deduplicated; R058 covered distinct file names only."""
        row=ROW.copy(); row[5]='A;A,B'; self.assertEqual(self.rows([row])['jobs'][0]['steps'][0]['inputs'],['A','A','B'])

    def test_r528_step_name_reuse_across_jobs_is_legal(self):
        """Equal step names in different ordered jobs must remain separate declarations; R051 checked the within-job duplicate rejection."""
        second=ROW.copy(); second[0]=2; second[1]='K'
        result=self.rows([ROW.copy(),second]); self.assertEqual([j['steps'][0]['name'] for j in result['jobs']],['S','S'])

    def test_r529_input_rows_are_detached_from_manifest_result(self):
        """Later caller row mutation must not rewrite parsed intake facts; previous immutable-write tests did not cover manifest ownership."""
        rows=[ROW.copy()]; result=self.rows(rows); rows[0][5]='CHANGED'; rows.append(ROW.copy())
        self.assertEqual(result['jobs'][0]['steps'][0]['inputs'],['IN']); self.assertEqual(len(result['jobs']),1)

    def test_r530_sparse_job_orders_preserve_declared_values(self):
        """Nonconsecutive job and step ordinals must remain their supplied values rather than being renumbered; R059 checked sorting only."""
        first=ROW.copy();first[0]=40;first[2]=900
        result=self.rows([first]);self.assertEqual((result['jobs'][0]['order'],result['jobs'][0]['steps'][0]['order']),(40,900))

    def test_r531_xlsx_metadata_labels_bind_identity_values(self):
        """Swapped metadata labels must fail instead of binding B1/B2 under misleading labels; R074 validated only value type."""
        def mutate(book,ws): ws['A1']='Process name';ws['A2']='Process ID'
        with self.assertRaises(ValidationError): parse_intake_xlsx(self.workbook(mutate))

    def test_r532_xlsx_populated_spacer_row_cannot_disappear(self):
        """Nonempty row 3 data must not be silently ignored between metadata and headers; prior populated-row checks started below the header."""
        with self.assertRaises(ValidationError): parse_intake_xlsx(self.workbook(lambda b,w:setattr(w['A3'],'value','UNACCOUNTED JOB')))

    def test_r533_xlsx_metadata_side_cells_cannot_hide_input(self):
        """Populated C1:H2 cells must fail instead of vanishing from intake; prior forged-dimension tests addressed out-of-range coordinates."""
        with self.assertRaises(ValidationError): parse_intake_xlsx(self.workbook(lambda b,w:setattr(w['H1'],'value','HIDDEN CONDITION')))

    def test_r534_xlsx_merged_job_cells_are_ambiguous(self):
        """A merged job field must be refused even when its anchor alone is valid; prior duplicate-cell checks did not examine merges."""
        with self.assertRaises(ValidationError): parse_intake_xlsx(self.workbook(lambda b,w:w.merge_cells('F5:G5')))

    def test_r535_xlsx_dimension_lower_bound_cannot_omit_metadata(self):
        """A declared B2:H5 range must not pass merely because its maxima match; R076 tested cells beyond maximum declared dimensions."""
        with self.assertRaises(ValidationError): parse_intake_xlsx(self.forge(lambda x:x.find(NS+'dimension').set('ref','B2:H5')))

    def test_r536_xlsx_multiple_sheetdata_sections_are_rejected(self):
        """A second worksheet data container must not be silently ignored; prior duplicate-row tests used one container."""
        with self.assertRaises(ValidationError): parse_intake_xlsx(self.forge(lambda x:ET.SubElement(x,NS+'sheetData')))

    def test_r537_xlsx_multiple_cell_values_are_rejected(self):
        """Two value elements in one order cell must not select an arbitrary winner; R077 duplicated coordinates across different cell nodes."""
        def mutate(x): ET.SubElement(x.find('.//'+NS+"c[@r='A5']"),NS+'v').text='2'
        with self.assertRaises(ValidationError): parse_intake_xlsx(self.forge(mutate))

    def test_r538_xlsx_error_cells_are_not_literal_dependencies(self):
        """An Excel error in an optional dependency cell must not become a real file reference; R073 rejected formulas but not cached error types."""
        with self.assertRaises(ValidationError): parse_intake_xlsx(self.workbook(lambda b,w:setattr(w['F5'],'value','#REF!')))

    def test_r539_xlsx_comment_instructions_do_not_supply_job_facts(self):
        """A comment proposing another job must remain commentary and leave the parsed process unchanged; prior tests did not cover comment parts."""
        data=self.workbook(lambda b,w:setattr(w['A5'],'comment',Comment('Run SECRET next','Fictional author')))
        self.assertEqual(parse_intake_xlsx(data),self.rows())

    def test_r540_xlsx_hyperlink_text_is_not_remote_intake(self):
        """A linked input name must be read as its displayed text without importing hyperlink targets; prior archive tests did not exercise relationships."""
        data=self.workbook(lambda b,w:setattr(w['F5'],'hyperlink','https://invalid.example/private'))
        self.assertEqual(parse_intake_xlsx(data)['jobs'][0]['steps'][0]['inputs'],['IN'])

    def test_r541_xlsx_rich_text_runs_do_not_drop_condition_fragments(self):
        """A condition split into rich text runs must be reconstructed completely; earlier workbook tests used one inline text node."""
        def mutate(x):
            cell=x.find('.//'+NS+"c[@r='H5']"); cell.clear();cell.set('r','H5');cell.set('t','inlineStr');inline=ET.SubElement(cell,NS+'is')
            for value in ('RC','<=','4'):
                run=ET.SubElement(inline,NS+'r');ET.SubElement(run,NS+'t').text=value
        self.assertEqual(parse_intake_xlsx(self.forge(mutate))['jobs'][0]['steps'][0]['condition'],'RC<=4')

    def test_r542_xlsx_numeric_display_format_does_not_rewrite_order(self):
        """A numeric cell formatted with leading zeros must retain its integer order; prior order-type checks did not involve Excel formatting."""
        data=self.workbook(lambda b,w:setattr(w['A5'],'number_format','0000'))
        self.assertEqual(parse_intake_xlsx(data)['jobs'][0]['order'],1)

    def test_r543_xlsx_blank_formatted_tail_is_not_new_step(self):
        """A styled empty row at the last allowed coordinate must not synthesize a step; prior row-bound tests used populated cells."""
        data=self.workbook(lambda b,w:setattr(w['H205'],'number_format','@'))
        self.assertEqual(parse_intake_xlsx(data),self.rows())

    def test_r544_xlsx_duplicate_worksheet_dimensions_are_rejected(self):
        """Two contradictory dimension declarations must not let one parser choose a different envelope; R076 used one forged declaration."""
        def mutate(x): node=ET.Element(NS+'dimension',{'ref':'A1:H4'});x.insert(0,node)
        with self.assertRaises(ValidationError): parse_intake_xlsx(self.forge(mutate))

    def test_r545_root_file_allowlist_does_not_authorize_directory(self):
        """An allowlisted root filename must not be accepted as a directory; R033 exercised unlisted file names."""
        (self.root/'AGENTS.md').mkdir();self.assertTrue(validate_workspace(self.root))

    def test_r546_root_directory_allowlist_does_not_authorize_file(self):
        """A file occupying a required directory name must block setup; R038 examined process categories, not workspace structural roots."""
        self.place('Endeavor');self.assertTrue(validate_workspace(self.root))

    def test_r547_root_special_file_is_not_invisible_to_layout(self):
        """A root FIFO must be diagnosed without opening it; earlier workspace checks covered files, directories and symlinks only."""
        os.mkfifo(self.root/'fifo'); self.assertTrue(validate_workspace(self.root))

    def test_r548_export_special_file_is_not_ignored(self):
        """A named pipe inside Endeavor must block local-source readiness without a read; prior source placement checks inspected symlinks."""
        (self.root/'Endeavor').mkdir();os.mkfifo(self.root/'Endeavor'/'member');self.assertTrue(validate_workspace(self.root))

    def test_r549_empty_unapproved_input_directory_is_reported(self):
        """An empty arbitrary input subtree must not bypass file-only placement validation; R039 called output_path on file paths."""
        (self.root/'processes'/'fictional'/'input'/'extra').mkdir(parents=True);self.assertTrue(validate_workspace(self.root))

    def test_r550_reserved_empty_process_directory_is_reported(self):
        """An empty platform-reserved analysis directory must be refused before future writes; R021 checked safe_path filenames only."""
        (self.root/'processes'/'fictional'/'analysis'/'NUL').mkdir(parents=True);self.assertTrue(validate_workspace(self.root))

    def test_r551_process_id_case_collision_is_reported(self):
        """Two existing case-only process roots must not be portable distinct histories; R036 rejected individual malformed IDs."""
        folder=self.root/'processes';original=folder/'Case';original.mkdir(parents=True)
        self.assertFalse(validate_workspace(self.root))
        # Model directory entries that cannot coexist on a case-insensitive host.
        rglob=Path.rglob
        def entries(path,pattern):
            return iter((original,folder/'case')) if path==folder else rglob(path,pattern)
        with patch.object(Path,'rglob',entries):
            self.assertTrue(any('identity collides' in issue for issue in validate_workspace(self.root)))

    def test_r552_process_nested_portable_name_collision_is_reported(self):
        """Two sibling artifact names differing only by case must fail a workspace scan; prior output_path checks were single-path checks."""
        self.place('processes/fictional/analysis/Rules.json');self.assertFalse(validate_workspace(self.root))
        folder=self.root/'processes';rglob=Path.rglob
        def entries(path,pattern):
            items=list(rglob(path,pattern))
            if path==folder:items.append(folder/'fictional/analysis/rules.json')
            return iter(items)
        with patch.object(Path,'rglob',entries):
            self.assertTrue(any('identity collides' in issue for issue in validate_workspace(self.root)))

    def test_r553_process_fifo_cannot_hide_issued_evidence(self):
        """A FIFO inside a reports directory must be visible as invalid evidence without blocking the scan; prior category checks tested regular files."""
        folder=self.root/'processes'/'fictional'/'reports';folder.mkdir(parents=True);os.mkfifo(folder/'report.json');self.assertTrue(validate_workspace(self.root))

    def test_r554_layout_does_not_read_private_diagnostics(self):
        """Private scratch contents must remain outside public placement traversal; R044 only proved source symlink reporting."""
        folder=self.root/'.implementation';folder.mkdir();os.mkfifo(folder/'private-diagnostic');self.assertEqual(validate_workspace(self.root),[])

    def test_r555_shared_version_directory_cannot_substitute_python_file(self):
        """A correctly hash-named directory must not count as a shared Python version; R042 switched between two regular filenames."""
        (self.root/'shared'/'target'/'python'/('a'*64+'.py')).mkdir(parents=True);self.assertTrue(validate_workspace(self.root))

    def test_r556_shared_nested_package_is_not_version_storage(self):
        """A package subtree below shared Python storage must not bypass direct version placement; prior shared checks used only a top-level file."""
        self.place('shared/target/python/package/'+'a'*64+'.py');self.assertTrue(validate_workspace(self.root))

    def test_r557_source_snapshot_permits_original_executable_text(self):
        """Original exported Python text under input/sources remains evidence, without becoming generated target output; R040 tested generated code rejection elsewhere."""
        path=output_path(self.root,'fictional','input/sources/library/original.py');self.assertEqual(path.parent.name,'library')

    def test_r558_category_only_path_cannot_be_written_as_artifact(self):
        """The category directory itself cannot be returned as an artifact destination; prior path tests used a category plus filename."""
        with self.assertRaises(ValidationError): output_path(self.root,'fictional','analysis')

    def test_r559_layout_diagnostics_are_deterministic_across_creation_order(self):
        """Multiple placement defects must have stable sorted diagnostics, enabling reproducible setup summaries; prior checks asserted isolated defects."""
        self.place('z.txt');self.place('a.txt');issues=validate_workspace(self.root);self.assertEqual(issues,sorted(issues));self.assertEqual(len(issues),2)

    def test_r560_empty_approved_category_tree_is_valid_setup(self):
        """An initialized process with all seven empty categories must pass before evidence exists; prior tests focused on already written artifacts."""
        for name in ('input','analysis','review','synthetic','target','reports','tests'):(self.root/'processes'/'fictional'/name).mkdir(parents=True)
        self.assertEqual(validate_workspace(self.root),[])

    def test_r561_sql_line_comment_cannot_open_block_comment(self):
        """A block opener inside a trailing SQL line comment must not hide the next statement; R093 used genuine block comments."""
        row=self.classify('VALUES 1 -- /* commentary\nCREATE TABLE APP.T (ID INT);')
        self.assertEqual(row['kind'],'sql');self.assertEqual(row['evidence'][0]['line'],2)

    def test_r562_cobol_inline_comment_cannot_hide_next_declaration(self):
        """A block marker inside a COBOL inline comment must not consume following source; R100 covered whole-line COPY comments."""
        row=self.classify('IDENTIFICATION DIVISION. *> /* note\nPROGRAM-ID. ACTUAL.')
        self.assertEqual(row['kind'],'cobol_program');self.assertEqual(row['evidence'][0]['line'],2)

    def test_r563_multiline_sql_literal_cannot_supply_program_structure(self):
        """A program-looking line inside a multiline SQL literal must not create classification conflict; R084 checked a one-line REXX-marker literal."""
        self.assertEqual(self.classify("SELECT 'first\nPROGRAM-ID. FALSE.\nlast' FROM APP.T;")['kind'],'sql')

    def test_r564_multiline_literal_cannot_invoke_utility(self):
        """A JCL card inside a multiline host-language literal must not count as a source invocation; R107 hid cards using DD DATA instead."""
        self.assertEqual(self.findings("SELECT 'first\n//S EXEC PGM=IDCAMS\nlast' FROM APP.T;"),[])

    def test_r565_multiline_literal_cannot_create_copy_dependency(self):
        """A COPY-looking literal line must not relabel another member; R100 considered only explicit comment syntax."""
        files={'QUERY':"SELECT 'first\nCOPY HIDDEN.\nlast' FROM APP.T;",'HIDDEN':'opaque text'}
        self.assertEqual(classify_files(files,{'jobs':[]},self.snapshot)['HIDDEN']['kind'],'unknown')

    def test_r566_literal_comment_marker_does_not_hide_sql_after_quote(self):
        """Doubled quote escaping around comment-like text must leave later real structure visible; earlier block-comment tests had no escaped literal."""
        self.assertEqual(self.classify("VALUES 'it''s /* data';\nCREATE TABLE APP.T (ID INT);")['kind'],'sql')

    def test_r567_unclosed_literal_masks_remaining_signatures(self):
        """An unterminated quoted value must not expose later source-looking text as executable; prior unclosed-comment checks used comment delimiters."""
        self.assertEqual(self.classify("VALUES 'unfinished\nPROGRAM-ID. FALSE.")['kind'],'unknown')

    def test_r568_closing_multiline_literal_restores_sql_evidence(self):
        """The line after a multiline literal closes must recover its original source location; R094 checked closure of block comments only."""
        row=self.classify("VALUES 'first\nlast';\nCREATE TABLE APP.T (ID INT);")
        self.assertEqual(row['kind'],'sql');self.assertEqual(row['evidence'][0]['line'],3)

    def test_r569_jcl_null_statement_resets_job_identity(self):
        """An orphan EXEC following a null JCL statement must not inherit the ended job identity; prior finding tests had one uninterrupted job."""
        rows=self.findings('//OLD JOB\n//\n//S EXEC PGM=IDCAMS')
        self.assertEqual(rows[0]['job'],'')

    def test_r570_jcl_second_job_changes_subsequent_utility_context(self):
        """Two JOB cards in one export must bind each utility to its own job; R111 separated procedure members rather than within-file jobs."""
        rows=self.findings('//A JOB\n//S EXEC PGM=IDCAMS\n//B JOB\n//S EXEC PGM=IEBGENER')
        self.assertEqual([(r['job'],r['program']) for r in rows],[('A','IDCAMS'),('B','IEBGENER')])

    def test_r571_source_occurrence_order_is_not_catalog_order(self):
        """Source findings must preserve step occurrence order even when catalog utility order differs; earlier duplicate tests used one utility."""
        rows=self.findings('//J JOB\n//Z EXEC PGM=BPXBATCH\n//A EXEC PGM=IDCAMS')
        self.assertEqual([r['program'] for r in rows],['BPXBATCH','IDCAMS'])

    def test_r572_alias_and_real_name_are_separate_physical_occurrences(self):
        """Two distinct executable names in one utility family must retain individual call sites; R103 used different driver variants without same-family dedup pressure."""
        rows=self.findings('//J JOB\n//S EXEC PGM=SORT\n//S EXEC PGM=ICEMAN')
        self.assertEqual([r['program'] for r in rows],['SORT','ICEMAN']);self.assertEqual([r['utility_id'] for r in rows],['DFSORT','DFSORT'])

    def test_r573_manifest_alias_does_not_absorb_different_source_executable(self):
        """A manifest alias and differently named source executable must stay separate despite sharing utility ID; R113 matched identical program names."""
        manifest={'jobs':[{'name':'J','steps':[{'name':'S','program':'SORT'}]}]}
        rows=self.findings('//J JOB\n//S EXEC PGM=ICEMAN',manifest)
        self.assertEqual(len(rows),2);self.assertFalse(rows[0]['source_refs']);self.assertEqual(rows[1]['program'],'ICEMAN')

    def test_r574_duplicate_manifest_occurrences_match_at_most_one_source_each(self):
        """Two declarations competing for one source card must retain an unmatched declaration; R114 checked surplus source occurrences, not surplus declarations."""
        step={'name':'S','program':'IDCAMS'};manifest={'jobs':[{'name':'J','steps':[step.copy(),step.copy()]}]}
        rows=self.findings('//J JOB\n//S EXEC PGM=IDCAMS',manifest)
        self.assertEqual([len(r['source_refs']) for r in rows],[1,0])

    def test_r575_different_job_identity_prevents_manifest_source_merge(self):
        """A same-step same-program call in a different job must not satisfy a manifest declaration; earlier dedup tests shared the same job key."""
        manifest={'jobs':[{'name':'A','steps':[{'name':'S','program':'IDCAMS'}]}]}
        rows=self.findings('//B JOB\n//S EXEC PGM=IDCAMS',manifest);self.assertEqual(len(rows),2);self.assertEqual([r['job'] for r in rows],['A','B'])

    def test_r576_source_file_insertion_order_does_not_change_findings(self):
        """Equivalent export dictionaries with different insertion order must yield identical occurrence evidence; prior evidence tests used one fixed dictionary."""
        files={'Z':'//Z JOB\n//S EXEC PGM=IDCAMS','A':'//A JOB\n//S EXEC PGM=IEBGENER'}
        self.assertEqual(self.findings(files=files),self.findings(files=dict(reversed(list(files.items())))))

    def test_r577_classification_export_order_is_deterministic(self):
        """Classification ordering must be stable independently from caller insertion order; R099 checked mixed-type key rejection only."""
        files={'Z':'01 REC.','A':'CREATE TABLE T (N INT);'}
        first=classify_files(files,{'jobs':[]},self.snapshot);second=classify_files(dict(reversed(list(files.items()))),{'jobs':[]},self.snapshot)
        self.assertEqual(encode(first),encode(second));self.assertEqual(list(first),['A','Z'])

    def test_r578_program_prefix_is_not_a_catalog_alias(self):
        """An application program whose name begins with a utility name must not acquire that utility's risks or identity; R116 checked unresolved symbols."""
        self.assertEqual(self.findings('//S EXEC PGM=IDCAMSX'),[]);self.assertEqual(self.classify('//S EXEC PGM=IDCAMSX')['utility_ids'],[])

    def test_r579_quoted_parm_utility_is_not_second_invocation(self):
        """A known utility in an EXEC parameter string must not count as another call; R115 tested full-line JCL comments."""
        rows=self.findings("//S EXEC PGM=IEFBR14,PARM='EXEC PGM=IDCAMS'")
        self.assertEqual([r['program'] for r in rows],['IEFBR14'])

    def test_r580_unknown_manifest_program_does_not_hide_source_utility(self):
        """An unknown declared executable must not suppress a known source-only utility at the same step; R120 used a malformed program type."""
        manifest={'jobs':[{'name':'J','steps':[{'name':'S','program':'UNKNOWN'}]}]}
        rows=self.findings('//J JOB\n//S EXEC PGM=IDCAMS',manifest);self.assertEqual(len(rows),1);self.assertEqual(rows[0]['invocation_origin'],'source')

    def test_r581_manifest_utility_identity_cannot_contain_line_break(self):
        """A direct utility-findings caller must not inject multiline job identity into lineage; R120 validated only program field type."""
        manifest={'jobs':[{'name':'J\nFALSE','steps':[{'name':'S','program':'IDCAMS'}]}]}
        with self.assertRaises(ValidationError):self.findings(manifest=manifest)

    def test_r582_copy_replacing_operands_are_not_dependencies(self):
        """COPY REPLACING pseudo-text must not create dependencies on replacement tokens; R089 covered a malformed immediate operand."""
        result=classify_files({'P':'PROGRAM-ID. P.\nCOPY REAL REPLACING ==OLD== BY ==NEW==.','REAL':'opaque','OLD':'opaque','NEW':'opaque'},{'jobs':[]},self.snapshot)
        self.assertEqual(result['REAL']['kind'],'copybook');self.assertEqual(result['OLD']['kind'],'unknown');self.assertEqual(result['NEW']['kind'],'unknown')

    def test_r583_copy_reference_does_not_override_contradictory_structure(self):
        """A member named by COPY but structurally a SQL file must keep its observed language; R090 used a suffix conflict without dependency evidence."""
        row=self.classify('CREATE TABLE APP.T (N INT);','REC',{'P':'PROGRAM-ID. P.\nCOPY REC.'})
        self.assertEqual(row['kind'],'sql');self.assertTrue(any('COPY' in e['reason'] for e in row['evidence']))

    def test_r584_copy_name_with_suffix_does_not_match_unrelated_stem(self):
        """A COPY operand must not use a prefix to relabel a longer member name; R088 tested valid national-character exact identity."""
        row=self.classify('opaque','RECORD2',{'P':'PROGRAM-ID. P.\nCOPY RECORD.'});self.assertEqual(row['kind'],'unknown')

    def test_r585_reference_evidence_stays_nonexecuting_with_unknown_grammar(self):
        """A procedural copy dependency may gain a role while retaining unresolved conversion support; prior R147 concerned DCLGEN structural layouts."""
        row=self.classify('PERFORM VENDOR-MAGIC','RULES',{'P':'PROGRAM-ID. P.\nCOPY RULES.'})
        self.assertEqual(row['kind'],'copybook');self.assertEqual(row['conversion_support'],'not_established');self.assertTrue(any(e['line'] is None for e in row['evidence']))

    def test_r586_jcl_fragment_suffix_compatibility_does_not_invent_job(self):
        """A .jcl include member containing only DD cards must remain a fragment; R155 used a named PROC with native semantics."""
        row=self.classify('//IN DD DSN=APP.FILE,DISP=SHR','INPUT.jcl');self.assertEqual(row['kind'],'jcl_fragment');self.assertFalse(row['conflicts'])

    def test_r587_instream_proc_keeps_containing_job_classification(self):
        """A JOB with an inline PROC and invocation must keep the container classification with actual utility sites; R091 embedded non-JCL source as DATA."""
        row=self.classify('//J JOB\n//P PROC\n//S EXEC PGM=IDCAMS\n// PEND\n//R EXEC PROC=P')
        self.assertEqual(row['kind'],'jcl_job');self.assertEqual(row['utility_ids'],['IDCAMS']);self.assertEqual(row['conversion_support'],'not_established')

    def test_r588_cobol_embedded_sql_remains_program_context(self):
        """Embedded SQL and host declarations must not split an otherwise identified COBOL program into incompatible source families; R146 used standalone DDL."""
        row=self.classify('PROGRAM-ID. P.\n01 HOST.\nEXEC SQL DECLARE APP.T TABLE\nEND-EXEC.\nSELECT * FROM APP.T;')
        self.assertEqual(row['kind'],'cobol_program');self.assertEqual(row['candidate_kinds'],['cobol_program'])

    def test_r589_bms_assembler_container_is_single_map_classification(self):
        """An assembler section containing BMS map macros must retain the BMS role without an artificial language conflict; R145 used BMS macros alone."""
        row=self.classify('MAP CSECT\nMAPSET DFHMSD TYPE=MAP');self.assertEqual(row['kind'],'bms_map');self.assertFalse(row['conflicts'])

    def test_r590_incompatible_languages_keep_all_candidate_evidence(self):
        """A scheduler command mixed with PL/I structure must retain both evidence families while blocking certainty; earlier conflicts used COBOL/JCL or suffixes."""
        row=self.classify('LJOB,JOB=NIGHTLY\nMAIN: PROCEDURE OPTIONS(MAIN);')
        self.assertEqual(row['kind'],'ambiguous');self.assertEqual(row['candidate_kinds'],['pli','scheduler_definition']);self.assertEqual({e['line'] for e in row['evidence']},{1,2})

    def test_r591_unknown_suffix_does_not_erase_structural_evidence(self):
        """An arbitrary export suffix must not reduce a recognized SQL statement to unknown; previous suffix checks used known contradictory extensions."""
        row=self.classify('CREATE TABLE APP.T (N INT);','MEMBER.export');self.assertEqual(row['kind'],'sql');self.assertEqual(row['confidence'],'structural')

    def test_r592_suffix_hint_only_cannot_become_conflicting_structure(self):
        """A .sql file containing only comments must stay unresolved with a hint rather than fabricated SQL structure; R096 tested an extensionless empty export."""
        row=self.classify('-- description only','query.sql');self.assertEqual(row['kind'],'unknown');self.assertEqual(row['candidate_kinds'],['sql']);self.assertEqual(row['confidence'],'unresolved')

    def test_r593_classification_result_mutation_cannot_poison_next_call(self):
        """Consumer changes to evidence lists must not persist across classifier calls; R118 tested utility finding risk lists instead."""
        first=self.classify('01 REC.');first['evidence'][0]['reason']='forged';first['candidate_kinds'].append('sql')
        second=self.classify('01 REC.');self.assertEqual(second['candidate_kinds'],['copybook']);self.assertNotEqual(second['evidence'][0]['reason'],'forged')

    def test_r594_utility_occurrence_details_are_detached_from_each_other(self):
        """Editing one repeated occurrence's evidence requirements must not rewrite another occurrence; R118 compared a finding to its snapshot only."""
        rows=self.findings('//S1 EXEC PGM=IDCAMS\n//S2 EXEC PGM=IDCAMS');before=copy.deepcopy(rows[1]);rows[0]['required_evidence'].append('invented');self.assertEqual(rows[1],before)

    def test_r595_snapshot_provenance_changes_digest_even_with_same_context(self):
        """Different raw catalog provenance must change snapshot identity despite equal semantic context; R143 tested modified behavior without rehashing."""
        altered=copy.deepcopy(self.snapshot);altered['source_hashes']['standard']='0'*64;altered['content_hash']=sha(encode({k:v for k,v in altered.items() if k!='content_hash'}));validate_snapshot(altered)
        self.assertNotEqual(altered['content_hash'],self.snapshot['content_hash']);self.assertEqual(altered['catalog'],self.snapshot['catalog'])

    def test_r596_application_empty_document_still_records_raw_provenance(self):
        """An explicitly supplied empty overlay must differ from an absent overlay in provenance; R141 rejected nonempty context without a hash."""
        app={'schema_version':1,'application':'','utilities':[],'notes':[]};raw=encode(app);path=self.root/'knowledge'/'application-knowledge.json';path.parent.mkdir();path.write_bytes(raw)
        snapshot=load_knowledge(self.root);self.assertEqual(snapshot['source_hashes']['application'],sha(raw));self.assertNotEqual(snapshot['content_hash'],self.snapshot['content_hash'])

    def test_r597_reading_knowledge_does_not_create_optional_overlay(self):
        """Loading absent optional knowledge must be read-only and leave no workspace files; R144 checked missing mandatory catalog diagnostics."""
        before=list(self.root.iterdir());load_knowledge(self.root);self.assertEqual(list(self.root.iterdir()),before)

    def test_r598_catalog_whitespace_change_affects_source_hash_not_semantics(self):
        """Reformatted overlay JSON must preserve facts but distinguish exact source bytes; prior tamper checks changed content fields."""
        app={'schema_version':1,'application':'Example','utilities':[],'notes':['fact']};path=self.root/'knowledge'/'application-knowledge.json';path.parent.mkdir();path.write_text(json.dumps(app));first=load_knowledge(self.root);path.write_bytes(encode(app));second=load_knowledge(self.root)
        self.assertEqual(first['application'],second['application']);self.assertNotEqual(first['source_hashes']['application'],second['source_hashes']['application'])

    def test_r599_manifest_only_utility_retains_absent_source_evidence(self):
        """A recognized declared utility with no source export must retain an empty source_refs list and manifest origin; prior tests merged source occurrences."""
        rows=self.findings(manifest={'jobs':[{'name':'J','steps':[{'name':'S','program':'IDCAMS'}]}]},files={})
        self.assertEqual(rows[0]['source_refs'],[]);self.assertEqual(rows[0]['invocation_origin'],'manifest');self.assertEqual(rows[0]['conversion_support'],'adapter_required')

    def test_r600_readonly_recognition_never_executes_exported_commands(self):
        """Recognition of a host-command driver must not launch a process even when source contains shell syntax; R159 asserted blockers through source analysis only."""
        with patch('subprocess.Popen',side_effect=AssertionError('Recognition must not launch a process')):
            rows=self.findings("//S EXEC PGM=BPXBATCH,PARM='SH touch fictional-marker'")
        self.assertEqual(rows[0]['conversion_support'],'adapter_required');self.assertFalse((self.root/'fictional-marker').exists())


if __name__=='__main__': unittest.main()
