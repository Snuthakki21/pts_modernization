import csv
import json
import unittest
from io import StringIO
from pathlib import Path
from openpyxl import load_workbook
import test_workflow
from test_source import COBOL
from workbench.coverage import build_coverage, write_coverage


class CoverageTests(unittest.TestCase):
    setUp = test_workflow.WorkflowTests.setUp
    answer = test_workflow.WorkflowTests.answer

    def verified(self, files=None, manifest=None):
        self.c.create(manifest or test_workflow.MANIFEST, files or {'ELIGIBLE.cbl': COBOL})
        self.c.start('process-a'); self.c.advance('process-a')
        self.c.import_answers('process-a', self.answer(), 'Example reviewer')
        self.c.advance('process-a')
        return self.c.ledger.get('process-a')

    def test_inventory_preserves_every_original_line_and_explicit_scope(self):
        extra = COBOL.replace('ELIGIBLE.', 'UNUSED.')
        files = {'ELIGIBLE.cbl': COBOL, 'UNUSED.cbl': extra, 'ORPHAN.cpy': '05 UNUSED PIC X.\n',
                 'screen.bms': 'SCREEN DFHMDI SIZE=(24,80)\n', 'query.sql': 'SELECT * FROM A;\n',
                 'README.export': 'UNKNOWN STATEMENT\n'}
        doc = self.verified(files)
        model = build_coverage(doc, self.root)
        self.assertEqual({f['path'] for f in model['files']}, set(files))
        self.assertEqual(len(model['rows']), sum(len(t.splitlines()) for t in files.values()))
        for row in model['rows']:
            self.assertEqual(row['source_text'], files[row['source_path']].splitlines()[row['source_line']-1])
            self.assertEqual(len(row['source_hash']), 64)
            self.assertTrue(row['unit_id']); self.assertTrue(row['reason'])
        unused = [r for r in model['rows'] if r['source_path'] in ('UNUSED.cbl', 'ORPHAN.cpy')]
        self.assertTrue(all(r['disposition'] == 'out_of_scope' for r in unused))
        self.assertTrue(all(r['reason'] for r in unused))
        other = [r for r in model['rows'] if r['source_path'] in ('screen.bms', 'query.sql', 'README.export')]
        self.assertTrue(all(r['disposition'] == 'blocked' for r in other))
        self.assertFalse(model['summary']['completion_eligible'])
        self.assertEqual(model['summary']['out_of_scope_lines'], len(unused))

    def test_rule_lines_have_target_spans_and_return_has_concrete_verified_adapter(self):
        doc = self.verified()
        model = build_coverage(doc, self.root)
        self.assertTrue(model['summary']['completion_eligible'], model['summary'])
        executable = [r for r in model['rows'] if r['source_line'] >= 10]
        self.assertTrue(all(r['target_file'] and r['target_version'] and r['target_start'] and r['tests'] for r in executable))
        terminal = model['rows'][-1]
        self.assertEqual(terminal['disposition'], 'platform_replaced_verified')
        self.assertIn('return', terminal['replacement'].lower())
        self.assertTrue(terminal['evidence'])
        self.assertEqual(model['summary']['semantic_units']['verified'], model['summary']['semantic_units']['applicable'])

    def test_used_copybook_has_mapping_unreferenced_copybook_has_explicit_exclusion(self):
        files = {'ELIGIBLE.cbl': COBOL.replace('  05 AGE PIC 9(3).', 'COPY AGECOPY.'),
                 'AGECOPY.cpy': '05 AGE PIC 9(3).\n', 'NOTUSED.cpy': '05 UNUSED PIC X.\n'}
        model = build_coverage(self.verified(files), self.root)
        used = next(r for r in model['rows'] if r['source_path'] == 'AGECOPY.cpy')
        self.assertEqual(used['disposition'], 'mapped_verified')
        self.assertTrue(used['target_file']); self.assertTrue(used['tests'])
        self.assertEqual(next(r for r in model['rows'] if r['source_path'] == 'NOTUSED.cpy')['disposition'], 'out_of_scope')

    def test_dd_and_scheduling_are_blocked_without_verified_native_adapter(self):
        files = {'ELIGIBLE.cbl': COBOL, 'JOBA.jcl': '//JOBA JOB TIME=5\n//S010 EXEC PGM=ELIGIBLE\n//IN DD DSN=PROD.INPUT,DISP=SHR\n'}
        model = build_coverage(self.verified(files), self.root)
        cards = [r for r in model['rows'] if r['source_path'] == 'JOBA.jcl']
        self.assertEqual(cards[0]['disposition'], 'blocked')
        self.assertEqual(cards[-1]['disposition'], 'blocked')
        self.assertFalse(cards[-1]['replacement'])
        self.assertFalse(model['summary']['completion_eligible'])

    def test_declared_job_dispatch_has_verified_replacement_without_scheduler_claim(self):
        files={'ELIGIBLE.cbl':COBOL,'JOBA.jcl':'//JOBA JOB\n//S010 EXEC PGM=ELIGIBLE\n'}
        model=build_coverage(self.verified(files),self.root)
        cards=[r for r in model['rows'] if r['source_path']=='JOBA.jcl']
        self.assertEqual(cards[0]['disposition'],'platform_replaced_verified')
        self.assertIn('named',cards[0]['replacement'])
        self.assertIn('scheduler',cards[0]['reason'])
        self.assertTrue(cards[0]['target_file']);self.assertTrue(cards[0]['evidence'])
        self.assertTrue(model['summary']['completion_eligible'], model['summary'])

    def test_jcl_exec_maps_its_exact_job_step_when_program_is_reused(self):
        manifest=test_workflow.MANIFEST+'| 2 | JOBB | 1 | S020 | ELIGIBLE | INPUT | OUTPUT | Always |\n'
        files={'ELIGIBLE.cbl':COBOL,'jobs.jcl':'//JOBA JOB\n//S010 EXEC PGM=ELIGIBLE\n//JOBB JOB\n//S020 EXEC PGM=ELIGIBLE\n'}
        model=build_coverage(self.verified(files,manifest),self.root)
        executions=[r for r in model['rows'] if r['source_path']=='jobs.jcl' and ' EXEC ' in r['source_text']]
        self.assertTrue(all(len(r['target_mappings'])==1 for r in executions))
        self.assertNotEqual(executions[0]['target_start'],executions[1]['target_start'])
        self.assertTrue(model['summary']['completion_eligible'])

    def test_unsupported_never_becomes_mainframe_only(self):
        doc = self.verified({'ELIGIBLE.cbl': COBOL.replace('GOBACK.', 'PERFORM NEXT-PARA.')})
        model = build_coverage(doc, self.root)
        row = model['rows'][-1]
        self.assertEqual(row['disposition'], 'blocked')
        self.assertFalse(row['replacement'])
        self.assertNotIn('mainframe-only', row['reason'])

    def test_manifest_changes_and_missing_baseline_revoke_credit_without_execution(self):
        from unittest.mock import patch
        doc=self.verified()
        manifest=self.root/'processes/process-a/input/process-input.md'
        frozen=manifest.read_bytes()
        manifest.write_bytes(frozen.replace(b'Always',b'RC=999'))
        self.assertFalse(build_coverage(doc,self.root)['summary']['completion_eligible'])
        with patch('workbench.fixtures.verify_program') as program,patch('workbench.fixtures.adversarial_review') as adversarial,patch('workbench.orchestration.verify_jobs') as jobs:
            model=build_coverage(doc,self.root)
            self.assertFalse(model['summary']['completion_eligible'])
            self.assertEqual(model['summary']['verified_applicable_lines'],0)
            self.assertTrue(any('manifest' in error.lower() for error in model['summary']['integrity_errors']))
            program.assert_not_called();adversarial.assert_not_called();jobs.assert_not_called()
        manifest.write_bytes(frozen)
        changed_doc=dict(doc);changed_doc['jobs']=[dict(job) for job in doc['jobs']]
        changed_doc['jobs'][0]['steps']=[dict(step) for step in doc['jobs'][0]['steps']]
        changed_doc['jobs'][0]['steps'][0]['condition']='RC=999'
        self.assertFalse(build_coverage(changed_doc,self.root)['summary']['completion_eligible'])
        missing=dict(doc);missing.pop('manifest_hash')
        missing_model=build_coverage(missing,self.root)
        self.assertFalse(missing_model['summary']['completion_eligible'])
        self.assertTrue(any('baseline' in error.lower() for error in missing_model['summary']['integrity_errors']))
        self.assertTrue(build_coverage(doc,self.root)['summary']['completion_eligible'])

    def test_comment_grammar_is_source_kind_aware_and_unknown_behavior_is_blocked(self):
        files={'ELIGIBLE.cbl':COBOL,'unknown.export':'--EXEC UNKNOWN\n*>EXEC OTHER\n//*EXEC UNKNOWN\n','query.sql':'--SQL comment\nSELECT * FROM ACTUAL_TABLE;\n','screen.bms':'--NOT VALIDATED BMS COMMENT\n'}
        model=build_coverage(self.verified(files),self.root)
        unknown=[row for row in model['rows'] if row['source_path']=='unknown.export']
        self.assertTrue(all(row['disposition']=='blocked' for row in unknown))
        self.assertEqual(next(row for row in model['rows'] if row['source_path']=='query.sql')['disposition'],'non_executable')
        self.assertEqual(next(row for row in model['rows'] if row['source_path']=='screen.bms')['disposition'],'blocked')
        self.assertFalse(model['summary']['completion_eligible'])
        cobol=COBOL.replace('  GOBACK.','--EXEC UNKNOWN\n  GOBACK.')
        other=self.root/'another'
        from workbench.coordinator import Coordinator
        coordinator=Coordinator(other)
        try:
            coordinator.create(test_workflow.MANIFEST,{'ELIGIBLE.cbl':cobol});coordinator.start('process-a');coordinator.advance('process-a')
            doc=coordinator.ledger.get('process-a');accounting=build_coverage(doc,other)
            unsupported=next(row for row in accounting['rows'] if row['source_text']=='--EXEC UNKNOWN')
            self.assertEqual(unsupported['disposition'],'blocked')
            self.assertGreaterEqual(accounting['summary']['applicable_lines'],16)
        finally:coordinator.close()

    def test_frozen_target_database_tampering_cannot_receive_credit(self):
        import sqlite3
        from unittest.mock import patch
        doc=self.verified();run=doc['runs'][-1]['id']
        database=self.root/'processes/process-a/target'/run/'target.sqlite'
        with sqlite3.connect(database) as db:db.execute('DELETE FROM results')
        self.assertFalse(build_coverage(doc,self.root)['summary']['completion_eligible'])
        with patch('workbench.fixtures.verify_program') as execute:
            model=build_coverage(doc,self.root)
            self.assertFalse(model['summary']['completion_eligible'])
            self.assertEqual(model['summary']['verified_applicable_lines'],0)
            self.assertTrue(any('database' in error.lower() for error in model['summary']['integrity_errors']))
            execute.assert_not_called()
        # Row reconciliation remains independent even if a caller supplies a wrong pin.
        from workbench.domain import sha
        forged=dict(doc);forged['artifact_hashes']=dict(doc['artifact_hashes'])
        relative=database.relative_to(self.root/'processes/process-a').as_posix()
        forged['artifact_hashes'][relative]=sha(database.read_bytes())
        reconciled=build_coverage(forged,self.root)
        self.assertFalse(reconciled['summary']['completion_eligible'])
        self.assertTrue(any('results are missing' in error for error in reconciled['summary']['integrity_errors']))
        database.unlink();missing=build_coverage(doc,self.root)
        self.assertFalse(missing['summary']['completion_eligible'])
        self.assertFalse(database.exists())

    def test_each_tampered_evidence_surface_revokes_credit(self):
        doc = self.verified()
        run = doc['runs'][-1]['id']
        base = self.root/'processes/process-a'
        paths = [base/'synthetic'/run/'ELIGIBLE/expected.json',
                 base/'synthetic'/run/'ELIGIBLE/actual-and-comparison.json',
                 self.root/'shared/target/python'/(doc['program_versions']['ELIGIBLE']+'.py'),
                 base/'analysis/source-analysis.json']
        for path in paths:
            with self.subTest(path=path.name):
                raw = path.read_bytes(); path.write_bytes(raw+b' ')
                model = build_coverage(doc, self.root)
                self.assertFalse(model['summary']['completion_eligible'])
                self.assertTrue(model['summary']['integrity_errors'])
                self.assertEqual(model['summary']['verified_applicable_lines'], 0)
                path.write_bytes(raw)
        self.assertTrue(build_coverage(doc, self.root)['summary']['completion_eligible'])

    def test_forged_adversarial_result_cannot_receive_credit(self):
        from workbench.domain import encode
        doc = self.verified(); run = doc['runs'][-1]
        run['programs']['ELIGIBLE']['adversarial']['mutations'] = []
        path=self.root/'processes/process-a/synthetic'/run['id']/'ELIGIBLE/actual-and-comparison.json'
        path.write_bytes(encode(run['programs']['ELIGIBLE']))
        model=build_coverage(doc,self.root)
        self.assertFalse(model['summary']['completion_eligible'])
        self.assertEqual(model['summary']['verified_applicable_lines'], 0)

    def test_changed_source_preserves_frozen_original_text_and_revokes_credit(self):
        doc=self.verified(); source=self.root/'processes/process-a/input/sources/ELIGIBLE.cbl'
        source.write_text(COBOL.replace('>= 18','>= 99'))
        model=build_coverage(doc,self.root)
        self.assertFalse(model['summary']['fully_accounted'])
        self.assertEqual(model['files'][0]['source_text'], COBOL)
        self.assertEqual(model['rows'][9]['source_text'], COBOL.splitlines()[9])
        self.assertEqual(model['summary']['verified_applicable_lines'], 0)

    def test_cancelled_accountability_never_executes_targets(self):
        from unittest.mock import patch
        doc=self.verified();doc['cancel_requested']=True
        with patch('workbench.fixtures.verify_program',side_effect=AssertionError('must not execute')) as program, patch('workbench.fixtures.adversarial_review',side_effect=AssertionError('must not execute')) as adversarial, patch('workbench.orchestration.verify_jobs',side_effect=AssertionError('must not execute')) as jobs:
            model=build_coverage(doc,self.root)
            self.assertEqual(program.call_count,0)
            self.assertEqual(adversarial.call_count,0)
            self.assertEqual(jobs.call_count,0)
        self.assertFalse(model['summary']['completion_eligible'])
        self.assertEqual(model['summary']['verified_applicable_lines'],0)

    def test_incomplete_verification_never_receives_credit(self):
        doc = self.verified(); doc['verification_finished'] = False
        model = build_coverage(doc, self.root)
        self.assertFalse(model['summary']['completion_eligible'])
        self.assertEqual(model['summary']['verified_applicable_lines'], 0)

    def test_workbook_preserves_long_original_lines_across_reversible_chunks(self):
        text='UNRECOGNIZED ' + 'x'*40000 + '\n'
        model=build_coverage(self.verified({'ELIGIBLE.cbl':COBOL,'long.txt':text}),self.root)
        root=self.root/'long-coverage';write_coverage(model,root)
        book=load_workbook(root/'coverage.xlsx')
        try:
            sheet=book['Text chunks']
            chunks=[r for r in sheet.iter_rows(min_row=2,values_only=True) if r[0]=='Source lines' and r[2]=='source_text']
            self.assertEqual(''.join(r[4] for r in chunks),text.rstrip('\n'))
        finally:book.close()

    def test_exports_match_inventory_and_escape_source_text(self):
        doc = self.verified({'ELIGIBLE.cbl': COBOL, 'payload.txt': '=HYPERLINK("bad") <script>alert(1)</script>\n'})
        model = build_coverage(doc, self.root); root = self.root/'coverage-out'
        paths = write_coverage(model, root)
        self.assertEqual({p.name for p in paths}, {'coverage.json','coverage.csv','coverage.xlsx','coverage.html'})
        self.assertEqual(json.loads((root/'coverage.json').read_text())['rows'], model['rows'])
        rows = list(csv.DictReader(StringIO((root/'coverage.csv').read_text())))
        self.assertEqual(len(rows), len(model['rows']))
        self.assertTrue(rows[-1]['source_text'].startswith("'="))
        book = load_workbook(root/'coverage.xlsx'); sheet=book['Source lines']
        self.assertEqual(sheet.max_row, len(model['rows'])+1)
        self.assertTrue(all(c.data_type != 'f' for row in sheet for c in row)); book.close()
        html=(root/'coverage.html').read_text()
        self.assertNotIn('<script>', html); self.assertIn('&lt;script&gt;', html)
        self.assertIn('<pre>=HYPERLINK', html)


if __name__ == '__main__': unittest.main()
