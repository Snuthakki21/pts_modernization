"""Effort, budget and scaling must keep measured facts separate from assumptions."""
import unittest
from pathlib import Path
import tempfile
from workbench.domain import ValidationError

class EconomicsTests(unittest.TestCase):
    def test_forecast_does_not_mix_processes_and_procs_or_invent_rates(self):
        from workbench.economics import forecast
        result=forecast([],[],{'cohorts':[{'id':'procs','unit':'jcl_procs','total':100}], 'capacity_hours_per_week':40}, {})
        self.assertIsNone(result['remaining_effort_hours']['base'])
        self.assertEqual(result['unestimated_units'],100)

    def test_twenty_hours_times_estate_is_explicit_scenario(self):
        from workbench.economics import forecast
        plan={'cohorts':[{'id':'processes','unit':'process','total':100,'hours_per_unit':{'low':15,'base':20,'high':30}}], 'capacity_hours_per_week':40}
        result=forecast([],[],plan,{})
        self.assertEqual(result['remaining_effort_hours']['base'],'2000')
        self.assertEqual(result['capacity_weeks']['base'],'50')
        self.assertFalse(result['full_completion_commitment'])
        self.assertEqual(result['cohorts'][0]['basis'],'ASSUMPTION')

    def test_inline_accepted_report_supports_all_declared_evidence_links(self):
        from fastapi.testclient import TestClient
        from workbench.api import create_app
        from test_workflow import MANIFEST
        from test_source import COBOL
        from openpyxl import load_workbook
        from io import BytesIO
        from urllib.parse import urlencode
        with tempfile.TemporaryDirectory() as temp:
            app=create_app(temp);c=app.state.coordinator
            try:
                c.create(MANIFEST,{'ELIGIBLE.cbl':COBOL},demo=True)
                c.start('process-a');c.advance('process-a')
                book=load_workbook(BytesIO(c.artifact('process-a','review/sme-checklist.xlsx').read_bytes()))
                for row in range(2,book['Checklist'].max_row+1):
                    book['Checklist'].cell(row,5,'Yes');book['Checklist'].cell(row,7,'Fictional test reviewer')
                out=BytesIO();book.save(out);book.close()
                c.import_answers('process-a',out.getvalue(),'Fictional test reviewer');c.advance('process-a');c.advance('process-a')
                doc=c.ledger.get('process-a');self.assertTrue(doc['report_verified'])
                path=next(p for p in doc['report_hashes'] if p.endswith('/executive-report.html'))
                response=TestClient(app,base_url='http://127.0.0.1:8765').get('/api/process/process-a/artifact?'+urlencode({'path':path,'inline':'true'}))
                self.assertEqual(response.status_code,200,response.text)
                self.assertIn('rules.html',response.text)
            finally:c.close()

    def work(self,**changes):
        return {'id':'pilot-1','kind':'work','process_id':'process-a','stage':'pilot_total','hours':'20','complete':True,
                'started_at':'2026-01-01T00:00:00Z','ended_at':'2026-01-02T00:00:00Z','actor':'fixture-analyst',
                'recorded_by':'fictional receipt','evidence':'Fictional observed timesheet',**changes}

    def pilot(self,**changes):
        return {'id':'process-a','status':'COMPLETED','report_verified':True,'measurement_report_intact':True,
                'coverage_summary':{'completion_eligible':True},'blockers':[],**changes}

    def test_completed_pilot_refines_matching_process_forecast_only(self):
        from workbench.economics import forecast
        p={'cohorts':[{'id':'processes','unit':'process','total':100}],'capacity_hours_per_week':40}
        model=forecast([self.pilot()],[self.work()],p,{})
        self.assertEqual(model['remaining_effort_hours']['base'],'1980')
        self.assertEqual(model['capacity_weeks']['base'],'49.5')
        self.assertEqual(model['cohorts'][0]['samples'],[{'process_id':'process-a','hours':'20'}])
        for changes in ({'demo':True},{'fixture_only':True},{'status':'COMPLETED_WITH_BLOCKERS'},
                        {'report_verified':False},{'measurement_report_intact':False},{'blockers':[{}]}):
            with self.subTest(changes=changes):
                self.assertIsNone(forecast([self.pilot(**changes)],[self.work()],p,{})['remaining_effort_hours']['base'])
        p['cohorts'][0]['unit']='jcl_procs'
        self.assertIsNone(forecast([self.pilot()],[self.work()],p,{})['remaining_effort_hours']['base'])

    def test_unknown_subset_does_not_turn_into_zero_or_full_estate_estimate(self):
        from workbench.economics import forecast
        p={'cohorts':[{'id':'jobs','unit':'jcl_jobs','total':10,'hours_per_unit':{'low':1,'base':2,'high':3}},
                      {'id':'procs','unit':'jcl_procs','total':4}], 'allocation_note':'Jobs exclude all separately counted PROC implementation work'}
        m=forecast([],[],p,{})
        self.assertEqual(m['estimated_subset_effort_hours']['base'],'20')
        self.assertIsNone(m['remaining_effort_hours']['base']);self.assertEqual(m['unestimated_units'],4)

    def test_plan_validation_rejects_mixed_grains_double_rates_and_nan(self):
        from workbench.economics import validate_plan
        cases=[{'cohorts':[{'id':'p','unit':'process','total':True}]},
               {'cohorts':[{'id':'p','unit':'process','total':100,'hours_per_unit':{'low':3,'base':2,'high':1}}]},
               {'cohorts':[{'id':'p','unit':'process','total':100}], 'capacity_hours_per_week':'NaN'},
               {'cohorts':[{'id':'p','unit':'process','total':100},{'id':'a','unit':'jcl_jobs','total':100}], 'allocation_note':'Not disjoint'},
               {'cohorts':[{'id':'p','unit':[],'total':100}]},
               {'cohorts':[{'id':'p','unit':'jcl_procs','total':100,'sample_process_ids':['process-a']}]},
               {'cohorts':[{'id':'p','unit':'process','total':-1}]}]
        for case in cases:
            with self.subTest(case=case),self.assertRaises(ValidationError):validate_plan(case)

    def test_partial_or_conflicting_effort_never_calibrates(self):
        from workbench.economics import work_summary,forecast
        detail=self.work(id='d',stage='conversion',hours=8,complete=False)
        records=[self.work(hours=7),detail]
        self.assertFalse(work_summary(records,'process-a')['complete'])
        records=[detail,{**detail,'id':'d2','stage':'validation'}]
        summary=work_summary(records,'process-a');self.assertTrue(summary['issues'])
        self.assertIsNone(summary['total_effort_hours'])
        self.assertIsNone(forecast([self.pilot()],records,{'cohorts':[{'id':'p','unit':'process','total':10}]},{})['remaining_effort_hours']['base'])

    def test_budget_credits_expiry_separation_and_straddling_receipts(self):
        from workbench.economics import usage_summary,validate_receipt
        budget={'id':'budget','kind':'budget','process_id':None,'provider':'Claude','account':'team','unit':'credits',
                'allowance':'100','consumed':'20','period_start':'2026-01-01T00:00:00Z','period_end':'2026-02-01T00:00:00Z',
                'snapshot_at':'2026-01-02T00:00:00Z','evidence':'Fictional billing receipt','recorded_by':'fixture'}
        usage={'id':'usage','kind':'usage','process_id':'process-a','provider':'Claude','account':'team','model':'model',
               'started_at':'2026-01-03T00:00:00Z','ended_at':'2026-01-03T01:00:00Z','quantities':{'credits':'1.25','input_tokens':500},
               'evidence':'Fictional actual counters','recorded_by':'fixture'}
        validate_receipt(budget);validate_receipt(usage)
        result=usage_summary([budget,usage,{**usage,'id':'other','provider':'Copilot'}],'2026-01-04T00:00:00Z')
        self.assertEqual(result['primary_budget_unit'],'credits');self.assertEqual(result['budgets'][0]['conditional_remaining'],'78.75')
        self.assertEqual(len(result['quantities']),4)
        self.assertIsNone(usage_summary([budget,usage],'2026-02-01T00:00:00Z')['budgets'][0]['conditional_remaining'])
        crossing={**usage,'started_at':'2026-01-01T00:00:00Z'}
        self.assertIsNone(usage_summary([budget,crossing],'2026-01-04T00:00:00Z')['budgets'][0]['conditional_remaining'])
        with self.assertRaises(ValidationError):validate_receipt({**budget,'unit':'input_tokens'})

    def test_twenty_distinct_credit_scenarios_use_exact_decimals_without_token_conversion(self):
        from workbench.economics import forecast
        for i in range(1,21):
            rate=str(i/100)
            p={'cohorts':[{'id':'p','unit':'process','total':100,'usage_per_unit':[{'provider':'Copilot','account':'team','unit':'credits','low':rate,'base':rate,'high':rate}]}]}
            result=forecast([],[],p,{})
            self.assertEqual(result['cohorts'][0]['remaining_ai_usage'][0]['remaining']['base'],str(i))
            self.assertIsNone(result['remaining_effort_hours']['base'])

    def test_historical_missing_times_and_crash_durations_remain_unknown(self):
        from workbench.economics import timing
        doc={'id':'p','created':'2026-01-01T00:00:00Z','status':'COMPLETED'}
        self.assertIsNone(timing(doc,[],[],'2026-01-10T00:00:00Z')['elapsed_hours'])
        transitions=[{'process_id':'p','status':'READY','created':doc['created']},
                     {'process_id':'p','status':'WAITING_SME','created':'2026-01-01T01:00:00Z'},
                     {'process_id':'p','status':'COMPLETED','created':'2026-01-01T05:00:00Z'}]
        attempts=[{'process_id':'p','stage':'QUEUED_ANALYSIS','elapsed_seconds':None,'outcome':None}]
        t=timing(doc,transitions,attempts,'2026-02-01T00:00:00Z')
        self.assertEqual(t['elapsed_hours'],'5');self.assertEqual(t['status_hours']['sme_wait'],'4');self.assertEqual(t['unfinished_attempts'],1)
        self.assertIsNone(t['service_hours']);self.assertFalse(t['service_timings_complete'])

    def test_receipts_are_idempotent_correctable_and_tamper_evident(self):
        from workbench.coordinator import Coordinator
        from test_workflow import MANIFEST
        from test_source import COBOL
        with tempfile.TemporaryDirectory() as temp:
            c=Coordinator(temp)
            try:
                c.create(MANIFEST,{'ELIGIBLE.cbl':COBOL})
                first=c.record_measurement(self.work());self.assertEqual(c.record_measurement(self.work()),first)
                self.assertEqual(len(c.ledger.measurements()),1)
                with self.assertRaises(ValidationError):c.record_measurement(self.work(hours=21))
                c.record_measurement(self.work(id='correction',hours=21,supersedes='pilot-1'))
                self.assertEqual(c.economics('process-a')['processes'][0]['work']['total_effort_hours'],'21')
                with self.assertRaises(ValidationError):c.record_measurement(self.work(id='fork',supersedes='pilot-1'))
                path=Path(temp)/'processes/process-a'/first['path'];path.write_text('{}')
                self.assertIn('pilot-1',c.economics()['integrity_errors'])
            finally:c.close()

    def test_session_clocks_capture_work_and_preserve_unknown_restart(self):
        from workbench.coordinator import Coordinator
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as temp:
            c=Coordinator(temp)
            try:
                session={'id':'work-1','process_id':None,'actor':'Claude-session','stage':'framework'}
                start=c.begin_work(session);self.assertEqual(c.begin_work(session),start)
                with self.assertRaises(ValidationError):c.begin_work({**session,'id':'work-2'})
                closed=c.end_work('work-1');self.assertTrue(closed['recorded'])
                self.assertTrue(c.end_work('work-1')['receipt_id']);self.assertEqual(len(c.ledger.measurements()),1)
                self.assertIsNotNone(c.economics()['framework']['recorded_detail_hours'])
                c.begin_work({**session,'id':'work-2'});c.close();c=Coordinator(temp)
                with self.assertRaisesRegex(ValidationError,'lost on restart'):c.end_work('work-2')
                self.assertEqual(c.end_work('work-2',abandon=True)['state'],'UNMEASURED')
                self.assertTrue(c.economics()['work_sessions'])
            finally:c.close()

    def test_http_retains_receipt_writer_while_copilot_mcp_cannot_mutate_economics(self):
        from fastapi.testclient import TestClient
        from workbench.api import create_app
        from tools.workbench_mcp import WorkflowBridge
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as temp:
            app=create_app(temp);c=app.state.coordinator
            try:
                client=TestClient(app,base_url='http://127.0.0.1:8765');token=client.get('/api/state').json()['token']
                record=self.work(process_id=None,stage='framework',complete=False,hours=8)
                self.assertEqual(client.post('/api/economics/receipts',json=record).status_code,403)
                headers={'Origin':'http://127.0.0.1:8765','X-Workbench-Token':token}
                self.assertEqual(client.post('/api/economics/receipts',json=record,headers=headers).status_code,200)
                self.assertEqual(client.get('/api/economics').json()['framework']['recorded_detail_hours'],'8')
                bridge=WorkflowBridge()
                self.assertTrue(client.post('/api/economics/receipts',json=record,headers=headers).json()['recorded'])
                self.assertEqual(client.get('/api/economics').json()['receipt_count'],1)
                start={'id':'w','process_id':None,'actor':'Claude-local-session','stage':'framework'}
                self.assertEqual(client.post('/api/economics/work/start',json=start,headers=headers).status_code,200)
                self.assertTrue(client.post('/api/economics/work/stop',json={'id':'w'},headers=headers).json()['recorded'])
                with patch.object(bridge,'request') as request:
                    for name,args in [('workbench_record_measurement',{'receipt':record}),('workbench_economics',{}),
                                      ('workbench_begin_work',start),('workbench_end_work',{'id':'w'})]:
                        with self.subTest(name=name), self.assertRaises(ValidationError):bridge.call(name,args)
                    request.assert_not_called()
            finally:c.close()

    def credit(self,**changes):
        return {'id':'credit-1','kind':'usage','process_id':'process-a','provider':'Copilot','account':'team','model':'all',
                'started_at':'2026-01-01T00:00:00Z','ended_at':'2026-01-02T00:00:00Z','quantities':{'credits':'5'},
                'coverage':'complete_process','evidence':'Fictional complete pilot credit statement','recorded_by':'fixture',**changes}

    def test_completed_migration_learns_actual_credits_without_reentering_rates(self):
        from workbench.economics import forecast
        plan={'cohorts':[{'id':'p','unit':'process','total':100}]}
        result=forecast([self.pilot()],[self.work(),self.credit()],plan,{})
        ai=result['cohorts'][0]['remaining_ai_usage'][0]
        self.assertEqual(ai['remaining']['base'],'495');self.assertEqual(ai['sample_size'],1)
        self.assertEqual(ai['basis'],'ATTRIBUTED_COMPLETE_PILOT_CREDITS');self.assertEqual(ai['unit'],'credits')
        self.assertIsNone(forecast([self.pilot()],[self.credit(coverage='partial')],plan,{})['cohorts'][0]['remaining_ai_usage'])

    def test_calibration_selection_does_not_erase_completed_estate_work(self):
        from workbench.economics import forecast
        processes=[self.pilot()]+[self.pilot(id='p'+str(i)) for i in range(9)]
        plan={'cohorts':[{'id':'p','unit':'process','total':100,'sample_process_ids':['process-a']}]}
        m=forecast(processes,[self.work()],plan,{})
        self.assertEqual(m['cohorts'][0]['completed_units'],10);self.assertEqual(m['remaining_effort_hours']['base'],'1800')

    def test_duplicate_credit_ids_and_cross_process_actor_overlap_are_visible(self):
        from workbench.economics import usage_summary,forecast,work_summary
        records=[self.credit(coverage='partial'),self.credit(id='duplicate',coverage='partial')]
        usage=usage_summary(records,'2026-01-10T00:00:00Z')
        self.assertTrue(usage['issues']);self.assertEqual(usage['quantities'],[])
        records=[self.work(),self.work(stage='conversion',complete=False,hours=8,id='work-detail'),
                 self.work(id='other-detail',stage='conversion',complete=False,hours=8,process_id='other')]
        self.assertFalse(work_summary(records,'process-a')['complete'])
        self.assertTrue(work_summary(records,'other')['issues'])

    def test_whole_pilot_credit_total_contains_details_without_double_counting(self):
        from workbench.economics import usage_summary
        details=self.credit(id='d',coverage='partial',quantities={'credits':'2'},ended_at='2026-01-01T01:00:00Z')
        usage=usage_summary([self.credit(),details],'2026-01-10T00:00:00Z')
        self.assertEqual(usage['quantities'][0]['quantity'],'5')
        self.assertFalse(usage['issues'])
        usage=usage_summary([self.credit(quantities={'credits':'1'}),details],'2026-01-10T00:00:00Z')
        self.assertTrue(usage['issues']);self.assertFalse(usage['quantities'])

    def test_forecast_numeric_bounds_and_disjoint_process_cohorts(self):
        from workbench.economics import forecast
        plan={'cohorts':[{'id':'p','unit':'process','total':100000000,'hours_per_unit':{'low':'1e15','base':'1e15','high':'1e15'}}], 'capacity_hours_per_week':'1e-30'}
        result=forecast([],[],plan,{})
        self.assertEqual(result['remaining_effort_hours']['base'],'100000000000000000000000')
        self.assertIsNotNone(result['capacity_weeks']['base'])
        with self.assertRaises(ValidationError):forecast([],[],{**plan,'capacity_hours_per_week':'1e-999'},{})
        cohorts={'cohorts':[{'id':'simple','unit':'process','total':10,'process_ids':['process-a']},
                            {'id':'complex','unit':'process','total':20,'process_ids':['b']}],'allocation_note':'Disjoint declared processes'}
        m=forecast([self.pilot(),self.pilot(id='b')],[self.work(),self.work(id='b',process_id='b',hours=40)],cohorts,{})
        self.assertEqual(m['remaining_effort_hours']['base'],'940')

    def test_preupgrade_partial_service_durations_cannot_calibrate_full_process(self):
        from workbench.economics import timing
        doc={'id':'p','created':'2025-01-01T00:00:00Z','status':'COMPLETED'}
        states=[{'process_id':'p','status':'VERIFYING','created':'2026-01-01T00:00:00Z'},
                {'process_id':'p','status':'COMPLETED','created':'2026-01-01T02:00:00Z'}]
        attempts=[{'process_id':'p','stage':'QUEUED_VERIFY','elapsed_seconds':'7200','outcome':'returned'}]
        m=timing(doc,states,attempts,'2026-01-10T00:00:00Z')
        self.assertEqual(m['service_hours'],'2');self.assertFalse(m['service_timings_complete']);self.assertIsNone(m['elapsed_hours'])

    def test_usage_persists_failure_and_cache_hit_without_extra_charge(self):
        from workbench.coordinator import Coordinator
        from workbench.provider import StructuredProvider
        from test_workflow import MANIFEST
        from test_source import COBOL
        from unittest.mock import patch
        import json
        with tempfile.TemporaryDirectory() as temp:
            c=Coordinator(temp)
            try:
                doc=c.create(MANIFEST,{'ELIGIBLE.cbl':COBOL},demo=True)
                c.provider=StructuredProvider('https://model.test/endpoint','fixture-model','not-a-real-token',True)
                response={'choices':[{'message':{'content':json.dumps({'summary':'fixture','assumptions':[],'questions':[]})}}], 'usage':{'prompt_tokens':123,'completion_tokens':12}}
                with patch('workbench.provider.post_json',return_value=(response,{})):
                    c.observed_analysis(doc,'Fictional source');c.observed_analysis(doc,'Fictional source')
                failed={**response,'choices':[{'message':{'content':'{}'}}]}
                with patch('workbench.provider.post_json',return_value=(failed,{})),self.assertRaises(ValidationError):c.observed_analysis(doc,'Different fictional source')
                before=c.economics('process-a')['usage'];self.assertEqual(before['provider_calls']['requests'],2)
                self.assertEqual(before['provider_calls']['cache_hits'],1)
                self.assertEqual(next(r['quantity'] for r in before['quantities'] if r['unit']=='input_tokens'),'246')
                c.close();c=Coordinator(temp)
                self.assertEqual(c.economics('process-a')['usage'],before)
                self.assertEqual(before['billing_state'],'UNKNOWN')
            finally:c.close()

    def test_work_close_retry_after_receipt_commit_recovers_without_double_count(self):
        from workbench.coordinator import Coordinator
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as temp:
            c=Coordinator(temp)
            try:
                c.begin_work({'id':'clock','process_id':None,'actor':'fixture','stage':'framework'})
                with patch.object(c.ledger,'close_work_session',side_effect=OSError('crash after receipt')),self.assertRaises(OSError):c.end_work('clock')
                self.assertEqual(len(c.ledger.measurements()),1)
                c.close();c=Coordinator(temp);self.assertTrue(c.end_work('clock')['recorded'])
                self.assertEqual(len(c.ledger.measurements()),1)
            finally:c.close()

    def test_stale_optional_plan_does_not_block_existing_process_reporting(self):
        from workbench.coordinator import Coordinator
        from workbench.economics import economics_view
        from unittest.mock import patch
        from test_workflow import MANIFEST
        from test_source import COBOL
        with tempfile.TemporaryDirectory() as temp:
            c=Coordinator(temp)
            try:
                c.create(MANIFEST,{'ELIGIBLE.cbl':COBOL})
                c.record_measurement({'id':'plan','kind':'plan','process_id':None,'evidence':'Fictional estate','recorded_by':'fixture',
                                      'plan':{'cohorts':[{'id':'p','unit':'process','total':0}]}})
                with patch('workbench.economics.forecast',side_effect=ValidationError('Forecast scope is smaller than recorded completed scope')):
                    model=c.economics('process-a')
                self.assertIsNone(model['forecast']);self.assertTrue(model['forecast_errors'])
                self.assertEqual(c.ledger.get('process-a')['status'],'READY')
            finally:c.close()

class CopilotCountTests(unittest.TestCase):
    def receipt(self,**kw):
        return {'id':'credits','kind':'usage','process_id':'p','provider':'GitHub Copilot','account':'team','model':'all',
                'started_at':'2026-01-01T00:00:00Z','ended_at':'2026-01-02T00:00:00Z','quantities':{'credits':'5'},'evidence':'Fictional billing record','recorded_by':'test',**kw}
    def test_only_actual_credit_receipts_count_and_zero_remains_known(self):
        from workbench.economics import usage_summary
        result=usage_summary([self.receipt(quantities={'credits':'0'}),self.receipt(id='tokens',provider='Claude',quantities={'input_tokens':2000})],'2026-10-06T00:00:00Z','p')
        self.assertEqual(result['copilot_credits']['accounts'][0]['credits_used'],'0')
        self.assertEqual(usage_summary([],'2026-10-06T00:00:00Z')['copilot_credits']['status'],'UNKNOWN')
    def test_overlapping_copilot_credits_are_withheld_without_token_fallback(self):
        from workbench.economics import usage_summary
        result=usage_summary([self.receipt(),self.receipt(id='duplicate')],'2026-10-06T00:00:00Z','p')
        self.assertEqual(result['copilot_credits']['status'],'UNKNOWN');self.assertTrue(result['issues'])
    def test_excluded_scope_cannot_calibrate_full_estate_conversion(self):
        from workbench.economics import forecast
        pilot={'id':'p','status':'COMPLETED','report_verified':True,'measurement_report_intact':True,'coverage_summary':{'completion_eligible':True},'requirements':{'excluded_ids':['omitted']}}
        result=forecast([pilot],[self.receipt(coverage='complete_process')],{'cohorts':[{'id':'all','unit':'process','total':100}]},{})
        self.assertEqual(result['cohorts'][0]['remaining_units'],100);self.assertIsNone(result['cohorts'][0]['remaining_ai_usage'])
