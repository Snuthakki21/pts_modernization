"""Evidence-based reverse engineering, distinct from conversion support."""
import copy
from pathlib import Path
import unittest

from workbench.domain import ValidationError, sha
from workbench.program_insights import program_insights, render_program_insights
from workbench.source import analyze_sources
from workbench.mainframe import load_knowledge
from test_source import COBOL


def fixture():
    files = {'ELIGIBLE.cbl': COBOL, 'AGE.dclgen': 'EXEC SQL DECLARE APP.AGE TABLE (AGE INTEGER) END-EXEC.\n01 HOST-AGE PIC 9(3).\n',
             'CARDS.cntl': 'SORT FIELDS=(1,3,CH,A)\n', 'SCREENS.bms': 'MAPSET DFHMSD TYPE=MAP\nSCREEN1 DFHMDI SIZE=(24,80)\nSCREEN2 DFHMDI SIZE=(24,80)\n',
             'PROC.sql': "-- CREATE PROCEDURE FALSE1()\nVALUES 'CREATE PROCEDURE FALSE2()';\n/* CREATE PROCEDURE FALSE3() */\nCREATE\n PROCEDURE APP.REALPROC(IN P INTEGER) LANGUAGE SQL BEGIN VALUES P; END;\nCALL APP.NOT_A_DEFINITION();\n",
             'MYSTERY.bin': 'unclassified opaque source\n', 'OTHER.cbl': COBOL.replace('ELIGIBLE', 'OTHER')}
    doc = {'id': 'insight-test', 'status': 'WAITING_COPILOT', 'jobs': [{'name': 'JOBA', 'steps': [{'name': 'S1', 'program': 'ELIGIBLE', 'inputs': [], 'outputs': [], 'condition': 'Always'}]}],
           'transactions': [{'id': 'ELIG', 'program': 'ELIGIBLE', 'mapset': 'MAPSET', 'map': 'SCREEN1'}],
           'mainframe_knowledge': load_knowledge(Path('.')), 'lineage_scope': [p for p in files if p != 'OTHER.cbl'], 'sme_packet_version': 4,
           'artifacts': [], 'blockers': []}
    doc['analysis'] = analyze_sources(files, doc)
    return doc


class ProgramInsightsTests(unittest.TestCase):
    def test_no_analysis_is_unknown_not_empty_estate(self):
        view = program_insights({'id': 'p'})
        self.assertIsNone(view['inventory']['retained_files'])
        self.assertTrue(all(row['file_count'] is None for row in view['inventory']['categories']))
        self.assertEqual(view['basis'], 'NOT_ANALYZED')

    def test_definition_counts_are_not_file_counts_or_conversion_claims(self):
        view = program_insights(fixture())
        kinds = {r['kind']: r for r in view['inventory']['categories']}
        self.assertEqual(kinds['cobol_program']['file_count'], 2)
        self.assertEqual(kinds['cobol_program']['selected_file_count'], 1)
        self.assertEqual(kinds['bms_map']['definition_count'], 2)
        self.assertEqual(kinds['bms_map']['file_count'], 1)
        self.assertEqual(kinds['dclgen']['file_count'], 1)
        self.assertEqual(kinds['utility_control']['file_count'], 1)
        self.assertEqual(view['inventory']['sql_procedures']['observed_definitions'], 1)
        self.assertIsNone(view['inventory']['sql_procedures']['estate_total'])
        self.assertEqual(view['inventory']['sql_procedures']['evidence'][0]['name'], 'APP.REALPROC')
        self.assertEqual(view['inventory']['sql_procedures']['evidence'][0]['start_line'], 4)
        self.assertGreater(kinds['unknown']['file_count'], 0)
        self.assertIn('Recognition', view['boundary'])

    def test_program_documents_trace_job_transaction_rules_and_hashes(self):
        doc = fixture(); before = copy.deepcopy(doc)
        view = program_insights(doc)
        program = next(p for p in view['programs'] if p['name'] == 'ELIGIBLE')
        self.assertEqual(program['rule_counts']['identified'], 2)
        self.assertIsNone(program['rule_counts']['converted_verified'])
        self.assertEqual(program['rules'][0]['category'], 'unclassified')
        self.assertEqual(program['rules'][0]['selection'], 'Yes (default, not saved)')
        self.assertEqual(program['source']['source_hash'], sha(COBOL))
        self.assertEqual({m['kind'] for m in program['entrypoints']}, {'job_step', 'transaction'})
        self.assertIn('Business purpose', program['functional_summary'])
        self.assertEqual(doc, before)

    def test_excluded_rule_has_exact_reason_and_preserved_logic(self):
        doc = fixture(); p = doc['analysis']['programs']['ELIGIBLE']; rule = p['rules'].pop()
        p['omitted_rules'] = [rule]
        view = program_insights(doc)
        item = next(r for p in view['programs'] if p['name'] == 'ELIGIBLE' for r in p['rules'] if r['id'] == rule['id'])
        self.assertEqual(item['selection'], 'No')
        self.assertEqual(item['reason'], 'Not converted because selected No in requirements.')
        self.assertTrue(item['description'])

    def test_technical_fields_preserve_frozen_record_layout(self):
        doc = fixture()
        program = next(p for p in program_insights(doc)['programs'] if p['name'] == 'ELIGIBLE')
        fields = {field['name']: field for field in program['technical_spec']['fields']}
        self.assertEqual(fields['AGE']['width'], 3)
        self.assertEqual(fields['ACTIVE']['width'], 1)
        self.assertEqual(fields['DECISION']['width'], 1)
        for name, field in fields.items():
            frozen = doc['analysis']['programs']['ELIGIBLE']['fields'][name]
            self.assertEqual(field['group'], frozen['group'])
            self.assertEqual(field['storage_section'], frozen['storage_section'])
            self.assertEqual(field['source_ref'], frozen['source_ref'])

    def test_entry_mappings_follow_program_calls_not_colocated_declarations(self):
        from workbench.lineage import map_lineage
        doc = fixture(); doc.pop('lineage_scope'); doc['transactions'] = []
        entry = COBOL.replace('  GOBACK.', '  CALL "CALLED".\n  GOBACK.')
        files = {'MULTI.cbl': entry + 'END PROGRAM ELIGIBLE.\n' + COBOL.replace('ELIGIBLE', 'UNRELATED'),
                 'CALLED.cbl': COBOL.replace('ELIGIBLE', 'CALLED'),
                 'JOBA.jcl': '//JOBA JOB\n//S1 EXEC PGM=ELIGIBLE\n'}
        doc['analysis'] = analyze_sources(files, doc)
        doc['lineage'] = map_lineage(files, doc)
        programs = {p['name']: p for p in program_insights(doc)['programs']}
        self.assertEqual(programs['UNRELATED']['entrypoints'], [])
        for name in ('ELIGIBLE', 'CALLED'):
            self.assertEqual([(entry['label'], entry['entry_program']) for entry in programs[name]['entrypoints']],
                             [('JOBA.S1', 'ELIGIBLE')])

    def test_unknowns_out_of_scope_and_conflicts_remain_visible(self):
        doc = fixture(); original_conflicts = sum(c['kind'] == 'ambiguous' for c in doc['analysis']['classifications'].values())
        doc['analysis']['classifications']['CARDS.cntl']['kind'] = 'ambiguous'
        view = program_insights(doc)
        self.assertEqual(next(p for p in view['programs'] if p['name'] == 'OTHER')['scope'], 'OUTSIDE_SELECTED_PROCESS')
        self.assertEqual(next(c for c in view['inventory']['categories'] if c['kind'] == 'ambiguous')['file_count'], original_conflicts + 1)
        self.assertTrue(view['inventory']['unresolved_classifications'])

    def test_cursors_reject_bools_negatives_and_overflow(self):
        for cursor in (True, -1, 500):
            with self.subTest(cursor=cursor), self.assertRaises(ValidationError):program_insights(fixture(), after=cursor)

    def test_html_escapes_source_and_links_only_local_reports(self):
        doc = fixture(); doc['analysis']['programs']['ELIGIBLE']['rules'][0]['plain'] = '<script>alert(1)</script>'
        html = render_program_insights(program_insights(doc))
        self.assertNotIn('<script>', html)
        self.assertIn('&lt;script&gt;', html)
        self.assertIn('href="rules.html"', html)
        self.assertIn('program-', html)
        self.assertIn('SHA256', html)

    def test_factory_projection_keeps_existing_capabilities_and_adds_insights(self):
        from workbench.factory import bounded_view
        view = bounded_view(fixture(), program_after=1)
        self.assertEqual(view['program_insights']['after'], 1)
        self.assertIn('capabilities', view)


    def test_twenty_runtime_randomized_definition_states_ignore_nondefinitions(self):
        import random
        import secrets
        seed = secrets.randbits(63); rng = random.Random(seed)
        for index in range(20):
            name = 'P' + str(rng.randrange(10**9, 10**10))
            fake = 'FAKE' + str(index)
            gap = rng.choice([' ', '\n', '\n  ', ' \n\t'])
            text = (f"-- CREATE PROCEDURE {fake}()\n"
                    f"VALUES 'CREATE PROCEDURE {fake}()';\n"
                    f"/* outer /* CREATE PROCEDURE {fake}() */ still comment */\n"
                    f"CREATE{gap}PROCEDURE APP.{name}(IN P INTEGER) LANGUAGE SQL BEGIN VALUES P; END;\n"
                    f"CALL APP.{fake}();\n")
            doc = fixture(); doc['analysis']['assets'] = [{'path':'probe.sql', 'source_text':text, 'source_hash':sha(text), 'kind':'other_source'}]
            doc['analysis']['programs'] = {}; doc['analysis']['classifications'] = {'probe.sql': {'kind':'unknown'}}
            with self.subTest(seed=seed, case=index):
                result = program_insights(doc)['inventory']['sql_procedures']
                self.assertEqual(result['observed_definitions'], 1)
                self.assertEqual(result['evidence'][0]['name'], 'APP.' + name)
                self.assertEqual(result['evidence'][0]['source_hash'], sha(text))
                self.assertEqual(result['evidence'][0]['start_line'], 4)

    def test_pages_preserve_same_named_programs_in_distinct_files(self):
        doc = fixture(); source = doc['analysis']['assets'][0]
        doc['analysis']['assets'] = [{'path':f'folder-{i}/ELIGIBLE.cbl', 'source_text':COBOL, 'source_hash':sha(COBOL), 'kind':'cobol_program', 'selected':False} for i in range(23)]
        first = program_insights(doc); second = program_insights(doc, after=first['next_after'])
        self.assertEqual(first['program_total'], 23)
        self.assertEqual(len(first['programs']), 20)
        self.assertEqual(len(second['programs']), 3)
        self.assertFalse(second['has_more'])
        self.assertEqual(len({p['id'] for p in first['programs'] + second['programs']}), 23)
        self.assertTrue(all(p['analysis_state'] == 'NOT_SEMANTICALLY_ANALYZED' for p in first['programs']))

    def test_issued_report_pins_consolidated_program_knowledge(self):
        import json
        import tempfile
        from io import BytesIO
        from openpyxl import load_workbook
        from workbench.coordinator import Coordinator
        from test_factory import ONLINE
        with tempfile.TemporaryDirectory() as temp:
            coordinator = Coordinator(temp)
            try:
                coordinator.create(ONLINE, {'ELIGIBLE.cbl': COBOL}, demo=True)
                coordinator.start('online-pilot'); coordinator.advance('online-pilot')
                book = load_workbook(coordinator.artifact('online-pilot', 'review/sme-checklist.xlsx'))
                for row in book['Checklist'].iter_rows(min_row=2): row[4].value = 'Yes'
                output = BytesIO(); book.save(output); book.close()
                coordinator.import_answers('online-pilot', output.getvalue(), 'Fictional documentation test reviewer')
                coordinator.advance('online-pilot'); coordinator.advance('online-pilot')
                root = 'reports/report-0001/'
                model = json.loads(coordinator.artifact('online-pilot', root + 'program-insights.json').read_text())
                receipt = json.loads(coordinator.artifact('online-pilot', root + 'inspection.json').read_text())
                self.assertEqual(model['verification_basis'], 'REPLAYED_COVERAGE')
                self.assertEqual(model['programs'][0]['rule_counts']['identified'], 2)
                for name in ('program-insights.json', 'program-insights.html'):
                    self.assertEqual(receipt['sha256'][name], sha(coordinator.artifact('online-pilot', root + name).read_bytes()))
                self.assertIn('program-insights.html', coordinator.artifact('online-pilot', root + 'factory.html').read_text())
            finally: coordinator.close()


if __name__ == '__main__': unittest.main()
