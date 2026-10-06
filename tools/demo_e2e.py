"""Fictional-only full flow, including returned SME file and reuse metrics."""
import argparse
from pathlib import Path
from io import BytesIO
import json
from openpyxl import load_workbook
from workbench.coordinator import Coordinator
from workbench.reports import portfolio
from workbench.domain import require

def run(root):
    example=Path(__file__).parent.parent/'examples';manifest=(example/'process-input.md').read_text()
    sources={p.name:p.read_text() for p in (example/'Endeavor').iterdir()}
    c=Coordinator(root)
    try:
        for pid in ['example-referral','example-reuse']:
            doc=c.create(manifest.replace('example-referral',pid),sources,False)
            # Counting is intentionally exercised in this isolated fixture workspace;
            # carry an explicit disclosure into every generated presentation slide.
            doc['fixture_only']=True;c.ledger.save(doc)
            c.start(pid);c.advance(pid)
            book=load_workbook(c.artifact(pid,'review/sme-checklist.xlsx'))
            for row in book['Checklist'].iter_rows(min_row=2):row[4].value='Yes';row[6].value='Fictional fixture reviewer'
            out=BytesIO();book.save(out);book.close();c.import_answers(pid,out.getvalue(),'Fictional fixture reviewer')
            c.advance(pid);c.advance(pid);doc=c.ledger.get(pid)
            require(doc['status']=='COMPLETED','Fictional supported process did not complete: '+str(doc['blockers']))
        pf=portfolio(c.ledger);require(pf['unique_program_versions']==1 and pf['program_memberships']==2,'Shared program deduplication failed')
        return {'fixture_only':True,'portfolio':pf,'processes':[{'id':p['id'],'status':p['status'],'cases':sum(len(v['actual']) for v in p['runs'][-1]['programs'].values())} for p in c.ledger.list()]}
    finally:c.close()

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--root',required=True);args=p.parse_args();print(json.dumps(run(args.root),indent=2))
